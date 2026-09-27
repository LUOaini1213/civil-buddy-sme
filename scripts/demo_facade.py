#!/usr/bin/env python3
"""Façade demo: one curtain-wall subcontractor's jobs, run on the SYNTHETIC files in examples/facade-demo.

  1 linked     the tender and the packing as ONE run: the ITT's logistics clauses, a loading plan from the panel
               list under the clause's container type, the English statements tied to clause and plan figure, and
               the link record; then a revised panel list (rev B) re-run, and the statements that changed are named
  2 tender     解析招标 facade_itt_doc.md, then 技术标 and 废标检查 read the same session's hand-off
  3 packing    按 facade_panels_zh.xlsx 装柜 (and the English list): container plan + the conservation line
  4 site docs  项目日报 from daily_report_input.txt; 安全交底 from wah_briefing_input.txt, a high-risk post
               that writes nothing until a licensed person types the sign-off sentence. This script never
               types it: pass --sign "<the sentence>" yourself, the way `civil exec --confirm` is yours to pass.

Every turn goes through civil.run_task (what `civil exec` calls) in steps mode: offline, no model key.
It runs inside a job folder this script creates (default: a new temp folder, kept so the drafts can be
opened) and writes nothing outside it. Exit 1 when a flow errors, 2 on bad arguments.

  python scripts/demo_facade.py [--job NEW_DIR] [--sign SENTENCE]
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional

sys.dont_write_bytecode = True          # not even __pycache__ outside the job folder
ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples" / "facade-demo"
INPUTS = ("facade_itt_doc.md", "facade_panels.xlsx", "facade_panels_zh.xlsx", "facade_panels_rev_b.xlsx", "daily_report_input.txt",
          "wah_briefing_input.txt")
PROJECT = "合成示例办公楼幕墙分包"
SESSION = "civil-cli"
TENDER_ROWS = ("招标编号", "注册资格/工作类别", "投标截止", "工期", "投标有效期", "履约担保", "缺陷责任期/质保期", "备选投标方案")
# Section 4 of the ITT, and two appendix rows, by a token only that clause carries.
FACADE_CLAUSES = (("PMU mock-up", "PMU"), ("VMU mock-up", "VMU"), ("heat soak", "14179"), ("site water test", "hose"),
                  ("PE-endorsed calcs/shop drawings", "Professional Engineer"), ("warranty", "water tightness"),
                  ("A-frame delivery", "A-frame"), ("40HQ containers", "high cube"), ("container gross mass", "gross mass"),
                  ("CTU Code securing", "CTU Code"), ("delivery sequence", "installation programme"), ("insurance", "All Risks"))
APPENDIX_ROWS = (("liquidated damages", "per day of delay"), ("retention", "progress payment"))
PANELS = ("facade_panels_zh.xlsx", "facade_panels.xlsx")
LINK_SESSION = "civil-link"
LINK_ASK = "Link the tender facade_itt_doc.md to the packing list facade_panels.xlsx and write the logistics response"
LINK_REV_B = "招标装柜联动：按招标 facade_itt_doc.md 和改版装箱单 facade_panels_rev_b.xlsx 重出物流应答"
DAILY_ROWS = ("日期", "天气", "部位", "形象进度", "出勤", "机械材料", "安全质量记事", "明日计划")


def tables(md: str) -> List[Dict[str, str]]:
    """Every row of every Markdown table, as {column: cell}."""
    rows: List[Dict[str, str]] = []
    header: Optional[List[str]] = None
    for raw in md.splitlines():
        line = raw.strip()
        if not (line.startswith("|") and line.endswith("|")):
            header = None
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        if header is None:
            header = cells
            continue
        rows.append(dict(zip(header, cells)))
    return rows


def first_cells(md: str) -> Dict[str, List[str]]:
    """First cell -> the second cell of each row that starts with it."""
    found: Dict[str, List[str]] = {}
    for row in tables(md):
        values = list(row.values())
        if len(values) > 1:
            found.setdefault(values[0], []).append(values[1])
    return found


def section(md: str, title: str) -> str:
    match = re.search(rf"^## {re.escape(title)}[^\n]*\n(.*?)(?=^## |\Z)", md, re.S | re.M)
    return match.group(1).strip() if match else ""


def draft(out: Dict[str, Any]) -> str:
    path = next((Path(f["path"]) for f in out.get("files") or [] if str(f.get("path", "")).endswith(".md")), None)
    return path.read_text(encoding="utf-8") if path and path.is_file() else ""


class Demo:
    def __init__(self, job: Path, sign: str = "") -> None:
        self.job, self.sign = job, sign
        self.errors: List[str] = []
        self.result: Dict[str, Any] = {"job": str(job)}

    def say(self, line: str = "") -> None:
        print(line, flush=True)

    def rel(self, path: Any) -> str:
        try:
            return Path(str(path)).resolve().relative_to(self.job).as_posix()
        except (OSError, ValueError):
            return str(path)

    def turn(self, flow: str, text: str, skill: str, *, session: str = SESSION, confirm: bool = False,
             command: str = "") -> Dict[str, Any]:
        from packing_assistant.civil import run_task

        self.say(f"  $ civil exec {command or repr(text)}")
        out = run_task(text, session_id=session, confirm=confirm)
        if out.get("skill") != skill:
            self.errors.append(f"{flow}: routed to {out.get('skill')!r}, expected {skill!r}")
        self.say("  reply: " + str(out.get("reply") or "").splitlines()[0][:160] if out.get("reply") else "  reply: (none)")
        files = [self.rel(f.get("path")) for f in out.get("files") or [] if isinstance(f, dict) and f.get("path")]
        if files:
            self.say("  wrote: " + ", ".join(files))
        return out

    def expect_written(self, flow: str, out: Dict[str, Any]) -> str:
        md = draft(out)
        if not (out.get("ok") and out.get("wrote") and md):
            self.errors.append(f"{flow}: no draft written (ok={out.get('ok')}, error={out.get('error_code')})")
        return md

    # 1 ---------------------------------------------------------------------------------------------
    @staticmethod
    def figure(statement: Dict[str, Any]) -> str:
        """The plan figure a statement rests on, in one line."""
        f = statement.get("figures") or {}
        kind = statement.get("kind")
        if kind == "container_type":
            return f"clause names {f.get('clause_names') or '-'}; plan made in {f.get('plan_container_type') or f.get('container_type')}"
        if kind == "containers_used" and f.get("containers_used") is not None:
            return (f"{f.get('containers_used')} x {f.get('container_type')} (N0 {f.get('n0')}) for {f.get('pieces')} pieces / "
                    f"{f.get('cargo_net_kg'):,.0f} kg net from {f.get('panel_list')}")
        if kind == "gross_mass" and f.get("max_gross_kg") is not None:
            return (f"heaviest container {f.get('max_gross_kg'):,} kg gross ({f.get('max_cargo_kg'):,} cargo + {f.get('container_tare_kg'):,} tare)"
                    f" vs limit {f.get('limit_kg'):,.0f} kg" + (f", margin {f.get('margin_kg'):,}" if f.get("margin_kg") is not None else ""))
        if kind == "crate_structure":
            return f"{f.get('pending_design')} of {f.get('n_boxes')} crates pending detailed design"
        return f"not modelled -> {statement.get('owner')}"

    def link_turn(self, text: str, label: str) -> Dict[str, Any]:
        out = self.turn("linked", text, "bid-parse", session=LINK_SESSION)
        link = out.get("tender_packing_link") or {}
        if not (out.get("ok") and link.get("statements")):
            self.errors.append(f"linked ({label}): no linked response (ok={out.get('ok')}, error={out.get('error_code')})")
            return {}
        record = next((Path(f["path"]) for f in out.get("files") or [] if str(f.get("path", "")).endswith("tender-packing-link.json")), None)
        if record is None or not record.is_file():
            self.errors.append(f"linked ({label}): no link record written")
        inputs = link.get("inputs") or {}
        self.say(f"    container type: {(link.get('container') or {}).get('reason')}")
        self.say("    inputs: " + " · ".join(f"{k} {(inputs.get(k) or {}).get('name')} sha256 {str((inputs.get(k) or {}).get('sha256'))[:12]}"
                                            for k in ("tender", "panel_list", "plan")))
        for s in link["statements"]:
            self.say(f"    {s['id']} {s.get('cite') or 'Clause ' + (s.get('clause') or '-')} · {s['kind']} · {s['status']} · {self.figure(s)}")
        if record is not None:
            self.say(f"    link record: {self.rel(record)} (statement -> clause -> plan figures -> sha256 of tender, list, plan)")
        return {"out": out, "link": link, "record": record}

    def linked(self) -> None:
        from packing_assistant.tender_packing_link import KIND_TITLE

        self.say("\n== 1 Tender <-> packing, linked: the ITT's logistics clauses, the plan under them, the English statements")
        first = self.link_turn(LINK_ASK, "first run")
        if not first:
            return
        link = first["link"]
        clauses = link.get("clauses") or []
        self.say("    logistics clauses read from the ITT: " + "; ".join(f"{c['clause']} {'/'.join(c['kinds'])}" for c in clauses))
        by_kind = {s["kind"]: s for s in link["statements"]}
        plan = link.get("plan") or {}
        clause_type = next((c for c in clauses if "container_type" in c["kinds"]), None)
        named = str((by_kind.get("container_type") or {}).get("figures", {}).get("clause_names") or "").split(", ")
        if clause_type and plan.get("container_type") not in named:
            self.errors.append(f"linked: plan in {plan.get('container_type')}, the ITT's Clause {clause_type['clause']} names {named}")
        for kind in ("securing", "handling", "delivery_sequence"):
            if by_kind.get(kind, {}).get("status") == "covered":
                self.errors.append(f"linked: {kind} marked covered, the plan does not model it")
        s2 = by_kind.get("containers_used") or {}
        self.say(f"    {KIND_TITLE['containers_used']} as the English bid-book states it ({s2.get('id')}): {s2.get('text')}")
        bidbook = next((self.rel(f["path"]) for f in first["out"].get("files") or [] if str(f.get("path", "")).endswith("bidbook.en.md")), "")
        self.say(f"    English bid-book (qualifications and price stay [TO FILL] for people): {bidbook}")

        self.say("  A revised panel list arrives (rev B: level L9 added, L8 panels heavier). Same request, new list:")
        second = self.link_turn(LINK_REV_B, "rev B")
        if not second:
            return
        changes = second["link"].get("changes_since_previous") or {}
        self.say(f"    since the previous run: {changes.get('summary')}")
        for item in changes.get("changed") or []:
            moved = "; ".join(f"{k} {a} -> {b}" for k, (a, b) in item["figures"].items() if k != "rows_per_container")
            self.say(f"      {item['id']} changed: {moved or 'rows per container moved'} -> re-confirm")
        if changes.get("unchanged"):
            self.say("      re-derived with the same figures: " + ", ".join(c["id"] for c in changes["unchanged"]))
        if not changes.get("needs_reconfirmation"):
            self.errors.append("linked: the revised panel list changed no statement")
        self.result["linked"] = {"first": {k: first["link"].get(k) for k in ("container", "inputs", "statements", "plan")},
                                 "rev_b": {k: second["link"].get(k) for k in ("container", "inputs", "statements", "plan")},
                                 "changes": changes, "record": str(second["record"]) if second.get("record") else "",
                                 "submit_blocked": second["out"].get("submit_blocked")}
        self.say("  sign-off: nothing here is booked or submitted (submit_blocked stays true). A person confirms the loading plan"
                 " before booking, and re-confirms every statement the re-run names.")

    # 2 ---------------------------------------------------------------------------------------------
    def tender(self) -> None:
        self.say("\n== 2 Tender review: parse the ITT, then the bid posts read its hand-off")
        out = self.turn("tender", "解析招标 facade_itt_doc.md", "bid-parse")
        md = self.expect_written("tender", out)
        cells = first_cells(md)
        rows = {name: cells.get(name, ["(no row)"])[0] for name in TENDER_ROWS}
        for name, value in rows.items():
            self.say(f"    {name}: {value}")
        scoring = [v[0] for k, v in cells.items() if k.startswith("评分点：")]
        rejections = [v[0] for k, v in cells.items() if k.startswith("否决条款")]
        blanks = sum(cell.count("未在原文检出") for values in cells.values() for cell in values)
        shown = [name for name, token in FACADE_CLAUSES if token in md]
        self.say(f"    scoring points: {len(scoring)} ({'; '.join(scoring)})")
        self.say(f"    rejection clauses listed one by one: {len(rejections)}")
        self.say(f"    rows left 未在原文检出 for a person to check against the ITT: {blanks}")
        self.say(f"    façade specification clauses (ITT section 4) shown in tender.parse.md: {len(shown)} of {len(FACADE_CLAUSES)}"
                 + (f"; not yet: {', '.join(n for n, _ in FACADE_CLAUSES if n not in shown)}" if len(shown) < len(FACADE_CLAUSES) else ""))
        self.say("    appendix rows shown: " + ", ".join(f"{n} {'yes' if t in md else 'no'}" for n, t in APPENDIX_ROWS))
        self.result["tender"] = {"ok": out.get("ok"), "skill": out.get("skill"), "rows": rows, "scoring": scoring,
                                 "rejections": rejections, "facade_clauses_shown": shown, "submit_blocked": out.get("submit_blocked")}
        tech = self.turn("bid-tech", "出一份技术标提纲", "bid-tech")
        tech_md = self.expect_written("bid-tech", tech)
        self.say(f"    chapters from scoring points: {len(re.findall(r'^### ', tech_md, re.M))} · "
                 f"cells left [A001] for the bid team: {tech_md.count('[A001]')}")
        comp = self.turn("bid-compliance", "出一份废标检查表", "bid-compliance")
        comp_md = self.expect_written("bid-compliance", comp)
        states = [row.get("三态", "") for row in tables(comp_md)]
        self.say(f"    requirements with no response yet (未响应): {states.count('未响应')} · rejection clauses to tick by hand: "
                 f"{len(tables(section(comp_md, '10 否决与拒收条款自查清单')))}")
        self.result["bid_tech"] = {"ok": tech.get("ok"), "skill": tech.get("skill"), "written": bool(tech_md)}
        self.result["bid_compliance"] = {"ok": comp.get("ok"), "skill": comp.get("skill"), "written": bool(comp_md),
                                         "unanswered": states.count("未响应")}
        self.say("  sign-off: bid posts are low risk, so no sentence is asked; submit_blocked stays true and every"
                 " P0 (qualification / rejection) row waits for a person.")

    # 3 ---------------------------------------------------------------------------------------------
    def packing(self) -> None:
        self.say("\n== 3 Packing and shipping: 24 unitised panels into 40HQ containers")
        plans = {}
        for name, session in zip(PANELS, (SESSION, SESSION + "-en")):
            out = self.turn("packing", f"按 {name} 装柜，柜型 40HQ", "pack-ship", session=session)
            plan = (out.get("pack_ship") or {}).get("plan") or {}
            cons = plan.get("conservation") or {}
            if not (out.get("ok") and plan.get("source") == "solver" and plan.get("can_fit") is True):
                self.errors.append(f"packing {name}: no usable plan (ok={out.get('ok')}, error={out.get('error_code')})")
            if cons.get("ok") is not True:
                self.errors.append(f"packing {name}: conservation check did not pass: {cons}")
            structure = plan.get("structure") or {}
            self.say(f"    plan: {plan.get('containers_used')} × {plan.get('container_type')} · N0 {plan.get('n0')}"
                     f" · can_fit {plan.get('can_fit')} · binding {plan.get('binding_constraint')} · crates {plan.get('n_boxes')} for {cons.get('pieces_in')} panels"
                     f" · space {plan.get('utilization')} · weight {plan.get('weight_utilization')}")
            self.say(f"    conservation (list -> crates): pieces {cons.get('pieces_in')} -> {cons.get('pieces_out')} · "
                     f"kg {cons.get('kg_in')} -> {cons.get('kg_out')} · {'ok' if cons.get('ok') else 'FAILED'}")
            self.say(f"    crate structure: pass {structure.get('pass')} · reinforce {structure.get('needs_reinforcement')} · fail "
                     f"{structure.get('fail')} · pending detailed design {structure.get('pending_design')} (the engine does not invent a pass)")
            plans[name] = plan
        control = self.without_notes(self.job / "inputs" / "facade_panels.xlsx")
        out = self.turn("packing", f"按 {control.name} 装柜，柜型 40HQ", "pack-ship", session=SESSION + "-control")
        plans[control.name] = (out.get("pack_ship") or {}).get("plan") or {}
        keys = ("containers_used", "n0", "n_boxes", "utilization", "weight_utilization", "binding_constraint")
        changed = [name for name in PANELS if any(plans[name].get(k) != plans[control.name].get(k) for k in keys)]
        self.say("    handling notes (glass / upright / no stack), against the same list with the notes removed: "
                 + (f"they change the plan for {', '.join(changed)}" if changed else "no effect on the plan, in either language"))
        library = (ROOT / "knowledge" / "packing_knowledge_base.json").read_text(encoding="utf-8").lower()
        if "stillage" not in library and "a-frame" not in library:
            self.say("    not modelled: A-frame stillages (the ITT asks for them). That needs the contractor's"
                     " stillage size, tare and capacity.")
        zh = plans["facade_panels_zh.xlsx"]
        self.result["packing"] = {"plan": {k: zh.get(k) for k in (*keys, "container_type", "can_fit")},
                                  "conservation": zh.get("conservation"), "notes_change_plan": changed}
        self.say("  sign-off: pack-plan.md is an internal draft (不可直接订舱); a person confirms the plan in the workbench"
                 " HITL step, and lashing / VGM are signed separately.")

    @staticmethod
    def without_notes(source: Path) -> Path:
        """The control list: the same rows with the note column emptied, next to the source in the job folder."""
        import openpyxl

        target = source.with_name("panels_no_notes.xlsx")
        wb = openpyxl.load_workbook(source)
        ws = wb["materials"]
        column = [c.value for c in ws[1]].index("note") + 1
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=column).value = None
        wb.save(target)
        return target

    # 4 ---------------------------------------------------------------------------------------------
    def daily(self) -> None:
        self.say("\n== 4 Site documents: daily report, then the work-at-height briefing")
        daily_text = (self.job / "inputs" / "daily_report_input.txt").read_text(encoding="utf-8").strip()
        out = self.turn("daily", daily_text, "pm-daily", command="- < inputs/daily_report_input.txt")
        md = self.expect_written("daily", out)
        cells = first_cells(md)
        for name in DAILY_ROWS:
            self.say(f"    {name}: {cells.get(name, ['(no row)'])[0][:90]}")
        blanks = [who for who in ("填报人", "审核人") if f"{who}空栏" in md]
        if blanks:
            self.say(f"  sign-off: {' / '.join(blanks)} left blank for a person; low-risk post, no sentence asked.")
        self.result["daily"] = {"ok": out.get("ok"), "skill": out.get("skill"), "written": bool(md),
                                "rows": {k: cells.get(k, [""])[0] for k in DAILY_ROWS}}

    def briefing(self) -> None:
        from packing_assistant.civil import CONFIRM, CONFIRM_EN

        brief_text = (self.job / "inputs" / "wah_briefing_input.txt").read_text(encoding="utf-8").strip()
        out = self.turn("briefing", brief_text, "safety-brief", command="- < inputs/wah_briefing_input.txt")
        refused = bool(out.get("hitl_pending")) and not out.get("wrote") and not out.get("files")
        if not refused:
            self.errors.append("briefing: the high-risk post wrote without the sign-off sentence")
        self.say(f"  SIGN-OFF NEEDED HERE: 安全交底 is a high-risk post. It wrote nothing: a licensed person must type"
                 f" 「{CONFIRM}」 or \"{CONFIRM_EN}\" (civil desktop / TUI dialog, or `civil exec --confirm` run by that person).")
        brief = {"refused_without_sentence": refused, "written": False}
        if not self.sign:
            self.say("  This demo does not supply the sentence. To see the draft, the person reruns with --sign and types it.")
        else:
            out = self.turn("briefing", brief_text, "safety-brief", confirm=True, command="--confirm - < inputs/wah_briefing_input.txt")
            md = self.expect_written("briefing (signed)", out)
            sign_row = next((row for row in tables(section(md, "11 签字栏"))), {})
            self.say("    cover: " + section(md, "1 封面")[:90])
            self.say("    signature row: " + " · ".join(f"{k} {v}" for k, v in sign_row.items()) + " (names are never pre-filled)")
            brief.update(written=bool(md), file=next((self.rel(f["path"]) for f in out.get("files") or []
                                                      if str(f.get("path", "")).endswith(".md")), ""))
        self.result["briefing"] = brief

    def run(self) -> Dict[str, Any]:
        for flow in (self.linked, self.tender, self.packing, self.daily, self.briefing):
            try:
                flow()
            except Exception as exc:  # a flow that raises is reported, and the next one still runs
                self.errors.append(f"{flow.__name__}: {type(exc).__name__}: {exc}")
                self.say(f"  ERROR {flow.__name__}: {type(exc).__name__}: {exc}")
        self.result["errors"] = list(self.errors)
        return self.result


def make_job(job: Optional[Path]) -> Path:
    """A new job folder holding the fixtures under inputs/ and a CIVIL.md with the two stated facts."""
    from packing_assistant.runtime.project_instructions import TEMPLATE

    if job is None:
        job = Path(tempfile.mkdtemp(prefix="civil-facade-demo-"))
    elif job.exists():
        raise FileExistsError(f"{job} already exists; the demo only writes into a folder it creates")
    job = job.resolve()
    (job / "inputs").mkdir(parents=True)
    for name in INPUTS:
        shutil.copyfile(FIXTURES / name, job / "inputs" / name)
    civil_md = re.sub(r"^- 辖区：.*$", "- 辖区：SG", TEMPLATE.replace("- 项目：\n", f"- 项目：{PROJECT}\n", 1), count=1, flags=re.M)
    (job / "CIVIL.md").write_text(civil_md, encoding="utf-8")
    return job


def run_demo(job: Optional[Path] = None, sign: str = "") -> Dict[str, Any]:
    """Create the job folder, enter it and run the flows.

    Leaves the process inside the job folder, with tempfile.tempdir pointing into it."""
    os.environ["PYTHON_DOTENV_DISABLED"] = "1"
    os.environ["CIVIL_AGENT_MODE"] = "steps"
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from packing_assistant.civil import enter_workspace

    job = make_job(job)
    scratch = job / ".civil-buddy" / "tmp"
    scratch.mkdir(parents=True, exist_ok=True)
    tempfile.tempdir = str(scratch)                             # a library's temp files (openpyxl) stay in the job folder
    os.environ["MPLCONFIGDIR"] = str(scratch / "matplotlib")    # and so does matplotlib's config folder
    enter_workspace(str(job))
    print(f"SYNTHETIC façade demo · job folder {job}\n"
          "All inputs are invented fixtures (examples/facade-demo); no contractor's data.", flush=True)
    return Demo(job, sign).run()


def main(argv: Optional[List[str]] = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--job", default="", help="a folder that does not exist yet, created as the job folder; default: a new temp folder")
    parser.add_argument("--sign", default="", help="the licensed person types the sign-off sentence here; never filled in for you")
    args = parser.parse_args(argv)
    sys.path.insert(0, str(ROOT))
    from packing_assistant.civil import CONFIRM, CONFIRM_EN
    from packing_assistant.runtime.civil_config import is_confirmation

    if args.sign and not is_confirmation(args.sign):
        print(f"--sign must be one of the two sentences exactly: {CONFIRM} | {CONFIRM_EN}", file=sys.stderr)
        return 2
    try:
        result = run_demo(Path(args.job).expanduser().absolute() if args.job else None, sign=args.sign.strip())
    except (FileExistsError, OSError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print()
    if result["errors"]:
        print("FAIL demo_facade: " + " | ".join(result["errors"]))
        return 1
    print(f"PASS demo_facade · drafts under {result['job']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
