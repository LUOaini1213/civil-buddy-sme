#!/usr/bin/env python3
"""Interim measurement identity, confirmation, amount boundaries and exports."""
from __future__ import annotations

import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PYTHON_DOTENV_DISABLED"] = "1"

from packing_assistant import expert_turn
from packing_assistant.office_job import tables_from_md
from packing_assistant.post_drafts import build_draft as dispatch
from packing_assistant.post_drafts.interim import MISSING, TOOL, _owned, _with_units, build_draft, parse
from packing_assistant.runtime.scheduler import Scheduler
from packing_assistant.runtime.tool_engine import default_engine


def draft(text):
    return build_draft("interim", TOOL, text)


def table(md, title):
    return next(rows for name, rows in tables_from_md(md) if name.startswith(title))


class Interim(unittest.TestCase):
    def test_exclusive_dispatch_and_no_country_tax_or_payment_default(self):
        md = draft("")
        self.assertEqual(md, dispatch("interim", TOOL, ""))
        self.assertIsNone(build_draft("subcontract", TOOL, ""))
        self.assertIsNone(build_draft("interim", "subcontract__settle", ""))
        self.assertEqual(12, len([s for s in md.splitlines() if s.startswith("## ")]))
        for text in ("AI 草稿", "submit_blocked=true", "无业主", "签字：（空）", "TBD"):
            self.assertIn(text, md)
        for text in ("CN：", "SG：", "GB 50500", "财建〔", "Security of Payment", "80%", "3%", "## 用户原文", "可以付款", "已获业主确认"):
            self.assertNotIn(text, md)

    def test_all_development_facts_have_correct_semantic_fields(self):
        from scripts.eval_post_content import classify, load, spots, split_echo
        for case in load("cases"):
            if case["post"] != "interim":
                continue
            with self.subTest(case=case["id"]):
                md = draft(case["request"])
                body, echo = split_echo(md)
                for fact in case["facts"]:
                    self.assertEqual("placed", classify(fact, body, echo, spots(body)), fact)
                self.assertTrue(all(t not in md for t in case.get("forbidden", [])))

    def test_narrative_completed_declared_previous_and_supervisor_are_not_owner(self):
        row = parse("合同编号HT-2025-018，计量期间2026-08-26到2026-09-25，清单项C35承台混凝土本期申报1286.5立方，上期末开累3420立方，监理已经审了1250立方，业主那边还没确认。").rows[0]
        self.assertEqual(("HT-2025-018", "2026-08-26", "2026-09-25", "承台混凝土", "C35", "1286.5立方", "3420立方", "1250立方"),
                         tuple(row[k] for k in ("contract", "begin", "end", "name", "model", "declared", "previous", "supervisor")))
        for key in ("owner_qty", "owner_amount", "claimed", "price", "completed", "cumulative"):
            self.assertNotIn(key, row)

    def test_owner_signed_document_does_not_copy_supervisor_or_declared_quantity(self):
        row = parse("3#楼主体钢筋HRB400这期报186.4吨，监理核下来182.75，业主计量单JL-031也签了").rows[0]
        self.assertEqual("JL-031", row["owner_source"])
        self.assertEqual("182.75", row["supervisor"])
        self.assertNotIn("owner_qty", row)
        self.assertNotIn("owner_amount", row)
        data = parse("名称：钢筋\n监理审：182.75\n确认来源：JL-031")
        events = _owned(data, data.events)
        self.assertEqual(MISSING, events[-1]["quantity"])

    def test_two_narrative_items_do_not_share_confirmations_sources_or_units(self):
        data = parse("3#楼主体钢筋HRB400这期报186.4吨，监理核下来182.75，业主计量单JL-031也签了；5#楼加气块砌体报了930方，监理还没看，业主更没确认。")
        self.assertEqual(2, len(data.rows))
        first, second = _with_units(data.rows)
        self.assertEqual(("钢筋", "HRB400", "吨"), tuple(first[k] for k in ("name", "model", "unit")))
        self.assertEqual(("加气块砌体", "930方", "方"), tuple(second[k] for k in ("name", "declared", "unit")))
        for key in ("supervisor", "owner_qty", "owner_source", "model"):
            self.assertNotIn(key, second)

    def test_explicit_later_unique_material_reference_returns_to_its_own_item(self):
        data = parse("3#楼钢筋HRB400本期申报35吨；5#楼加气块砌体报了930方；上期末开累我只记得钢筋是412.6吨")
        self.assertEqual(2, len(data.rows))
        self.assertEqual("412.6吨", data.rows[0]["previous"])
        self.assertNotIn("previous", data.rows[1])

    def test_ambiguous_material_reference_remains_unassigned_not_last_item(self):
        data = parse("3#楼钢筋HRB400本期申报35吨；5#楼钢筋HRB500本期申报40吨；上期末开累钢筋是412.6吨")
        self.assertEqual(3, len(data.rows))
        self.assertNotIn("name", data.rows[-1])
        self.assertEqual("412.6吨", data.rows[-1]["previous"])
        self.assertTrue(all("previous" not in r for r in data.rows[:2]))

    def test_later_reference_cannot_cross_file_or_repeated_filename_epoch(self):
        data = parse("### A.txt\n名称：钢筋\n本期申报：35吨\n### A.txt\n上期末开累钢筋是412.6吨")
        self.assertEqual(2, len(data.rows))
        self.assertNotIn("previous", data.rows[0])
        self.assertNotIn("name", data.rows[1])

    def test_leading_zero_boq_code_is_a_string_not_quantity(self):
        row = parse("清单编码 040203006001\n名称：沥青混凝土面层 AC-13C").rows[0]
        self.assertEqual("040203006001", row["code"])
        self.assertEqual("AC-13C", row["model"])
        for key in ("contract_qty", "previous", "declared", "supervisor", "owner_qty", "price"):
            self.assertNotIn(key, row)

    def test_short_command_keeps_unlabelled_quantity_without_declaring_its_role(self):
        for text in ("写一份验工计价 模板 120m2", "模板 120m2"):
            with self.subTest(text=text):
                row = parse(text).rows[0]
                self.assertEqual("模板", row["name"])
                self.assertEqual("120m2", row["unassigned"])
                for key in ("declared", "completed", "contract_qty", "supervisor", "owner_qty", "previous"):
                    self.assertNotIn(key, row)
                rows = table(draft(text), "5 计量")
                self.assertEqual("120m2", rows[1][rows[0].index("用户未分派数量原文（口径待核）")])
                self.assertEqual(["TBD"]*5, rows[1][-5:])

    def test_spaced_narrative_review_quantity_is_not_a_new_unassigned_item(self):
        row = parse("名称：钢筋\n监理已经审了 1250立方").rows[0]
        self.assertEqual("钢筋", row["name"])
        self.assertEqual("1250立方", row["supervisor"])
        self.assertNotIn("unassigned", row)

    def test_all_quantity_roles_remain_separate_without_cumulative_or_reduction_math(self):
        row = parse("名称：钢筋\n合同量：100吨\n上期末开累：20吨\n本期完成：35吨\n本期申报：30吨\n本期末累计：52吨\n监理审：28吨\n业主核：25吨").rows[0]
        self.assertEqual(["100吨", "20吨", "35吨", "30吨", "52吨", "28吨", "25吨"],
                         [row[k] for k in ("contract_qty", "previous", "completed", "declared", "cumulative", "supervisor", "owner_qty")])
        rows = table(draft("名称：钢筋\n本期申报：30吨\n上期末开累：20吨\n监理审：28吨"), "5 计量")
        self.assertEqual(["TBD"]*5, rows[1][-5:])
        self.assertEqual(MISSING, rows[1][rows[0].index("用户本期末累计（未核）")])

    def test_pending_hearsay_future_and_not_reviewed_are_records_not_quantities(self):
        for text, key in (("预计监理审核量35吨", "supervisor"), ("业主还没确认35吨", "owner_qty"),
                          ("明天本期申报35吨", "declared"), ("听说监理审了35吨", "supervisor"),
                          ("计划本期完成35吨", "completed"), ("业主确认量：预计35吨", "owner_qty"),
                          ("本期申报：没报35吨", "declared")):
            with self.subTest(text=text):
                row = parse("名称：钢筋\n"+text).rows[0]
                self.assertEqual(MISSING, row.get(key, MISSING))
                self.assertIn("35吨", row["record"])

    def test_unknown_rate_does_not_erase_explicit_user_owner_quantity(self):
        row = parse("名称：钢筋\n业主确认30吨但单价没定").rows[0]
        self.assertEqual("30吨", row["owner_qty"])
        self.assertNotIn("price", row)
        self.assertEqual("单价没定", row["record"])
        self.assertEqual(["TBD"]*5, table(draft("名称：钢筋\n业主确认30吨但单价没定"), "5 计量")[1][-5:])

    def test_forecast_declaration_does_not_erase_actual_user_supervisor_quantity(self):
        row = parse("名称：钢筋\n预计本期申报35吨但监理已经审30吨").rows[0]
        self.assertNotIn("declared", row)
        self.assertEqual("30吨", row["supervisor"])
        self.assertEqual("预计本期申报35吨", row["record"])
        self.assertNotIn("owner_qty", row)

    def test_quantity_fields_cannot_extract_model_code_date_version_or_question(self):
        for label, key in (("合同量", "contract_qty"), ("上期末开累", "previous"), ("本期申报", "declared"),
                           ("监理审", "supervisor"), ("业主核", "owner_qty")):
            for value in ("HRB400", "C35", "AC-13C", "JL-031", "V2.0", "2026-10-20", "35吨？", "不是35吨", "例如35吨"):
                with self.subTest(label=label, value=value):
                    row = parse("名称：钢筋\n"+label+"："+value).rows[0]
                    self.assertEqual(MISSING, row.get(key, MISSING))

    def test_negated_unknown_or_question_name_does_not_leak_an_embedded_model(self):
        for value in ("不是钢筋 HRB400", "钢筋 HRB400？", "未定 HRB400"):
            with self.subTest(value=value):
                data = parse("部位：3#楼\n名称："+value)
                self.assertTrue(all(r.get("model", MISSING) == MISSING for r in data.rows))
                self.assertTrue(all(r.get("name", MISSING) == MISSING for r in data.rows))

    def test_raw_owner_quantity_kept_without_confirmed_or_payable_claim(self):
        md = draft("名称：钢筋\n本期申报：35吨\n业主确认量：30吨\n确认来源：JL-031\n业主代表：周海峰\n业主确认记录：用户说已签")
        rows = table(md, "5 计量")
        self.assertEqual("30吨", rows[1][rows[0].index("用户业主确认量（未核）")])
        self.assertEqual("JL-031", rows[1][rows[0].index("业主确认来源（未核）")])
        self.assertEqual(["TBD"]*5, rows[1][-5:])
        self.assertEqual("未执行", table(md, "用户确认记录")[1][-1])

    def test_unique_boq_reference_only_reuses_its_own_item(self):
        data = parse("名称：钢筋\n清单编码：A\n本期申报：35吨\n名称：砌体\n清单编码：B\n本期申报：930方\n清单编码：A\n监理审：30吨")
        self.assertEqual(2, len(data.rows))
        self.assertEqual("30吨", data.rows[0]["supervisor"])
        self.assertNotIn("supervisor", data.rows[1])

    def test_duplicate_boq_reference_without_identity_stays_unassigned(self):
        data = parse("名称：钢筋\n清单编码：DUP\n本期申报：35吨\n名称：砌体\n清单编码：DUP\n清单编码：DUP\n监理审：30吨")
        self.assertEqual(3, len(data.rows))
        self.assertNotIn("name", data.rows[-1])
        self.assertEqual("30吨", data.rows[-1]["supervisor"])
        self.assertTrue(all("supervisor" not in r for r in data.rows[:2]))

    def test_explicit_period_qualifier_resolves_only_its_own_existing_boq_row(self):
        data = parse("| 名称 | 清单编码 | 期间 | 本期申报 |\n| --- | --- | --- | --- |\n"
                     "| 钢筋 | A | 9月 | 35吨 |\n| 钢筋 | A | 10月 | 40吨 |\n\n"
                     "期间：9月\n清单编码：A\n监理审：30吨")
        self.assertEqual(2, len(data.rows))
        self.assertEqual("30吨", data.rows[0]["supervisor"])
        self.assertNotIn("supervisor", data.rows[1])

    def test_explicit_contract_qualifier_resolves_only_its_own_existing_boq_row(self):
        data = parse("| 名称 | 清单编码 | 合同 | 本期申报 |\n| --- | --- | --- | --- |\n"
                     "| 钢筋 | A | HT-A | 35吨 |\n| 砌体 | A | HT-B | 930方 |\n\n"
                     "合同：HT-A\n清单编码：A\n业主核：30吨")
        self.assertEqual(2, len(data.rows))
        self.assertEqual("30吨", data.rows[0]["owner_qty"])
        self.assertNotIn("owner_qty", data.rows[1])

    def test_period_qualifier_with_duplicate_boq_in_two_contracts_is_still_ambiguous(self):
        data = parse("| 名称 | 清单编码 | 合同 | 期间 | 本期申报 |\n| --- | --- | --- | --- | --- |\n"
                     "| 钢筋 | A | HT-A | 9月 | 35吨 |\n| 砌体 | A | HT-B | 9月 | 930方 |\n\n"
                     "期间：9月\n清单编码：A\n监理审：30吨")
        self.assertEqual(3, len(data.rows))
        self.assertNotIn("name", data.rows[-1])
        self.assertEqual("30吨", data.rows[-1]["supervisor"])
        self.assertTrue(all("supervisor" not in r for r in data.rows[:2]))

    def test_explicit_contract_cannot_resolve_boq_in_another_file(self):
        data = parse("### A.txt\n合同：HT-A\n清单编码：A\n名称：钢筋\n本期申报：35吨\n"
                     "### B.txt\n合同：HT-A\n清单编码：A\n监理审：30吨")
        self.assertEqual(2, len(data.rows))
        self.assertNotIn("name", data.rows[1])
        self.assertNotIn("supervisor", data.rows[0])

    def test_new_model_and_location_do_not_copy_quantities_prices_or_sources(self):
        for label, value in (("规格", "HRB500"), ("部位", "5#楼")):
            with self.subTest(label=label):
                data = parse("名称：钢筋 HRB400\n部位：3#楼\n本期申报：35吨\n申报来源：D.csv\n单价：SGD 800/t\n"+label+"："+value)
                self.assertEqual(2, len(data.rows))
                self.assertEqual("钢筋", data.rows[1]["name"])
                for key in ("declared", "declared_source", "price", "unit", "owner_qty", "code"):
                    self.assertNotIn(key, data.rows[1])

    def test_scope_changes_do_not_copy_previous_item_or_amount(self):
        for label, first, second, key in (("合同", "HT-A", "HT-B", "contract"), ("标段", "A", "B", "lot"),
                                           ("期间", "9月", "10月", "period"), ("辖区", "CN", "SG", "jurisdiction"),
                                           ("版本", "V1", "V2", "version"), ("币种", "SGD", "CNY", "currency")):
            with self.subTest(label=label):
                data = parse(label+"："+first+"\n名称：钢筋\n本期申报：35吨\n单价：800/t\n"+label+"："+second+"\n本期申报：12吨")
                self.assertEqual(2, len(data.rows))
                self.assertEqual([first, second], [r[key] for r in data.rows])
                self.assertNotIn("name", data.rows[1])
                self.assertNotIn("price", data.rows[1])

    def test_blank_table_scope_field_does_not_mask_conflicting_file_contract(self):
        data = parse("合同：HT-A\n| 名称 | 合同 | 本期申报 |\n| --- | --- | --- |\n| 钢筋 | | 35吨 |\n\n合同：HT-B\n本期申报：12吨")
        self.assertEqual(2, len(data.rows))
        self.assertNotIn("name", data.rows[1])
        self.assertEqual("HT-B", data.rows[1]["contract"])

    def test_full_blank_row_and_table_tail_do_not_borrow_previous_identity(self):
        data = parse("| 名称 | 清单编码 | 本期申报 |\n| --- | --- | --- |\n| 钢筋 | A | 35吨 |\n| | | |\n\n监理审：30吨")
        self.assertEqual(3, len(data.rows))
        self.assertEqual(MISSING, data.rows[1]["name"])
        self.assertNotIn("name", data.rows[2])
        self.assertNotIn("supervisor", data.rows[0])

    def test_multi_declarations_keep_each_source_without_adding_or_losing_repeats(self):
        data = parse("名称：钢筋\n本期申报：10吨\n申报来源：D1.csv\n本期申报：10吨\n申报来源：D2.csv\n监理审：9吨\n监理来源：S.csv")
        self.assertEqual([("10吨", "D1.csv"), ("10吨", "D2.csv"), ("9吨", "S.csv")], [(e["quantity"], e["source"]) for e in data.events])
        self.assertNotIn("cumulative", data.rows[0])
        self.assertNotIn("owner_qty", data.rows[0])

    def test_reordered_role_sources_attach_to_correct_quantity(self):
        data = parse("| 确认来源 | 监理来源 | 申报来源 | 名称 | 本期申报 | 监理审 | 业主核 |\n| --- | --- | --- | --- | --- | --- | --- |\n| O.csv | S.csv | D.csv | 钢筋 | 35吨 | 32吨 | 30吨 |")
        events = {e["field"]: (e["quantity"], e.get("source")) for e in data.events}
        self.assertEqual(("35吨", "D.csv"), events["本期申报量"])
        self.assertEqual(("32吨", "S.csv"), events["监理审核量"])
        self.assertEqual(("30吨", "O.csv"), events["业主确认量"])

    def test_general_source_does_not_become_owner_supervisor_or_price_source(self):
        row = parse("名称：钢筋\n来源：A.pdf\n本期申报：35吨").rows[0]
        self.assertEqual("A.pdf", row["source"])
        for key in ("owner_source", "supervisor_source", "price_source", "declared_source"):
            self.assertNotIn(key, row)

    def test_source_only_does_not_supply_the_previous_roles_quantity(self):
        data = parse("名称：钢筋\n本期申报：35吨\n监理来源：S.csv")
        self.assertEqual(MISSING, data.events[-1]["quantity"])
        self.assertEqual("S.csv", data.events[-1]["source"])
        self.assertNotIn("supervisor", data.rows[0])

    def test_same_named_files_and_regions_have_independent_rows(self):
        data = parse("### A.txt\n辖区：CN\n名称：钢筋 HRB400\n本期申报：35吨\n### A.txt\n名称：钢筋\n监理审：30吨")
        self.assertEqual(["CN", "UNSPECIFIED"], [r["jurisdiction"] for r in data.rows])
        for key in ("model", "declared", "previous", "price"):
            self.assertNotIn(key, data.rows[1])

    def test_units_are_literal_and_mixed_units_not_converted(self):
        rows = _with_units(parse("名称：钢筋\n本期申报：35吨\n监理审：120kg").rows)
        self.assertEqual(MISSING, rows[0].get("unit", MISSING))
        self.assertIn("待核", rows[0]["unit_note"])
        self.assertEqual("120kg", rows[0]["supervisor"])
        row = _with_units(parse("名称：钢筋\n单位：kg\n本期申报：35吨").rows)[0]
        self.assertEqual("kg", row["unit"])
        self.assertIn("待核", row["unit_note"])

    def test_bare_review_number_remains_bare_and_does_not_guess_a_unit(self):
        row = _with_units(parse("名称：钢筋\n监理审：182.75").rows)[0]
        self.assertEqual("182.75", row["supervisor"])
        self.assertNotIn("unit", row)

    def test_prices_and_amounts_preserve_roles_and_never_generate_payable(self):
        md = draft("名称：钢筋\n本期申报：35吨\n合同单价：SGD 800/t\n单价来源：P.csv\n申报金额：SGD 28000\n申报金额来源：D.csv\n业主核准金额：SGD 26000\n核准金额来源：O.csv\n价税口径：不含税")
        rows = table(md, "5 计量")
        for header, value in (("用户合同单价原文（未核）", "SGD 800/t"), ("单价来源", "P.csv"), ("用户申报金额（未核）", "SGD 28000"),
                              ("用户业主核准金额（未核）", "SGD 26000"), ("核准金额来源", "O.csv"), ("用户价税口径", "不含税")):
            self.assertEqual(value, rows[1][rows[0].index(header)])
        self.assertEqual(["TBD"]*5, rows[1][-5:])

    def test_price_and_amount_cannot_come_from_model_date_quantity_or_percent(self):
        for label, key in (("合同单价", "price"), ("申报金额", "claimed"), ("业主核准金额", "owner_amount")):
            for value in ("HRB400", "HRB 400", "JL-031", "2026-10-20", "35吨", "3%", "V2.0", "不是800元", "预计800元"):
                with self.subTest(label=label, value=value):
                    self.assertEqual(MISSING, parse("名称：钢筋\n"+label+"："+value).rows[0].get(key, MISSING))

    def test_variations_adjustment_and_claims_are_records_not_amounts_added_to_payment(self):
        md = draft("名称：钢筋\n变更项目：替换规格\n变更文号：BG-A\n变更状态：用户说已批\n变更金额：SGD 500\n调价依据：X.pdf\n索赔资料：C.pdf")
        rows = table(md, "6 变更")
        self.assertEqual("SGD 500", rows[1][rows[0].index("用户变更金额原文（未核）")])
        self.assertEqual("BG-A", rows[1][rows[0].index("变更文号 / 来源")])
        self.assertEqual("用户说已批", rows[1][rows[0].index("用户变更状态（未核）")])
        self.assertEqual("TBD", rows[1][-1])

    def test_payment_certificate_without_quantity_never_supplies_confirmation_or_payment(self):
        row = parse("名称：钢筋\n付款证书：PC-01\n付款证书来源：PC.pdf\n支付资料：待核PC签字\n附件目录：计量表、试验资料").rows[0]
        self.assertEqual("PC-01", row["certificate"])
        for key in ("owner_qty", "owner_amount", "claimed", "price", "declared"):
            self.assertNotIn(key, row)

    def test_deductions_and_wages_are_separate_raw_fields_not_net_calculation(self):
        md = draft("名称：钢筋\n农民工工资款：200元\n预付款抵扣：30元\n甲供材扣款：40元\n质保金：50元\n违约金：60元\n扣减依据：D.pdf")
        rows = table(md, "9 扣减")
        self.assertEqual(["30元", "40元", "50元", "60元"], [rows[1][rows[0].index(h)] for h in
                         ("用户预付款抵扣（未核）", "用户甲供材扣款（未核）", "用户质保金原文（未核）", "用户违约金原文（未核）")])
        self.assertEqual("TBD", rows[1][-1])
        wages = table(md, "8 农民工")
        self.assertEqual("200元", wages[1][wages[0].index("用户工资款原文（未核）")])

    def test_each_deduction_and_wage_source_binds_only_its_explicit_amount_role(self):
        data = parse("名称：钢筋\n农民工工资款：200元\n工资款来源：W.csv\n预付款抵扣：30元\n预付款抵扣来源：A.csv\n"
                     "甲供材扣款：40元\n甲供材扣款来源：S.csv\n质保金：50元\n质保金来源：R.csv\n违约金：60元\n违约金来源：P.csv\n扣减依据：G.pdf")
        self.assertEqual([("200元", "W.csv"), ("30元", "A.csv"), ("40元", "S.csv"), ("50元", "R.csv"), ("60元", "P.csv")],
                         [(e["quantity"], e["source"]) for e in data.events])
        self.assertEqual("G.pdf", data.rows[0]["deduction_source"])
        self.assertTrue(all("G.pdf" != e["source"] for e in data.events))

    def test_common_owner_and_supervisor_labels_still_keep_roles_separate(self):
        for label in ("业主审核量", "甲方确认量", "建设单位核", "发包人核准量"):
            with self.subTest(label=label):
                row = parse("名称：钢筋\n"+label+"：30吨").rows[0]
                self.assertEqual("30吨", row["owner_qty"])
                self.assertNotIn("supervisor", row)
        self.assertEqual("32吨", parse("名称：钢筋\n监理核量32吨").rows[0]["supervisor"])
        self.assertEqual("30吨", parse("名称：钢筋\n甲方已核准30吨").rows[0]["owner_qty"])

    def test_dates_are_literal_not_codes_or_assumed_quarter_calendar(self):
        row = parse("期次：2026年第3季度\n名称：钢筋").rows[0]
        self.assertEqual(MISSING, row["begin"])
        self.assertEqual(MISSING, row["end"])
        for value in ("JL-20261020", "V2.0", "不是2026-10-20", "2026-10-20？"):
            self.assertEqual(MISSING, parse("名称：钢筋\n截止日期："+value).rows[0]["end"])

    def test_english_fields_keep_raw_units_and_separate_source_roles(self):
        row = parse("Project: Synthetic A\nContract: HT-A\nJurisdiction: SG\nBOQ code: 000012\nName: rebar\nModel: HRB400\nDeclared quantity: 1,200kg\nSupervisor quantity: 1100\nOwner source: O.pdf").rows[0]
        self.assertEqual(("000012", "rebar", "HRB400", "1,200kg", "1100", "O.pdf"), tuple(row[k] for k in ("code", "name", "model", "declared", "supervisor", "owner_source")))
        self.assertNotIn("owner_qty", row)

    def test_evidence_keeps_physical_location_separate_from_input_position(self):
        data = parse("### A.txt\n名称：钢筋\n部位：3#楼\n本期申报：35吨")
        record = _owned(data, data.events)[0]
        self.assertEqual("3#楼", record["location"])
        self.assertEqual("A.txt · 本轮输入第4行（非原件页码）", record["input_position"])

    def test_examples_masked_but_explicit_actual_reset_is_preserved(self):
        data = parse("例如清单项钢筋本期申报35吨，监理审30吨\n实际情况：名称：卷材\n本期申报：76卷")
        self.assertEqual(1, len(data.rows))
        self.assertEqual("76卷", data.rows[0]["declared"])
        self.assertNotIn("35吨", repr(data))


class InterimRuntime(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="interim-regression-")
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        for context in (patch.object(expert_turn, "_OUT", self.root), patch("packing_assistant.runtime.memory._OUT", self.root),
                        patch("packing_assistant.runtime.scheduler.get_scheduler", return_value=Scheduler()),
                        patch.dict(os.environ, {"CIVIL_JOB_ROOT": "", "CIVIL_SANDBOX_ROOTS": str(self.root), "CIVIL_SANDBOX": "workspace-write", "CIVIL_APPROVAL": "on-request"})):
            context.start()
            self.addCleanup(context.stop)

    def test_actual_word_excel_match_markdown_and_preserve_codes_and_formula_literals(self):
        from xml.etree import ElementTree as ET
        from zipfile import ZipFile
        import openpyxl
        result = default_engine().execute(TOOL, {"text": "清单编码：000012\n名称：=rebar|A\n部位：A & B\n本期申报：35吨", "session_id": "export-interim"}, expert_id="interim", intent="run")
        self.assertTrue(result["ok"] and result["wrote"] and result["submit_blocked"], result)
        paths = {Path(f["path"]).suffix: Path(f["path"]) for f in result["files"]}
        self.assertTrue({".md", ".docx", ".xlsx"}.issubset(paths), paths)
        expected = table(paths[".md"].read_text(encoding="utf-8"), "5 计量")
        book = openpyxl.load_workbook(paths[".xlsx"])
        try:
            sheet = next(book[n] for n in book.sheetnames if n.startswith("5 计量"))
            self.assertEqual(expected, [["" if c.value is None else c.value for c in row] for row in sheet])
            self.assertEqual("000012", sheet.cell(2, 1).value)
            self.assertTrue(all(c.data_type != "f" for sheet in book for row in sheet for c in row))
        finally:
            book.close()
        with ZipFile(paths[".docx"]) as archive:
            document = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        tables = [[["".join(t.text or "" for t in c.findall(".//w:t", ns)) for c in r.findall("w:tc", ns)]
                   for r in tbl.findall("w:tr", ns)] for tbl in document.findall(".//w:tbl", ns)]
        self.assertIn(expected, tables)

    def test_questions_and_readonly_cannot_write(self):
        chat = expert_turn.run_expert_turn("监理核量和业主确认怎么区分？", "interim", session_id="chat-interim")
        self.assertFalse(chat["wrote"])
        with patch.dict(os.environ, {"CIVIL_SANDBOX": "read-only"}):
            denied = expert_turn.run_named_exclusive(TOOL, {"text": "写验工计价草稿", "session_id": "ro-interim"})
        self.assertFalse(denied["wrote"])
        self.assertFalse(list(self.root.rglob("*.md")))

    def test_all_hostile_inputs_survive_without_ragged_columns_or_verdicts(self):
        from scripts.test_post_robustness import INPUTS, VERDICTS, table_rows
        import re
        for name, source in INPUTS.items():
            with self.subTest(input=name):
                result = expert_turn.run_named_exclusive(TOOL, {"text": source, "session_id": "hostile-interim"})
                self.assertTrue(result["wrote"] and result["submit_blocked"], result)
                md = next(Path(f["path"]) for f in result["files"] if f["path"].endswith(".md")).read_text(encoding="utf-8")
                parsed = table_rows(md)
                for i in range(1, len(parsed)):
                    line, cells = parsed[i]
                    previous_line, previous = parsed[i-1]
                    if line != previous_line+1 or all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                        continue
                    self.assertEqual(len(previous), len(cells))
                self.assertTrue(all(p.search(md) is None for p in VERDICTS))


if __name__ == "__main__":
    unittest.main()
