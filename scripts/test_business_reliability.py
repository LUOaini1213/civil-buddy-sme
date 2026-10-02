#!/usr/bin/env python3
"""Offline regressions for KB mutations and job-folder Office interchange."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "demo"))

import kbio
import openpyxl

from packing_assistant import office_job


class BusinessReliabilityTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory(prefix="civil-business-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        environment = patch.dict(os.environ, {
            "CIVIL_JOB_ROOT": str(self.root),
            "CIVIL_SANDBOX_ROOTS": str(self.root),
        })
        environment.start()
        self.addCleanup(environment.stop)

    def test_kb_write_and_delete_refresh_index(self) -> None:
        with patch.object(kbio, "KB_ROOT", self.root), patch(
            "packing_assistant.kb_search.reindex_kb_file", return_value=True
        ) as reindex:
            metadata = kbio.write_text("company/notes.md", "# Updated notes\n")
            self.assertEqual(metadata["path"], "company/notes.md")
            self.assertEqual(reindex.call_count, 1)
            self.assertEqual(reindex.call_args.args, ("company/notes.md",))
            kbio.delete_file("company/notes.md")
            self.assertFalse((self.root / "company/notes.md").exists())
            self.assertEqual(reindex.call_count, 2)

    def test_guard_failure_never_falls_back_to_direct_write(self) -> None:
        with patch.object(kbio, "KB_ROOT", self.root), patch(
            "packing_assistant.sandbox.guarded_write_text", side_effect=RuntimeError("guard unavailable")
        ), patch.object(kbio, "_kb_index_hook") as reindex:
            with self.assertRaisesRegex(RuntimeError, "guard unavailable"):
                kbio.write_text("company/notes.md", "must not be written")
            self.assertFalse((self.root / "company/notes.md").exists())
            reindex.assert_not_called()

    def test_kb_delete_obeys_guard(self) -> None:
        target = self.root / "notes.md"
        target.write_text("keep", encoding="utf-8")
        with patch.object(kbio, "KB_ROOT", self.root), patch(
            "packing_assistant.sandbox.assert_write", side_effect=PermissionError("denied")
        ):
            with self.assertRaises(PermissionError):
                kbio.delete_file("notes.md")
        self.assertEqual(target.read_text(encoding="utf-8"), "keep")

    def test_index_failure_does_not_undo_a_successful_write(self) -> None:
        with patch.object(kbio, "KB_ROOT", self.root), patch(
            "packing_assistant.kb_search.reindex_kb_file", side_effect=RuntimeError("index offline")
        ), self.assertLogs("civil.demo_kbio", level="WARNING"):
            kbio.write_text("notes.md", "saved content")
        self.assertEqual((self.root / "notes.md").read_text(encoding="utf-8"), "saved content")

    def test_draft_cells_remain_literal_and_sheet_names_are_unique(self) -> None:
        title = "A" * 31
        path = self.root / "draft.xlsx"
        office_job.write_xlsx(path, [
            (title, [["=1+1", "#N/A"]]),
            (title.lower(), [["plain text"]]),
        ])
        workbook = openpyxl.load_workbook(path, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(len({name.casefold() for name in workbook.sheetnames}), 2)
        self.assertTrue(all(len(name) <= 31 for name in workbook.sheetnames))
        for coordinate, value in (("A1", "=1+1"), ("B1", "#N/A")):
            cell = workbook.worksheets[0][coordinate]
            self.assertEqual(cell.value, value)
            self.assertEqual(cell.data_type, "s")

    def test_patch_preserves_owner_sheets_and_formulas(self) -> None:
        path = self.root / "ledger.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "Owner"
        workbook.active["A1"] = "=SUM(1,2)"
        workbook.create_sheet("CB草稿审批记录")["A1"] = "owner content"
        workbook.create_sheet("CB草稿-旧表")["A1"] = "outdated draft"
        workbook.save(path)
        workbook.close()
        office_job.patch_xlsx(path, [("新表", [["=1+1"]])])
        patched = openpyxl.load_workbook(path, data_only=False)
        try:
            self.assertEqual(patched["Owner"]["A1"].data_type, "f")
            self.assertEqual(patched["CB草稿审批记录"]["A1"].value, "owner content")
            self.assertNotIn("CB草稿-旧表", patched.sheetnames)
            self.assertEqual(patched["CB草稿-新表"]["A1"].data_type, "s")
        finally:
            patched.close()
        before = path.read_bytes()
        office_job.patch_xlsx(path, [])
        self.assertEqual(path.read_bytes(), before)

    def test_bad_office_files_do_not_abort_job_context(self) -> None:
        (self.root / "broken.xlsx").write_bytes(b"not a workbook")
        with ZipFile(self.root / "broken.docx", "w") as archive:
            archive.writestr("unrelated.xml", "<root />")
        (self.root / "valid.md").write_text("usable context", encoding="utf-8")
        blob = office_job.job_files_blob()
        self.assertIn("usable context", blob)
        self.assertEqual(blob.count("（读失败）"), 2)

    def test_bad_named_workbook_keeps_sibling_export(self) -> None:
        (self.root / "ledger.xlsx").write_bytes(b"not a workbook")
        draft = self.root / "draft.md"
        draft.write_text("| Item |\n| --- |\n| Draft |\n", encoding="utf-8")
        paths = office_job.export_md_to_xlsx(draft, query="update ledger.xlsx")
        self.assertEqual(paths, [draft.with_suffix(".xlsx")])
        self.assertEqual((self.root / "ledger.xlsx").read_bytes(), b"not a workbook")

    def test_automatic_export_never_patches_named_source_or_overwrites_previous_copy(self) -> None:
        source = self.root / "ledger.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "=SUM(1,2)"
        workbook.save(source)
        workbook.close()
        before = source.read_bytes()
        draft = self.root / "ledger.md"
        draft.write_text("## 用户原文\nupdate ledger.xlsx\n\n| Item |\n| --- |\n| First draft |\n", encoding="utf-8")
        first = office_job.export_md_to_xlsx(draft, query="update ledger.xlsx")
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(first, [self.root / "ledger-2.xlsx"])
        first_bytes = first[0].read_bytes()
        draft.write_text("| Item |\n| --- |\n| Second draft |\n", encoding="utf-8")
        second = office_job.export_md_to_xlsx(draft)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(first[0].read_bytes(), first_bytes)
        self.assertEqual(second, [self.root / "ledger-3.xlsx"])
        book = openpyxl.load_workbook(second[0])
        try:
            self.assertEqual(book.active["A2"].value, "Second draft")
        finally:
            book.close()

    def test_job_files_exclude_sensitive_paths_and_reject_other_types(self) -> None:
        (self.root / "private_key_notes.txt").write_text("synthetic fixture", encoding="utf-8")
        (self.root / "valid.md").write_text("abcdef", encoding="utf-8")
        (self.root / "program.py").write_text("pass", encoding="utf-8")
        self.assertEqual([row["name"] for row in office_job.list_job_files()], ["valid.md"])
        with self.assertRaises(PermissionError):
            office_job.read_job_file(self.root / "private_key_notes.txt")
        with self.assertRaises(ValueError):
            office_job.read_job_file(self.root / "program.py")
        self.assertEqual(office_job.read_job_file(self.root / "valid.md", 3), "abc")
        self.assertEqual(office_job.read_job_file(self.root / "valid.md", -1), "")

    def test_job_context_excludes_links_outside_job_root(self) -> None:
        outside = self.root / "outside"
        outside.mkdir()
        target = outside / "notes.md"
        target.write_text("outside job root", encoding="utf-8")
        job = self.root / "job"
        job.mkdir()
        link = job / "linked.md"
        try:
            link.symlink_to(target)
        except OSError:
            self.skipTest("symbolic links unavailable on this host")
        with patch.dict(os.environ, {"CIVIL_JOB_ROOT": str(job)}):
            self.assertEqual(office_job.list_job_files(), [])
            with self.assertRaises(PermissionError):
                office_job.read_job_file(link)


if __name__ == "__main__":
    unittest.main()
