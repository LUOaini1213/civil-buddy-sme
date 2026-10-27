#!/usr/bin/env python3
"""Infrastructure drafts through real HTTP, selected uploads and Office copies."""
from io import BytesIO
from pathlib import Path
import sys
import unittest
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import test_finale_http_acceptance as finale
from scripts import test_workbench_flow as flow
from packing_assistant.office_job import tables_from_md
from packing_assistant.post_drafts.design_infrastructure import PROFILES, UNKNOWN

TITLES = {"municipal": "道路区段专业核对", "bridge": "桥梁比选与输入来源",
          "port": "港工对象与资料接口", "tunnel": "区间支护与来源待核",
          "traffic": "交通调查与实验待核", "design-coord": "会审行动与证据归属"}
ACK = "我明白，将由持证人员签认"


class InfrastructureHttpAcceptance(unittest.TestCase):
    setUp = flow.WorkbenchFlowTests.setUp
    post = flow.WorkbenchFlowTests.post
    markdown = finale.FinaleHttpAcceptance.markdown
    spreadsheet_rows = finale.FinaleHttpAcceptance.spreadsheet_rows
    word_table_rows = finale.FinaleHttpAcceptance.word_table_rows

    def test_all_six_natural_inputs_reach_word_and_excel_professional_fields(self):
        cases = (
            ("municipal", "生成市政道路原则稿。东环路按城市次干路做，桩号K0+010到K1+720，红线宽24米，设计速度60km/h，现状地面标高21.45米。", "设计速度", "60km/h"),
            ("bridge", "生成桥梁设计提纲。东河大桥跨径3×27米，桥宽23.6米，下部双柱墩配钻孔灌注桩。", "桥宽", "23.6米"),
            ("port", "生成港航设计提纲。1#泊位，高水位4.23m，低水位-0.51m，H1%波高1.63m。", "波浪要素", "H1%波高1.63m"),
            ("tunnel", "生成隧道设计提纲。青岭隧道帮忙，最大埋深38.9米，勘察报告里是IV级围岩，准备上CRD法。", "埋深", "38.9米"),
            ("traffic", "生成交通报告骨架。节点：路口甲；调查时段：07:20-08:20；东进口左转301pcu/h；流量文件：调查甲.csv；信号周期：109秒", "交通量", "301pcu/h"),
            ("design-coord", "生成图纸会审纪要。一个是地下室入口，设计院刘工说月底前出变更单，问题编号记HS-211，责任专业定建筑。", "责任人", "刘工"),
        )
        for expert, message, field, expected in cases:
            with self.subTest(expert=expert):
                self.sid = "infra-natural-"+uuid4().hex[:16]
                md, done = self.markdown(message, expert, confirm_text=ACK)
                title = TITLES[expert]
                table = dict(tables_from_md(md))[title]
                self.assertEqual(expected, table[1][table[0].index(field)])
                self.assertEqual("", table[1][-1])
                self.assertEqual(table, self.spreadsheet_rows(done, title))
                self.assertEqual(table, self.word_table_rows(done, table[0], title=title))
                self.assertTrue(done["submit_blocked"])

    def test_selected_reordered_csvs_keep_blank_rows_and_original_bytes(self):
        cases = (
            ("municipal", "红线宽,道路名称,设计速度\n24m,东路,60km/h\n18m,,\n", "红线宽", "24m", "18m", "道路名称"),
            ("bridge", "跨径,桥名,桥宽\n3×27米,一桥,23.6米\n40m,,\n", "跨径组合", "3×27米", "40m", "桥梁名称"),
            ("port", "高水位,泊位名称,波浪\n4.23m,1#泊位,H1%波高1.63m\n4.57m,,\n", "设计高水位", "4.23m", "4.57m", "码头名称"),
            ("tunnel", "埋深,区间名称,围岩\n38.9m,甲区间,IV级\n13.1m,,\n", "埋深", "38.9m", "13.1m", "隧道名称"),
            ("traffic", "交通量,节点,方向,调查时段\n301pcu/h,路口甲,东进口左转,07:20-08:20\n209辆,,,\n", "交通量", "301pcu/h", "209辆", "研究名称"),
            ("design-coord", "问题,问题编号,责任人\n管线碰撞,Q-A,张工\n净高不足,,\n", "问题描述", "管线碰撞", "净高不足", "问题编号"),
        )
        for expert, source, field, first, second, identity in cases:
            with self.subTest(expert=expert):
                self.sid = "infra-csv-"+uuid4().hex[:16]
                data = source.encode("utf-8")
                uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                            files={"files": ("合成基础设施.csv", data, "text/csv")})
                self.assertEqual(200, uploaded.status_code, uploaded.text)
                identifier = uploaded.json()["files"][0]["id"]
                original, _ = flow.uploads.upload_file(self.sid, identifier)
                md, done = self.markdown("请按选中资料生成本岗内部草稿，缺项保持待填。", expert,
                                         attachments=[identifier], confirm_text=ACK)
                title = TITLES[expert]
                tables = [(name, table) for name, table in tables_from_md(md) if name.startswith(title)]
                self.assertEqual(2, len(tables))
                self.assertEqual([first, second], [table[1][table[0].index(field)] for _, table in tables])
                self.assertEqual(UNKNOWN, tables[1][1][1][tables[1][1][0].index(identity)])
                self.assertEqual(tables[0][1], self.spreadsheet_rows(done, title))
                self.assertEqual(tables[0][1], self.word_table_rows(done, tables[0][1][0], title=title))
                artifact = next(item for item in done["deliverables"] if item["name"].endswith(".xlsx"))
                import openpyxl
                book = openpyxl.load_workbook(BytesIO(self.client.get("/api/file", params={"path": artifact["path"]}).content))
                try:
                    actual = [[value for value in r] for r in book[tables[1][0]].iter_rows(values_only=True)]
                    self.assertEqual(tables[1][1], [[v if v is not None else "" for v in r] for r in actual])
                finally:
                    book.close()
                self.assertEqual(data, original.read_bytes())

    def test_two_movements_and_two_pipes_are_independent_downloadable_rows(self):
        for expert, message, title, key, expected, missing in (
            ("traffic", "生成交通报告骨架。节点：路口甲；调查时段：07:20-08:20；东进口左转301pcu/h；流量文件：a.csv；现状信号周期：109秒；西进口右转172辆", TITLES["traffic"], "交通量", ["301pcu/h", "172辆"], "流量文件"),
            ("municipal", "生成市政道路原则稿。道路名称：东路；现状管线：DN300供水；管线来源：图甲；管线权属单位：单位甲；现状管线：DN200燃气", "管线接口逐项归属", "现状管线", ["DN300供水", "DN200燃气"], "管线权属单位"),
        ):
            with self.subTest(expert=expert):
                self.sid = "infra-child-"+uuid4().hex[:16]
                md, done = self.markdown(message, expert, confirm_text=ACK)
                table = dict(tables_from_md(md))[title]
                self.assertEqual(expected, [r[table[0].index(key)] for r in table[1:]])
                self.assertEqual(UNKNOWN, table[2][table[0].index(missing)])
                self.assertEqual(table, self.spreadsheet_rows(done, title))
                self.assertEqual(table, self.word_table_rows(done, table[0], title=title))

    def test_multiple_selected_files_cannot_lend_span_basis_or_source(self):
        identifiers, originals = [], []
        for name, text in (("甲.csv", "桥名,跨径,资料来源,辖区\n一桥,31m,资料甲,CN\n"),
                           ("乙.csv", "桥名,桥宽\n二桥,16m\n")):
            data = text.encode("utf-8")
            uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                        files={"files": (name, data, "text/csv")})
            self.assertEqual(200, uploaded.status_code, uploaded.text)
            identifier = uploaded.json()["files"][0]["id"]
            identifiers.append(identifier)
            original, _ = flow.uploads.upload_file(self.sid, identifier)
            originals.append((original, data))
        md, done = self.markdown("请根据选中的两个附件生成桥梁提纲，不借用其他对象资料。", "bridge",
                                 attachments=identifiers, confirm_text=ACK)
        second = md.split("# 独立材料记录 2", 1)[1]
        self.assertIn("| 桥梁名称 | 二桥 |", second)
        self.assertIn("| 跨径组合 | UNSPECIFIED |", second)
        self.assertIn("- 对象辖区：UNSPECIFIED", second)
        self.assertNotIn("资料甲", second)
        self.assertNotIn("31m", second)
        self.assertTrue(done["submit_blocked"])
        for original, data in originals:
            self.assertEqual(data, original.read_bytes())

    def test_all_six_question_turns_never_write_and_high_risk_needs_exact_ack(self):
        for expert in PROFILES:
            with self.subTest(expert=expert):
                self.sid = "infra-chat-"+uuid4().hex[:16]
                done, _ = self.post("这个岗位需要哪些资料？只解释，不生成文件。", expert_ids=[expert])
                self.assertFalse(done["wrote"])
                self.assertEqual([], done["deliverables"])
        for expert in ("bridge", "port", "tunnel"):
            for fields in ({}, {"confirm_ok": True}, {"confirm_text": ACK[:-1]}):
                with self.subTest(expert=expert, fields=fields):
                    self.sid = "infra-gate-"+uuid4().hex[:16]
                    done, _ = self.post("请生成本岗内部提纲。", expert_ids=[expert], **fields)
                    self.assertFalse(done["wrote"], done)
                    self.assertEqual([], done["deliverables"])
                    self.assertIn(ACK, done["text"])
        self.no_model.assert_not_called()

    def test_formula_like_identifier_and_markup_remain_literal_in_copies(self):
        md, done = self.markdown("生成桥梁提纲\n|桥名|跨径|\n|---|---|\n|=A&#124;B &lt; 9|字面&amp;lt;|",
                                 "bridge", confirm_text=ACK)
        table = dict(tables_from_md(md))[TITLES["bridge"]]
        self.assertEqual("=A|B < 9", table[1][0])
        self.assertEqual("字面&lt;", table[1][1])
        self.assertEqual(table, self.word_table_rows(done, table[0], title=TITLES["bridge"]))
        artifact = next(item for item in done["deliverables"] if item["name"].endswith(".xlsx"))
        import openpyxl
        book = openpyxl.load_workbook(BytesIO(self.client.get("/api/file", params={"path": artifact["path"]}).content))
        try:
            cells = [cell for sheet in book for row in sheet for cell in row]
            self.assertIn("=A|B < 9", [cell.value for cell in cells])
            self.assertFalse(any(cell.data_type == "f" for cell in cells))
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
