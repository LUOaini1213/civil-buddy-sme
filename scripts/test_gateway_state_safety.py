#!/usr/bin/env python3
"""Gateway state safety: the packing approval runs Team B once, exports never collide, busy says when to retry,
and a restart never resumes a half-run session silently.

Measured on main be54f0c before the fix (the same setups as below):

  * /api/confirm twice in a row: 200 done, 200 done, Team B ran 2 times. Twice at once: Team B ran 2 times.
    Cancel, then confirm: 200 done, the cancel was overturned.
  * Two sessions exporting in the same second: both got shipment_20260928_100000.xlsx; session A's download
    link served session B's workbook (plan PLAN-B).
  * A busy /api/tender/link and CAD: 429 with no Retry-After. A model endpoint answering 429: "模型接口返回 429，
    请检查模型名、Key 与额度" - a rate limit reported as a configuration problem.
  * A session killed in team_b_running was still team_b_running after a restart and a confirm ran Team B from
    it; a Team A crash left a sessions placeholder and a runs row with phase and status NULL for ever.
  * 8 concurrent /api/pipeline calls: 8 accepted, no cap.

Everything runs in a temp directory against the in-process gateway. No network, no model key.
Usage: python scripts/test_gateway_state_safety.py   (about 20 s)
"""
from __future__ import annotations

import datetime as dt
import http.server
import io
import json
import os
import sqlite3
import sys
import tempfile
import threading
import time
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TMP = Path(tempfile.mkdtemp(prefix="cb-state-safety-"))
os.environ.update({
    "PYTHON_DOTENV_DISABLED": "1", "PACKING_LLM_AGENT": "0", "PACKING_SKIP_SKJOLBER": "1", "CB_STORAGE": "sqlite",
    "PACKING_OUTPUT_DIR": str(TMP / "out"), "PACKING_TRACE_DIR": str(TMP / "out" / "traces"),
    "CB_DB_PATH": str(TMP / "cb.db"), "PACKING_LG_CHECKPOINT_PATH": str(TMP / "lg.db"),
})
for _key in [k for k in os.environ if k.endswith("_API_KEY")]:
    os.environ.pop(_key)
os.environ.pop("CIVIL_TOKEN", None)

import openpyxl  # noqa: E402
from fastapi import HTTPException  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import gateway.app as G  # noqa: E402
import packing_assistant.harness as H  # noqa: E402

CLIENT = TestClient(G.app, raise_server_exceptions=False)
MATS = [{"id": "C1", "name": "crate", "quantity": 2, "weight_kg": 50,
         "length_mm": 1200, "width_mm": 800, "height_mm": 600}]
TEAM_B = {"runs": 0}
_REAL_TEAM_B = H.run_team_b


def _counted_team_b(*args, **kwargs):
    TEAM_B["runs"] += 1
    time.sleep(0.4)  # two confirms sent together really overlap
    return _REAL_TEAM_B(*args, **kwargs)


def _paused(sid: str) -> None:
    r = CLIENT.post("/api/pipeline", json={"user_input": "state safety", "session_id": sid, "materials": MATS,
                                           "preset": "", "enable_auto_confirm": False, "save_artifacts": False,
                                           "packing_options": {"crate_passthrough": True, "multi_start": False}})
    assert r.status_code == 200, r.text[:300]
    assert (r.json().get("public") or r.json()).get("phase") == "await_user_confirm", r.text[:300]


def _confirm(sid: str, path: str = "/api/confirm"):
    return CLIENT.post(path, json={"session_id": sid, "action": "confirm", "container_type": "40HQ"})


def test_sequential_double_confirm_runs_team_b_once() -> None:
    _paused("ss-seq")
    with patch.object(H, "run_team_b", _counted_team_b):
        before = TEAM_B["runs"]
        first, second = _confirm("ss-seq"), _confirm("ss-seq")
        via_resume = _confirm("ss-seq", "/api/resume/ss-seq/team-b")
        runs = TEAM_B["runs"] - before
    assert first.status_code == 200 and first.json()["phase"] == "done" and first.json()["replayed"] is False, first.text[:300]
    assert second.status_code == 200 and second.json()["replayed"] is True, second.text[:300]
    assert second.json()["phase"] == "done" and second.json().get("container_plan") == first.json().get("container_plan")
    assert via_resume.status_code == 200 and via_resume.json()["replayed"] is True, via_resume.text[:300]
    assert runs == 1, f"Team B ran {runs} times (main be54f0c: 2)"


def test_concurrent_double_confirm_runs_team_b_once() -> None:
    _paused("ss-con")
    out = []
    with patch.object(H, "run_team_b", _counted_team_b):
        before = TEAM_B["runs"]
        threads = [threading.Thread(target=lambda: out.append(_confirm("ss-con"))) for _ in range(2)]
        [t.start() for t in threads]
        [t.join() for t in threads]
        runs = TEAM_B["runs"] - before
    assert sorted(r.status_code for r in out) == [200, 200], [r.text[:200] for r in out]
    assert sorted(bool(r.json()["replayed"]) for r in out) == [False, True], [r.json().get("replayed") for r in out]
    assert runs == 1, f"Team B ran {runs} times at once (main be54f0c: 2)"


def test_cancel_is_final() -> None:
    _paused("ss-cancel")
    with patch.object(H, "run_team_b", _counted_team_b):
        before = TEAM_B["runs"]
        assert CLIENT.post("/api/confirm", json={"session_id": "ss-cancel", "action": "cancel"}).status_code == 200
        again = _confirm("ss-cancel")
        resumed = _confirm("ss-cancel", "/api/resume/ss-cancel/team-b")
        runs = TEAM_B["runs"] - before
    assert again.status_code == 409 and "cancelled" in again.json()["detail"], again.text[:300]
    assert resumed.status_code == 409, resumed.text[:300]
    assert runs == 0, runs
    assert CLIENT.get("/api/session/ss-cancel").json()["phase"] == "cancelled"
    G._SESSIONS.clear()  # and on disk, after a restart
    assert CLIENT.get("/api/session/ss-cancel").json()["phase"] == "cancelled"


def test_cancel_while_team_b_runs_wins() -> None:
    _paused("ss-race")
    started = threading.Event()

    def slow_b(*args, **kwargs):
        started.set()
        time.sleep(1.0)
        return _REAL_TEAM_B(*args, **kwargs)

    out = []
    with patch.object(H, "run_team_b", slow_b):
        t = threading.Thread(target=lambda: out.append(_confirm("ss-race")))
        t.start()
        assert started.wait(30)
        assert CLIENT.get("/api/session/ss-race").json()["phase"] == "team_b_running"
        assert CLIENT.post("/api/confirm", json={"session_id": "ss-race", "action": "cancel"}).status_code == 200
        t.join()
    assert out[0].status_code == 409 and "cancelled while Team B" in out[0].json()["detail"], out[0].text[:300]
    assert CLIENT.get("/api/session/ss-race").json()["phase"] == "cancelled"
    G._SESSIONS.clear()
    assert CLIENT.get("/api/session/ss-race").json()["phase"] == "cancelled", "the cancel must win on disk too"


def test_a_cancel_that_stops_team_b_is_409_not_500() -> None:
    from packing_assistant.runtime import cancel

    _paused("ss-stop")
    started = threading.Event()

    def stopped_b(*args, **kwargs):
        started.set()
        for _ in range(100):
            if cancel.is_cancelled("ss-stop"):
                raise cancel.RunCancelled("stopped")
            time.sleep(0.05)
        raise AssertionError("the cancel never arrived")

    out = []
    with patch.object(H, "run_team_b", stopped_b):
        t = threading.Thread(target=lambda: out.append(_confirm("ss-stop")))
        t.start()
        assert started.wait(30)
        assert CLIENT.post("/api/confirm", json={"session_id": "ss-stop", "action": "cancel"}).status_code == 200
        t.join()
    assert out[0].status_code == 409, out[0].text[:300]
    assert CLIENT.get("/api/session/ss-stop").json()["phase"] == "cancelled"


def test_a_failed_team_b_can_be_retried() -> None:
    _paused("ss-fail")

    def broken_b(*args, **kwargs):
        raise RuntimeError("solver crashed")

    with patch.object(H, "run_team_b", broken_b):
        assert _confirm("ss-fail").status_code == 500
    assert CLIENT.get("/api/session/ss-fail").json()["phase"] == "await_user_confirm"
    r = _confirm("ss-fail")
    assert r.status_code == 200 and r.json()["phase"] == "done" and r.json()["replayed"] is False, r.text[:300]


def test_nothing_to_confirm_is_409() -> None:
    G._store_session("ss-none", {"session_id": "ss-none", "phase": "materials_ready", "materials": MATS})
    r = _confirm("ss-none")
    assert r.status_code == 409 and "phase=materials_ready" in r.json()["detail"], r.text[:300]


class _Frozen(dt.datetime):
    @classmethod
    def now(cls, tz=None):
        return dt.datetime(2026, 9, 28, 10, 0, 0, tzinfo=tz)


def test_exports_in_the_same_second_never_collide() -> None:
    import packing_assistant.export_pack as ep

    for sid, plan in (("exp-A", "PLAN-A"), ("exp-B", "PLAN-B")):
        G._store_session(sid, {"session_id": sid, "packing_plan_id": plan, "container_plan": {}, "boxes": [],
                               "materials": []})
    with patch.object(ep, "datetime", _Frozen):
        a = CLIENT.post("/api/export/shipment", json={"session_id": "exp-A"}).json()
        b = CLIENT.post("/api/export/shipment", json={"session_id": "exp-B"}).json()
    name_a, name_b = Path(a["xlsx_path"]).name, Path(b["xlsx_path"]).name
    assert name_a != name_b, name_a
    assert name_a.startswith("shipment_exp-A_20260928T100000Z_") and len(name_a) == len("shipment_exp-A_20260928T100000Z_abcdef.xlsx"), name_a
    for link, plan in ((a["download_url"], "PLAN-A"), (b["download_url"], "PLAN-B")):
        got = CLIENT.get(link)
        assert got.status_code == 200, got.text[:200]
        assert openpyxl.load_workbook(io.BytesIO(got.content))["摘要"].cell(1, 2).value == plan
    # exclusive create: even a forced name clash refuses rather than overwrites
    with patch.object(ep, "_unique_name", lambda state: Path(name_a).stem):
        try:
            ep.export_shipment_xlsx({"packing_plan_id": "PLAN-C", "container_plan": {}, "boxes": [], "materials": []},
                                    output_dir=Path(a["xlsx_path"]).parent)
        except FileExistsError:
            pass
        else:
            raise AssertionError("an existing export was overwritten")
    assert openpyxl.load_workbook(a["xlsx_path"])["摘要"].cell(1, 2).value == "PLAN-A"


class _Full:
    def acquire(self, *args, **kwargs):
        return False

    def release(self):
        pass


def test_every_busy_429_says_when_to_retry() -> None:
    import demo.cad_api as cad
    import gateway.web_link as wl

    with patch.object(wl, "_RUNS", _Full()):
        r = CLIENT.post("/api/tender/link", data={"session_id": "busy"},
                        files={"tender": ("t.md", b"# t\n", "text/markdown"),
                               "panel_list": ("p.csv", b"a,b\n1,2\n", "text/csv")})
    assert r.status_code == 429 and r.headers.get("retry-after") == "5", (r.status_code, dict(r.headers))
    with patch.object(cad, "COMPUTE", _Full()):
        for fn in (cad.compute, cad.project_compute):
            try:
                fn(lambda: 1)
            except HTTPException as exc:
                assert exc.status_code == 429 and exc.headers == {"Retry-After": "5"}, (fn.__name__, exc.headers)
            else:
                raise AssertionError(fn.__name__)
    # global cap on the pipeline endpoints (CIVIL_PIPELINE_CONCURRENCY, default 4)
    codes, retry = [], []

    def slow_pipeline(*args, **kwargs):
        time.sleep(1.0)
        return {"phase": "done", "session_id": kwargs.get("session_id")}

    def one(i):
        r = CLIENT.post("/api/pipeline", json={"session_id": f"cap{i}", "materials": MATS, "preset": ""})
        codes.append(r.status_code)
        if r.status_code == 429:
            retry.append(r.headers.get("retry-after"))

    with patch.object(G, "run_agent_pipeline", slow_pipeline):
        threads = [threading.Thread(target=one, args=(i,)) for i in range(8)]
        [t.start() for t in threads]
        [t.join() for t in threads]
    assert sorted(codes) == [200] * 4 + [429] * 4, codes
    assert retry == ["5"] * 4, retry
    assert CLIENT.post("/api/pipeline", json={"session_id": "cap-after", "materials": MATS, "preset": "",
                                              "enable_auto_confirm": True}).status_code == 200, "slots come back"


def test_a_model_429_is_reported_as_rate_limiting() -> None:
    class Throttled(http.server.BaseHTTPRequestHandler):
        def do_POST(self):
            self.rfile.read(int(self.headers.get("content-length") or 0))
            self.send_response(429)
            self.send_header("Retry-After", "5")
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(b'{"error": "rate limited"}')

        def log_message(self, *args):
            pass

    from packing_assistant.runtime import model_client

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Throttled)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    env = {"CIVIL_API_BASE": f"http://127.0.0.1:{server.server_port}/v1", "CIVIL_API_KEY": "local-fake", "CIVIL_MODEL": "m"}
    try:
        with patch.dict(os.environ, env):
            model_client.complete([{"role": "user", "content": "hi"}])
    except model_client.ModelError as exc:
        message = str(exc)
    else:
        raise AssertionError("a 429 must fail the call")
    finally:
        server.shutdown()
        server.server_close()
    assert "rate-limiting (429); retry later" in message and "限流" in message and "5 秒" in message, message
    assert "Key" not in message, message


def test_restart_marks_running_sessions_interrupted() -> None:
    from packing_assistant import storage as S
    from packing_assistant.session_store import load_checkpoint_meta, recover_interrupted

    _paused("ss-src")
    src = G._get_session("ss-src")
    G.save_session("rc-b", {**src, "session_id": "rc-b", "run_id": "run-rc-b", "phase": "team_b_running",
                            "user_action": "confirm"})
    G.save_session("rc-ok", {**src, "session_id": "rc-ok", "run_id": "run-rc-ok"})  # paused: stays confirmable
    st = S.get_storage()
    st.ensure_run({"run_id": "run-rc-a", "session_id": "rc-a", "started_at": "2026-09-28T00:00:00+00:00"})
    st.ensure_run({"run_id": "run-rc-fin", "session_id": "rc-fin", "started_at": "2026-09-28T00:00:00+00:00"})
    st.insert_event({"run_id": "run-rc-fin", "session_id": "rc-fin", "type": "done", "ts": "2026-09-28T00:00:01+00:00",
                     "status": "ok", "schema": "packing.stream.v1"})
    # before the sweep the crashed Team B reads as done: Team A's final_response is still in its state
    assert (load_checkpoint_meta("rc-b") or {}).get("status") == "resumed"
    G._SESSIONS.clear()
    with TestClient(G.app, raise_server_exceptions=False) as client:  # runs the startup handlers
        b, a = client.get("/api/session/rc-b").json(), client.get("/api/session/rc-a").json()
        assert b["phase"] == "interrupted" and a["phase"] == "interrupted", (b.get("phase"), a.get("phase"))
        cb = client.post("/api/confirm", json={"session_id": "rc-b", "action": "confirm"})
        ca = client.post("/api/confirm", json={"session_id": "rc-a", "action": "confirm"})
        assert cb.status_code == 409 and "team_b_running" in cb.json()["detail"], cb.text[:300]
        assert ca.status_code == 409 and "interrupted" in ca.json()["detail"], ca.text[:300]
        ok = client.post("/api/confirm", json={"session_id": "rc-ok", "action": "confirm"})
        assert ok.status_code == 200 and ok.json()["phase"] == "done", ok.text[:300]
        pending = client.get("/api/checkpoints?pending_hitl=true&limit=100").json()["checkpoints"]
        assert not {"rc-a", "rc-b"} & {p["session_id"] for p in pending}, "interrupted is not awaiting a person"
    con = sqlite3.connect(os.environ["CB_DB_PATH"])
    try:
        runs = dict((r[0], (r[1], r[2])) for r in con.execute(
            "SELECT run_id, phase, status FROM runs WHERE run_id LIKE 'run-rc-%'"))
        fin = con.execute("SELECT status FROM sessions WHERE session_id='rc-fin'").fetchone()
    finally:
        con.close()
    assert runs["run-rc-a"] == ("interrupted", "aborted") and runs["run-rc-b"] == ("interrupted", "aborted"), runs
    assert fin == ("placeholder",), "a run that reached done is not a crash"
    assert recover_interrupted() == {"sessions": 0, "runs": 0}, "a second sweep finds nothing"


def test_json_store_is_swept_too() -> None:
    from packing_assistant.session_store import load_session, recover_interrupted, save_session

    with patch.dict(os.environ, {"CB_STORAGE": "json"}):
        save_session("rc-json", {"session_id": "rc-json", "run_id": "run-rc-json", "phase": "team_a_running"})
        assert recover_interrupted()["sessions"] == 1
        state = load_session("rc-json")
    assert state["phase"] == "interrupted" and state["interrupted_phase"] == "team_a_running", state


def main() -> int:
    tests = [
        test_sequential_double_confirm_runs_team_b_once,
        test_concurrent_double_confirm_runs_team_b_once,
        test_cancel_is_final,
        test_cancel_while_team_b_runs_wins,
        test_a_cancel_that_stops_team_b_is_409_not_500,
        test_a_failed_team_b_can_be_retried,
        test_nothing_to_confirm_is_409,
        test_exports_in_the_same_second_never_collide,
        test_every_busy_429_says_when_to_retry,
        test_a_model_429_is_reported_as_rate_limiting,
        test_restart_marks_running_sessions_interrupted,
        test_json_store_is_swept_too,
    ]
    for test in tests:
        test()
        print(f"[OK] {test.__name__}")
    print("GATEWAY STATE SAFETY PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
