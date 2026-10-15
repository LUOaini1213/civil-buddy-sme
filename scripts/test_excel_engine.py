"""Offline adversarial admission and saved-cache verification, no Office startup."""
from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
import sys
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from packing_assistant.documents.common import DocumentError
from packing_assistant.documents.excel_engine import inspect_engine_input, verify_engine_output
from packing_assistant.documents.ooxml import Spreadsheet, S, R


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Quantities"
    sheet.append(["Item", "Qty", "Rate", "Amount"])
    sheet.append(["Concrete", 6, 12.5, "=B2*C2"])
    sheet.append(["Steel", 4, 20, "=B3*C3"])
    sheet["D4"] = "=SUM(D2:D3)"
    sheet["A6"] = "Internal engineering draft"
    sheet.merge_cells("A6:D6")
    sheet["A1"].font = Font(name="Calibri", bold=True, color="FF112233")
    sheet["A1"].fill = PatternFill(patternType="solid", fgColor="FFABCDEF")
    sheet["D2"].number_format = "0.00"
    sheet.column_dimensions["A"].width = 24
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "B2"
    sheet.print_area = "A1:D6"
    other = book.create_sheet("Rates and quantities")
    other["A1"] = "=IF(Quantities!D4>0,ROUND(Quantities!D4,2),0)"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def parts(data):
    with ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def rebuild(members):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        for name, data in reversed(list(members.items())):
            archive.writestr(name, data)
    return stream.getvalue()


def change_xml(data, path, mutate):
    members = parts(data)
    root = ET.fromstring(members[path])
    mutate(root)
    members[path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return rebuild(members)


def formula(data, value):
    def mutate(root):
        root.find(f".//{{{S}}}c[@r='D2']/{{{S}}}f").text = value
    return change_xml(data, "xl/worksheets/sheet1.xml", mutate)


def saved(data):
    results = {("Quantities", "D2"): 75, ("Quantities", "D3"): 80,
               ("Quantities", "D4"): 155, ("Rates and quantities", "A1"): 155}
    members = parts(data)
    captured = []
    for part, name in (("xl/worksheets/sheet1.xml", "Quantities"), ("xl/worksheets/sheet2.xml", "Rates and quantities")):
        root = ET.fromstring(members[part])
        for cell in root.iter(f"{{{S}}}c"):
            if cell.find(f"{{{S}}}f") is not None:
                value = results[(name, cell.get("r"))]
                cell.set("t", "n")
                existing = cell.find(f"{{{S}}}v")
                if existing is None:
                    existing = ET.SubElement(cell, f"{{{S}}}v")
                existing.text = str(value)
                captured.append({"sheet": name, "cell": cell.get("r"), "type": "number", "value": value})
        members[part] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    after = rebuild(members)
    receipt = {"schema_version": 1, "engine": "Microsoft Excel", "engine_version": "16.0",
               "calculation_method": "CalculateFullRebuild", "calculation_state": "done",
               "owned_instance": True, "workbook_count": 1,
               "source_sha256": hashlib.sha256(data).hexdigest(), "output_sha256": hashlib.sha256(after).hexdigest(),
               "formula_results": captured}
    return after, receipt


class ExcelEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.original = workbook()

    def inspect(self, data):
        return inspect_engine_input(Spreadsheet(data))

    def assert_refused(self, data):
        with self.assertRaises(DocumentError):
            self.inspect(data)

    def verify_changed(self, change, *, code="excel_semantic_conflict"):
        after, receipt = saved(self.original)
        after = change(after)
        receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
        with self.assertRaises(DocumentError) as raised:
            verify_engine_output(self.original, after, receipt)
        self.assertEqual(raised.exception.code, code)

    def test_useful_selfcontained_workbook_and_full_cache_receipt(self):
        report = self.inspect(self.original)
        self.assertTrue(report["eligible"])
        self.assertEqual(report["formula_count"], 4)
        self.assertFalse(report["engine_started"])
        after, receipt = saved(self.original)
        self.assertNotEqual(self.original, after)
        result = verify_engine_output(self.original, after, receipt)
        self.assertEqual(result["verification"]["status"], "pass")
        self.assertEqual(result["verification"]["engineering_facts"], "not_checked")

    def test_supported_builtins_references_and_literal_labels(self):
        for text in ('SUM(B2:C3)', 'IF(B2>0,"https://example.invalid",0)', 'SUMIF(B2:B3,">0",C2:C3)',
                     'COUNTIF(B2:B3,">=4")', 'AND(B2>0,C2>0)', 'OR(B2=0,C2=0)',
                     'VLOOKUP(B2,B2:C3,2,FALSE)', 'INDEX(B2:C3,1,2)', 'MATCH(B2,B2:B3,0)',
                     'ROUND(-1.475,2)', "'Rates and quantities'!$A$1+1", '-B2^2+10%'):
            with self.subTest(text=text):
                self.assertTrue(self.inspect(formula(self.original, text))["eligible"])

    def test_formula_executable_external_and_dynamic_surfaces_refused(self):
        for text in ('CALL("evil.dll","run","J")', 'EXEC("cmd")', 'REGISTER.ID("evil","run")',
                     'WEBSERVICE("https://example.invalid")', 'HYPERLINK("https://example.invalid")',
                     'RTD("server",,"value")', 'IMAGE("https://example.invalid")',
                     'INDIRECT("[evil.xlsx]Sheet1!A1")', 'OFFSET(A1,1,1)', 'MyAddin(1)',
                     '_xlfn.LAMBDA(x,x)(1)', 'SUM([evil.xlsx]Sheet1!A1)', "cmd|' /C calc'!A1",
                     'SUM(SecretName)', 'SUM(A:A)', 'SUM(Quantities:Other!A1)', 'SUM(A1 B1)',
                     '{1,2}', 'A1#', 'SUM(Table1[Qty])', 'SUM(A1,)', 'IF(TRUE,1,EXEC("bad"))'):
            with self.subTest(text=text):
                self.assert_refused(formula(self.original, text))

    def test_unknown_parts_even_orphaned_are_refused(self):
        for path in ('xl/vbaProject.bin', 'xl/embeddings/oleObject1.bin', 'xl/activeX/activeX1.xml',
                     'xl/connections.xml', 'xl/queryTables/queryTable1.xml', 'xl/externalLinks/externalLink1.xml',
                     'xl/drawings/drawing1.xml', 'xl/charts/chart1.xml', 'xl/tables/table1.xml',
                     'customXml/item1.xml', '_xmlsignatures/sig1.xml', 'xl/printerSettings/printerSettings1.bin'):
            with self.subTest(path=path):
                members = parts(self.original)
                members[path] = b"opaque"
                self.assert_refused(rebuild(members))

    def test_external_and_unknown_relationships_refused(self):
        for attributes in ({"Id": "external", "Type": R + "/externalLink", "Target": "file:///secret.xlsx", "TargetMode": "External"},
                           {"Id": "custom", "Type": R + "/customXml", "Target": "../docProps/app.xml"},
                           {"Id": "strange", "Type": R + "/worksheet", "Target": "https://example.invalid/book.xlsx"}):
            self.assert_refused(change_xml(self.original, "xl/_rels/workbook.xml.rels", lambda root: ET.SubElement(root, root[0].tag, attributes)))

    def test_duplicate_relationship_ids_and_sheet_names_refused(self):
        self.assert_refused(change_xml(self.original, "xl/_rels/workbook.xml.rels", lambda root: root[1].set("Id", root[0].get("Id"))))
        self.assert_refused(change_xml(self.original, "xl/workbook.xml", lambda root: root.find(f"{{{S}}}sheets")[1].set("name", "quantities")))

    def test_static_print_names_supported_but_expressions_and_custom_names_refused(self):
        def name(value, text, **attributes):
            def mutate(root):
                container = root.find(f"{{{S}}}definedNames")
                container.clear()
                node = ET.SubElement(container, f"{{{S}}}definedName", {"name": value, "localSheetId": "0", **attributes})
                node.text = text
            return change_xml(self.original, "xl/workbook.xml", mutate)
        self.assertTrue(self.inspect(name("_xlnm.Print_Titles", "'Quantities'!$1:$2,'Quantities'!$A:$B"))["eligible"])
        for value, text, attributes in (("Secret", 'WEBSERVICE("https://example.invalid")', {}),
                                        ("_xlnm.Print_Area", 'INDIRECT("A1:D6")', {}),
                                        ("_xlnm.Print_Area", "'Unknown'!$A$1:$D$6", {}),
                                        ("_xlnm.Print_Area", "'Quantities'!A1:D6", {}),
                                        ("_xlnm.Print_Area", "'Quantities'!$A$1:$D$6", {"xlm": "1"}),
                                        ("Auto_Open", 'EXEC("cmd")', {})):
            self.assert_refused(name(value, text, **attributes))

    def test_noncell_formula_surfaces_and_extensions_refused(self):
        for tag, inner in (("conditionalFormatting", "formula"), ("dataValidations", "formula1"), ("extLst", "f")):
            def mutate(root):
                outer = ET.SubElement(root, f"{{{S}}}{tag}")
                ET.SubElement(outer, f"{{{S}}}{inner}").text = 'EXEC("cmd")'
            self.assert_refused(change_xml(self.original, "xl/worksheets/sheet1.xml", mutate))
        self.assert_refused(change_xml(self.original, "xl/worksheets/sheet1.xml", lambda root: ET.SubElement(root, f"{{{S}}}f").__setattr__("text", "B2*C2")))

    def test_shared_array_datatable_formulas_and_iteration_refused(self):
        for kind in ("shared", "array", "dataTable"):
            self.assert_refused(change_xml(self.original, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}f").set("t", kind)))
        for key, value in (("iterate", "1"), ("fullPrecision", "0"), ("refMode", "R1C1")):
            self.assert_refused(change_xml(self.original, "xl/workbook.xml", lambda root: root.find(f"{{{S}}}calcPr").set(key, value)))

    def test_duplicate_or_misplaced_cells_refused(self):
        def duplicate(root):
            row = root.find(f".//{{{S}}}row[@r='2']")
            ET.SubElement(row, f"{{{S}}}c", {"r": "D2"})
        self.assert_refused(change_xml(self.original, "xl/worksheets/sheet1.xml", duplicate))
        self.assert_refused(change_xml(self.original, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='B2']").set("r", "B3")))

    def test_ordinary_input_and_formula_changes_refused(self):
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='B2']/{{{S}}}v").__setattr__("text", "7")))
        self.verify_changed(lambda data: formula(data, "B2+C2"))

    def test_style_merge_and_layout_loss_refused(self):
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='A1']").set("s", "0")))
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.remove(root.find(f"{{{S}}}mergeCells"))))
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f"{{{S}}}cols/{{{S}}}col").set("width", "12")))

    def test_missing_error_stale_and_wrong_typed_caches_refused(self):
        for kind, value in (("n", "76"), ("e", "#DIV/0!"), ("b", "1"), ("n", "NaN")):
            def mutate(root):
                cell = root.find(f".//{{{S}}}c[@r='D2']")
                cell.set("t", kind)
                cell.find(f"{{{S}}}v").text = value
            after, receipt = saved(self.original)
            after = change_xml(after, "xl/worksheets/sheet1.xml", mutate)
            receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
            with self.assertRaises(DocumentError):
                verify_engine_output(self.original, after, receipt)
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='D2']").remove(root.find(f".//{{{S}}}c[@r='D2']/{{{S}}}v"))), code="unsupported_excel_engine")

    def test_receipt_versions_bindings_duplicates_and_incomplete_values_refused(self):
        for field, value in (("schema_version", 2), ("schema_version", True), ("engine", "Other"), ("calculation_state", "pending"),
                             ("source_sha256", "0" * 64), ("output_sha256", "0" * 64),
                             ("owned_instance", False), ("workbook_count", 2), ("workbook_count", True), ("engine_version", ""), ("engine_version", "unknown")):
            after, receipt = saved(self.original)
            receipt[field] = value
            with self.subTest(field=field), self.assertRaises(DocumentError):
                verify_engine_output(self.original, after, receipt)
        after, receipt = saved(self.original)
        receipt["formula_results"][1] = receipt["formula_results"][0]
        with self.assertRaises(DocumentError):
            verify_engine_output(self.original, after, receipt)
        after, receipt = saved(self.original)
        receipt["formula_results"].pop()
        with self.assertRaises(DocumentError):
            verify_engine_output(self.original, after, receipt)

    def test_formula_casing_and_equivalent_numeric_serialization_preserved(self):
        after, receipt = saved(self.original)
        after = change_xml(after, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='D4']/{{{S}}}f").__setattr__("text", "sum(d2:d3)"))
        after = change_xml(after, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='B2']/{{{S}}}v").__setattr__("text", "6.000"))
        receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
        self.assertEqual(verify_engine_output(self.original, after, receipt)["verification"]["status"], "pass")

    def test_real_default_row_height_change_is_not_silently_accepted(self):
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f"{{{S}}}sheetFormatPr").set("defaultRowHeight", "14.4")))

    def test_theme_whitespace_normalizes_but_actual_font_or_accent_change_fails(self):
        after, receipt = saved(self.original)
        def whitespace(root):
            for node in root.iter():
                if len(node):
                    node.text = "\n  "
        after = change_xml(after, "xl/theme/theme1.xml", whitespace)
        receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
        self.assertEqual(verify_engine_output(self.original, after, receipt)["verification"]["status"], "pass")
        from packing_assistant.documents.excel_engine import A
        self.verify_changed(lambda data: change_xml(data, "xl/theme/theme1.xml", lambda root: root.find(f"{{{A}}}themeElements/{{{A}}}clrScheme/{{{A}}}accent1/{{{A}}}srgbClr").set("val", "FF0000")))
        self.verify_changed(lambda data: change_xml(data, "xl/theme/theme1.xml", lambda root: root.find(f"{{{A}}}themeElements/{{{A}}}fontScheme/{{{A}}}minorFont/{{{A}}}latin").set("typeface", "Arial")))

    def test_empty_protection_default_omission_and_theme_cached_names_only(self):
        after, receipt = saved(self.original)
        after = change_xml(after, "xl/workbook.xml", lambda root: root.remove(root.find(f"{{{S}}}workbookProtection")))
        after = change_xml(after, "xl/styles.xml", lambda root: root.find(f"{{{S}}}fonts/{{{S}}}font/{{{S}}}name").set("val", "宋体"))
        receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
        self.assertEqual(verify_engine_output(self.original, after, receipt)["verification"]["status"], "pass")
        # A1's explicit (non-theme) Calibri font cannot be changed.
        self.verify_changed(lambda data: change_xml(data, "xl/styles.xml", lambda root: root.find(f"{{{S}}}fonts")[1].find(f"{{{S}}}name").set("val", "Arial")))

    def test_known_engine_metadata_is_exact_shape_and_cannot_hide_formulas(self):
        from packing_assistant.documents.excel_engine import CALCF
        def extension(root, hostile=False):
            outer = ET.SubElement(root, f"{{{S}}}extLst")
            ext = ET.SubElement(outer, f"{{{S}}}ext", {"uri": "{B58B0392-4F1F-4190-BB64-5DF3571DCE5F}"})
            features = ET.SubElement(ext, f"{{{CALCF}}}calcFeatures")
            ET.SubElement(features, f"{{{CALCF}}}feature", {"name": "microsoft.com:RD"})
            if hostile:
                ET.SubElement(features, f"{{{S}}}definedName", {"name": "Auto_Open"}).text = 'EXEC("cmd")'
        self.assertTrue(self.inspect(change_xml(self.original, "xl/workbook.xml", extension))["eligible"])
        self.assert_refused(change_xml(self.original, "xl/workbook.xml", lambda root: extension(root, True)))

    def test_tiny_wrong_cache_and_malformed_result_identity_do_not_pass(self):
        after, receipt = saved(self.original)
        receipt["formula_results"][0]["sheet"] = ["Quantities"]
        with self.assertRaises(DocumentError):
            verify_engine_output(self.original, after, receipt)
        after, receipt = saved(self.original)
        after = change_xml(after, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='D2']/{{{S}}}v").__setattr__("text", "1e-15"))
        receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
        receipt["formula_results"][0]["value"] = 0
        with self.assertRaises(DocumentError):
            verify_engine_output(self.original, after, receipt)

    def test_numeric_alignment_and_outline_attributes_are_not_boolean_flags(self):
        for key, value in (("indent", "2"), ("relativeIndent", "2"), ("textRotation", "45"), ("readingOrder", "2")):
            def align(root):
                node = root.find(f"{{{S}}}cellXfs")[1]
                ET.SubElement(node, f"{{{S}}}alignment", {key: value})
            with self.subTest(key=key):
                self.verify_changed(lambda data: change_xml(data, "xl/styles.xml", align))
        for path, selector, key in (("xl/worksheets/sheet1.xml", f".//{{{S}}}row[@r='1']", "outlineLevel"),
                                     ("xl/worksheets/sheet1.xml", f"{{{S}}}cols/{{{S}}}col", "outlineLevel"),
                                     ("xl/worksheets/sheet1.xml", f"{{{S}}}sheetFormatPr", "outlineLevelRow")):
            self.verify_changed(lambda data: change_xml(data, path, lambda root: root.find(selector).set(key, "2")))

    def test_inherited_formatting_and_explicit_local_suppression_are_preserved(self):
        def inherited(root):
            ET.SubElement(root.find(f"{{{S}}}cellStyleXfs/{{{S}}}xf"), f"{{{S}}}alignment", {"textRotation": "45"})
        self.verify_changed(lambda data: change_xml(data, "xl/styles.xml", inherited))
        self.verify_changed(lambda data: change_xml(data, "xl/styles.xml", lambda root: root.find(f"{{{S}}}cellXfs")[1].set("applyFont", "0")))

    def test_print_fit_outline_and_tab_properties_are_preserved(self):
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f"{{{S}}}sheetPr/{{{S}}}pageSetUpPr").set("fitToPage", "1")))
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: root.find(f"{{{S}}}sheetPr/{{{S}}}outlinePr").set("summaryBelow", "0")))
        self.verify_changed(lambda data: change_xml(data, "xl/worksheets/sheet1.xml", lambda root: ET.SubElement(root.find(f"{{{S}}}sheetPr"), f"{{{S}}}tabColor", {"rgb": "FF112233"})))

    def test_resource_limits_rich_strings_and_long_input_precision_are_strict(self):
        for value in ("(" * 65 + "B2" + ")" * 65, "SUM(A1:XFD1048576)", "1+" * 1100 + "1"):
            self.assert_refused(formula(self.original, value))
        def rich(root):
            inline = root.find(f".//{{{S}}}c[@r='A2']/{{{S}}}is")
            inline.clear()
            run = ET.SubElement(inline, f"{{{S}}}r")
            ET.SubElement(run, f"{{{S}}}t").text = "Concrete"
        self.assert_refused(change_xml(self.original, "xl/worksheets/sheet1.xml", rich))
        # Preserve raw input precision even beyond Decimal's default 28 digits.
        before = change_xml(self.original, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='B2']/{{{S}}}v").__setattr__("text", "1.123456789012345678901234567891"))
        after, receipt = saved(before)
        after = change_xml(after, "xl/worksheets/sheet1.xml", lambda root: root.find(f".//{{{S}}}c[@r='B2']/{{{S}}}v").__setattr__("text", "1.123456789012345678901234567892"))
        receipt["output_sha256"] = hashlib.sha256(after).hexdigest()
        with self.assertRaises(DocumentError) as raised:
            verify_engine_output(before, after, receipt)
        self.assertEqual(raised.exception.code, "excel_semantic_conflict")

    def test_worksheet_phonetic_default_is_preserved_and_inline_phonetic_denied(self):
        def default(root):
            ET.SubElement(root, f"{{{S}}}phoneticPr", {"fontId": "0", "type": "noConversion"})
        before = change_xml(self.original, "xl/worksheets/sheet1.xml", default)
        self.assertTrue(self.inspect(before)["eligible"])
        after, receipt = saved(before)
        self.assertEqual(verify_engine_output(before, after, receipt)["verification"]["status"], "pass")
        for key, value in (("fontId", "1"), ("alignment", "center")):
            changed = change_xml(after, "xl/worksheets/sheet1.xml", lambda root: root.find(f"{{{S}}}phoneticPr").set(key, value))
            receipt["output_sha256"] = hashlib.sha256(changed).hexdigest()
            with self.assertRaises(DocumentError) as raised:
                verify_engine_output(before, changed, receipt)
            self.assertEqual(raised.exception.code, "excel_semantic_conflict")
        for key, value in (("fontId", "999"), ("type", "unknown"), ("formula", "EXEC(1)")):
            self.assert_refused(change_xml(before, "xl/worksheets/sheet1.xml", lambda root: root.find(f"{{{S}}}phoneticPr").set(key, value)))
        self.assert_refused(change_xml(before, "xl/worksheets/sheet1.xml", lambda root: ET.SubElement(root.find(f".//{{{S}}}is"), f"{{{S}}}phoneticPr", {"fontId": "0", "type": "noConversion"})))


if __name__ == "__main__":
    unittest.main()
