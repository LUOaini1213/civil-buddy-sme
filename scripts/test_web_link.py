#!/usr/bin/env python3
"""The tender <-> packing link from a browser (gateway/web_link.py) and the page a visitor without a token lands on.

  guard      POST /api/tender/link, POST /api/tender/link/demo, GET /demo and the deliverable route answer 401 without
             CIVIL_TOKEN (a browser gets an HTML page, not the app) and create no job folder
  upload     a tender (.md / .docx / .pdf) and a panel list (.xlsx / .csv) in a job folder of their own under the data
             volume: the same tender.packing_link statements as the CLI (S1-S7, 1 covered / 2 partial / 4 for a person,
             6 x 40HQ from Clause 4.8), sha256 of the bytes sent, submit_blocked, the four deliverables
  stale      a second upload in the same session (rev B) names the statements that went stale (S2, S3, S6, S7) and the
             previous job; another session compares with nothing
  refusals   type allow-list, bytes that do not match the extension, empty, too large (file, request, unzipped),
             no Content-Length, not multipart, bad session / container type, missing file, too many files, busy
  paths      a client file name never becomes a path; the job folder is bound per context (office_job.job_root_scope),
             CIVIL_JOB_ROOT is untouched and the sandbox is not widened (a scoped folder outside its roots stays closed)
  demo       one call copies the SYNTHETIC examples/facade-demo files, runs rev A then rev B, leaves the examples as they were
  prune      10 jobs per session, 5 demo sessions, and nothing this module did not name is removed; active,
             queued and draining sessions are kept until their last request (and its tool workers) leave
  hardening  a timed-out run holds its slot until its worker exits (never more than 2 workers), the worker stops at
             the next merge step, row / cell / 20 MB caps on panel lists, English failure text (raw reason in the
             job's log), a session of its own without session_id, no device names, no '.' session folder
  review     an independent reviewer's probes: every app route closed without the token, foreign Origin, ?token=
             on a POST, hostile file names, a lying Content-Length, inputs not downloadable
  landing    / and /workbench without the token: what this is and how to get access, in English, not "gateway down";
             index.html / workbench.html handle 401 and show the :8765 link only on this machine
No model and no network.

Numeric upload/merge-timeout cases use separately named geometry-only synthetic
workbooks. They do not satisfy the original A-frame/upright/no-stack instructions.
Original uploads and the public demo retain those instructions, return no plan,
and have explicit refusal checks. Source examples and historical scores stay intact.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import unittest
import zipfile
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PYTHON_DOTENV_DISABLED"] = "1"
for _key in [k for k in os.environ if k.endswith("_API_KEY")] + ["CIVIL_SANDBOX", "CIVIL_APPROVAL", "CIVIL_JOB_ROOT"]:
    os.environ.pop(_key, None)

from fastapi.testclient import TestClient  # noqa: E402

from gateway import app as gateway  # noqa: E402
from gateway import web_link  # noqa: E402
from packing_assistant import storage  # noqa: E402

EX = ROOT / "examples" / "facade-demo"
TOKEN = "t6-web-link-token"
BEARER = {"Authorization": "Bearer " + TOKEN}
REMOTE = {"base_url": "http://remote.invalid", "client": ("192.0.2.1", 12345)}
ITT = (EX / "facade_itt_doc.md").read_bytes()
REV_A = (EX / "facade_panels.xlsx").read_bytes()
REV_B = (EX / "facade_panels_rev_b.xlsx").read_bytes()
DELIVERABLES = {"tender-packing-link.md", "bidbook.en.md", "tender-packing-link.json", "pack-plan.json"}


def geometry_control(data: bytes) -> bytes:
    """Independent synthetic transport-free input for numeric and timeout coverage."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["materials"]
    column = next(cell.column for cell in ws[1] if cell.value == "note")
    for row in ws.iter_rows(min_row=2):
        row[column - 1].value = "Geometry-only SYNTHETIC control; handling tested separately"
    stream = io.BytesIO(); wb.save(stream); wb.close()
    return stream.getvalue()


GEOMETRY_REV_A = geometry_control(REV_A)
GEOMETRY_REV_B = geometry_control(REV_B)


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def docx_bytes(text: str) -> bytes:
    """A minimal Word file: one paragraph per line (what word/document.xml needs, nothing more)."""
    from xml.sax.saxutils import escape

    paras = "".join(f"<w:p><w:r><w:t xml:space=\"preserve\">{escape(line)}</w:t></w:r></w:p>" for line in text.splitlines())
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f"<w:body>{paras}</w:body></w:document>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Override '
                   'PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                   'wordprocessingml.document.main+xml"/></Types>')
        z.writestr("word/document.xml", doc)
    return buf.getvalue()


def pdf_bytes(text: str) -> bytes:
    """A text-layer PDF written by hand (Helvetica, 55 lines a page). Non-Latin-1 punctuation is transliterated."""
    table = {"\u2264": "<=", "\u00d7": "x", "\u2014": "-", "\u2013": "-", "\u2019": "'", "\u2018": "'", "\u201c": '"',
             "\u201d": '"', "\u00e7": "c", "\u2192": "->"}
    lines = ["".join(table.get(ch, ch if ord(ch) < 128 else "?") for ch in line) for line in text.splitlines()]
    pages = [lines[i:i + 55] for i in range(0, len(lines), 55)] or [[]]
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for page in pages:
        ops = ["BT", "/F1 7 Tf", "9 TL", "30 810 Td"]
        for line in page:
            esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({esc}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream")
        content_no = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> "
                       f"/Contents {content_no} 0 R >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        body = obj if isinstance(obj, bytes) else obj.encode("latin-1")
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def csv_bytes(xlsx: bytes) -> bytes:
    import csv
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True)
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in wb["materials"].iter_rows(values_only=True):
        writer.writerow(["" if v is None else v for v in row])
    return buf.getvalue().encode("utf-8")


class Case(unittest.TestCase):
    def setUp(self) -> None:
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.dict(os.environ))
        for key in ("CIVIL_ALLOW_OPEN_LAN", "CIVIL_JOB_ROOT", "UVICORN_HOST", "CIVIL_LINK_MAX_TENDER_MB",
                    "CIVIL_LINK_MAX_PANEL_MB", "CIVIL_LINK_MAX_PANEL_ROWS"):
            os.environ.pop(key, None)
        self.tmp = Path(tempfile.mkdtemp(prefix="test-web-link-"))
        self.addCleanup(shutil.rmtree, self.tmp, True)
        os.environ["PACKING_OUTPUT_DIR"] = str(self.tmp / "output")
        os.environ["CIVIL_TOKEN"] = TOKEN
        # startup maintenance must not back up a real database while the gateway starts in a test
        self.stack.enter_context(patch.object(storage, "storage_mode", return_value="json"))
        self.client = TestClient(gateway.app)
        self.root = self.tmp / "output" / "web-link"

    def upload(self, tender=("facade_itt_doc.md", ITT), panel=("facade_panels.xlsx", REV_A), headers=BEARER, **data):
        files = {}
        if tender is not None:
            files["tender"] = tender
        if panel is not None:
            files["panel_list"] = panel
        return self.client.post("/api/tender/link", headers=headers, files=files, data=data)

    def jobs(self) -> list:
        return sorted(p for p in self.root.glob("*/*") if p.is_dir()) if self.root.is_dir() else []


class GuardTests(Case):
    def test_every_new_route_needs_the_token(self) -> None:
        self.assertEqual(401, self.upload(headers={}).status_code)
        self.assertEqual(401, self.upload(headers={"Authorization": "Bearer wrong"}).status_code)
        self.assertEqual(401, self.client.post("/api/tender/link/demo").status_code)
        self.assertEqual(401, self.client.get("/api/tender/link/file/web/20260926T000000000000Z-abcdef/bidbook.en.md").status_code)
        page = self.client.get("/demo", headers={"Accept": "text/html"})
        self.assertEqual(401, page.status_code)
        self.assertIn("text/html", page.headers["content-type"])
        self.assertIn("Access token required", page.text)
        self.assertNotIn("Run the linked demo", page.text)
        self.assertEqual(401, self.client.get("/demo").status_code)
        self.assertEqual([], self.jobs(), "a refused request created a job folder")
        del os.environ["CIVIL_TOKEN"]
        remote = TestClient(gateway.app, **REMOTE)
        self.assertEqual(403, remote.post("/api/tender/link/demo").status_code)
        self.assertEqual(403, remote.get("/demo").status_code)
        self.assertEqual([], self.jobs())

    def test_demo_page_is_english_and_builds_no_html_from_data(self) -> None:
        page = self.client.get("/demo", headers=BEARER)
        self.assertEqual(200, page.status_code)
        self.assertIn('<html lang="en">', page.text)
        self.assertIn("Run the linked demo", page.text)
        self.assertIn("SYNTHETIC", page.text)
        self.assertNotIn("innerHTML", page.text, "server data (clause text, file names) must go in as text")
        self.assertIsNone(re.search(r"[\u4e00-\u9fff]", page.text), "the demo page is English")


class UploadTests(Case):
    def upload(self, tender=("facade_itt_doc.md", ITT), panel=("geometry_panels.xlsx", GEOMETRY_REV_A), headers=BEARER, **data):
        return super().upload(tender=tender, panel=panel, headers=headers, **data)

    def test_original_transport_instructions_remain_a_supplement_request(self) -> None:
        original = self.upload(panel=("facade_panels.xlsx", REV_A), session_id="original-handling").json()
        self.assertTrue(original["ok"])
        self.assertIsNone(original["plan"])
        self.assertEqual(original["plan_refusal"]["error"], "unsupported_transport_requirements")
        self.assertEqual(original["counts"]["covered"], 0)
        self.assertEqual(original["inputs"]["panel_list"]["sha256"], sha(REV_A))
        self.assertTrue(all("A-frame" in row["requirements"]["note"] for row in original["plan_refusal"]["needs_human"]))
        self.assertTrue(original["submit_blocked"])

    def test_upload_links_the_two_files(self) -> None:
        r = self.upload(session_id="judge-1")
        self.assertEqual(200, r.status_code, r.text[:400])
        d = r.json()
        self.assertTrue(d["ok"])
        self.assertEqual([f"S{i}" for i in range(1, 8)], [s["id"] for s in d["statements"]])
        self.assertEqual({"covered": 1, "partial": 2, "gap": 0, "human_required": 4}, d["counts"])
        self.assertEqual("40HQ", d["container"]["type"])
        self.assertEqual("4.8", d["container"]["clause"])
        self.assertEqual({"container_type": "40HQ", "containers_used": 6}, {k: d["plan"][k] for k in ("container_type", "containers_used")})
        self.assertEqual(5, len(d["clauses"]))
        self.assertEqual(sha(ITT), d["inputs"]["tender"]["sha256"])
        self.assertEqual(sha(GEOMETRY_REV_A), d["inputs"]["panel_list"]["sha256"])
        self.assertEqual(d["inputs"]["tender"]["sha256"], d["inputs"]["tender"]["uploaded_as"])
        self.assertRegex(d["inputs"]["plan"]["sha256"], r"^[0-9a-f]{64}$")
        self.assertTrue(d["submit_blocked"])
        self.assertFalse(d["confirmed_by_person"])
        self.assertIsNone(d["changes_since_previous"])
        self.assertEqual([], d["stale_statements"])
        self.assertEqual(DELIVERABLES, {f["name"] for f in d["files"]})
        for s in d["statements"]:
            if s["status"] == "covered":
                self.assertNotEqual("—", s["figure"], f"{s['id']} covered without a plan figure")
        job = self.root / "judge-1" / d["job_id"]
        self.assertTrue((job / "facade_itt_doc.md").is_file())
        self.assertEqual(DELIVERABLES, {p.name for p in (job / "out").iterdir()})
        for f in d["files"]:
            got = self.client.get(f["url"], headers=BEARER)
            self.assertEqual(200, got.status_code, f["url"])
            self.assertEqual((job / "out" / f["name"]).read_text(encoding="utf-8"), got.text)
        record = json.loads((job / "out" / "tender-packing-link.json").read_text(encoding="utf-8"))
        self.assertIs(True, record["submit_blocked"])
        self.assertIs(False, record["confirmed_by_person"])
        written = [p for p in self.tmp.rglob("*") if p.is_file()]
        self.assertTrue(all(str(p).startswith(str(job)) for p in written), [str(p) for p in written])
        self.assertNotIn("CIVIL_JOB_ROOT", os.environ)
        # the deliverable route serves only the four names, only for well-formed ids
        for bad in (f"/api/tender/link/file/judge-1/{d['job_id']}/facade_itt_doc.md",
                    f"/api/tender/link/file/judge-1/{d['job_id']}/..%2F..%2Fx.json",
                    "/api/tender/link/file/judge-1/not-a-job/bidbook.en.md"):
            self.assertEqual(404, self.client.get(bad, headers=BEARER).status_code, bad)

    def test_second_upload_names_the_stale_statements(self) -> None:
        first = self.upload(session_id="judge-2").json()
        second = self.upload(panel=("geometry_panels_rev_b.xlsx", GEOMETRY_REV_B), session_id="judge-2").json()
        self.assertTrue(second["ok"], second)
        self.assertEqual(first["job_id"], second["previous_job_id"])
        self.assertEqual(["S2", "S3", "S6", "S7"], second["stale_statements"])
        self.assertEqual(8, second["plan"]["containers_used"])
        self.assertIn("panel list", second["changes_since_previous"]["summary"])
        self.assertEqual(first["inputs"]["tender"]["sha256"], second["inputs"]["tender"]["sha256"])
        self.assertNotEqual(first["inputs"]["panel_list"]["sha256"], second["inputs"]["panel_list"]["sha256"])
        other = self.upload(panel=("geometry_panels_rev_b.xlsx", GEOMETRY_REV_B), session_id="someone-else").json()
        self.assertIsNone(other["previous_job_id"])
        self.assertEqual([], other["stale_statements"])

    def test_word_pdf_and_csv_inputs(self) -> None:
        text = ITT.decode("utf-8")
        docx = self.upload(tender=("itt.docx", docx_bytes(text)), session_id="docx").json()
        self.assertTrue(docx["ok"], docx)
        self.assertEqual(5, len(docx["clauses"]))
        self.assertEqual({"covered": 1, "partial": 2, "gap": 0, "human_required": 4}, docx["counts"])
        pdf = self.upload(tender=("itt.pdf", pdf_bytes(text)), session_id="pdf").json()
        self.assertTrue(pdf["ok"], pdf)
        self.assertEqual(5, len(pdf["clauses"]))
        self.assertEqual(6, pdf["plan"]["containers_used"])
        csv = self.upload(panel=("geometry_panels.csv", csv_bytes(GEOMETRY_REV_A)), session_id="csv").json()
        self.assertTrue(csv["ok"], csv)
        self.assertEqual({"container_type": "40HQ", "containers_used": 6}, {k: csv["plan"][k] for k in ("container_type", "containers_used")})

    def test_a_client_file_name_never_becomes_a_path(self) -> None:
        for name in ("../../escape.md", "..\\..\\escape.md", "/etc/passwd.md", "C:\\Windows\\x.md", ".env.md"):
            d = self.upload(tender=(name, ITT), session_id="names").json()
            self.assertTrue(d["ok"], (name, d))
            saved = d["inputs"]["tender"]["name"]
            self.assertRegex(saved, r"^[A-Za-z0-9][A-Za-z0-9._-]*\.md$", name)
            self.assertTrue((self.root / "names" / d["job_id"] / saved).is_file(), name)
        stray = [p for p in self.tmp.rglob("*escape*") if not str(p).startswith(str(self.root))]
        self.assertEqual([], stray)
        same = self.upload(tender=("list.md", ITT), panel=("list.csv", csv_bytes(REV_A)), session_id="same").json()
        self.assertTrue(same["ok"], same)


class RefusalTests(Case):
    def refused(self, response, status: int, code: str) -> None:
        self.assertEqual(status, response.status_code, response.text[:300])
        self.assertEqual(code, response.json()["error_code"], response.text[:300])

    def test_refusals(self) -> None:
        self.refused(self.upload(tender=("itt.exe", b"MZ...")), 415, "type_not_allowed")
        self.refused(self.upload(tender=("itt.xlsx", REV_A)), 415, "type_not_allowed")
        self.refused(self.upload(panel=("panels.pdf", b"%PDF-1.4")), 415, "type_not_allowed")
        self.refused(self.upload(tender=("itt.docx", ITT)), 415, "type_mismatch")
        self.refused(self.upload(tender=("itt.pdf", ITT)), 415, "type_mismatch")
        self.refused(self.upload(tender=("itt.md", b"abc\x00def")), 415, "type_mismatch")
        self.refused(self.upload(tender=("itt.md", "caf\u00e9".encode("latin-1"))), 415, "type_mismatch")
        self.refused(self.upload(panel=("panels.xlsx", docx_bytes("not a workbook"))), 415, "type_mismatch")
        self.refused(self.upload(tender=("itt.md", b"")), 400, "empty_file")
        self.refused(self.upload(tender=None), 422, "missing_file")
        self.refused(self.upload(session_id="../x"), 400, "bad_session")
        self.refused(self.upload(session_id="demo-12345678"), 400, "bad_session")
        self.refused(self.upload(container_type="40HQ; rm -rf /"), 400, "bad_container_type")
        self.refused(self.client.post("/api/tender/link", headers=BEARER, json={"tender": "x"}), 415, "multipart_required")
        three = self.client.post("/api/tender/link", headers=BEARER, files=[
            ("tender", ("a.md", ITT)), ("panel_list", ("b.xlsx", REV_A)), ("extra", ("c.md", ITT))])
        self.refused(three, 400, "bad_form")
        chunked = self.client.post("/api/tender/link", headers={**BEARER, "Content-Type": "multipart/form-data; boundary=x"},
                                   content=iter([b"--x--\r\n"]))
        self.refused(chunked, 411, "length_required")
        self.assertEqual([], self.jobs(), "a refused upload created a job folder")

    def test_size_limits(self) -> None:
        os.environ["CIVIL_LINK_MAX_TENDER_MB"] = "0.002"     # 2 KB; the ITT is 4 KB
        self.refused(self.upload(), 413, "too_large")
        os.environ["CIVIL_LINK_MAX_TENDER_MB"] = "10"
        os.environ["CIVIL_LINK_MAX_PANEL_MB"] = "0.001"
        self.refused(self.upload(), 413, "too_large")        # the whole request is over the sum as well
        del os.environ["CIVIL_LINK_MAX_PANEL_MB"]
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/workbook.xml", "<workbook/>")
            with z.open("xl/worksheets/sheet1.xml", "w") as member:
                chunk = b"\0" * (1024 * 1024)
                for _ in range(101):
                    member.write(chunk)
        self.assertLess(len(bomb.getvalue()), 1024 * 1024)
        self.refused(self.upload(panel=("bomb.xlsx", bomb.getvalue())), 413, "too_large")
        self.assertEqual([], self.jobs())

    def test_busy_server_says_so(self) -> None:
        taken = []
        while web_link._RUNS.acquire(blocking=False):
            taken.append(1)
        try:
            self.refused(self.upload(), 429, "busy")
        finally:
            for _ in taken:
                web_link._RUNS.release()


class ScopeTests(Case):
    def test_job_root_scope_is_per_context_and_does_not_widen_the_sandbox(self) -> None:
        from packing_assistant import office_job, sandbox

        inside = self.tmp / "output" / "web-link" / "s" / "job"
        inside.mkdir(parents=True)
        (inside / "a.md").write_text("x", encoding="utf-8")
        (inside / ".env").write_text("SECRET=1", encoding="utf-8")
        default = office_job.job_root()
        seen = {}
        with office_job.job_root_scope(inside):
            self.assertEqual(inside, office_job.job_root())
            self.assertTrue(office_job.job_root_granted())
            self.assertEqual((inside / "a.md").resolve(), office_job._resolve_job_file(inside / "a.md"))
            with self.assertRaises(PermissionError):
                office_job._resolve_job_file(inside / ".env")                        # secret names stay closed
            with self.assertRaises(PermissionError):
                office_job._resolve_job_file(self.tmp / "output" / "elsewhere.md")   # outside the scoped folder
            worker = threading.Thread(target=lambda: seen.setdefault("plain", office_job.job_root()))
            worker.start()
            worker.join()
            roots = [Path(r).resolve() for r in sandbox.default_profile().allowed_write_roots]
            self.assertNotIn(inside.resolve(), roots, "the scope must not become a sandbox root")
        self.assertEqual(default, office_job.job_root())
        self.assertEqual(default, seen["plain"], "a thread started without copy_context() must not inherit the scope")
        self.assertNotIn("CIVIL_JOB_ROOT", os.environ)
        outside = Path(tempfile.mkdtemp(prefix="test-web-link-outside-"))
        self.addCleanup(shutil.rmtree, outside, True)
        (outside / "t.md").write_text("x", encoding="utf-8")
        with office_job.job_root_scope(outside):
            with self.assertRaises(PermissionError) as caught:
                office_job._resolve_job_file(outside / "t.md")
        self.assertIn("outside allowed root", str(caught.exception))


class DemoTests(Case):
    def test_demo_runs_rev_a_then_rev_b_on_copies(self) -> None:
        before = {p.name: sha(p.read_bytes()) for p in EX.iterdir() if p.is_file()}
        r = self.client.post("/api/tender/link/demo", headers=BEARER)
        self.assertEqual(200, r.status_code, r.text[:400])
        d = r.json()
        self.assertTrue(d["synthetic"])
        self.assertTrue(d["session_id"].startswith("demo-"))
        self.assertIsNone(d["rev_a"]["previous_job_id"])
        self.assertEqual(d["rev_a"]["job_id"], d["rev_b"]["previous_job_id"])
        # Historical 6/8-container results ignored these source requirements.
        # The current demo must retain them and ask for actual packaged inputs.
        for revision in ("rev_a", "rev_b"):
            self.assertIsNone(d[revision]["plan"])
            self.assertEqual(d[revision]["plan_refusal"]["error"], "unsupported_transport_requirements")
            self.assertEqual(d[revision]["counts"]["covered"], 0)
            self.assertTrue(d[revision]["submit_blocked"])
            self.assertTrue(all("A-frame" in row["requirements"]["note"] for row in d[revision]["plan_refusal"]["needs_human"]))
        self.assertEqual(["S1", "S2", "S3", "S6"], d["stale_statements"])
        self.assertIn("panel list", d["rev_b"]["changes_since_previous"]["summary"])
        self.assertEqual(sha(REV_B), d["rev_b"]["inputs"]["panel_list"]["sha256"])
        self.assertEqual(before, {p.name: sha(p.read_bytes()) for p in EX.iterdir() if p.is_file()})
        self.assertEqual(2, len(list((self.root / d["session_id"]).iterdir())))


class PruneTests(Case):
    def test_prune_preserves_active_jobs_at_every_retention_limit(self) -> None:
        ready, finish = threading.Barrier(2, timeout=10), threading.Barrier(2, timeout=10)
        errors, made = [], []
        session = self.root / "demo-active"
        for day in range(1, 13):
            (session / f"202608{day:02d}T000000000000Z-abcdef").mkdir(parents=True)

        def work():
            made.append(web_link._new_job(session))
            os.utime(session, (1, 1))
            ready.wait()
            finish.wait()
            self.assertTrue(made[0].is_dir())

        def run():
            try:
                web_link._locked_run(session.name, work)
            except BaseException as exc:
                errors.append(exc)

        with patch.object(web_link, "KEEP_JOBS_PER_SESSION", 1), patch.object(web_link, "KEEP_JOBS_TOTAL", 1):
            thread = threading.Thread(target=run)
            thread.start()
            try:
                ready.wait()
                for i in range(5):
                    newer = self.root / f"demo-new-{i}"
                    (newer / f"2099010{i + 1}T000000000000Z-abcdef").mkdir(parents=True)
                    os.utime(newer, (10 + i, 10 + i))
                web_link._locked_run("other", lambda: None)  # another request's finally prunes
                self.assertEqual(13, len(web_link._jobs(session)))
                self.assertTrue(made[0].exists())
                self.assertTrue(web_link._session_lock(session.name).locked())
            finally:
                finish.wait()
                thread.join(timeout=10)
        self.assertFalse(thread.is_alive())
        self.assertEqual([], errors)
        self.assertNotIn(session.name, web_link._SESSION_USERS)
        self.assertFalse(session.exists(), "an idle old demo becomes eligible again")

    def test_a_queued_request_keeps_its_previous_jobs_until_it_enters(self) -> None:
        queued = threading.Barrier(2, timeout=10)
        underlying = threading.Lock()
        underlying.acquire()
        session = self.root / "queued"
        for day in range(1, 13):
            (session / f"202608{day:02d}T000000000000Z-abcdef").mkdir(parents=True)
        errors, seen = [], []

        class QueuedLock:
            def acquire(self, **kwargs):
                queued.wait()  # registration has happened; the request has not acquired the session
                return underlying.acquire(**kwargs)

            def release(self):
                underlying.release()

        def run():
            try:
                web_link._locked_run(session.name, lambda: seen.append(len(web_link._jobs(session))))
            except BaseException as exc:
                errors.append(exc)

        with patch.object(web_link, "_session_lock", return_value=QueuedLock()):
            thread = threading.Thread(target=run)
            thread.start()
            try:
                queued.wait()
                web_link._prune(self.root)
                self.assertEqual(12, len(web_link._jobs(session)))
                self.assertEqual(1, web_link._SESSION_USERS.get(session.name))
            finally:
                underlying.release()
                thread.join(timeout=10)
        self.assertFalse(thread.is_alive())
        self.assertEqual([], errors)
        self.assertEqual([12], seen)
        self.assertEqual(10, len(web_link._jobs(session)))
        self.assertNotIn(session.name, web_link._SESSION_USERS)

    def test_timed_out_tool_keeps_its_real_worker_resources_until_exit(self) -> None:
        from packing_assistant.runtime.tool_engine import ToolEngine

        entered, finish = threading.Barrier(2, timeout=10), threading.Barrier(2, timeout=10)
        released = threading.Event()
        engine = ToolEngine()
        errors = []
        session = "demo-timeout"

        def handler(args):
            entered.wait()
            finish.wait()
            return {"ok": True}

        engine.register("tender.packing_link", handler, expert_id="bid-parse", timeout_s=0.01)
        when_idle = engine.when_idle

        def tracked_idle(run_id, callback):
            def done():
                try:
                    callback()
                finally:
                    released.set()
            when_idle(run_id, done)

        def run():
            try:
                web_link._locked_run(session, lambda: web_link._upload_job(
                    session, (".md", "t.md", ITT), (".xlsx", "p.xlsx", REV_A), "", ""))
            except BaseException as exc:
                errors.append(exc)

        slots = threading.BoundedSemaphore(2)
        with patch.object(web_link, "_RUNS", slots), patch.object(engine, "when_idle", tracked_idle), \
                patch("packing_assistant.runtime.tool_engine.default_engine", return_value=engine):
            thread = threading.Thread(target=run)
            thread.start()
            try:
                entered.wait()
                thread.join(timeout=5)  # response returns while the actual handler is still blocked
                self.assertFalse(thread.is_alive())
                self.assertEqual(1, len(errors))
                self.assertIsInstance(errors[0], web_link.Refusal)
                self.assertEqual("timeout", errors[0].code)
                self.assertEqual(1, sum(len(v) for v in engine._workers.values()))
                self.assertTrue(web_link._session_lock(session).locked())
                self.assertFalse(released.is_set())
                with self.assertRaises(web_link.Refusal) as busy:
                    web_link._locked_run(session, lambda: self.fail("overlapping request entered"))
                self.assertEqual(429, busy.exception.status)
                self.assertTrue(slots.acquire(blocking=False))
                try:
                    self.assertFalse(slots.acquire(blocking=False), "the timed-out worker released its slot early")
                finally:
                    slots.release()
                active = self.root / session
                os.utime(active, (1, 1))
                for i in range(5):
                    newer = self.root / f"demo-after-{i}"
                    newer.mkdir()
                    os.utime(newer, (10 + i, 10 + i))
                web_link._locked_run("other", lambda: None)
                self.assertTrue(active.is_dir(), "prune deleted a timed-out worker's job")
            finally:
                finish.wait()
                thread.join(timeout=10)
                self.assertTrue(released.wait(timeout=10))
        self.assertFalse(web_link._session_lock(session).locked())
        self.assertNotIn(session, web_link._SESSION_USERS)
        self.assertNotIn(session, web_link._DRAINING_SESSIONS)
        self.assertEqual({}, engine._workers)
        self.assertTrue(slots.acquire(blocking=False))
        self.assertTrue(slots.acquire(blocking=False))
        self.assertFalse(slots.acquire(blocking=False))
        slots.release()
        slots.release()

    def test_prune_keeps_the_latest_and_touches_only_its_own_folders(self) -> None:
        session = self.root / "s1"
        names = [f"202609{d:02d}T000000000000Z-abcdef" for d in range(1, 13)]
        for n in names:
            (session / n / "out").mkdir(parents=True)
        (session / "notes").mkdir()
        (self.root / "README.txt").parent.mkdir(parents=True, exist_ok=True)
        (self.root / "README.txt").write_text("keep", encoding="utf-8")
        for i in range(7):
            demo = self.root / f"demo-0000000{i}"
            demo.mkdir()
            os.utime(demo, (1_000_000 + i, 1_000_000 + i))
        web_link._prune(self.root)
        self.assertEqual(names[2:], sorted(p.name for p in session.iterdir() if p.name != "notes"))
        self.assertTrue((session / "notes").is_dir())
        self.assertTrue((self.root / "README.txt").is_file())
        self.assertEqual([f"demo-0000000{i}" for i in range(2, 7)], sorted(p.name for p in self.root.glob("demo-*")))


class HardeningTests(Case):
    """Round-3 review items (2026-09-27). Each failed on main 923ed38 (scratchpad probe.py, 0/9) and passes here."""

    def blocked_engine(self, gate: threading.Event, seen: dict, timeout: float = 0.05):
        """A ToolEngine whose link tool blocks until ``gate`` is set: a run the engine times out but whose worker
        is still alive, as a 60 s timeout on a big panel list is in production."""
        from packing_assistant.runtime import tool_engine

        guard = threading.Lock()

        def handler(args):
            with guard:
                seen["live"] = seen.get("live", 0) + 1
                seen["peak"] = max(seen.get("peak", 0), seen["live"])
            try:
                gate.wait(20)
            finally:
                with guard:
                    seen["live"] -= 1
            return {"ok": False, "error_code": "late"}

        def make():
            engine = tool_engine.ToolEngine()
            engine.register("tender.packing_link", handler, expert_id="bid-parse", timeout_s=timeout)
            return engine

        return patch.object(tool_engine, "default_engine", make)

    def wait_for_slots(self, slots: threading.BoundedSemaphore, n: int) -> None:
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            got = [slots.acquire(blocking=False) for _ in range(n)]
            for ok in got:
                if ok:
                    slots.release()
            if all(got):
                return
            time.sleep(0.05)
        self.fail(f"the {n} slots did not come back after the workers exited")

    def test_timed_out_runs_never_run_more_workers_than_the_limit(self) -> None:
        """Main released the slot when the request timed out, while the tool's worker kept parsing: 3 rounds of 2
        uploads left 6 workers alive and still admitted a seventh call. Now a slot is held until its worker exits."""
        gate, seen, codes = threading.Event(), {}, []
        slots = threading.BoundedSemaphore(2)

        def upload(sid):
            try:
                web_link._locked_run(sid, lambda: web_link._upload_job(
                    sid, (".md", "t.md", ITT), (".xlsx", "p.xlsx", REV_A), "", ""))
                codes.append("ran")
            except web_link.Refusal as r:
                codes.append(r.code)

        with patch.object(web_link, "_RUNS", slots), self.blocked_engine(gate, seen):
            try:
                for rnd in range(3):
                    threads = [threading.Thread(target=upload, args=(f"round{rnd}-{k}",)) for k in range(2)]
                    for t in threads:
                        t.start()
                    for t in threads:
                        t.join(10)
                        self.assertFalse(t.is_alive(), "a request waited for its worker instead of answering")
                self.assertLessEqual(seen["peak"], 2, "more tool workers alive than CIVIL_LINK_CONCURRENCY")
                self.assertEqual(2, seen["live"])
                self.assertEqual(["busy"] * 4 + ["timeout"] * 2, sorted(codes))
                with self.assertRaises(web_link.Refusal) as busy:
                    web_link._locked_run("quick", lambda: self.fail("a call was admitted while both slots drain"))
                self.assertEqual(429, busy.exception.status)
            finally:
                gate.set()
            self.wait_for_slots(slots, 2)
        self.assertEqual(0, seen["live"])
        self.assertEqual({}, web_link._SESSION_USERS)
        self.assertEqual(set(), web_link._DRAINING_SESSIONS)

    def test_a_timed_out_link_stops_at_the_next_merge_step(self) -> None:
        """The pack step of a big list is quadratic in pieces (300 rows: 79 s). Main let a timed-out worker finish
        (it ran 69 s past a 10 s timeout); with the slot now held until the worker exits, it must stop: the merge
        search checks the engine's timeout event (tools/packing.py _can_merge)."""
        import openpyxl
        from packing_assistant.runtime import tool_engine

        self.assertTrue(self.upload(session_id="warm").json()["ok"])     # imports loaded, as on a running server
        wb = openpyxl.load_workbook(io.BytesIO(GEOMETRY_REV_A))
        ws = wb["materials"]
        row = [c.value for c in ws[2]]
        row[2], row[4] = 1, row[3]
        for i in range(300):
            ws.append([f"X{i:04d}"] + row[1:])
        big = io.BytesIO()
        wb.save(big)
        engine = tool_engine.default_engine()
        engine.tools["tender.packing_link"].timeout_s = 2.0
        released = threading.Event()
        idle = engine.when_idle
        slots = threading.BoundedSemaphore(2)

        def tracked(run_id, callback):
            def done():
                try:
                    callback()
                finally:
                    released.set()
            idle(run_id, done)

        with patch.object(web_link, "_RUNS", slots), patch.object(engine, "when_idle", tracked), \
                patch.object(tool_engine, "default_engine", return_value=engine):
            t0 = time.monotonic()
            with self.assertRaises(web_link.Refusal) as caught:
                web_link._locked_run("big", lambda: web_link._upload_job(
                    "big", (".md", "t.md", ITT), (".xlsx", "p.xlsx", big.getvalue()), "", ""))
            answered = time.monotonic() - t0
            self.assertEqual("timeout", caught.exception.code)
            self.assertTrue(released.wait(15), "the timed-out worker kept packing")
            stopped = time.monotonic() - t0
        self.assertLess(stopped - answered, 10, (answered, stopped))
        self.assertEqual({}, engine._workers)

    def test_panel_list_row_and_cell_caps(self) -> None:
        self.assertEqual(20 * 1024 * 1024, web_link.MAX_UNZIPPED)
        self.assertEqual(5000, web_link.max_panel_rows())
        # the demo workbook has 8 rows (5 on 'materials', 3 on 'README'); the CSV has 5 lines
        os.environ["CIVIL_LINK_MAX_PANEL_ROWS"] = "7"
        r = self.upload(session_id="rows")
        self.assertEqual((413, "too_many_rows"), (r.status_code, r.json()["error_code"]), r.text[:300])
        self.assertIn("more than 7 rows", r.json()["detail"])
        os.environ["CIVIL_LINK_MAX_PANEL_ROWS"] = "4"
        r = self.upload(panel=("p.csv", csv_bytes(REV_A)), session_id="rows")
        self.assertEqual((413, "too_many_rows"), (r.status_code, r.json()["error_code"]))
        self.assertEqual([], self.jobs(), "a refused list created a job folder")
        os.environ["CIVIL_LINK_MAX_PANEL_ROWS"] = "8"
        self.assertTrue(self.upload(session_id="rows").json()["ok"], "a list at the limit is linked")
        os.environ["CIVIL_LINK_MAX_PANEL_ROWS"] = "5"
        self.assertTrue(self.upload(panel=("p.csv", csv_bytes(REV_A)), session_id="rows").json()["ok"])
        del os.environ["CIVIL_LINK_MAX_PANEL_ROWS"]
        before = len(self.jobs())

        def workbook(sheet_xml: bytes) -> bytes:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                z.writestr("xl/workbook.xml", "<workbook/>")
                z.writestr("xl/worksheets/sheet1.xml", sheet_xml)
            return buf.getvalue()

        cell = b'<c t="n"><v>1</v></c>'
        rows = workbook(b"<worksheet><sheetData>" + b"".join(b'<row r="%d">' % i + cell + b"</row>" for i in range(1, 5002))
                        + b"</sheetData></worksheet>")
        prefixed = workbook(b"<x:worksheet><x:sheetData>" + b"<x:row>" * 5001 + b"</x:sheetData></x:worksheet>")
        wide = workbook(b"<worksheet><sheetData><row>" + cell * 100_001 + b"</row></sheetData></worksheet>")
        inflating = io.BytesIO()
        with zipfile.ZipFile(inflating, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/workbook.xml", "<workbook/>")
            with z.open("xl/worksheets/sheet1.xml", "w") as member:
                for _ in range(21):
                    member.write(b" " * (1024 * 1024))
        csv = b"id,name,quantity\n" + b"".join(b"X%05d,panel,1\n" % i for i in range(5000))
        for name, data, code in (("rows.xlsx", rows, "too_many_rows"), ("prefixed.xlsx", prefixed, "too_many_rows"),
                                 ("wide.xlsx", wide, "too_many_rows"), ("inflating.xlsx", inflating.getvalue(), "too_large"),
                                 ("rows.csv", csv, "too_many_rows")):
            t0 = time.perf_counter()
            r = self.upload(panel=(name, data), session_id="caps")
            self.assertEqual((413, code), (r.status_code, r.json().get("error_code")), (name, r.text[:300]))
            self.assertLess(time.perf_counter() - t0, 10, f"{name}: refused only after parsing")
        self.assertEqual(before, len(self.jobs()), "a refused list created a job folder")

    def test_tool_failures_are_told_in_english_and_logged_raw(self) -> None:
        cjk = re.compile(r"[\u3000-\u9fff\uff00-\uffef]")
        broken = self.upload(tender=("broken.pdf", b"%PDF-1.4\n not a pdf"), session_id="errors")
        self.assertEqual(422, broken.status_code, broken.text[:300])
        body = broken.json()
        self.assertEqual("invalid_args", body["error_code"])
        self.assertTrue(body["detail"].startswith("the link did not run: the tender broken.pdf could not be read. A file "),
                        body["detail"])
        self.assertIsNone(cjk.search(broken.text), broken.text)
        self.assertNotIn(str(self.tmp), broken.text.replace("\\\\", "\\"))
        [job] = self.jobs()
        log = (job / "run-error.log").read_text(encoding="utf-8")
        self.assertIn("invalid_args", log)
        self.assertRegex(log, cjk, "the raw reason is kept in the job's log")
        self.assertEqual(404, self.client.get(f"/api/tender/link/file/errors/{job.name}/run-error.log",
                                              headers=BEARER).status_code)
        gate, seen = threading.Event(), {}
        slots = threading.BoundedSemaphore(2)
        with patch.object(web_link, "_RUNS", slots), self.blocked_engine(gate, seen):
            try:
                late = self.upload(session_id="errors")
            finally:
                gate.set()
            self.wait_for_slots(slots, 2)
        self.assertEqual((422, "timeout"), (late.status_code, late.json()["error_code"]), late.text[:300])
        self.assertIn("took longer than 0.05 s", late.json()["detail"])
        self.assertIsNone(cjk.search(late.text), late.text)
        odd = web_link._failure_text("weird<code>", {"reason": "\u5931\u8d25 C:\\x"}, None)
        self.assertIsNone(cjk.search(odd), odd)
        self.assertNotIn("<", odd)
        self.assertNotIn("C:", odd)

    def test_uploads_without_a_session_id_do_not_share_one(self) -> None:
        first = self.upload().json()
        second = self.upload(panel=("facade_panels_rev_b.xlsx", REV_B)).json()
        self.assertTrue(first["ok"] and second["ok"], (first, second))
        for d in (first, second):
            self.assertRegex(d["session_id"], r"^web-[0-9a-f]{16}$")
            self.assertIsNone(d["previous_job_id"], "another caller's job was used as 'previous'")
            self.assertEqual([], d["stale_statements"])
        self.assertNotEqual(first["session_id"], second["session_id"])
        self.assertNotIn(first["session_id"], web_link._SESSION_LOCKS, "the lock map grows with every upload")
        self.assertNotIn(second["session_id"], web_link._SESSION_LOCKS)
        again = self.upload(panel=("facade_panels_rev_b.xlsx", REV_B), session_id=first["session_id"]).json()
        self.assertEqual(first["job_id"], again["previous_job_id"], "a returned session_id continues that session")

    def test_no_link_route_computes_without_the_token(self) -> None:
        """The token gate still comes first: without it no route reaches _locked_run, the tool engine or the row
        count, whatever the body (the new session default and the row cap run only after the guard)."""
        from packing_assistant.runtime import tool_engine

        calls = []
        remote = TestClient(gateway.app, **REMOTE)
        big = b"id\n" + b"x\n" * 6000
        with patch.object(web_link, "_locked_run", side_effect=lambda *a, **k: calls.append("run")), \
                patch.object(web_link, "_check_rows", side_effect=lambda *a, **k: calls.append("rows")), \
                patch.object(tool_engine, "default_engine", side_effect=lambda: calls.append("engine")):
            for headers in ({}, {"Authorization": "Bearer wrong"}, {"Cookie": "cb_token=wrong"}):
                for method, path, kwargs in (
                        ("POST", "/api/tender/link", {"files": {"tender": ("t.md", ITT), "panel_list": ("p.csv", big)}}),
                        ("POST", "/api/tender/link", {"files": {"tender": ("t.md", ITT), "panel_list": ("p.xlsx", REV_A)},
                                                      "data": {"session_id": "."}}),
                        ("POST", "/api/tender/link/demo", {}),
                        ("GET", "/api/tender/link/file/web-0123456789abcdef/20260927T000000000000Z-abcdef/bidbook.en.md", {}),
                        ("GET", "/demo", {})):
                    status = remote.request(method, path, headers=headers, **kwargs).status_code
                    self.assertEqual(401, status, (method, path, headers))
        self.assertEqual([], calls, "a request without the token reached the link")
        self.assertEqual([], self.jobs())

    def test_windows_device_names_are_not_used_as_file_names(self) -> None:
        for name, saved in (("CON.md", "CON_.md"), ("nul.md", "nul_.md"), ("lpt1.md", "lpt1_.md"), ("aux.x.md", "aux.x_.md"),
                            ("CONSOLE.md", "CONSOLE.md"), ("com10.md", "com10.md")):
            self.assertEqual(saved, web_link._safe_name(name, ".md", "tender"), name)

    def test_a_session_id_never_names_the_output_root_or_a_parent(self) -> None:
        """agent_loop._safe_sid('.') was '.', so read_link_record rglobbed every session's folder and answered a
        question in one session from another session's link record."""
        from packing_assistant.runtime import agent_loop, tool_engine
        from packing_assistant.runtime.session_packing import _path
        from packing_assistant.tender_packing_link import LINK_FILE

        for sid in (".", "..", "...", "a/../..", "..\\..", "C:", "x:y", " ", "CON", "nul", "a.", "-x", "\u6807", "x" * 200):
            got = agent_loop._safe_sid(sid)
            self.assertRegex(got, r"^sid-[0-9a-f]{16}$", sid)
            self.assertEqual(got, agent_loop._safe_sid(sid), "the same id always maps to the same folder")
            self.assertEqual(got, _path(sid).parent.name)
        for sid in ("judge-1", "t-1234abcd", "sess-ab12", "eval-case.1", "default"):
            self.assertEqual(sid, agent_loop._safe_sid(sid))
        self.assertEqual("default", agent_loop._safe_sid(""))
        self.assertNotEqual(agent_loop._safe_sid("."), agent_loop._safe_sid(".."))
        d = self.upload(session_id="rec").json()
        record = self.root / "rec" / d["job_id"] / "out" / LINK_FILE
        out = self.tmp / "agent-out"
        (out / "other" / "job").mkdir(parents=True)
        shutil.copy(record, out / "other" / "job" / LINK_FILE)
        with patch.object(agent_loop, "_OUT", out):
            self.assertTrue(tool_engine._read_link_record({"session_id": "other"})["ok"])
            for sid in (".", "..", "./", "other/.."):
                got = tool_engine._read_link_record({"session_id": sid})
                self.assertEqual("no_link_record", got.get("error_code"), (sid, str(got)[:200]))


def sheet_xlsx(rows_xml: str, *, part: str = "xl/worksheets/sheet1.xml", head: str = "", tail: str = "") -> bytes:
    """A workbook openpyxl opens, with its one sheet at ``part`` (the relationships may name any path)."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" '
                   'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                   f'<Override PartName="/{part}" '
                   'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        z.writestr("_rels/.rels", f'<Relationships xmlns="{pkg}"><Relationship Id="rId1" '
                   f'Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<workbook xmlns="{main}" xmlns:r="{rel}"><sheets>'
                   '<sheet name="materials" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{pkg}"><Relationship Id="rId1" '
                   f'Type="{rel}/worksheet" Target="/{part}"/></Relationships>')
        z.writestr(part, f'<worksheet xmlns="{main}">{head}<sheetData>{rows_xml}</sheetData>{tail}</worksheet>')
    return buf.getvalue()


def sheet_row(r: int, values, first_col: str = "A") -> str:
    cells = "".join(f'<c r="{chr(ord(first_col) + i)}{r}" t="inlineStr"><is><t>{v}</t></is></c>' if isinstance(v, str)
                    else f'<c r="{chr(ord(first_col) + i)}{r}"><v>{v}</v></c>' for i, v in enumerate(values))
    return f'<row r="{r}">{cells}</row>'


PANEL_HEAD = sheet_row(1, ["Mark", "Description", "Qty", "Weight (kg)", "Length (mm)", "Width (mm)", "Height (mm)"])


def panel_rows(n: int, start: int = 2) -> str:
    return "".join(sheet_row(start + k, [f"P{k:05d}", "Unitised panel", 1, 450, 4200, 1500, 250]) for k in range(n))


class ReviewRound3Tests(Case):
    """The independent review of #77 (2026-09-28): each upload below passed the row and cell caps and then cost the
    reader seconds to hours, and with the slot now held until the worker stops, two of them kept the link busy."""

    def test_caps_count_every_part_the_row_numbers_the_columns_and_merges(self) -> None:
        far_cell = '<c r="XFD1" t="inlineStr"><is><t>x</t></is></c>'
        wide_head = PANEL_HEAD.replace("</row>", far_cell + "</row>")
        inputs = {
            # a sheet anywhere the relationships point: 60,000 rows took the reader 17 s, then the pack step
            "rows outside xl/worksheets": sheet_xlsx(PANEL_HEAD + panel_rows(5001), part="xl/sheets/data.xml"),
            # 2 KB: openpyxl pads the gap to row 1,048,576 (10 s, 410 MB)
            "last row numbered 1,048,576": sheet_xlsx(PANEL_HEAD + panel_rows(10) + sheet_row(1048576, ["END"])),
            # one header cell in column XFD: 4,000 rows took 11 s
            "a header cell in column XFD": sheet_xlsx(wide_head + panel_rows(4000)),
            # both: the reader was still running at 90 s and the slot was never released
            "column XFD and row 1,048,576": sheet_xlsx(wide_head + panel_rows(10) + sheet_row(1048576, ["END"])),
            "stated dimension A1:XFD1048576": sheet_xlsx(PANEL_HEAD + panel_rows(10), head='<dimension ref="A1:XFD1048576"/>'),
            # openpyxl makes one object per merged cell when a two-row header makes it read merges: 1 M took 27 s
            "one merge over A40:Z40000": sheet_xlsx(PANEL_HEAD + panel_rows(10), tail='<mergeCells count="1">'
                                                    '<mergeCell ref="A40:Z40000"/></mergeCells>'),
            "5,001 rows numbered without r": sheet_xlsx("<row><c><v>1</v></c></row>" * 5001),
        }
        csv_cr = b"Mark,Qty,Weight (kg)\r" + b"P1,1,450\r" * 5001        # the csv reader ends a line at a bare CR too
        before = len(self.jobs())
        for name, data in [*inputs.items(), ("CR-only csv", csv_cr)]:
            t0 = time.perf_counter()
            ext = ".csv" if name.endswith("csv") else ".xlsx"
            r = self.upload(panel=("p" + ext, data), session_id="caps3")
            self.assertEqual((413, "too_many_rows"), (r.status_code, r.json().get("error_code")), (name, r.text[:300]))
            self.assertLess(time.perf_counter() - t0, 5, f"{name}: refused only after parsing")
        self.assertEqual(before, len(self.jobs()), "a refused list created a job folder")
        # at the limits: 5,000 rows, a grid of 5,000 x 200, a small header merge
        head_200 = PANEL_HEAD.replace("</row>", '<c r="GR1" t="inlineStr"><is><t>x</t></is></c></row>')
        for name, data in (("5,000 rows", sheet_xlsx(PANEL_HEAD + panel_rows(4999))),
                           ("5,000 x 200", sheet_xlsx(head_200 + panel_rows(4999))),
                           ("header merge", sheet_xlsx(PANEL_HEAD + panel_rows(10), tail='<mergeCells count="1">'
                                                       '<mergeCell ref="E1:G1"/></mergeCells>'))):
            web_link._check_rows(data, "xlsx", name)
        # every table in the repository still passes
        tables = [p for top in ("docs", "examples", "test") for pattern in ("*.xlsx", "*.csv")
                  for p in (ROOT / top).glob(f"**/{pattern}")]
        self.assertGreater(len(tables), 20)
        for path in tables:
            kind = "xlsx" if path.suffix == ".xlsx" else "text"
            if kind == "xlsx" and not zipfile.is_zipfile(path):
                continue
            web_link._check_rows(path.read_bytes(), kind, path.name)

    def test_a_timed_out_link_stops_while_reading_a_padded_sheet(self) -> None:
        """Behind the caps as well: the reader checks the engine's timeout event while openpyxl pads rows, so a sheet
        that gets past the upload check (a workbench file, a future reader) cannot hold a slot for hours."""
        from packing_assistant.runtime import tool_engine

        self.assertTrue(self.upload(session_id="warm3").json()["ok"])
        far = sheet_xlsx(PANEL_HEAD.replace("</row>", '<c r="XFD1" t="inlineStr"><is><t>x</t></is></c></row>')
                         + panel_rows(10) + sheet_row(1048576, ["END"]))
        engine = tool_engine.default_engine()
        engine.tools["tender.packing_link"].timeout_s = 2.0
        released = threading.Event()
        idle = engine.when_idle

        def tracked(run_id, callback):
            def done():
                try:
                    callback()
                finally:
                    released.set()
            idle(run_id, done)

        with patch.object(web_link, "_RUNS", threading.BoundedSemaphore(2)), patch.object(engine, "when_idle", tracked), \
                patch.object(tool_engine, "default_engine", return_value=engine):
            t0 = time.monotonic()
            with self.assertRaises(web_link.Refusal) as caught:
                web_link._locked_run("far", lambda: web_link._upload_job(
                    "far", (".md", "t.md", ITT), (".xlsx", "p.xlsx", far), "", ""))
            answered = time.monotonic() - t0
            self.assertEqual("timeout", caught.exception.code)
            self.assertTrue(released.wait(15), "the timed-out worker kept reading padded rows")
        self.assertLess(time.monotonic() - t0 - answered, 10)

    def test_sessions_of_their_own_leave_no_empty_folders(self) -> None:
        """Every upload without a session_id has a folder of its own; pruning removed its jobs and left the folder,
        so the web-link root grew by one folder per upload (8 uploads, 3 jobs kept: 13 folders, 10 empty)."""
        with patch.object(web_link, "KEEP_JOBS_TOTAL", 2):
            for _ in range(4):
                self.assertTrue(self.upload().json()["ok"])
        folders = [p for p in self.root.iterdir() if p.is_dir()]
        self.assertEqual([], [p.name for p in folders if not any(p.iterdir())], "empty session folders left behind")
        self.assertEqual(2, len(folders))
        self.assertEqual(2, len(self.jobs()))

    def test_every_session_file_uses_one_folder_rule(self) -> None:
        """memory.py and session_handoff.py kept the old rule: session '.' wrote its summary and hand-off into
        demo/out itself, and an id such as 'x y' split one session's files across two folders."""
        from packing_assistant.runtime import agent_loop, memory, session_handoff
        from packing_assistant.runtime.session_packing import _path

        for sid in (".", "..", "x y", "CON", "C:", "default", "web-0123", ""):
            folder = agent_loop._safe_sid(sid)
            self.assertEqual({folder}, {memory.summary_path(sid).parent.name, session_handoff.handoff_path(sid).parent.name,
                                        _path(sid).parent.name}, sid)
            self.assertNotEqual("out", memory.summary_path(sid).parent.name, sid)


class LandingTests(Case):
    def test_visitor_without_token_is_told_what_this_is(self) -> None:
        for path in ("/", "/workbench"):
            page = self.client.get(path)
            self.assertEqual(200, page.status_code)
            self.assertIn("This server is private", page.text)
            self.assertIn("?token=", page.text)
            self.assertNotIn("uvicorn", page.text)
            self.assertNotIn("网关未启动", page.text)
            self.assertIsNone(re.search(r"https?://127\.0\.0\.1:8765", page.text))
        self.assertIn("投标应答", self.client.get("/", headers=BEARER).text)
        self.assertIn("gatewayDown", self.client.get("/workbench", headers=BEARER).text)
        del os.environ["CIVIL_TOKEN"]
        self.assertIn("投标应答", self.client.get("/").text)                       # this machine, no token set
        self.assertIn("This server is private", TestClient(gateway.app, **REMOTE).get("/").text)

    def test_pages_handle_401_and_hide_the_local_workbench_link(self) -> None:
        index = (ROOT / "frontend" / "index.html").read_text(encoding="utf-8")
        workbench = (ROOT / "frontend" / "workbench.html").read_text(encoding="utf-8")
        self.assertIn("this.needToken = res.status === 401", index)
        self.assertIn('v-if="gatewayDown && !needToken"', index)
        self.assertIn('data-cb-need-token="true"', index)
        self.assertIn('data-cb-need-token="true"', workbench)
        self.assertIn("gated.status === 401", workbench)
        self.assertIn('href="/demo"', index)
        # The standalone gateway additions must keep the Rust host's existing
        # same-origin packing proxy, navigation and single SSE observer.
        self.assertIn("const API = CB_UNIFIED_PACKING ? '/packing' : '';", workbench)
        self.assertIn('v-if="unifiedWorkbench || isLocalHost"', workbench)
        self.assertIn('href="/static/agent.html"', workbench)
        self.assertIn("this.wsStatus = 'SSE';", workbench)
        for name, page in (("index.html", index), ("workbench.html", workbench)):
            for line in page.splitlines():
                if re.search(r"href=\"http://127\.0\.0\.1:8765", line):
                    self.assertIn("isLocalHost", line, f"{name}: {line.strip()}")


class ReviewProbeTests(Case):
    """An independent reviewer's adversarial probes (2026-09-27), through the public routes only."""

    def test_whole_app_is_closed_without_the_token(self) -> None:
        """Every route the app registers, not only the new ones, from a remote client without the token: each HTTP
        route answers 401 and each WebSocket is refused, except /, /workbench (the landing page, not the app),
        /api/health and the /static mount. A route added later without the guard fails here."""
        from starlette.routing import Mount, WebSocketRoute

        remote = TestClient(gateway.app, **REMOTE)
        public = {"/", "/workbench", "/api/health"}
        opened, http, sockets = [], 0, 0
        for route in gateway.app.routes:
            path = re.sub(r"\{[^}]+\}", "x", getattr(route, "path", "")) or "/"
            if isinstance(route, WebSocketRoute):
                sockets += 1
                try:
                    with remote.websocket_connect(path) as ws:
                        ws.receive_text()
                        opened.append(("WS", path))
                except Exception:  # noqa: BLE001 - the guard closes the handshake; any refusal is fine
                    pass
                continue
            if isinstance(route, Mount):
                self.assertEqual("/static", route.path, "a mount other than /static is not covered here")
                continue
            for method in sorted((getattr(route, "methods", None) or {"GET"}) - {"HEAD"}):
                http += 1
                status = remote.request(method, path).status_code
                if path not in public and status != 401:
                    opened.append((method, path, status))
        self.assertEqual([], opened, "routes reachable without the token")
        self.assertGreater(http, 50)
        self.assertGreaterEqual(sockets, 1)
        for path in ("/", "/workbench"):
            self.assertIn("This server is private", remote.get(path).text)
        self.assertEqual([], self.jobs(), "a refused request created a job folder")

    def test_no_token_foreign_origin_and_query_token_write_nothing(self) -> None:
        remote = TestClient(gateway.app, **REMOTE)
        for headers in ({}, {"Authorization": "Bearer wrong"}, {"Cookie": "cb_token=wrong"},
                        {"Origin": "https://evil.example"}):
            r = remote.post("/api/tender/link", headers=headers,
                            files={"tender": ("t.md", ITT), "panel_list": ("p.xlsx", REV_A)})
            self.assertEqual(401, r.status_code, headers)
        on_post = self.client.post("/api/tender/link?token=" + TOKEN,
                                   files={"tender": ("t.md", ITT), "panel_list": ("p.xlsx", REV_A)})
        self.assertEqual(401, on_post.status_code, "?token= is a GET-only link, never a POST credential")
        pre = self.client.options("/api/tender/link", headers={"Origin": "https://evil.example",
                                                               "Access-Control-Request-Method": "POST"})
        self.assertNotIn(pre.headers.get("access-control-allow-origin"), ("*", "https://evil.example"))
        self.assertEqual([], self.jobs())
        self.assertFalse(self.root.exists() and any(self.root.rglob("*")), "a refused request wrote under web-link")

    def test_hostile_file_names_stay_in_the_job_folder(self) -> None:
        # ".md" is a dot-file with no extension: refused by the allow-list, which is fine
        self.assertEqual(415, self.upload(tender=(".md", ITT), session_id="hostile").status_code)
        names = ("CON.md", "nul.md", "..md", "a/../../../b.md", "....//....//x.md", "%2e%2e%2fx.md",
                 "x" * 400 + ".md", "\u6807\u4e66.md", "e\u0301.md", "t.md\n.md")
        for name in names:
            d = self.upload(tender=(name, ITT), session_id="hostile").json()
            self.assertTrue(d["ok"], (name, d))
            saved = d["inputs"]["tender"]["name"]
            self.assertRegex(saved, r"^[A-Za-z0-9][A-Za-z0-9._-]{0,80}\.md$", name)
            self.assertTrue((self.root / "hostile" / d["job_id"] / saved).is_file(), name)
        outside = [p for p in self.tmp.rglob("*") if p.is_file() and self.root not in p.parents]
        self.assertEqual([], outside, "a client file name put a file outside the web-link root")

    def test_declared_length_and_inputs_are_not_downloadable(self) -> None:
        huge = self.client.post("/api/tender/link", content=b"", headers={
            **BEARER, "Content-Type": "multipart/form-data; boundary=x", "Content-Length": "999999999999"})
        self.assertEqual(413, huge.status_code)
        d = self.upload(session_id="inp", project_name="<img src=x onerror=alert(1)>").json()
        self.assertTrue(d["ok"], d)
        tender = d["inputs"]["tender"]["name"]
        self.assertEqual(404, self.client.get(f"/api/tender/link/file/inp/{d['job_id']}/{tender}", headers=BEARER).status_code)
        for bad in (f"/api/tender/link/file/inp/{d['job_id']}/..%2f{tender}",
                    f"/api/tender/link/file/..%2finp/{d['job_id']}/bidbook.en.md"):
            self.assertEqual(404, self.client.get(bad, headers=BEARER).status_code, bad)
        for f in d["files"]:
            got = self.client.get(f["url"], headers=BEARER)
            self.assertEqual(200, got.status_code, f)
            self.assertNotIn("text/html", got.headers["content-type"])
            self.assertEqual("nosniff", got.headers.get("x-content-type-options"))
        self.assertTrue(d["submit_blocked"])
        self.assertFalse(d["confirmed_by_person"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
