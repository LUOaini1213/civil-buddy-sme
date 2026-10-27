#!/usr/bin/env python3
"""Inspect actual daily-report and related draft artifacts for unsupported facts."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from packing_assistant import expert_turn


class DailyReportContentTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory(prefix="civil-daily-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        output = patch.object(expert_turn, "_OUT", self.root)
        output.start()
        self.addCleanup(output.stop)
        environment = patch.dict(os.environ, {"CIVIL_SANDBOX_ROOTS": str(self.root), "CIVIL_JOB_ROOT": ""})
        environment.start()
        self.addCleanup(environment.stop)

    def _draft(self, tool: str, text: str, **fields) -> tuple[str, dict]:
        result = expert_turn.run_named_exclusive(tool, {
            "text": text, "session_id": "content-test", "confirm_ok": True, **fields,
        })
        self.assertTrue(result["wrote"], result.get("reply"))
        self.assertTrue(result["submit_blocked"])
        path = next(Path(row["path"]) for row in result["files"] if row["path"].endswith(".md"))
        return path.read_text(encoding="utf-8"), result

    def _chapter(self, text: str, title: str) -> str:
        return text.split(f"## {title}\n", 1)[1].split("\n## ", 1)[0]

    def test_template_request_does_not_become_site_or_assume_singapore(self) -> None:
        text, result = self._draft("pm-daily__log", "@项目日报 写一份项目日报模板")
        self.assertIn("- 辖区：UNSPECIFIED", text)
        self.assertIn("| 部位 | [A001] 待填部位 |", text)
        self.assertIn("| 天气 | 天气待填 |", text)
        self.assertIn("| 出勤 | 出勤待填 |", text)
        self.assertNotIn("| 部位 | @", text)
        self.assertNotIn("SG：BCA", text)
        import openpyxl

        path = next(Path(row["path"]) for row in result["files"] if row["path"].endswith(".xlsx"))
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            rows = list(workbook.worksheets[0].values)
            self.assertIn(("部位", "[A001] 待填部位"), rows)
            self.assertNotIn(("部位", "@项目日报 写一份项目日报模板"), rows)
        finally:
            workbook.close()

    def test_report_copies_only_supplied_site_weather_and_attendance(self) -> None:
        text, _ = self._draft("pm-daily__log", "@项目日报 写一份日报；辖区：CN；部位：1号楼三层；天气：暴雨；木工12人；钢筋工5人")
        self.assertIn("- 辖区：CN", text)
        self.assertIn("| 部位 | 1号楼三层 |", text)
        self.assertIn("用户口述：暴雨。", text)
        self.assertIn("木工12人；钢筋工5人", text)
        self.assertNotIn("17人", text)
        self.assertNotIn("SG：BCA", text)

    def test_short_daily_command_keeps_literal_weather_and_attendance(self) -> None:
        for command in ("写一份项目日报", "帮我写一份日报", "@项目日报 写一份项目日报"):
            with self.subTest(command=command):
                text, _ = self._draft("pm-daily__log", command+" 临边 晴 木工12人")
                self.assertIn("用户口述：晴。", self._chapter(text, "2 天气"))
                self.assertIn("木工12人", self._chapter(text, "5 出勤"))

    def test_short_daily_command_cannot_turn_forecasts_into_current_observations(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份项目日报 明天临边 晴 木工12人")
        self.assertIn("| 天气 | 天气待填 |", text)
        self.assertIn("| 出勤 | 出勤待填 |", text)

    def test_short_daily_command_does_not_interpret_rainwater_pipes_as_weather(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份项目日报 临边 雨水管 木工12人")
        self.assertIn("| 天气 | 天气待填 |", text)
        self.assertIn("木工12人", self._chapter(text, "5 出勤"))

    def test_pipe_names_and_trade_names_do_not_invent_weather_or_people(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份项目日报；部位：地下室；雨水管安装；防雪措施；钢筋工；出勤待填")
        self.assertIn("| 天气 | 天气待填 |", text)
        self.assertIn("| 出勤 | 出勤待填 |", text)
        self.assertIn("| 部位 | 地下室 |", text)

    def test_forecasts_and_plans_are_not_reported_as_actual_conditions(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份项目日报；明日预报晴；安排木工20人；部位：待填")
        self.assertIn("| 天气 | 天气待填 |", text)
        self.assertIn("| 出勤 | 出勤待填 |", text)
        self.assertIn("| 部位 | [A001] 待填部位 |", text)

    def test_structured_tool_fields_preserve_explicit_project_data(self) -> None:
        text, _ = self._draft("pm-daily__log", "生成项目日报模板", site="东侧作业面", weather="晴转多云", attendance="普工8人", jurisdiction="SG")
        self.assertIn("- 辖区：SG", text)
        self.assertIn("| 部位 | 东侧作业面 |", text)
        self.assertIn("用户口述：晴转多云", text)
        self.assertIn("普工8人", text)
        self.assertIn("SG：BCA", text)

    def test_named_facts_reach_report_chapters_and_excel_cells(self) -> None:
        text, result = self._draft("pm-daily__log", "@项目日报 写一份日报；项目名称：东岸改造工程；日期：2026-09-12；部位：1号楼三层；安全质量记事：围挡已完成复查；明日计划：复核东侧围挡并核对材料清单")
        header = self._chapter(text, "1 报头")
        safety = self._chapter(text, "7 安全质量记事")
        tomorrow = self._chapter(text, "8 明日拟安排")
        self.assertIn("项目名称：东岸改造工程", header)
        self.assertIn("日期：2026-09-12", header)
        self.assertNotIn("项目名称待填", header)
        self.assertIn("用户提供的安全质量记事：围挡已完成复查", safety)
        self.assertNotIn("用户未提供巡查事实", safety)
        self.assertIn("用户提供的明日计划：复核东侧围挡并核对材料清单", tomorrow)
        import openpyxl

        path = next(Path(row["path"]) for row in result["files"] if row["path"].endswith(".xlsx"))
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            cells = dict(list(workbook.worksheets[0].values)[1:])
            self.assertEqual(cells["项目名称"], "东岸改造工程")
            self.assertEqual(cells["日期"], "2026-09-12")
            self.assertEqual(cells["安全质量记事"], "围挡已完成复查")
            self.assertEqual(cells["明日计划"], "复核东侧围挡并核对材料清单")
        finally:
            workbook.close()

    def test_unlabelled_sentences_and_placeholder_fields_stay_unfilled(self) -> None:
        for request in (
            "写一份日报；东岸改造工程于2026-09-12复查围挡，明天核对材料",
            "写一份日报；项目名称：待填；日期：未提供；安全质量记事：未知；明日计划：TBD",
        ):
            with self.subTest(request=request):
                text, _ = self._draft("pm-daily__log", request)
                self.assertIn("项目名称待填", self._chapter(text, "1 报头"))
                self.assertIn("日期待填", self._chapter(text, "1 报头"))
                self.assertIn("安全质量记事待填", self._chapter(text, "7 安全质量记事"))
                self.assertIn("明日计划待填", self._chapter(text, "8 明日拟安排"))

    def test_adjacent_labels_do_not_bleed_into_other_chapters(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份日报\n工程名称:东岸改造工程 报告日期:2026年9月12日\n安全记事:围挡已复查，整改证据未提供 明日安排:整理整改证据")
        header = self._chapter(text, "1 报头")
        safety = self._chapter(text, "7 安全质量记事")
        tomorrow = self._chapter(text, "8 明日拟安排")
        self.assertIn("项目名称：东岸改造工程。日期：2026年9月12日", header)
        self.assertIn("围挡已复查，整改证据未提供", safety)
        self.assertNotIn("整理整改证据", safety)
        self.assertIn("整理整改证据", tomorrow)

    def test_narrative_weather_resources_and_next_day_reach_body_fields(self) -> None:
        text, _ = self._draft("pm-daily__log", "10月8日阴转小雨，试验楼2#楼七层剪力墙钢筋绑扎完成到C轴，"
                             "木工班13人、钢筋班17人，泵车一台停在南门院内，"
                             "下午监理查出七层西侧临边防护缺一段，已要求当天整改，"
                             "明天准备支七层顶板模板，帮我出今天的日报")
        body = text.split("## 日报字段", 1)[1]
        for label, value in (("日期", "10月8日"), ("天气", "阴转小雨"), ("部位", "2#楼七层剪力墙"),
                             ("出勤", "13人"), ("出勤", "17人"), ("机械材料", "泵车"),
                             ("安全质量记事", "临边防护缺一段"), ("明日计划", "支七层顶板模板")):
            self.assertTrue(any(label in line and value in line for line in body.splitlines()), (label, value))
        self.assertNotIn("支七层顶板模板", self._chapter(text, "4 形象进度"))
        self.assertNotIn("隐患已消除", body)

    def test_multiple_actual_work_locations_have_separate_literal_rows(self) -> None:
        text, result = self._draft("pm-daily__log", "今天两头忙，帮我记个日报：6#楼九层柱子浇了，"
                                  "来了9车商砼，泥工班到了19个人；西侧3#楼那边地下室还在挖，"
                                  "挖机两台早上坏了一台修到中午，人数班长没报上来，晚上下了阵雨停了一阵子")
        body = text.split("## 当日口述记录", 1)[1]
        self.assertIn("| 形象进度 | 6#楼九层柱子浇了 | 6#楼九层 |", body)
        self.assertIn("| 形象进度 | 西侧3#楼那边地下室还在挖 | 西侧3#楼 |", body)
        self.assertIn("| 出勤 | 泥工班到了19个人 | 未说明 |", body)
        self.assertIn("| 机械材料 | 来了9车商砼 | 未说明 |", body)
        self.assertIn("| 天气 | 阵雨 | 未说明 |", body)
        # The second building cannot inherit the first building's headcount.
        self.assertFalse(any("19" in line and "西侧3#楼" in line for line in body.splitlines()))
        import openpyxl
        path = next(Path(row["path"]) for row in result["files"] if row["path"].endswith(".xlsx"))
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            rows = [row for sheet in workbook for row in sheet.values]
            self.assertIn(("形象进度", "西侧3#楼那边地下室还在挖", "西侧3#楼"), rows)
        finally:
            workbook.close()

    def test_future_clause_context_does_not_leak_to_actual_resources_or_attendance(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份日报；明天准备安装北立面，安排木工23人，"
                             "泵车两台进场，预报晴；今天东立面幕墙安装完成，钢筋工4人")
        self.assertIn("| 出勤 | 钢筋工4人 |", text)
        self.assertIn("| 机械材料 | 待填 |", text)
        self.assertIn("| 天气 | 天气待填 |", text)
        self.assertNotIn("23人", self._chapter(text, "5 出勤"))

    def test_summary_table_is_not_grouped_under_original_input_heading(self) -> None:
        from scripts.eval_post_content import split_echo
        text, _ = self._draft("pm-daily__log", "写一份日报；日期：2026-10-08；形象进度：顶板完成两个流水段")
        body, echo = split_echo(text)
        self.assertIn("| 日期 | 2026-10-08 |", body)
        self.assertNotIn("| 日期 |", echo)

    def test_explicit_future_context_survives_commas_newlines_and_semicolons(self) -> None:
        from packing_assistant.tools.daily_observations import extract
        for text in (
            "明日计划：安装北立面，木工23人，泵车两台进场",
            "明天准备安装北立面；木工23人\n泵车两台进场",
            "明日安排：安装北立面，木工23人；预计泵车两台进场",
        ):
            with self.subTest(text=text):
                records = extract(text + "；今天东立面幕墙安装完成，钢筋工4人")
                self.assertFalse(any("23" in row.text or "泵车" in row.text for row in records if row.field != "tomorrow"))
                self.assertTrue(any(row.field == "labor" and row.text == "钢筋工4人" for row in records))

    def test_resource_requests_and_weather_assets_are_not_actual_observations(self) -> None:
        from packing_assistant.tools.daily_observations import extract
        records = extract("请出日报；泵车两台需进场；今天雨水管安装完成，防雪措施待填；"
                          "人数尚未记录；明天预报晴，泥工班28人")
        self.assertFalse(any(row.field in {"resources", "weather", "labor"} for row in records), records)

    def test_absent_and_hypothetical_workers_are_not_counted_as_attending(self) -> None:
        from packing_assistant.tools.daily_observations import extract
        for text in ("木工12人请假", "缺勤5人", "钢筋班14人未到场", "例如泥工班18人", "出勤上限30人"):
            with self.subTest(text=text):
                self.assertFalse(any(row.field == "labor" for row in extract(text)))

    def test_narrative_counts_are_whole_literal_values_not_suffixes_or_visitors(self) -> None:
        from packing_assistant.tools.daily_observations import extract
        for text in ("工人12.5人", "监理来访6人", "会议室容量30人", "请写成木工13人，钢筋班8人"):
            with self.subTest(text=text):
                self.assertFalse(any(row.field == "labor" for row in extract(text)))
        source = "今日普工1,200人，木工班16人"
        records = extract(source)
        self.assertEqual([row.text for row in records if row.field == "labor"], ["普工1,200人", "木工班16人"])
        self.assertTrue(all(row.text in source and (not row.location or row.location in source) for row in records))

    def test_english_presentation_keeps_the_same_literal_chinese_observations(self) -> None:
        text, _ = self._draft("pm-daily__log", "输出语言：English；10月8日阴转小雨；"
                             "6#楼九层柱子浇了，泥工班到了19个人；西侧3#楼那边地下室还在挖")
        self.assertIn("# Daily report draft", text)
        self.assertIn("| Date | 10月8日 |", text)
        self.assertIn("| Weather | 阴转小雨 |", text)
        self.assertIn("| Workers | 泥工班到了19个人 |", text)
        self.assertIn("| Work completed | 西侧3#楼那边地下室还在挖 | 西侧3#楼 |", text)
        self.assertIn("| Workers | 泥工班到了19个人 | Not provided |", text)

    def test_labelled_next_day_counts_require_an_actual_scope_reset(self) -> None:
        text, _ = self._draft("pm-daily__log", "写一份日报；明日计划：安装北立面；出勤：23人；"
                             "机械材料：泵车两台进场；今日出勤：钢筋工4人")
        self.assertIn("| 出勤 | 钢筋工4人 |", text)
        self.assertIn("| 机械材料 | 待填 |", text)
        self.assertIn("出勤：23人", self._chapter(text, "8 明日拟安排"))
        self.assertNotIn("23人", self._chapter(text, "5 出勤"))
        en = expert_turn._pm_daily_md("Generate a daily report. Tomorrow: install north facade. "
                                      "Workers: 23. Materials: two pumps expected. Actual workers: 4.")
        self.assertIn("| Workers | 4 |", en)
        self.assertIn("Workers：23", en.split("## Tomorrow", 1)[1])

    def test_unprovided_jurisdiction_is_unknown_in_other_written_drafts(self) -> None:
        for tool, request in (
            ("survey__record", "写一份测量记录"),
            ("construction__scheme_draft", "写一份施工方案讨论提纲"),
            ("method-hazard__judge_hazard", "写一份危大判定书 临边开挖"),
            ("lab-mix__report", "写一份配合比报告 C40"),
        ):
            with self.subTest(tool=tool):
                text, _ = self._draft(tool, request)
                if tool == "survey__record":
                    from packing_assistant.office_job import tables_from_md
                    regions = [row[header.index("辖区")] for _, table in tables_from_md(text)
                               if table and "辖区" in (header := table[0]) for row in table[1:]]
                    self.assertTrue(regions)
                    self.assertEqual({"UNSPECIFIED"}, set(regions))
                else:
                    self.assertIn("辖区：UNSPECIFIED", text)
                self.assertNotIn("辖区：SG", text)


if __name__ == "__main__":
    unittest.main()
