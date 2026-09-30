#!/usr/bin/env python3
"""The tender and the packing as one linked run (packing_assistant/tender_packing_link.py), on SYNTHETIC files.

  container   the plan is made in the type the tender's clause names (40HQ, and 40GP for a variant), 40HQ by default
              when it names none (said so); a type the planner cannot model (40OT, 40FR), a size with no type
              ("40-foot") or a type in a sentence that also says "not" gets no plan and goes to a person, who may name
              the type in the request (planned as asked; a clause naming another type is then never covered)
  no fit      24 panels in 20GP do not fit (18 of 24 crates placed): no statement is covered, no count or mass stated
  gates       a blank weight or a "10/12" quantity stops the plan and the row is named in the statements
  mass        a per-container mass clause is checked against the plan's per-container figures: within the limit,
              over it (gap), and on a cargo-only basis
  never       securing / lashing (CTU Code), A-frame stillages / upright / no stacking and delivery sequencing are never
              "covered"; crate structure pending detailed design stays with a person
  linked      a changed panel list (rev B) re-run names the statements that changed (containers 6 -> 8) and the inputs
              that moved; the same inputs change nothing; a clause removed from the tender withdraws its statement
  entry       English and Chinese trigger phrases reach it in steps mode (no model key) and write the matrix, the
              English bid-book and tender-packing-link.json; pack-ship takes the container type typed in the request
  gateway     /api/tender/delivery says materials_source = "sample" when it packs its canned materials, and takes the
              container type from the tender when the request gives none
  demo        scripts/demo_facade.py exits 0 and prints the link
No model and no network.

Container-figure cases use explicitly named geometry-only synthetic copies. The
original facade list's A-frame/upright/no-stack instructions now correctly block
automatic boxing and have their own refusal regression below; no source fixture
is edited, and a geometry copy is never presented as satisfying those instructions.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PYTHON_DOTENV_DISABLED"] = "1"
for _key in [k for k in os.environ if k.endswith("_API_KEY")] + ["CIVIL_SANDBOX", "CIVIL_APPROVAL", "CIVIL_JOB_ROOT"]:
    os.environ.pop(_key, None)
os.environ["CIVIL_AGENT_MODE"] = "steps"

FIXTURES = ROOT / "examples" / "facade-demo"
ITT = (FIXTURES / "facade_itt_doc.md").read_text(encoding="utf-8")
CONTAINER_CLAUSE = "4.8 Containers: panels fabricated overseas shall be shipped and delivered to site in 40HQ (40 ft high cube) containers."
MASS_CLAUSE = ("4.9 Container gross mass: the gross mass of each loaded container, including the container tare, shall not exceed "
               "20,000 kg to suit the site hoisting and road haulage arrangements.")
SEQUENCE_CLAUSE = ("4.11 Delivery sequence: deliveries shall be sequenced to the approved installation programme, floor by floor, "
                   "and each delivery shall be notified to the Main Contractor 48 hours in advance.")
HANDLING_CLAUSE = ("4.7 Packing and delivery: unitised panels shall be transported upright on steel A-frame stillages with the "
                   "glass faces protected; glazed panels shall not be stacked.")


def variant(**replace: str) -> str:
    text = ITT
    for old, new in replace.items():
        clause = {"container": CONTAINER_CLAUSE, "mass": MASS_CLAUSE, "sequence": SEQUENCE_CLAUSE, "handling": HANDLING_CLAUSE}[old]
        assert clause in text, old
        text = text.replace(clause, new)
    return text


class Link(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from packing_assistant.runtime import workspace

        cls.tmp = tempfile.TemporaryDirectory(prefix="tender-link-test-")
        cls.job = Path(cls.tmp.name).resolve() / "job"
        (cls.job / "inputs").mkdir(parents=True)
        for name in ("facade_itt_doc.md", "facade_panels.xlsx", "facade_panels_rev_b.xlsx"):
            shutil.copyfile(FIXTURES / name, cls.job / name)
        variants = {
            "itt_40gp.md": variant(container=CONTAINER_CLAUSE.replace("40HQ (40 ft high cube)", "40GP")),
            "itt_open_top.md": variant(container=CONTAINER_CLAUSE.replace("40HQ (40 ft high cube)", "40 ft open top")),
            "itt_no_type.md": variant(container="4.8 Containers: panels fabricated overseas shall be shipped to site in containers."),
            "itt_mass_6t.md": variant(mass=MASS_CLAUSE.replace("20,000 kg", "6,000 kg")),
            "itt_payload.md": variant(mass="4.9 Container payload: the maximum payload of each loaded container shall not exceed 2,500 kg."),
            "itt_plain.md": variant(handling="4.7 Packing: panels shall be packed in crates."),
            "itt_no_sequence.md": variant(sequence="4.11 Deliveries shall be notified to the Main Contractor in advance."),
            "itt_20gp.md": variant(container=CONTAINER_CLAUSE.replace("40HQ (40 ft high cube)", "20GP")),
            "itt_40fr.md": variant(container=CONTAINER_CLAUSE.replace("40HQ (40 ft high cube)", "40FR")),
            "itt_size_only.md": variant(container=CONTAINER_CLAUSE.replace("40HQ (40 ft high cube)", "40-foot")),
            "itt_not_stacked.md": variant(container="4.8 Containers: panels shall be shipped in 40HQ containers and shall not be stacked."),
        }
        for name, text in variants.items():
            (cls.job / name).write_text(text, encoding="utf-8")
        import openpyxl

        for original, geometry in (("facade_panels.xlsx", "geometry_panels.xlsx"),
                                   ("facade_panels_rev_b.xlsx", "geometry_panels_rev_b.xlsx")):
            wb = openpyxl.load_workbook(cls.job / original)
            ws = wb.active
            note_col = next(cell.column for cell in ws[1] if cell.value == "note")
            for row in ws.iter_rows(min_row=2):
                row[note_col - 1].value = "Geometry-only synthetic fixture; transport requirements tested separately"
            wb.save(cls.job / geometry)
            wb.close()

        head = ["id", "name", "quantity", "weight_kg", "total_weight_kg", "length_mm", "width_mm", "height_mm", "note"]
        good = ["P01", "Unitised panel L5 (SYNTHETIC)", 6, 450, 2700, 4200, 1500, 250, "glass"]
        for name, bad in (("panels_blank_weight.xlsx", ["P02", "Unitised panel L6 (SYNTHETIC)", 6, None, None, 4200, 1500, 250, "glass"]),
                          ("panels_qty_10_12.xlsx", ["P02", "Unitised panel L6 (SYNTHETIC)", "10/12", 450, None, 4200, 1500, 250, "glass"])):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "materials"
            for row in (head, good, bad):
                ws.append(row)
            wb.save(cls.job / name)
        (cls.job / "CIVIL.md").write_text("# CIVIL.md\n\n- 项目：合成示例办公楼幕墙分包\n- 辖区：SG\n", encoding="utf-8")
        cls.cwd = Path.cwd()
        home = patch.object(Path, "home", return_value=Path(cls.tmp.name) / "no-home")
        home.start()
        cls.addClassCleanup(home.stop)
        os.chdir(cls.job)
        workspace.activate(cls.job)

    @classmethod
    def tearDownClass(cls):
        from packing_assistant.runtime import workspace

        workspace.deactivate()
        os.chdir(cls.cwd)
        cls.tmp.cleanup()

    def link(self, tender: str = "facade_itt_doc.md", panels: str = "geometry_panels.xlsx", previous=None, container_type=None):
        from packing_assistant.tender_packing_link import run_link

        return run_link(str(self.job / tender), str(self.job / panels), previous=previous, container_type=container_type)

    @staticmethod
    def by_kind(out):
        return {s["kind"]: s for s in out["statements"]}

    # container type ------------------------------------------------------------------------------------------
    def test_original_facade_handling_requirements_block_automatic_boxing(self):
        out = self.link(panels="facade_panels.xlsx")
        self.assertIsNone(out["record"]["plan"])
        self.assertIsNone(out["record"]["inputs"]["plan"]["sha256"])
        refusal = out["record"]["plan_refusal"]
        self.assertEqual(refusal["error"], "unsupported_transport_requirements")
        self.assertTrue(any("A-frame" in str(item.get("requirements")) for item in refusal["needs_human"]))
        self.assertTrue(all(item["status"] != "covered" for item in out["statements"]))
        self.assertTrue(out["submit_blocked"])

    def test_plan_is_made_in_the_type_the_clause_names(self):
        out = self.link()
        record = out["record"]
        self.assertEqual(record["container"], {"type": "40HQ", "source": "tender_clause", "clause": "4.8",
                                               "reason": "Container type 40HQ taken from Clause 4.8."})
        self.assertEqual(out["plan"]["container_type"], "40HQ")
        self.assertEqual(self.by_kind(out)["container_type"]["status"], "covered")
        other = self.link("itt_40gp.md")
        self.assertEqual(other["record"]["container"]["type"], "40GP")
        self.assertEqual(other["plan"]["container_type"], "40GP")          # not the 40HQ default
        self.assertEqual(self.by_kind(other)["container_type"]["figures"]["plan_container_type"], "40GP")

    def test_no_type_in_the_tender_is_the_default_said_so(self):
        out = self.link("itt_no_type.md")
        self.assertEqual(out["record"]["container"]["source"], "default")
        self.assertIn("No container type (20GP / 40GP / 40HQ / 45HQ / OT / FR) was found in the ITT", out["record"]["container"]["reason"])
        row = self.by_kind(out)["container_type"]
        self.assertEqual(row["status"], "human_required")
        self.assertIn("planner's default", row["text"])
        self.assertIsNone(row["clause"])

    def test_a_type_the_planner_cannot_model_gets_no_plan(self):
        out = self.link("itt_open_top.md")
        record = out["record"]
        self.assertIsNone(record["container"]["type"])
        self.assertIn("40OT", record["container"]["reason"])
        self.assertIsNone(out["plan"])
        self.assertIsNone(record["plan"])
        self.assertIsNone(record["inputs"]["plan"]["sha256"])
        self.assertTrue(all(s["status"] != "covered" for s in out["statements"]), out["statements"])
        self.assertTrue(self.by_kind(out)["containers_used"]["placeholder"])
        self.assertIn("No plan", out["reply"])
        self.assertNotIn("pack-plan.json", [d["name"] for d in out["deliverables"]])
        self.assertIsNone(self.link("itt_40fr.md")["plan"])                  # 40FR: the planner has no flat rack

    def test_a_plan_that_does_not_fit_evidences_nothing(self):
        # 24 panels in 20GP: the engine stops at 9 containers holding 18 of 24 crates (N0 = 12), can_fit False.
        # Before the review fix S1 read "covered" and S2 said every piece was placed in 9 x 20GP.
        out = self.link("itt_20gp.md")
        self.assertIs(out["plan"]["can_fit"], False)
        kinds = self.by_kind(out)
        self.assertEqual(kinds["container_type"]["status"], "gap")
        self.assertEqual(kinds["containers_used"]["status"], "gap")
        self.assertTrue(kinds["containers_used"]["text"].startswith("[TO CONFIRM"), kinds["containers_used"]["text"])
        self.assertIn("hold 18 of the 24 crates", kinds["containers_used"]["text"])
        self.assertNotIn("places the", kinds["containers_used"]["text"])
        self.assertEqual(kinds["gross_mass"]["status"], "human_required")
        self.assertNotIn("max_gross_kg", kinds["gross_mass"]["figures"])
        self.assertFalse([s for s in out["statements"] if s["status"] == "covered"], out["statements"])
        self.assertIn("DOES NOT FIT", out["reply"])

    def test_a_size_or_a_negated_sentence_is_not_a_type_to_plan_in(self):
        size = self.link("itt_size_only.md")
        self.assertIsNone(size["plan"])
        self.assertIn("40 ft containers without the type", size["record"]["container"]["reason"])
        self.assertEqual(self.by_kind(size)["container_type"]["clause"], "4.8")      # not "the ITT names no type"
        negated = self.link("itt_not_stacked.md")
        self.assertIsNone(negated["plan"])
        self.assertIn("does not decide whether that allows or excludes it", negated["record"]["container"]["reason"])
        for out in (size, negated):
            self.assertFalse([s for s in out["statements"] if s["status"] == "covered"])

    def test_a_type_named_in_the_request_is_planned_and_never_covers_another_clause(self):
        chosen = self.link("itt_size_only.md", container_type="40HQ")
        self.assertEqual((chosen["plan"]["container_type"], chosen["record"]["container"]["source"]), ("40HQ", "request"))
        kinds = self.by_kind(chosen)
        self.assertEqual(kinds["container_type"]["status"], "human_required")     # the clause says 40 ft, not 40HQ
        self.assertEqual(kinds["containers_used"]["status"], "human_required")
        self.assertIn("[TO CONFIRM", kinds["containers_used"]["text"])
        same = self.link(container_type="40HQ")                                   # the request agrees with Clause 4.8
        self.assertEqual(self.by_kind(same)["container_type"]["status"], "covered")
        self.assertEqual(self.by_kind(same)["containers_used"]["clause"], "4.8")
        other = self.link(container_type="40GP")                                  # the request overrides Clause 4.8
        self.assertEqual(other["plan"]["container_type"], "40GP")
        self.assertEqual(self.by_kind(other)["container_type"]["status"], "human_required")
        self.assertIsNone(self.link(container_type="40FR")["plan"])

    def test_needs_human_rows_stop_the_plan_and_are_named(self):
        for panels, reason in (("panels_blank_weight.xlsx", "missing_weight"), ("panels_qty_10_12.xlsx", "invalid_quantity")):
            out = self.link(panels=panels)
            self.assertIsNone(out["record"]["plan"], panels)
            self.assertEqual((out["plan"]["source"], out["record"]["plan_refusal"]["error"]), ("needs_human", reason))
            self.assertNotIn("pack-plan.json", [d["name"] for d in out["deliverables"]])
            kinds = self.by_kind(out)
            self.assertIn(f"P02 ({reason})", kinds["containers_used"]["text"])
            self.assertEqual(kinds["crate_structure"]["status"], "human_required")   # still one row, no plan behind it
            self.assertFalse([s for s in out["statements"] if s["status"] == "covered"], panels)

    # mass ----------------------------------------------------------------------------------------------------
    def test_uncached_formula_names_the_cell_in_reply_and_refuses_a_plan(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "materials"
        ws.append(["name", "quantity", "weight_kg", "length_mm", "width_mm", "height_mm"])
        ws.append(["SYNTHETIC formula panel", "=6+6", 450, 4200, 1500, 250])
        wb.save(self.job / "uncached_formula.xlsx")
        out = self.link(panels="uncached_formula.xlsx")
        self.assertIsNone(out["record"]["plan"])
        self.assertIn("'materials'!B2", out["reply"])
        self.assertIn("Panel-list reading stopped", out["reply"])
        self.assertNotIn("pack-plan.json", [d["name"] for d in out["deliverables"]])
        self.assertFalse([s for s in out["statements"] if s["status"] == "covered"])
        self.assertTrue(out["submit_blocked"])

    def test_mass_clause_is_checked_per_container(self):
        row = self.by_kind(self.link())["gross_mass"]
        f = row["figures"]
        self.assertEqual((f["limit_kg"], f["limit_basis"], f["max_cargo_kg"], f["container_tare_kg"], f["max_gross_kg"]),
                         (20000.0, "gross", 2582.8, 3890.0, 6472.8))
        self.assertEqual(f["margin_kg"], 13527.2)
        self.assertEqual(row["status"], "partial")            # A-frame stillage mass (4.7) is not in the figure
        self.assertIn("[TO CONFIRM", row["text"])
        over = self.by_kind(self.link("itt_mass_6t.md"))["gross_mass"]
        self.assertEqual(over["status"], "gap")
        self.assertIn("exceeds the limit by 472.8 kg", over["text"])
        payload = self.by_kind(self.link("itt_payload.md"))["gross_mass"]
        self.assertEqual((payload["figures"]["limit_basis"], payload["status"]), ("cargo", "gap"))     # 2,582.8 > 2,500
        plain = self.by_kind(self.link("itt_plain.md"))["gross_mass"]
        self.assertEqual(plain["status"], "covered")            # no unmodelled packaging: the figure is the evidence
        self.assertEqual(plain["figures"]["margin_kg"], 13527.2)

    # never covered -------------------------------------------------------------------------------------------
    def test_lashing_stillages_and_sequence_are_never_covered(self):
        for name in ("facade_itt_doc.md", "itt_40gp.md", "itt_plain.md", "itt_mass_6t.md", "itt_open_top.md"):
            out = self.link(name)
            for s in out["statements"]:
                if s["kind"] in ("securing", "handling", "delivery_sequence"):
                    self.assertEqual(s["status"], "human_required", (name, s))
                    self.assertTrue(s["placeholder"] and s["text"].startswith("[TO CONFIRM"), s)
        kinds = self.by_kind(self.link())
        self.assertEqual(kinds["securing"]["clause"], "4.10")
        self.assertIn("CTU", kinds["securing"]["text"])
        self.assertEqual(kinds["handling"]["clause"], "4.7")
        self.assertIn("A-frame stillage size, tare or capacity", kinds["handling"]["text"])
        self.assertEqual(kinds["crate_structure"]["status"], "human_required")
        self.assertEqual(kinds["crate_structure"]["figures"]["pending_design"], 24)
        self.assertEqual(kinds["containers_used"]["status"], "partial")       # crate model, not the A-frame of 4.7
        self.assertEqual(kinds["containers_used"]["figures"]["containers_used"], 6)
        self.assertEqual(kinds["containers_used"]["clause"], "4.8")          # the count is tied to the clause it satisfies

    def test_bidbook_logistics_section_cites_clause_and_figure(self):
        out = self.link()
        book = out["bidbook_markdown"]
        section = book.split("## 6. Logistics & Packing (linked to the loading plan)", 1)[1].split("\n## 7.", 1)[0]
        for s in out["statements"]:
            self.assertIn(f"**{s['id']} (Clause {s['clause']}).**", section)
        self.assertIn("6 x 40HQ", section)
        self.assertIn("6,472.8 kg", section)
        self.assertIn("sha256", section)
        self.assertIn("S$ [TO FILL]", book)                  # price stays a person's job
        self.assertIn("| ITT 4.9 | Gross mass per loaded container (S3) |", book)
        self.assertNotIn("Delivery packing was **not run**", book)
        self.assertTrue(out["submit_blocked"])

    def test_the_english_bidbook_reads_in_english_and_says_it_is_synthetic(self):
        book = self.link()["bidbook_markdown"]
        lines = book.splitlines()
        self.assertEqual([line for line in lines if re.search(r"[㐀-鿿]", line)], [])      # no Chinese row left
        self.assertEqual(lines[0], "# Contractor's Proposal (Draft) — SYNTHETIC EXAMPLE - Facade Works Subcontract for "
                                   "Synthetic Office Tower")                                     # not "INVITATION TO TENDER"
        self.assertIn("> **SYNTHETIC.** The tender this draft answers is marked SYNTHETIC", book)
        for raw in ("can_fit =", "n0 =", "N0", "mid50", "type_source", "clause_names", "readiness:", "human_required"):
            self.assertNotIn(raw, book)
        self.assertIn("| Specialist method statement named in the tender |", book)
        self.assertIn("pending detailed design", book)
        self.assertIn("lower bound 6; 24 pieces; 10,800 kg net", book)
        self.assertEqual((book.count("[TO FILL]"), book.count("S$ [TO FILL]")), (50, 12))       # qualifications and price
        from packing_assistant.bidbook.sg_facade import build_sg_facade_bidbook
        real = build_sg_facade_bidbook(tender_text="# INVITATION TO TENDER\n# Facade Works for Harbour Tower\n")
        self.assertNotIn("SYNTHETIC", real["markdown"])                    # the banner follows the tender, not the demo
        self.assertEqual((real["project_title"], real["synthetic"]), ("Facade Works for Harbour Tower", False))
        rubber = build_sg_facade_bidbook(tender_text="# INVITATION TO TENDER\n# Harbour Tower\n"
                                                    "All gaskets shall be EPDM or SYNTHETIC RUBBER.\n"
                                                    "SYNTHETIC RUBBER SETTING BLOCKS\n")
        self.assertFalse(rubber["synthetic"])
        self.assertNotIn("SYNTHETIC.**", rubber["markdown"])
        for label in ("# Facade demo pack (SYNTHETIC)\n", "SYNTHETIC: fictional data\n", "# SYNTHETIC EXAMPLE - Tower\n"):
            self.assertTrue(build_sg_facade_bidbook(tender_text=label)["synthetic"], label)

    # stay linked -----------------------------------------------------------------------------------------------
    def test_changed_panel_list_names_the_stale_statements(self):
        first = self.link()
        self.assertIsNone(first["record"]["changes_since_previous"])
        again = self.link(previous=first["record"])
        same = again["record"]["changes_since_previous"]
        self.assertEqual((same["inputs_changed"], same["changed"], same["needs_reconfirmation"]), ([], [], []))
        rev_b = self.link(panels="geometry_panels_rev_b.xlsx", previous=first["record"])
        changes = rev_b["record"]["changes_since_previous"]
        self.assertEqual([i["input"] for i in changes["inputs_changed"]], ["panel_list", "plan"])
        by_id = {c["id"]: c for c in changes["changed"]}
        self.assertEqual(by_id["S2"]["figures"]["containers_used"], [6, 8])
        self.assertEqual(by_id["S3"]["figures"]["max_cargo_kg"], [2582.8, 2862.8])
        self.assertIn("S2", changes["needs_reconfirmation"])
        self.assertIn("S3", changes["needs_reconfirmation"])
        self.assertIn("containers used 6 -> 8", changes["summary"])
        self.assertEqual([c["id"] for c in changes["unchanged"]], ["S1", "S4", "S5"])
        self.assertIn("Changed since the previous run", rev_b["bidbook_markdown"])
        dropped = self.link("itt_no_sequence.md", previous=first["record"])["record"]["changes_since_previous"]
        self.assertEqual([i["input"] for i in dropped["inputs_changed"]], ["tender"])
        self.assertEqual([w["key"] for w in dropped["withdrawn"]], ["delivery_sequence@4.11"])
        self.assertIn("withdrawn (remove from the bid): earlier S7 (Delivery sequence, Clause 4.11)", dropped["summary"])
        retyped = self.link("itt_20gp.md", previous=first["record"])["record"]["changes_since_previous"]    # tender: 40HQ -> 20GP
        self.assertEqual([i["input"] for i in retyped["inputs_changed"]], ["tender", "plan"])
        s1 = next(c for c in retyped["changed"] if c["key"] == "container_type@4.8")
        self.assertTrue(s1["clause_changed"])
        self.assertEqual(s1["status"], ["covered", "gap"])
        blocked = self.link(panels="panels_blank_weight.xlsx", previous=first["record"])["record"]["changes_since_previous"]
        self.assertEqual(blocked["withdrawn"], [])          # no plan now: every earlier statement is changed, none vanishes
        self.assertEqual(len(blocked["changed"]), 7 - len(blocked["unchanged"]))
        self.assertIn("S2", blocked["needs_reconfirmation"])

    def test_record_binds_statement_clause_figures_and_hashes(self):
        import hashlib
        import json

        out = self.link()
        record = json.loads(next(d["text"] for d in out["deliverables"] if d["name"] == "tender-packing-link.json"))
        self.assertEqual(record["inputs"]["tender"]["sha256"], hashlib.sha256((self.job / "facade_itt_doc.md").read_bytes()).hexdigest())
        self.assertEqual(record["inputs"]["panel_list"]["sha256"],
                         hashlib.sha256((self.job / "geometry_panels.xlsx").read_bytes()).hexdigest())
        self.assertEqual(len(record["inputs"]["plan"]["sha256"]), 64)
        self.assertEqual(record["materials_source"], "panel_list")
        self.assertIs(record["submit_blocked"], True)
        self.assertIs(record["confirmed_by_person"], False)
        for s in record["statements"]:
            self.assertTrue({"id", "clause", "clause_sha256", "figures", "sha256", "status"} <= set(s))

    def test_reads_stay_inside_the_job_folder(self):
        from packing_assistant.tender_packing_link import run_link

        with self.assertRaises(PermissionError):
            run_link(str(FIXTURES / "facade_itt_doc.md"), str(self.job / "geometry_panels.xlsx"))

    # entry points ----------------------------------------------------------------------------------------------
    def test_trigger_phrases_route_to_the_linked_run(self):
        from packing_assistant.runtime.task_router import route_task, wants_link

        for text in ("Link the tender facade_itt_doc.md to the packing list geometry_panels.xlsx and write the logistics response",
                     "Write the tender logistics response from facade_itt_doc.md and geometry_panels.xlsx",
                     "Plan the packing of geometry_panels.xlsx under the tender's clauses in facade_itt_doc.md",
                     "按招标 facade_itt_doc.md 和装箱单 geometry_panels.xlsx 出投标物流应答",
                     "招标装柜联动：facade_itt_doc.md、geometry_panels.xlsx"):
            route = route_task(text)
            self.assertTrue(wants_link(text), text)
            self.assertEqual((route["expert_ids"], route["intent"]), (["bid-parse"], "run"), text)
        for text in ("What is a logistics response?", "物流应答是什么？"):
            self.assertEqual(route_task(text)["intent"], "chat", text)
        for text in ("解析招标 facade_itt_doc.md", "按 geometry_panels.xlsx 装柜，柜型 40HQ", "Parse the tender facade_itt_doc.md",
                     "不要做物流应答"):
            self.assertFalse(wants_link(text), text)

    def test_an_estimator_s_own_wording_reaches_the_link(self):
        """A tender document and a panel / packing list named together, and a check / match / comply / clauses word."""
        from packing_assistant.runtime.task_router import route_task, wants_link

        for text in ("Check whether geometry_panels.xlsx meets the logistics clauses of facade_itt_doc.md",
                     "Match facade_itt_doc.md against geometry_panels.xlsx",
                     "Does our loading plan for geometry_panels.xlsx comply with the container clauses in facade_itt_doc.md?",
                     "Check geometry_panels.xlsx against the shipping requirements in facade_itt_doc.md",
                     "Link ITT_Block_C.docx with panel_schedule_rev3.xlsx",
                     "核对 geometry_panels.xlsx 是否满足 facade_itt_doc.md 的物流条款",
                     "对照 facade_itt_doc.md 的运输条款检查 geometry_panels.xlsx",
                     "装箱单 geometry_panels.xlsx 符合招标 facade_itt_doc.md 的装柜要求吗？"):
            route = route_task(text)
            self.assertTrue(wants_link(text), text)
            self.assertEqual((route["expert_ids"], route["intent"]), (["bid-parse"], "run"), text)
        for text in ("How do I check geometry_panels.xlsx against facade_itt_doc.md?",          # asks about it
                     "Why is the securing statement of the link for facade_itt_doc.md and geometry_panels.xlsx left for a person?",
                     "怎么核对 geometry_panels.xlsx 是否满足 facade_itt_doc.md 的物流条款？"):
            self.assertEqual(route_task(text)["intent"], "chat", text)
        for text in ("Match the BOQ boq_facade.xlsx against the tender facade_itt_doc.md",     # not a packing list
                     "Check the price schedule rates.xlsx against the tender facade_itt_doc.md",
                     "Check the tender facade_itt_doc.md against our response bid_response.docx",  # no list at all
                     "Check whether spec.pdf meets the clauses of contract.docx",
                     "核对工程量清单 boq.xlsx 是否满足招标 facade_itt_doc.md 的要求",
                     "Don't link facade_itt_doc.md to geometry_panels.xlsx",
                     "Check geometry_panels.xlsx for missing weights"):
            self.assertFalse(wants_link(text), text)

    def test_look_alikes_do_not_run_the_link(self):
        """Review of #67: advice, hypothetical, past-tense and negated requests that name both files ran the link and
        wrote 11 files each (test/benchmarks/link_routing/dev_round3.json)."""
        from packing_assistant.runtime.task_router import route_task, wants_link

        p, t = "geometry_panels.xlsx", "facade_itt_doc.md"
        for text in (f"Should I check {p} against {t} first, or price it first?",
                     f"Is it worth matching {p} to {t} before the site visit?",
                     f"If I check {p} against {t}, will it change my crate count?",
                     f"Would it make sense to link {t} and {p}?",
                     f"I already checked {p} against {t} last week.",
                     f"Do I need to check {p} against {t}?",
                     f"Has anyone checked {p} against {t}?",
                     f"My colleague checked {p} against {t}.",
                     f"The PM checked {p} against {t} yesterday.",
                     f"我应该先核对 {p} 和 {t} 吗？",
                     f"要不要把 {p} 和 {t} 对照一下？",
                     f"你们上周已经把 {p} 和 {t} 核对过了。"):
            with self.subTest(text=text):
                self.assertEqual(route_task(text)["intent"], "chat")
        for text in (f"Pack {p} without checking {t}", f"Pack {p} without linking it to {t}",
                     f"Pack {p} but don't check it against {t}."):
            with self.subTest(text=text):
                self.assertFalse(wants_link(text))
                self.assertEqual(route_task(text)["intent"], "run")     # the packing it asks for still runs
        for text in (f"Is {p} compliant with the container clauses of {t}?", f"Please link {p} and {t}.",
                     f"Don't forget to check {p} against {t}.", f"{p} 能满足招标 {t} 的物流条款吗？"):
            with self.subTest(text=text):
                self.assertTrue(wants_link(text))
                self.assertEqual((route_task(text)["expert_ids"], route_task(text)["intent"]), (["bid-parse"], "run"))

    def test_deferred_and_past_look_alikes_do_not_run_the_link(self):
        """Review of PR #75 (dev_round3.json origin 'review-r3'): at d455e57 each of these still ran the link, because
        a sentence that opens with a verb was read as an imperative ("Remind me ...", "Linked ...", "Checking ...") or
        the check came after a deferral ("Wait until Monday, then check ..."), and "w/o" / "No checking" did not
        exclude the link from a pack request."""
        from packing_assistant.runtime.task_router import route_task, wants_link

        p, t = "geometry_panels.xlsx", "facade_itt_doc.md"
        for text in (f"Remind me tomorrow to check {p} against {t}.",
                     f"Checking {p} against {t} was a waste of time.",
                     f"Linked {t} and {p} yesterday, all fine.",
                     f"Matched {p} to {t} last week; no issues.",
                     f"Note to self: check {p} against {t} after the addendum.",
                     f"Wait until Monday, then check {p} against {t}.",
                     f"提醒我明天核对 {p} 和 {t}。"):
            with self.subTest(text=text):
                self.assertEqual(route_task(text)["intent"], "chat")
        for text in (f"Pack {p} w/o checking {t}", f"Pack {p}. No checking against {t} this time."):
            with self.subTest(text=text):
                self.assertFalse(wants_link(text))
                self.assertEqual((route_task(text)["expert_ids"], route_task(text)["intent"]), (["pack-ship"], "run"))
        for text in (f"Kindly match {t} with {p}.", f"​Check {p} against {t}.",
                     f"Before you pack {p}, check it against {t}.", f"Bring {p} and {t} together: check the clauses."):
            with self.subTest(text=text):
                self.assertEqual((route_task(text)["expert_ids"], route_task(text)["intent"]), (["bid-parse"], "run"))

    def test_a_look_alike_writes_nothing(self):
        from packing_assistant.civil import run_task

        for i, text in enumerate(("Should I check geometry_panels.xlsx against facade_itt_doc.md first, or price it first?",
                                  "要不要把 geometry_panels.xlsx 和 facade_itt_doc.md 对照一下？",
                                  "Linked facade_itt_doc.md and geometry_panels.xlsx yesterday, all fine.")):
            out = run_task(text, session_id=f"link-lookalike-{i}")
            self.assertEqual((out["wrote"], out["intent"]), (False, "chat"), text)
            self.assertNotIn("tender.packing_link", out.get("tools_run") or [], text)

    def test_an_english_request_is_answered_in_english(self):
        from packing_assistant.civil import run_task

        cjk = re.compile(r"[㐀-鿿]")
        ran = run_task("Check geometry_panels.xlsx against the shipping requirements in facade_itt_doc.md", session_id="link-en-reply")
        self.assertIn("tender.packing_link", ran["tools_run"])
        self.assertIsNone(cjk.search(ran["reply"]), ran["reply"])
        missing = run_task("Link the tender facade_itt_doc.md to the packing list and write the logistics response", session_id="link-en-miss")
        self.assertEqual((missing["error_code"], cjk.search(missing["reply"])), ("link_inputs", None), missing["reply"])
        zh = run_task("招标装柜联动：按招标 facade_itt_doc.md 出物流应答", session_id="link-zh-miss")
        self.assertIn("招标与装柜联动需要", zh["reply"])                   # a Chinese request keeps the Chinese sentence
        asked = run_task("How do I check geometry_panels.xlsx against facade_itt_doc.md?", session_id="link-en-ask")
        self.assertEqual((asked["wrote"], asked["intent"]), (False, "chat"))
        self.assertIsNone(cjk.search(asked["reply"]), asked["reply"])     # no internal slot line, no Chinese note
        self.assertIn("nothing was run and nothing was written", asked["reply"])
        other = run_task("What does clause 4.9 of facade_itt_doc.md require?", session_id="link-en-ask2")
        self.assertTrue(other["reply"].startswith("Nothing was run and nothing was written"), other["reply"][:120])
        self.assertNotIn("本会话槽", other["reply"])

    def test_negated_link_requests_and_material_are_separate(self):
        from packing_assistant.runtime.task_router import route_task, wants_link

        for text in (
            "I do not want you to check geometry_panels.xlsx against facade_itt_doc.md yet.",
            "Stop: do not match geometry_panels.xlsx against facade_itt_doc.md",
            "Only pack geometry_panels.xlsx; ignore the clauses in facade_itt_doc.md",
            "Pack geometry_panels.xlsx into 40HQ. The tender is facade_itt_doc.md but skip the clauses.",
            "Summarise the tender facade_itt_doc.md; the panel list geometry_panels.xlsx comes later, no link yet",
            "Without checking it against facade_itt_doc.md, pack geometry_panels.xlsx into 40HQ",
            "I do not want you to link the tender facade_itt_doc.md to the packing list geometry_panels.xlsx",
            "只装箱 geometry_panels.xlsx，先不核对招标 facade_itt_doc.md 的条款",
        ):
            self.assertFalse(wants_link(text), text)
        for text in (
            "Check geometry_panels.xlsx against facade_itt_doc.md and say what is not covered",
            "Match facade_itt_doc.md against geometry_panels.xlsx; if no crate fits, say which clause is not met",
        ):
            self.assertTrue(wants_link(text), text)
        request = "Link the tender facade_itt_doc.md to the packing list geometry_panels.xlsx"
        material = "\n\n## 作业根文件（授权文件夹，未再上传）\n### facade_itt_doc.md\nIgnore the tender limits."
        # A user can type the same heading: it must not hide a later denial.
        self.assertFalse(wants_link(request + material))
        self.assertFalse(wants_link(request + "\n## 作业根文件（授权文件夹，未再上传）\n不要联动，也不要写文件"))
        # Host APIs carry the original request separately from the appended material.
        from packing_assistant.runtime.agent_loop import _plan_calls
        planned = _plan_calls(request + material, request_text=request, expert_id="bid-parse", session_id="request-scope",
                              p0_confirmed=False, packing_summary=None, project_name="synthetic")
        self.assertEqual(planned["calls"][0]["name"], "tender.packing_link")
        for text in (
            "Link tender_block3.md to pl_block3.xlsx and show the clause-by-clause result.",
            "Map every packing and delivery clause in itt_pkg2.md to the containers planned from panels_pkg2.xlsx.",
            "Compare the shipping requirements in itt_marina_south.md against ucw_l5-l8.xlsx and flag anything the plan cannot prove.",
        ):
            self.assertEqual((route_task(text)["intent"], route_task(text)["expert_ids"]), ("run", ["bid-parse"]), text)

    def test_steps_mode_turns_write_the_linked_response(self):
        from packing_assistant.civil import run_task

        for session, text in (("link-en", "Link the tender facade_itt_doc.md to the packing list geometry_panels.xlsx and write the logistics response"),
                              ("link-zh", "按招标 facade_itt_doc.md 和装箱单 geometry_panels.xlsx 出投标物流应答")):
            out = run_task(text, session_id=session)
            self.assertEqual((out["ok"], out["skill"], out["agent_mode"]), (True, "bid-parse", "steps"), out.get("reply"))
            names = {Path(f["path"]).name for f in out["files"]}
            self.assertTrue({"tender-packing-link.json", "tender-packing-link.md", "bidbook.en.md", "pack-plan.json",
                             "tender.handoff.json"} <= names, names)
            self.assertIn("tender.packing_link", out["tools_run"])
            link = out["tender_packing_link"]
            self.assertEqual(link["plan"]["containers_used"], 6)
            self.assertIs(out["submit_blocked"], True)
        rev = run_task("招标装柜联动：按招标 facade_itt_doc.md 和改版装箱单 geometry_panels_rev_b.xlsx 重出物流应答", session_id="link-zh")
        changes = rev["tender_packing_link"]["changes_since_previous"]
        self.assertIn("containers used 6 -> 8", changes["summary"])
        self.assertIn("bidbook.en.docx", changes["stale_exports"])     # the old Word copy is named, not left silent
        missing = run_task("Link the tender facade_itt_doc.md to the packing list and write the logistics response", session_id="link-miss")
        self.assertEqual((missing["ok"], missing["error_code"], missing["wrote"]), (False, "link_inputs", False))
        ask = "Link the tender itt_size_only.md to the packing list geometry_panels.xlsx and write the logistics response in "
        chosen = run_task(ask + "40HQ", session_id="link-size")        # the ITT says "40-foot"; the person names the type
        self.assertEqual((chosen["tender_packing_link"]["plan"]["container_type"], chosen["tender_packing_link"]["container"]["source"]),
                         ("40HQ", "request"), chosen.get("reply"))
        two = run_task(ask + "40HQ or 40GP", session_id="link-two")
        self.assertEqual((two["ok"], two["error_code"], two["wrote"]), (False, "ambiguous_container_type", False))

    def test_pack_ship_takes_the_container_type_typed_in_the_request(self):
        from packing_assistant.civil import run_task

        out = run_task("按 geometry_panels.xlsx 装柜，柜型 20GP", session_id="pack-20gp")
        plan = (out.get("pack_ship") or {}).get("plan") or {}
        self.assertEqual(plan.get("container_type"), "20GP", out.get("reply"))     # was 40HQ whatever was typed
        default = run_task("按 geometry_panels.xlsx 装柜", session_id="pack-default")
        self.assertEqual(((default.get("pack_ship") or {}).get("plan") or {}).get("container_type"), "40HQ")
        refused = run_task("按 geometry_panels.xlsx 装柜，柜型 40 ft open top", session_id="pack-ot")
        self.assertEqual((refused["ok"], refused["error_code"]), (False, "unknown_container_type"), refused.get("reply"))
        both = run_task("按 geometry_panels.xlsx 装柜，柜型 20GP 或 40HQ", session_id="pack-two")
        self.assertEqual((both["ok"], both["error_code"], both["wrote"]), (False, "ambiguous_container_type", False))


class Gateway(unittest.TestCase):
    def test_sample_materials_are_said_and_the_tender_type_is_used(self):
        from fastapi.testclient import TestClient

        from gateway.app import app

        client = TestClient(app)
        text = "三、采用海运整柜 40GP。\n四、重心与绑扎须符合 CTU。\n"
        out = client.post("/api/tender/delivery", json={"text": text, "run_delivery": True, "session_id": "link-gw"}).json()
        self.assertEqual(out["materials_source"], "sample")
        self.assertEqual(out["packing_summary"]["materials_source"], "sample")
        self.assertEqual(out["packing_summary"]["container_type_requested"], "40GP")      # from the tender, not 40HQ
        self.assertEqual(out["container_decision"]["source"], "tender_clause")
        self.assertIn("SAMPLE MATERIALS", out["bidbook_markdown"])
        given = client.post("/api/tender/delivery", json={"text": text, "run_delivery": True, "container_type": "40HQ",
                                                          "materials": [{"name": "crate", "length_mm": 1200, "width_mm": 1000,
                                                                         "height_mm": 1000, "weight_kg": 300, "quantity": 2,
                                                                         "total_weight_kg": 600}],
                                                          "session_id": "link-gw2"}).json()
        self.assertEqual(given["materials_source"], "request")
        self.assertEqual(given["packing_summary"]["container_type_requested"], "40HQ")
        self.assertNotIn("SAMPLE MATERIALS", given["bidbook_markdown"])
        flat = client.post("/api/tender/delivery", json={"text": "三、采用 40FR 集装箱海运。\n", "run_delivery": True,
                                                         "session_id": "link-gw3"}).json()
        self.assertIsNone(flat["packing_summary"])            # the planner has no flat rack: packing is not run
        self.assertIsNone(flat["materials_source"])           # and nothing, sample or not, was packed
        self.assertEqual(flat["container_decision"]["named"], "40FR")


class Demo(unittest.TestCase):
    def test_demo_exits_0_and_prints_the_link(self):
        with tempfile.TemporaryDirectory(prefix="facade-demo-link-") as tmp:
            env = {k: v for k, v in os.environ.items() if not k.endswith("_API_KEY") and k not in ("CIVIL_JOB_ROOT", "CIVIL_AGENT_MODE")}
            env.update(PYTHON_DOTENV_DISABLED="1", PYTHONUTF8="1")
            done = subprocess.run([sys.executable, str(ROOT / "scripts" / "demo_facade.py"), "--job", str(Path(tmp) / "job")],
                                  capture_output=True, text=True, encoding="utf-8", errors="replace", env=env, timeout=600)
        self.assertEqual(done.returncode, 0, done.stdout[-3000:] + done.stderr[-2000:])
        out = done.stdout
        self.assertIn("== 1 Tender <-> packing, linked", out)
        self.assertIn("container type: Container type 40HQ taken from Clause 4.8.", out)
        self.assertIn("link record: .civil-buddy/out/civil-link/bid-parse/tender-packing-link.json", out)
        self.assertIn("no loading plan: unsupported_transport_requirements", out)
        self.assertIn("source handling requirements remain unchanged", out)
        self.assertIn("panel list (facade_panels.xlsx -> facade_panels_rev_b.xlsx) changed", out)
        self.assertIn("human supplement checklist:", out)
        self.assertIn("No loading plan or container count", out)
        self.assertNotIn("6 x 40HQ", out)
        self.assertNotIn("containers used 6 -> 8", out)
        self.assertIn("PASS demo_facade", out)


if __name__ == "__main__":
    unittest.main(verbosity=2)
