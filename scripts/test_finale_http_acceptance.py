#!/usr/bin/env python3
"""Offline content acceptance through the actual workbench HTTP/export path."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
import unittest
from uuid import uuid4
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import test_workbench_flow as flow
from packing_assistant.office_job import tables_from_md


class FinaleHttpAcceptance(unittest.TestCase):
    setUp = flow.WorkbenchFlowTests.setUp
    post = flow.WorkbenchFlowTests.post

    def test_entry_flags_match_missing_routes_and_keep_public_logistics_closed(self):
        from fastapi.testclient import TestClient
        from unittest.mock import patch
        from packing_assistant import access_guard
        with patch.object(access_guard, "configured_token", return_value="synthetic-entry-test-token"):
            with TestClient(flow.workbench.app, base_url="https://example.invalid", client=("198.51.100.9", 12345)) as client:
                headers = {"Authorization": "Bearer synthetic-entry-test-token"}
                caps = client.get("/api/health", headers=headers).json()["capabilities"]
                self.assertFalse(caps["agent_workbench"])
                self.assertFalse(caps["packing_workbench"])
                self.assertFalse(caps["logistics"])
                self.assertEqual(client.get("/api/agent/capabilities", headers=headers).status_code, 404)
                self.assertEqual(client.get("/packing", headers=headers).status_code, 404)
                self.assertEqual(client.get("/logistics", headers=headers).status_code, 403)

    def markdown(self, message, expert=None, **fields):
        if expert:
            fields["expert_ids"] = [expert]
        done, _ = self.post(message, **fields)
        self.assertTrue(done["ok"] and done["wrote"], done)
        restored = self.client.get(f"/api/sessions/{self.sid}").json()
        self.assertEqual([f["path"] for f in done["deliverables"]],
                         [f["path"] for f in restored["deliverables"]])
        texts = []
        for item in done["deliverables"]:
            response = self.client.get("/api/file", params={"path": item["path"]})
            self.assertEqual(response.status_code, 200)
            if Path(item["name"]).suffix == ".md":
                texts.append(response.content.decode("utf-8"))
            elif Path(item["name"]).suffix == ".docx":
                with ZipFile(BytesIO(response.content)) as archive:
                    self.assertIn("word/document.xml", archive.namelist())
        self.no_model.assert_not_called()
        self.assertTrue(texts)
        return "\n".join(texts), done

    def spreadsheet_rows(self, done, title):
        import openpyxl
        artifact = next(item for item in done["deliverables"] if Path(item["name"]).suffix == ".xlsx")
        response = self.client.get("/api/file", params={"path": artifact["path"]})
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(BytesIO(response.content))
        try:
            sheet = next(book[name] for name in book.sheetnames if name.startswith(title))
            return [["" if value is None else value for value in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()

    def word_table_rows(self, done, headers, *, title=None):
        from xml.etree import ElementTree as ET
        artifact = next(item for item in done["deliverables"] if Path(item["name"]).suffix == ".docx")
        response = self.client.get("/api/file", params={"path": artifact["path"]})
        self.assertEqual(response.status_code, 200)
        with ZipFile(BytesIO(response.content)) as archive:
            document = ET.fromstring(archive.read("word/document.xml"))
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        heading = None
        for table in document.find("w:body", namespaces):
            if table.tag == "{"+namespaces["w"]+"}p":
                style = table.find("w:pPr/w:pStyle", namespaces)
                if style is not None and style.get("{"+namespaces["w"]+"}val", "").startswith("Heading"):
                    heading = "".join(text.text or "" for text in table.findall(".//w:t", namespaces))
                continue
            if table.tag != "{"+namespaces["w"]+"}tbl":
                continue
            rows = [["".join(text.text or "" for text in cell.findall(".//w:t", namespaces))
                     for cell in row.findall("w:tc", namespaces)] for row in table.findall("w:tr", namespaces)]
            if rows and rows[0] == headers and (title is None or heading == title):
                return rows
        self.fail(f"Downloaded Word document has no table with headers {headers}")

    def test_colloquial_daily_places_facts_in_report_fields(self):
        md, _ = self.markdown("请生成项目日报。项目：合成幕墙A；日期：2026-10-09；"
                              "今日完成：幕墙安装20块；问题：材料未到；明日计划：安装15块。缺失信息保持待填。", "pm-daily")
        self.assertIn("| 形象进度 | 幕墙安装20块 |", md)
        self.assertIn("| 机械材料 | 材料未到 |", md)
        self.assertIn("| 明日计划 | 安装15块 |", md)

    def test_english_daily_is_not_an_empty_chinese_template(self):
        md, _ = self.markdown("Generate a daily report. Project: Synthetic Facade A. "
                              "Date: 2026-10-09. Work completed: installed 20 panels. "
                              "Workers: 8. Tomorrow: install 15 panels. Leave missing information unspecified.", "pm-daily", locale="en")
        self.assertIn("| Project | Synthetic Facade A |", md)
        self.assertIn("| Work completed | installed 20 panels |", md)
        self.assertIn("| Workers | 8 |", md)

    def test_narrative_daily_exports_both_locations_without_copying_headcounts(self):
        md, _ = self.markdown("请记录今天的日报：6#楼九层柱子浇了，泥工班到了19个人；"
                              "西侧3#楼那边地下室还在挖，挖机两台早上坏了一台，"
                              "人数没报上来，晚上下了阵雨；明天准备支七层顶板模板", "pm-daily")
        self.assertIn("| 形象进度 | 6#楼九层柱子浇了 | 6#楼九层 |", md)
        self.assertIn("| 形象进度 | 西侧3#楼那边地下室还在挖 | 西侧3#楼 |", md)
        self.assertIn("| 出勤 | 泥工班到了19个人 | 未说明 |", md)
        self.assertFalse(any(line.startswith("|") and "19" in line and "西侧3#楼" in line
                             for line in md.splitlines()))
        self.assertIn("| 天气 | 阵雨 | 未说明 |", md)

    def test_short_daily_command_downloads_literal_weather_and_attendance(self):
        md, done = self.markdown("写一份项目日报 临边 晴 木工12人", "pm-daily")
        report = dict(tables_from_md(md))["日报字段"]
        values = dict(report[1:])
        self.assertEqual("用户口述：晴。非气象站记录。", values["天气"])
        self.assertEqual("木工12人", values["出勤"])
        self.assertEqual(report, self.spreadsheet_rows(done, "日报字段"))
        self.assertEqual(report, self.word_table_rows(done, report[0]))

    def test_it_ops_downloads_keep_two_permissions_and_vpn_incident_separate(self):
        from packing_assistant.post_drafts.it import _MISSING
        md, done = self.markdown("生成运维手册提纲。财务新来的出纳要 OA 的报销单查询，用到 2026-12-31 就收回；监控回放平台那头，项目经理王建军要能导出关键片段，谁审批还没定。另外截标那天 VPN 掉过一次线，按 P3 记。", "it-ops")
        tables = dict(tables_from_md(md))
        permissions = tables["权限依据与申请审批（未执行）"]
        headers = permissions[0]
        cashier = next(r for r in permissions[1:] if r[headers.index("角色")] == "出纳")
        manager = next(r for r in permissions[1:] if r[headers.index("角色")] == "项目经理")
        self.assertEqual("2026-12-31", cashier[headers.index("回收条件")])
        self.assertEqual(_MISSING, manager[headers.index("回收条件")])
        self.assertEqual(_MISSING, manager[headers.index("业务审批人")])
        self.assertEqual("", manager[-1])
        faults = tables["故障分级与升级路径"]
        self.assertEqual("VPN", faults[1][0])
        self.assertIn("P3", faults[1][1])
        self.assertEqual(permissions, self.spreadsheet_rows(done, "权限依据与申请审批"))
        self.assertEqual(permissions, self.word_table_rows(done, headers, title="权限依据与申请审批（未执行）"))
        self.assertTrue(done["submit_blocked"])

    def test_it_app_downloads_place_actor_and_gate_in_separate_fields(self):
        md, done = self.markdown("生成需求说明书草稿\n系统：物资台账\n用户角色：试验员、商务只读\n功能需求：材料员登记进场过磅记录，无磅单影像不能入库\n验收标准：能按月导出明细\n需求来源：业务单A", "it-app")
        tables = dict(tables_from_md(md))
        requirements = tables["角色与功能需求（逐项待确认）"]
        row = next(r for r in requirements[1:] if r[1] == "材料员")
        self.assertEqual("登记进场过磅记录", row[3])
        self.assertEqual("无磅单影像不能入库", row[5])
        self.assertEqual(requirements, self.spreadsheet_rows(done, "角色与功能需求"))
        self.assertEqual(requirements, self.word_table_rows(done, requirements[0]))
        source = tables["需求来源与系统模块归属"]
        row = next(r for r in source[1:] if r[source[0].index("角色")] == "材料员")
        self.assertEqual("业务单A", row[source[0].index("需求来源")])
        self.assertEqual("", row[-1])

    def test_it_app_downloads_keep_module_delivery_and_peer_ownership(self):
        from packing_assistant.post_drafts.it import _MISSING
        md, done = self.markdown("生成需求说明书草稿。两个需求一块提：一个是工人考勤，门口闸机的数据想直接进项目部报表，劳务员每天不用再导表；另一个是机械台班登记，司机手机上填台班，机管员审。考勤这个王总要国庆前看到，台班那个没说时间。", "it-app")
        tables = dict(tables_from_md(md))
        time_records = [dict(rows[1:]) for title, rows in tables.items() if title.startswith("非功能与数据来源")]
        self.assertTrue(any(r["系统 / 模块 / 数据归属"] == "工人考勤" and r["期望交付时间（用户期望，未承诺）"] == "国庆前" for r in time_records))
        self.assertTrue(all(r["期望交付时间（用户期望，未承诺）"] == _MISSING for r in time_records if r["系统 / 模块 / 数据归属"] == "机械台班登记"))
        interfaces = tables["接口与主数据归属"]
        self.assertEqual("工人考勤", interfaces[1][0])
        self.assertEqual("门口闸机", interfaces[1][2])
        self.assertEqual(interfaces, self.spreadsheet_rows(done, "接口与主数据归属"))
        self.assertEqual(interfaces, self.word_table_rows(done, interfaces[0]))

    def test_it_data_downloads_do_not_turn_targets_or_plans_into_recovery_results(self):
        from packing_assistant.post_drafts.it import _MISSING
        md, done = self.markdown("生成备份策略草稿\n系统：OA；数据对象：审批附件；RPO：15分钟；RPO来源：业务单A；RTO：240分钟；RTO来源：业务单B；演练日期：计划2026-11-30；演练记录：预计生成记录A；实际恢复时长：预计35分钟；用户演练结果：预计恢复成功", "it-data")
        tables = dict(tables_from_md(md))
        targets = tables["系统数据恢复目标与策略来源"]
        self.assertEqual(["15分钟", "业务单A", "240分钟", "业务单B"], targets[1][2:6])
        drills = tables["恢复演练记录（用户提供，未核验）"]
        self.assertEqual(_MISSING, drills[1][2])
        self.assertEqual(_MISSING, drills[1][-1])
        evidence = tables["恢复步骤与演练证据归属"]
        self.assertEqual(_MISSING, evidence[1][5])
        self.assertEqual(_MISSING, evidence[1][8])
        self.assertEqual("", evidence[1][-1])
        self.assertIn("预计恢复成功", md)
        self.assertEqual(evidence, self.spreadsheet_rows(done, "恢复步骤与演练证据归属"))
        self.assertEqual(evidence, self.word_table_rows(done, evidence[0]))

    def test_selected_it_csvs_keep_blank_fields_credentials_and_originals_safe(self):
        from packing_assistant.post_drafts.it import _MISSING
        cases = (
            ("it-app", "系统,角色,功能,验收标准,密码\nOA,材料员,判断 a|b,本人资料,synthetic-it-secret-alpha\n平台B,,查询,,synthetic-it-secret-bravo\n", "角色与功能需求（逐项待确认）", 1),
            ("it-ops", "系统,角色,权限,审批人,密码\nOA,出纳,报销单查询,王工,synthetic-it-secret-alpha\n平台B,项目经理,导出,,synthetic-it-secret-bravo\n", "权限依据与申请审批（未执行）", 5),
            ("it-data", "系统,数据对象,RPO,RTO,密码\nOA,工资,半小时,2小时,synthetic-it-secret-alpha\n平台B,视频,,,synthetic-it-secret-bravo\n", "业务恢复目标（拟议，非实测承诺）", 1),
        )
        for post, source, title, blank_index in cases:
            with self.subTest(post=post):
                self.sid = "flow-it-" + uuid4().hex[:16]
                data = source.encode("utf-8")
                uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成IT资料.csv", data, "text/csv")})
                self.assertEqual(uploaded.status_code, 200, uploaded.text)
                identifier = uploaded.json()["files"][0]["id"]
                original, _ = flow.uploads.upload_file(self.sid, identifier)
                md, done = self.markdown("请按选中附件生成内部草稿，未给字段保持待填。", post, attachments=[identifier])
                rows = dict(tables_from_md(md))[title]
                self.assertEqual(["OA", "平台B"], [r[0] for r in rows[1:]])
                self.assertEqual(_MISSING, rows[2][blank_index])
                self.assertEqual(rows, self.spreadsheet_rows(done, title))
                self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
                if post == "it-app":
                    self.assertEqual("判断 a|b", rows[1][3])
                for secret in ("synthetic-it-secret-alpha", "synthetic-it-secret-bravo"):
                    self.assertNotIn(secret, md)
                self.assertEqual(data, original.read_bytes())

    def test_claim_downloads_denied_send_draft_only_and_refused_receipt_are_not_actual_service(self):
        from packing_assistant.post_drafts.claim import MISSING
        md, done = self.markdown("生成索赔意向草稿\n事件：停电\n已发函：N-A（草稿）\n我们不曾发过函N-B\n监理张三拒绝签收", "claim")
        tables = dict(tables_from_md(md))
        title, docs = next((k, v) for k, v in tables.items() if k.startswith("4 意向"))
        self.assertEqual(["N-A", "N-B"], [r[docs[0].index("用户函 / 通知 / 停工令编号")] for r in docs[1:]])
        self.assertTrue(all(r[docs[0].index("用户发出记录（非本轮发送）")] == MISSING for r in docs[1:]))
        self.assertTrue(all(r[docs[0].index("用户签收人（未核）")] == MISSING for r in docs[1:]))
        self.assertIn("草稿", docs[1][docs[0].index("用户通知 / 拟发记录")])
        self.assertIn("不曾", docs[2][docs[0].index("用户通知 / 拟发记录")])
        self.assertIn("拒绝签收", docs[2][docs[0].index("用户签收记录（未核）")])
        self.assertEqual(docs, self.spreadsheet_rows(done, title))
        self.assertEqual(docs, self.word_table_rows(done, docs[0], title=title))

    def test_claim_short_command_downloads_internal_draft_without_high_risk_acknowledgement(self):
        md, done = self.markdown("出索赔意向草稿。事件：停电\n发生日期：2026-09-01\n人数：12人", "claim")
        tables = dict(tables_from_md(md))
        resources = tables["用户同期资源与影响记录"]
        self.assertEqual("12人", resources[1][resources[0].index("用户窝工人数原文（未核）")])
        self.assertEqual(resources, self.spreadsheet_rows(done, "用户同期资源与影响记录"))
        self.assertEqual(resources, self.word_table_rows(done, resources[0], title="用户同期资源与影响记录"))
        self.assertIn("未签认", md)
        self.assertTrue(done["submit_blocked"])

    def test_claim_downloads_keep_two_events_resources_letters_receivers_and_evidence_owners_apart(self):
        from packing_assistant.post_drafts.claim import MISSING
        text = ("帮我出一份索赔意向草稿。一个是2#楼转换层结构图业主拖到2026-08-11才发，木工班组31人干等着，"
                "这事发过函ZJ-HAN-022，监理何志强签的字。另一个是上月底片区停电，地下室防水停了，"
                "防水队14个人窝着，还有台HBT60地泵趴窝，哪天恢复供电我还得查，函还没发，证据先让资料员孙倩去补。")
        md, done = self.markdown(text, "claim")
        tables = dict(tables_from_md(md))
        events = next(v for k, v in tables.items() if k.startswith("2 事件"))
        self.assertEqual(3, len(events))
        self.assertEqual(["2026-08-11", MISSING], [r[events[0].index("用户实际发图日期")] for r in events[1:]])
        docs = next(v for k, v in tables.items() if k.startswith("4 意向"))
        self.assertEqual(["ZJ-HAN-022", MISSING], [r[docs[0].index("用户函 / 通知 / 停工令编号")] for r in docs[1:]])
        self.assertEqual(["何志强", MISSING], [r[docs[0].index("用户签收人（未核）")] for r in docs[1:]])
        self.assertIn("片区停电", docs[2][docs[0].index("用户索赔事件（未核）")])
        resources = tables["用户同期资源与影响记录"]
        self.assertEqual(["31人", "14个人", MISSING], [r[resources[0].index("用户窝工人数原文（未核）")] for r in resources[1:]])
        self.assertEqual("HBT60", resources[-1][resources[0].index("用户机械型号")])
        self.assertEqual(MISSING, resources[-1][resources[0].index("用户机械数量原文（未核）")])
        actions = tables["用户补证计划"]
        self.assertIn("片区停电", actions[1][actions[0].index("用户索赔事件（未核）")])
        self.assertEqual("孙倩", actions[1][actions[0].index("用户补证责任人")])
        for prefix in ("2 事件", "4 意向", "用户同期资源", "用户补证", "用户待核"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertTrue(done["submit_blocked"])

    def test_claim_downloads_preserve_contract_conditions_and_requested_money_without_deadline_or_approval(self):
        text = ("生成索赔意向草稿\n项目：合成索赔A\n辖区：SG\n合同编号：C-A\n"
                "合同原文：如果事件持续，用户约定每12日提供记录；指定代表接收。\n合同原文来源：contract-A.pdf第7页\n"
                "工期条款原文：用户工期程序另行通知。\n事件：停电\n发生日期：2026-09-01\n日期来源：D-A\n"
                "用户主张金额：SGD12500\n金额来源：M-A\n用户主张工期：7天\n工期来源：T-A\n证据：用户日志A\n"
                "证据位置：用户目录第3项\n该证据补证责任：王五\n拟收件人：张三\n"
                "拟发函：N-A\n补证负责人：李四")
        md, done = self.markdown(text, "claim")
        tables = dict(tables_from_md(md))
        contract = next(v for k, v in tables.items() if k.startswith("3 合同"))
        owned_contracts = [r for r in contract[1:] if r[contract[0].index("合同编号")] == "C-A"]
        self.assertEqual(1, len(owned_contracts))
        self.assertEqual("如果事件持续，用户约定每12日提供记录；指定代表接收。", owned_contracts[0][contract[0].index("用户合同要求 / 索赔条款原文（未核）")])
        self.assertEqual("UNSPECIFIED · 未计算", owned_contracts[0][contract[0].index("本轮通知时限 / 最后日期")])
        requested = tables["用户意向请求（非获批）"]
        self.assertEqual("SGD12500", requested[1][requested[0].index("用户主张金额原文（未核）")])
        self.assertEqual("金额 TBD / 工期天数 TBD · 未核定", requested[1][-1])
        evidence = next(v for k, v in tables.items() if k.startswith("5 证据"))
        self.assertEqual("D-A", evidence[1][evidence[0].index("用户事件日期来源")])
        self.assertEqual("王五", evidence[1][evidence[0].index("该证据补证责任（明确归属才登记）")])
        for prefix in ("3 合同", "4 意向", "5 证据", "用户意向", "用户补证"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertNotIn("2026-09-13", md)

    def test_selected_claim_csv_keeps_blank_identity_independent_and_original_bytes_immutable(self):
        from packing_assistant.post_drafts.claim import MISSING
        source = ("事件,合同编号,部位,停工开始,日期来源,班组,人数,函号,签收人,签收凭证,补证负责人\n"
                  "停电,C-A,A区,2026-09-01,D-A,甲班,12人,N-A,张三,R-A,李四\n"
                  ",,,2026-09-03,,,8人,N-B,,,\n"
                  "晚供,C-B,B区,2026-09-05,,乙班,5人,N-C,王五,,赵六\n")
        data = source.encode("utf-8")
        response = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成索赔资料.csv", data, "text/csv")})
        self.assertEqual(response.status_code, 200, response.text)
        identifier = response.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        md, done = self.markdown("根据所选附件生成索赔意向草稿，缺项待核。", "claim", attachments=[identifier])
        tables = dict(tables_from_md(md))
        events = next(v for k, v in tables.items() if k.startswith("2 事件"))
        self.assertEqual(["停电", MISSING, "晚供"], [r[events[0].index("用户索赔事件（未核）")] for r in events[1:]])
        self.assertEqual(["D-A", MISSING, MISSING], [r[events[0].index("用户日期来源")] for r in events[1:]])
        docs = next(v for k, v in tables.items() if k.startswith("4 意向"))
        self.assertEqual(["张三", MISSING, "王五"], [r[docs[0].index("用户签收人（未核）")] for r in docs[1:]])
        self.assertEqual(["R-A", MISSING, MISSING], [r[docs[0].index("用户送达回证 / 签收凭证")] for r in docs[1:]])
        for prefix in ("2 事件", "4 意向", "用户同期资源", "用户补证"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertEqual(data, original.read_bytes())

    def test_claim_downloads_multi_letters_and_resource_types_leave_shared_proof_unassigned(self):
        from packing_assistant.post_drafts.claim import MISSING
        text = ("生成索赔意向草稿\n事件：停电\n已发函：N-A、N-B，业主代表张三签收\n签收凭证：R-A\n"
                "| 事件 | 班组 | 人数 | 机械型号 | 机械数量 | 同期记录来源 |\n"
                "| --- | --- | --- | --- | --- | --- |\n| 晚供 | 甲班 | 12人 | PC220 | 3台 | S-A |")
        md, done = self.markdown(text, "claim")
        tables = dict(tables_from_md(md))
        docs = next(v for k, v in tables.items() if k.startswith("4 意向"))
        id_col = docs[0].index("用户函 / 通知 / 停工令编号")
        receiver_col = docs[0].index("用户签收人（未核）")
        receipt_col = docs[0].index("用户送达回证 / 签收凭证")
        named = [r for r in docs[1:] if r[id_col] in {"N-A", "N-B"}]
        self.assertEqual(2, len(named))
        self.assertTrue(all(r[receiver_col] == MISSING and r[receipt_col] == MISSING for r in named))
        resources = tables["用户同期资源与影响记录"]
        self.assertEqual(["S-A", MISSING, MISSING], [r[resources[0].index("用户同期记录来源")] for r in resources[1:]])
        self.assertEqual([MISSING, "12人", MISSING], [r[resources[0].index("用户窝工人数原文（未核）")] for r in resources[1:]])
        self.assertEqual([MISSING, MISSING, "3台"], [r[resources[0].index("用户机械数量原文（未核）")] for r in resources[1:]])
        for prefix in ("4 意向", "用户同期资源"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))

    def test_subcontract_downloads_multi_deduction_and_paid_columns_leave_generic_vouchers_unassigned(self):
        from packing_assistant.post_drafts.subcontract import MISSING
        text = ("生成分包结算表头\n| 分包单位 | 水电费 | 质量罚款 | 通知单 | 本期已付 | 累计已付 | 支付凭证 |\n"
                "| --- | --- | --- | --- | --- | --- | --- |\n| 甲班 | 3150元 | 2400元 | N-A | 2000元 | 10000元 | P-A.pdf |")
        md, done = self.markdown(text, "subcontract")
        tables = dict(tables_from_md(md))
        deductions = next(v for k, v in tables.items() if k.startswith("6 扣款"))
        self.assertEqual([MISSING, "水电费", "质量罚款"], [r[deductions[0].index("用户拟扣款项")] for r in deductions[1:]])
        self.assertEqual([MISSING, "3150元", "2400元"], [r[deductions[0].index("用户拟扣款金额原文（未核）")] for r in deductions[1:]])
        self.assertEqual(["N-A", MISSING, MISSING], [r[deductions[0].index("用户通知单标识（未核）")] for r in deductions[1:]])
        payments = tables["用户支付记录（非支付批准）"]
        self.assertEqual(["P-A.pdf", MISSING, MISSING], [r[payments[0].index("支付凭证 / 来源")] for r in payments[1:]])
        self.assertEqual([MISSING, "2000元", MISSING], [r[payments[0].index("用户本期已付原文（未核）")] for r in payments[1:]])
        self.assertEqual([MISSING, MISSING, "10000元"], [r[payments[0].index("用户前期累计已付原文（未核）")] for r in payments[1:]])
        for prefix in ("6 扣款", "用户支付记录"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))

    def test_subcontract_downloads_keep_measurement_deductions_paid_and_wage_roles_separate(self):
        from packing_assistant.post_drafts.subcontract import MISSING
        text = ("生成分包结算表头。项目：合成分包A\n辖区：SG\n币种：SGD\n结算期间：2026-07-26至2026-08-25\n"
                "分包单位：甲班\n合同编号：C-A\n分项：钢筋绑扎\n申请量：35吨\n工长核量：31吨\n确认量：30吨\n"
                "确认量来源：Q-A.pdf\n总包确认量：29吨\n单价：SGD920/t\n单价来源：R-A.pdf\n"
                "质量罚款：SGD2400\n通知单：N-A\n扣款依据：D-A.pdf\n累计已付：SGD118000\n支付凭证：P-A.pdf\n"
                "用户应付人工费：SGD30000\n工资表：W-A.xlsx\n生活费预支：SGD5000\n"
                "分包单位：乙班\n合同编号：C-B\n分项：防水\n确认量：2600平\n单价：SGD38.6/m2\n"
                "质量罚款：SGD1200\n通知单：N-B\n累计已付：SGD8000")
        # This low-risk internal table needs no high-risk acknowledgement,
        # but still cannot approve, calculate or execute any settlement/payment.
        md, done = self.markdown(text, "subcontract")
        tables = dict(tables_from_md(md))
        quantities = next(v for k, v in tables.items() if k.startswith("3 本期"))
        self.assertEqual(["35吨", MISSING], [r[quantities[0].index("用户申请工程量（非确认量）")] for r in quantities[1:]])
        self.assertEqual(["30吨", "2600平"], [r[quantities[0].index("用户验工 / 确认工程量原文（未核）")] for r in quantities[1:]])
        self.assertEqual(["Q-A.pdf", MISSING], [r[quantities[0].index("用户确认量来源")] for r in quantities[1:]])
        values = next(v for k, v in tables.items() if k.startswith("4 合同"))
        self.assertTrue(all(r[-2:] == ["TBD · 未计算", "TBD · 未核定"] for r in values[1:]))
        deductions = next(v for k, v in tables.items() if k.startswith("6 扣款"))
        self.assertEqual(["甲班", "乙班"], [r[deductions[0].index("合同分包 / 分包单位 / 班组")] for r in deductions[1:]])
        self.assertEqual(["N-A", "N-B"], [r[deductions[0].index("用户通知单标识（未核）")] for r in deductions[1:]])
        self.assertEqual(["D-A.pdf", MISSING], [r[deductions[0].index("用户扣款依据 / 凭证")] for r in deductions[1:]])
        payments = tables["用户支付记录（非支付批准）"]
        self.assertEqual(["P-A.pdf", MISSING, MISSING], [r[payments[0].index("支付凭证 / 来源")] for r in payments[1:]])
        self.assertEqual("SGD5000", payments[2][payments[0].index("用户生活费预支原文（非工资结清）")])
        self.assertEqual(MISSING, payments[2][payments[0].index("用户工资代发已付原文（未核）")])
        for prefix in ("3 本期", "4 合同", "6 扣款", "用户支付记录", "8 农民工", "9 会签"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertTrue(done["submit_blocked"])

    def test_subcontract_narrative_downloads_keep_two_crews_damage_and_living_advance_apart(self):
        from packing_assistant.post_drafts.subcontract import MISSING
        text = ("出分包结算表头。老李的木工班和鑫源架子队这个月一块结，木工干的是3#楼标准层模板，"
                "工长核完是4120平，单价46.5。架子队搭的7#楼外架，量还没核出来，只知道合同价19.5一平，"
                "他们上回摔坏扣件要扣1800，另外木工班预支过生活费30000，合同号我回头再找。")
        md, done = self.markdown(text, "subcontract")
        tables = dict(tables_from_md(md))
        quantities = next(v for k, v in tables.items() if k.startswith("3 本期"))
        self.assertEqual(["老李的木工班", "鑫源架子队"], [r[quantities[0].index("合同分包 / 分包单位 / 班组")] for r in quantities[1:]])
        self.assertEqual(["4120平", MISSING], [r[quantities[0].index("用户工长核量原文（非本轮核定）")] for r in quantities[1:]])
        self.assertTrue(all(r[quantities[0].index("用户验工 / 确认工程量原文（未核）")] == MISSING for r in quantities[1:]))
        deductions = next(v for k, v in tables.items() if k.startswith("6 扣款"))
        self.assertEqual("鑫源架子队", deductions[1][deductions[0].index("合同分包 / 分包单位 / 班组")])
        self.assertEqual("1800", deductions[1][deductions[0].index("用户拟扣款金额原文（未核）")])
        payments = tables["用户支付记录（非支付批准）"]
        self.assertEqual("老李的木工班", payments[1][payments[0].index("合同分包 / 分包单位 / 班组")])
        self.assertEqual("30000", payments[1][payments[0].index("用户生活费预支原文（非工资结清）")])
        self.assertEqual(MISSING, payments[1][payments[0].index("用户工资代发已付原文（未核）")])
        for prefix in ("3 本期", "6 扣款", "用户支付记录"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))

    def test_selected_subcontract_csv_keeps_empty_party_identity_and_original_bytes_immutable(self):
        from packing_assistant.post_drafts.subcontract import MISSING
        source = ("分包单位,合同编号,结算期间,确认量,确认量来源,单价,质量罚款,通知单,累计已付,支付凭证\n"
                  "甲班,A,2026-08,30吨,Q-A.pdf,920元/吨,2400元,N-A,10000元,P-A.pdf\n"
                  ",,,20吨,,,1200元,,8000元,\n"
                  "乙班,B,2026-09,120平,,38元/平,,,,\n")
        data = source.encode("utf-8")
        response = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成分包资料.csv", data, "text/csv")})
        self.assertEqual(response.status_code, 200, response.text)
        identifier = response.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        md, done = self.markdown("根据所选附件生成分包结算表头，缺项待核。", "subcontract", attachments=[identifier])
        tables = dict(tables_from_md(md))
        quantities = next(v for k, v in tables.items() if k.startswith("3 本期"))
        self.assertEqual(["甲班", MISSING, "乙班"], [r[quantities[0].index("合同分包 / 分包单位 / 班组")] for r in quantities[1:]])
        self.assertEqual(["Q-A.pdf", MISSING, MISSING], [r[quantities[0].index("用户确认量来源")] for r in quantities[1:]])
        deductions = next(v for k, v in tables.items() if k.startswith("6 扣款"))
        self.assertEqual(["甲班", MISSING, "乙班"], [r[deductions[0].index("合同分包 / 分包单位 / 班组")] for r in deductions[1:]])
        self.assertEqual(["N-A", MISSING, MISSING], [r[deductions[0].index("用户通知单标识（未核）")] for r in deductions[1:]])
        for prefix in ("3 本期", "6 扣款", "用户支付记录"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertEqual(data, original.read_bytes())

    def test_subcontract_downloads_repeat_quantity_without_reusing_source_rate_or_old_period(self):
        from packing_assistant.post_drafts.subcontract import MISSING
        text = ("生成分包结算表头。项目：甲\n辖区：SG\n结算期间：2026-07-26至2026-08-25\n"
                "分包单位：甲班\n分项：钢筋\n确认量：35吨\n确认量来源：Q-A.pdf\n单价：920元/吨\n"
                "确认量：20吨\n结算期间：这个月\n确认量：10吨\n累计已付：预计明天支付8000元")
        md, done = self.markdown(text, "subcontract")
        tables = dict(tables_from_md(md))
        quantities = next(v for k, v in tables.items() if k.startswith("3 本期"))
        self.assertEqual(["35吨", "20吨", "10吨"], [r[quantities[0].index("用户验工 / 确认工程量原文（未核）")] for r in quantities[1:]])
        self.assertEqual(["Q-A.pdf", MISSING, MISSING], [r[quantities[0].index("用户确认量来源")] for r in quantities[1:]])
        self.assertEqual("这个月", quantities[3][quantities[0].index("结算期间原文")])
        values = next(v for k, v in tables.items() if k.startswith("4 合同"))
        self.assertEqual(["920元/吨", MISSING, MISSING], [r[values[0].index("用户合同单价原文（未核）")] for r in values[1:]])
        payments = tables["用户支付记录（非支付批准）"]
        self.assertEqual(MISSING, payments[1][payments[0].index("用户前期累计已付原文（未核）")])
        for prefix in ("3 本期", "4 合同", "用户支付记录"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))

    def test_safety_brief_downloads_keep_steps_hazards_contacts_and_unsigned_fields_separate(self):
        from packing_assistant.post_drafts.safety_brief import MISSING
        text = ("生成安全交底草稿。项目：合成交底A\n辖区：SG\n交底日期：明天\n交底人：张三\n"
                "作业部位：3#楼电梯井口\n步骤：拆盖板\n班组：木工甲班\n高度：3.2米\n参数来源：P-A.pdf\n"
                "危险点：用户点名坠落\n措施：用户措施A | B & C\n措施依据：用户方案A.pdf\n喊停条件：用户条件A\n"
                "监护人：李四\n监护人电话：13800000000\n步骤：吊料\n危险点：用户点名落物\n"
                "用户实施记录：预计明天设置围挡\n作业部位：屋面女儿墙外侧\n工序：打胶\n个人防护：安全带\n"
                "危险点：用户点名高坠\n安全员：王五")
        md, done = self.markdown(text, "safety-brief", confirm_text="我明白，将由持证人员签认")
        tables = dict(tables_from_md(md))
        measures = next(v for k, v in tables.items() if k.startswith("6 防护"))
        self.assertEqual(["拆盖板", "吊料", "打胶"], [r[measures[0].index("用户作业步骤")] for r in measures[1:]])
        self.assertEqual(["用户措施A | B & C", MISSING, MISSING], [r[measures[0].index("用户措施原文（非已落实）")] for r in measures[1:]])
        self.assertEqual(["用户方案A.pdf", MISSING, MISSING], [r[measures[0].index("用户措施依据")] for r in measures[1:]])
        self.assertTrue(all(r[measures[0].index("用户实施记录（未核）")] == MISSING for r in measures[1:]))
        self.assertTrue(all("未执行" in r[-1] for r in measures[1:]))
        self.assertIn("预计明天设置围挡", measures[2])
        for prefix in ("6 防护", "用户材料及参数", "用户责任与联系接口", "11 签字"):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        contacts = tables["用户责任与联系接口"]
        self.assertEqual(["3#楼电梯井口", "屋面女儿墙外侧"], [r[contacts[0].index("作业部位 / 工作面")] for r in contacts[1:]])
        self.assertEqual(["13800000000", MISSING], [r[contacts[0].index("用户联系电话原文（未核）")] for r in contacts[1:]])
        self.assertEqual(["（空）"]*4+["未执行 · 未签发"], tables["11 签字栏"][1])
        self.assertNotIn("已交底完毕", md)
        self.assertTrue(done["submit_blocked"])

    def test_safety_brief_requires_exact_current_turn_typed_confirmation(self):
        for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
            with self.subTest(fields=fields):
                self.sid = "safetygate-"+uuid4().hex[:16]
                done, _ = self.post("生成安全交底草稿，作业部位3#楼，工序吊料。", expert_ids=["safety-brief"], **fields)
                self.assertFalse(done["wrote"])
                self.assertEqual([], done["deliverables"])
                self.assertIn("我明白，将由持证人员签认", done["text"])
        self.sid = "safetysame-"+uuid4().hex[:16]
        self.markdown("生成安全交底草稿，部位3#楼，工序吊料。", "safety-brief", confirm_text="我明白，将由持证人员签认")
        later, _ = self.post("重新生成安全交底草稿，部位5#楼，工序通风。", expert_ids=["safety-brief"])
        self.assertFalse(later["wrote"])
        self.assertEqual([], later["deliverables"])
        self.no_model.assert_not_called()

    def test_selected_safety_csv_keeps_unassigned_rows_contacts_and_source_bytes_immutable(self):
        from packing_assistant.post_drafts.safety_brief import MISSING
        source = ("作业部位,工序,危险点,措施,措施依据,监护人,安全员,电话,高度,参数来源\n"
                  "3#楼,吊料,用户点名落物,用户措施A,A.pdf,张三,李四,13800000000,3.2米,P-A.pdf\n"
                  ",,用户点名碰撞,用户措施X,,,,,,\n"
                  "5#楼,通风,用户点名缺氧,用户措施B,,王五,,,4.1米,\n")
        data = source.encode("utf-8")
        response = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成交底资料.csv", data, "text/csv")})
        self.assertEqual(response.status_code, 200, response.text)
        identifier = response.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        md, done = self.markdown("根据所选附件生成安全交底草稿，未给字段保留待核。", "safety-brief", attachments=[identifier], confirm_text="我明白，将由持证人员签认")
        tables = dict(tables_from_md(md))
        measures = next(v for k, v in tables.items() if k.startswith("6 防护"))
        self.assertEqual(["3#楼", MISSING, "5#楼"], [r[measures[0].index("作业部位 / 工作面")] for r in measures[1:]])
        self.assertEqual(["A.pdf", MISSING, MISSING], [r[measures[0].index("用户措施依据")] for r in measures[1:]])
        contacts = tables["用户责任与联系接口"]
        numbered = [r for r in contacts[1:] if r[contacts[0].index("用户联系电话原文（未核）")] != MISSING]
        self.assertEqual(1, len(numbered))
        self.assertEqual(MISSING, numbered[0][contacts[0].index("用户责任接口 / 联系人原文")])
        for title, rows in tables.items():
            if title.startswith(("6 防护", "用户责任与联系接口")):
                self.assertEqual(rows, self.spreadsheet_rows(done, title))
                self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertEqual(data, original.read_bytes())

    def test_dispatch_downloads_keep_faces_resources_orders_and_stops_separate(self):
        from packing_assistant.post_drafts.dispatch import MISSING
        text = ("生成调度日报。项目：合成调度A\n日期：2026-09-18\n班次：白班\n天气：小雨\n"
                "作业面：3#楼十二层顶板\n今日完成：浇了约246方还没收面\n实际开始：7:30\n实绩来源：F-A.csv\n"
                "工种：泥工\n人数：34人\n出勤来源：L-A.csv\n机械：汽车泵\n台数：一台\n机械来源：P-A.csv\n"
                "指令：生产经理周建国要求钢筋班长刘海涛9月19日中午前把十三层墙柱插筋补齐\n指令凭证：I-A.pdf\n"
                "作业面：地下室B2层4区底板\n今日完成：防水卷材铺贴约730平米\n"
                "停工：22:40停电，原因配电房跳闸，底板钢筋绑扎中断\n"
                "指令：项目副经理陈立新口头要求机电分包赵鹏次日8点前恢复供电")
        md, done = self.markdown(text, "dispatch")
        tables = dict(tables_from_md(md))
        for prefix, word_title in (("4 当日", "4 当日实际"), ("6 指令", "6 指令栏"), ("8 停复工", "8 停复工与异常")):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=word_title))
        actual = next(v for k, v in tables.items() if k.startswith("4 当日"))
        self.assertEqual(["约246方", "约730平米"], [r[actual[0].index("用户当日实际数量原文（未核）")] for r in actual[1:]])
        self.assertEqual(["F-A.csv", MISSING], [r[actual[0].index("实绩来源")] for r in actual[1:]])
        self.assertTrue(all(r[-1] == "未执行 / 待核" for r in actual[1:]))
        actions = next(v for k, v in tables.items() if k.startswith("6 指令"))
        self.assertEqual(["生产经理周建国", "项目副经理陈立新"], [r[actions[0].index("用户下达人 / 指令来源")] for r in actions[1:]])
        self.assertEqual(["钢筋班长刘海涛", "机电分包赵鹏"], [r[actions[0].index("用户接收人 / 接收方")] for r in actions[1:]])
        self.assertEqual(["9月19日中午前", "次日8点前"], [r[actions[0].index("用户要求期限")] for r in actions[1:]])
        self.assertEqual(["I-A.pdf", MISSING], [r[actions[0].index("指令凭证来源")] for r in actions[1:]])
        self.assertTrue(all(r[-2:] == ["草稿 · 待责任人核准 · 未下达", "（空）"] for r in actions[1:]))
        stops = next(v for k, v in tables.items() if k.startswith("8 停复工"))
        self.assertEqual("配电房跳闸", stops[1][stops[0].index("用户停工原因")])
        self.assertEqual(MISSING, stops[1][stops[0].index("用户实际恢复记录（未核）")])
        for title, rows in tables.items():
            if title.startswith("用户人员记录") or title.startswith("用户机械记录"):
                self.assertEqual(rows, self.spreadsheet_rows(done, title))
                self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        self.assertTrue(done["submit_blocked"])

    def test_dispatch_downloads_keep_next_day_plans_out_of_actual_counts_and_restoration(self):
        from packing_assistant.post_drafts.dispatch import MISSING
        text = ("生成调度日报。日期：2026-09-18\n天气：小雨\n天气预报：预计明天晴\n作业面：3#楼\n"
                "停工：断电\n用户恢复记录：预计明天恢复供电\n明日安排：计划浇筑35方\n"
                "工种：木工\n人数：34人\n机械：塔吊TC6013\n台数：一台\n"
                "今日完成：浇了约20方\n工种：泥工\n人数：12人")
        md, done = self.markdown(text, "dispatch")
        tables = dict(tables_from_md(md))
        title, rows = next((k, v) for k, v in tables.items() if k.startswith("4 当日"))
        self.assertEqual(2, len(rows))
        self.assertEqual("约20方", rows[1][rows[0].index("用户当日实际数量原文（未核）")])
        self.assertEqual("2026-09-18", rows[1][rows[0].index("调度日期")])
        self.assertEqual(rows, self.spreadsheet_rows(done, title))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="4 当日实际"))
        labor_title, labor = next((k, v) for k, v in tables.items() if "用户人数原文" in v[0])
        self.assertEqual(["34人", "12人"], [r[labor[0].index("用户人数原文")] for r in labor[1:]])
        self.assertEqual(["用户计划（非实到 / 非已执行）", "用户现场记录（未核）"], [r[-1] for r in labor[1:]])
        self.assertEqual(labor, self.spreadsheet_rows(done, labor_title))
        self.assertEqual(labor, self.word_table_rows(done, labor[0], title="用户人员记录"))
        stops = next(v for k, v in tables.items() if k.startswith("8 停复工"))
        self.assertEqual(MISSING, stops[1][stops[0].index("用户实际恢复记录（未核）")])
        self.assertEqual("预计明天恢复供电", stops[1][stops[0].index("用户拟复工 / 恢复条件")])
        self.assertTrue(all(w not in md for w in ("已恢复供电", "隐患已消除", "同意复工")))

    def test_selected_dispatch_csv_keeps_unassigned_faces_and_source_bytes_immutable(self):
        from packing_assistant.post_drafts.dispatch import MISSING
        source = ("作业面,日期,今日完成,工种,人数,出勤来源,机械,台数,机械来源,指令,责任人\n"
                  "3#楼,2026-09-18,浇了35方,木工,34人,L-A.csv,塔吊TC6013,一台,P-A.csv,补齐插筋,甲班\n"
                  ",,,,17人,L-X.csv,,,,,\n"
                  "5#楼,2026-09-19,铺贴约730平米,,,,汽车泵,,,,\n")
        data = source.encode("utf-8")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成调度资料.csv", data, "text/csv")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        identifier = uploaded.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        self.assertEqual(data, original.read_bytes())
        md, done = self.markdown("根据所选附件生成调度日报，缺项保持待核。", "dispatch", attachments=[identifier])
        tables = dict(tables_from_md(md))
        title, rows = next((k, v) for k, v in tables.items() if k.startswith("4 当日"))
        self.assertEqual(["3#楼", MISSING, "5#楼"], [r[rows[0].index("作业面 / 工作面")] for r in rows[1:]])
        self.assertEqual(["35方", MISSING, "约730平米"], [r[rows[0].index("用户当日实际数量原文（未核）")] for r in rows[1:]])
        self.assertEqual(rows, self.spreadsheet_rows(done, title))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="4 当日实际"))
        labor_title, labor = next((k, v) for k, v in tables.items() if "用户人数原文" in v[0])
        self.assertEqual(["木工", MISSING, MISSING], [r[labor[0].index("用户工种 / 班组")] for r in labor[1:]])
        self.assertEqual(["34人", "17人", MISSING], [r[labor[0].index("用户人数原文")] for r in labor[1:]])
        self.assertEqual(["L-A.csv", "L-X.csv", MISSING], [r[labor[0].index("出勤来源")] for r in labor[1:]])
        self.assertEqual(labor, self.spreadsheet_rows(done, labor_title))
        self.assertEqual(labor, self.word_table_rows(done, labor[0], title="用户人员记录"))
        self.assertEqual(data, original.read_bytes())

    def test_interim_downloads_keep_declared_reviewed_owner_and_payable_separate(self):
        from packing_assistant.post_drafts.interim import MISSING
        text = ("生成验工计价草稿。合同：HT-A\n计量期间：2026-08-26到2026-09-25\n"
                "清单编码：000012\n名称：钢筋 HRB400\n部位：3#楼\n合同量：100吨\n上期末开累：20吨\n"
                "本期完成：35吨\n本期申报：30吨\n申报来源：D-A.csv\n监理审：28吨\n监理来源：S-A.csv\n"
                "业主计量单JL-A也签了\n合同单价：SGD 800/t\n单价来源：P-A.csv\n申报金额：SGD 24000\n"
                "名称：加气块砌体\n部位：5#楼\n本期申报：930方\n监理还没看\n业主更没确认\n"
                "预计业主确认量900方")
        # Ordinary low-risk internal measurement drafting does not need the
        # high-risk acknowledgement; it still cannot issue/payment-certify.
        md, done = self.markdown(text, "interim")
        tables = dict(tables_from_md(md))
        for prefix, word_title in (("3 本期", "3 本期范围"), ("4 计量", "4 计量依据"), ("5 计量", "5 计量草表"),
                                   ("6 变更", "6 变更、物价、索赔"), ("7 过程", "7 过程结算与进度款")):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=word_title))
        title, rows = next((k, v) for k, v in tables.items() if k.startswith("5 计量"))
        self.assertEqual(["000012", MISSING], [r[rows[0].index("清单编码")] for r in rows[1:]])
        self.assertEqual(["钢筋", "加气块砌体"], [r[rows[0].index("清单名称")] for r in rows[1:]])
        self.assertEqual(["30吨", "930方"], [r[rows[0].index("用户本期申报量（未核）")] for r in rows[1:]])
        self.assertEqual(["28吨", MISSING], [r[rows[0].index("用户监理审核量（未核）")] for r in rows[1:]])
        self.assertTrue(all(r[rows[0].index("用户业主确认量（未核）")] == MISSING for r in rows[1:]))
        self.assertEqual(["JL-A", MISSING], [r[rows[0].index("业主确认来源（未核）")] for r in rows[1:]])
        self.assertEqual(["SGD 800/t", MISSING], [r[rows[0].index("用户合同单价原文（未核）")] for r in rows[1:]])
        self.assertEqual(["SGD 24000", MISSING], [r[rows[0].index("用户申报金额（未核）")] for r in rows[1:]])
        self.assertTrue(all(r[-5:] == ["TBD"]*5 for r in rows[1:]))
        event_title, events = next((k, v) for k, v in tables.items() if "用户单笔原文" in v[0])
        role, quantity, source = (events[0].index(h) for h in ("数量 / 价款记录项", "用户单笔原文", "本笔来源"))
        owner_events = [r for r in events[1:] if r[role] == "业主确认量"]
        self.assertEqual(1, len(owner_events))
        self.assertEqual((MISSING, "JL-A"), (owner_events[0][quantity], owner_events[0][source]))
        self.assertEqual(events, self.spreadsheet_rows(done, event_title))
        self.assertEqual(events, self.word_table_rows(done, events[0], title="5 计量草表"))
        self.assertTrue(done["submit_blocked"])

    def test_interim_downloads_keep_period_boq_units_and_confirmation_sources_independent(self):
        from packing_assistant.post_drafts.interim import MISSING
        text = ("生成验工计价草稿。\n| 确认来源 | 名称 | 清单编码 | 期间 | 合同量 | 上期末开累 | 本期申报 | 监理审 | 业主核 | 单价 |\n"
                "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |\n"
                "| O-A.pdf | 钢筋 | 000012 | 9月 | 100吨 | 20吨 | 35吨 | 32吨 | 30吨 | SGD 800/t |\n"
                "| | 钢筋 | 000012 | 10月 | | | 40吨 | | | |\n"
                "| O-C.pdf | 砌体 | 000013 | 9月 | | | 930方 | 900m3 | 880 | |")
        md, done = self.markdown(text, "interim")
        title, rows = next((k, v) for k, v in tables_from_md(md) if k.startswith("5 计量"))
        self.assertEqual(["9月", "10月", "9月"], [r[rows[0].index("计量期间 / 期次")] for r in rows[1:]])
        self.assertEqual(["100吨", MISSING, MISSING], [r[rows[0].index("用户合同量（未核）")] for r in rows[1:]])
        self.assertEqual(["20吨", MISSING, MISSING], [r[rows[0].index("用户上期末开累（未核）")] for r in rows[1:]])
        self.assertEqual(["30吨", MISSING, "880"], [r[rows[0].index("用户业主确认量（未核）")] for r in rows[1:]])
        self.assertEqual(["O-A.pdf", MISSING, "O-C.pdf"], [r[rows[0].index("业主确认来源（未核）")] for r in rows[1:]])
        self.assertEqual(["吨", "吨", MISSING], [r[rows[0].index("计量单位（用户原文）")] for r in rows[1:]])
        self.assertIn("待核", rows[3][rows[0].index("单位待核")])
        self.assertEqual(["SGD 800/t", MISSING, MISSING], [r[rows[0].index("用户合同单价原文（未核）")] for r in rows[1:]])
        self.assertEqual(rows, self.spreadsheet_rows(done, title))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="5 计量草表"))

    def test_selected_interim_csv_preserves_unassigned_reviewed_quantities_and_original(self):
        from packing_assistant.post_drafts.interim import MISSING
        source = ("名称,清单编码,期间,本期申报,申报来源,监理审,监理来源,业主核,确认来源\n"
                  "钢筋,000012,9月,35吨,D-A.csv,32吨,S-A.csv,30吨,O-A.pdf\n"
                  ",,,,,28吨,S-X.csv,,\n"
                  "砌体,000013,10月,930方,D-B.csv,,,,\n")
        data = source.encode("utf-8")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成计量资料.csv", data, "text/csv")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        identifier = uploaded.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        self.assertEqual(data, original.read_bytes())
        md, done = self.markdown("根据所选附件生成验工计价草稿，缺项保持待核。", "interim", attachments=[identifier])
        title, rows = next((k, v) for k, v in tables_from_md(md) if k.startswith("5 计量"))
        self.assertEqual(["000012", MISSING, "000013"], [r[rows[0].index("清单编码")] for r in rows[1:]])
        self.assertEqual(["钢筋", MISSING, "砌体"], [r[rows[0].index("清单名称")] for r in rows[1:]])
        self.assertEqual(["32吨", "28吨", MISSING], [r[rows[0].index("用户监理审核量（未核）")] for r in rows[1:]])
        self.assertEqual(["S-A.csv", "S-X.csv", MISSING], [r[rows[0].index("监理确认来源")] for r in rows[1:]])
        self.assertEqual(["30吨", MISSING, MISSING], [r[rows[0].index("用户业主确认量（未核）")] for r in rows[1:]])
        self.assertTrue(all(r[-5:] == ["TBD"]*5 for r in rows[1:]))
        self.assertEqual(rows, self.spreadsheet_rows(done, title))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="5 计量草表"))
        self.assertEqual(data, original.read_bytes())

    def test_site_material_downloads_keep_quantity_sources_stock_and_consumption_separate(self):
        from packing_assistant.post_drafts.site_materials import MISSING
        text = ("生成材料核算表头。项目：合成材料A\n期间：9月\n部位：3#楼\n材料：钢筋 HRB400\n"
                "已完工程量：40吨\n工程量来源：W.csv\n消耗指标：1.02\n指标来源：指标-A.pdf\n"
                "领料：10吨\n领料来源：I1.csv\n领料：20吨\n领料来源：I2.csv\n退料：3吨\n退料来源：R.csv\n"
                "盘点：实存5吨\n盘点来源：C.csv\n实耗：26吨\n实耗来源：A.csv\n"
                "单价：SGD 800/t\n单价来源：P.csv\n节超符号约定：超耗为正\n"
                "材料：商品混凝土 C35P8\n部位：地下室B2层底板\n已完工程量：1284.6方\n"
                "盘点：现场收方与小票差47.8方\n预计实耗1400方")
        # This post is low-risk; an ordinary explicit draft request needs no
        # high-risk typed acknowledgement, but remains an internal draft.
        md, done = self.markdown(text, "material-site")
        tables = dict(tables_from_md(md))
        for prefix, word_title in (("3 应耗", "3 应耗量口径 / 计划用量"), ("5 节超", "5 节超口径 / 盘点"),
                                   ("6 核算", "6 核算表头 / 节超待核")):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=word_title))
        title, rows = next((k, v) for k, v in tables.items() if k.startswith("6 核算"))
        self.assertEqual(["钢筋", "商品混凝土"], [r[rows[0].index("材料名称")] for r in rows[1:]])
        self.assertEqual(["10吨；20吨", MISSING], [r[rows[0].index("用户本期领料原文")] for r in rows[1:]])
        self.assertEqual(["5吨", MISSING], [r[rows[0].index("用户盘点实存原文（未核）")] for r in rows[1:]])
        self.assertEqual(["26吨", MISSING], [r[rows[0].index("用户实耗原文（未核）")] for r in rows[1:]])
        self.assertEqual("现场收方与小票差47.8方", rows[2][rows[0].index("用户盘点记录（未核）")])
        self.assertEqual("SGD 800/t", rows[1][rows[0].index("用户单价原文（未核）")])
        self.assertEqual("P.csv", rows[1][rows[0].index("单价来源")])
        self.assertTrue(all(r[-5:] == ["TBD"]*5 for r in rows[1:]))
        pending = next(v for k, v in tables.items() if k.startswith("7 原因"))
        self.assertEqual("预计实耗1400方", pending[2][pending[0].index("用户待核记录")])
        event_title, events = next((k, v) for k, v in tables.items() if "用户单笔数量 / 指标原文" in v[0])
        quantity, source, field = (events[0].index(h) for h in ("用户单笔数量 / 指标原文", "本笔来源", "数量记录项"))
        self.assertEqual([("10吨", "I1.csv"), ("20吨", "I2.csv"), ("3吨", "R.csv")],
                         [(r[quantity], r[source]) for r in events[1:] if r[field] in {"本期领料", "本期退料"}])
        self.assertEqual(events, self.spreadsheet_rows(done, event_title))
        self.assertEqual(events, self.word_table_rows(done, events[0], title="4 实耗量口径 / 领退记录"))
        self.assertTrue(done["submit_blocked"])

    def test_site_material_downloads_keep_models_periods_units_and_sources_independent(self):
        from packing_assistant.post_drafts.site_materials import MISSING
        text = ("生成材料核算表头。\n| 材料 | 部位 | 规格 | 期间 | 工程量 | 工程量来源 | 领料 | 领料来源 | 退料 | 盘点 |\n"
                "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |\n"
                "| 防水卷材 | 地下室外墙 | SBS | 9月 | 1893平 | W-A.pdf | 76卷 | I-A.csv | 14卷 | 实存12卷 |\n"
                "| 防水卷材 | 地下室外墙 | | 10月 | | | 20卷 | | | |\n"
                "| 钢筋 | 3#楼 | HRB400 | 9月 | | | 35吨 | I-B.csv | 120kg | |")
        md, done = self.markdown(text, "material-site")
        title, rows = next((k, v) for k, v in tables_from_md(md) if k.startswith("6 核算"))
        self.assertEqual(["SBS", MISSING, "HRB400"], [r[rows[0].index("规格")] for r in rows[1:]])
        self.assertEqual(["9月", "10月", "9月"], [r[rows[0].index("核算期间")] for r in rows[1:]])
        self.assertEqual(["平", MISSING, MISSING], [r[rows[0].index("工程量单位（用户原文）")] for r in rows[1:]])
        self.assertEqual(["卷", "卷", MISSING], [r[rows[0].index("材料单位（用户原文）")] for r in rows[1:]])
        self.assertIn("待核", rows[3][rows[0].index("材料单位待核")])
        self.assertEqual(["W-A.pdf", MISSING, MISSING], [r[rows[0].index("工程量来源")] for r in rows[1:]])
        self.assertEqual(["I-A.csv", MISSING, "I-B.csv"], [r[rows[0].index("领料来源")] for r in rows[1:]])
        self.assertEqual(["12卷", MISSING, MISSING], [r[rows[0].index("用户盘点实存原文（未核）")] for r in rows[1:]])
        self.assertEqual(rows, self.spreadsheet_rows(done, title))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="6 核算表头 / 节超待核"))

    def test_selected_site_material_csv_preserves_unassigned_quantities_and_original_bytes(self):
        from packing_assistant.post_drafts.site_materials import MISSING
        source = ("材料,部位,规格,期间,领料,领料来源,退料,退料来源,实存,盘点来源\n"
                  "钢筋,3#楼,HRB400,9月,35吨,I-A.csv,12吨,R-A.csv,5吨,C-A.csv\n"
                  ",,,,,,3吨,R-X.csv,,\n"
                  "卷材,地下室外墙,,10月,76卷,I-B.csv,,,,\n")
        data = source.encode("utf-8")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成材料资料.csv", data, "text/csv")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        identifier = uploaded.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        self.assertEqual(data, original.read_bytes())
        md, done = self.markdown("根据所选附件生成材料核算表头，缺项保持待核。", "material-site", attachments=[identifier])
        title, rows = next((k, v) for k, v in tables_from_md(md) if k.startswith("6 核算"))
        self.assertEqual(["钢筋", MISSING, "卷材"], [r[rows[0].index("材料名称")] for r in rows[1:]])
        self.assertEqual(["HRB400", MISSING, MISSING], [r[rows[0].index("规格")] for r in rows[1:]])
        self.assertEqual(["12吨", "3吨", MISSING], [r[rows[0].index("用户退料原文")] for r in rows[1:]])
        self.assertEqual(["R-A.csv", "R-X.csv", MISSING], [r[rows[0].index("退料来源")] for r in rows[1:]])
        self.assertEqual(["5吨", MISSING, MISSING], [r[rows[0].index("用户盘点实存原文（未核）")] for r in rows[1:]])
        self.assertTrue(all(r[rows[0].index("用户实耗原文（未核）")] == MISSING for r in rows[1:]))
        self.assertEqual(rows, self.spreadsheet_rows(done, title))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="6 核算表头 / 节超待核"))
        self.assertEqual(data, original.read_bytes())

    def test_quality_downloads_keep_lots_reports_correction_and_review_separate(self):
        from packing_assistant.post_drafts.quality import MISSING
        text = ("生成质量检查表。单位工程：合成质量A\n检验批：1#楼四层顶板\n检验批编号：LOT-A\n"
                "主控检查项：裂缝\n自检记录：收缩裂缝\n实测：0.3mm\n材料：混凝土\n规格：C35P8\n复试报告：GJ-A\n报告来源：A.pdf\n"
                "整改责任：甲班\n整改期限：10月20日\n整改完成：用户说已完成\n复查计划：明天复查\n"
                "检验批：2#楼三层砌体\n检验批编号：LOT-B\n一般检查项：拉结筋\n自检记录：漏放\n问题数量：12处\n整改责任：乙班")
        md, done = self.markdown(text, "quality", confirm_text="我明白，将由持证人员签认")
        tables = dict(tables_from_md(md))
        for prefix, word_title in (("2 划分说明", "2 划分说明 / 检验批"), ("4 主控", "4 主控项目检查栏"), ("5 一般", "5 一般项目检查栏"),
                                   ("6 隐蔽", "6 隐蔽专项"), ("7 通病", "7 通病防治核对 / 问题记录"),
                                   ("8 不符合", "8 不符合时的处理路径栏目 / 整改"), ("9 资料闭合", "9 资料闭合 / 复查与归档")):
            title, rows = next((k, v) for k, v in tables.items() if k.startswith(prefix))
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=word_title))
        for title, rows in tables.items():
            if title.startswith("3 进场"):
                self.assertEqual(rows, self.spreadsheet_rows(done, title))
                self.assertEqual(rows, self.word_table_rows(done, rows[0]))
        checks = next(v for k, v in tables.items() if k.startswith("7 通病"))
        self.assertEqual(["LOT-A", "LOT-B"], [r[checks[0].index("检验批编号")] for r in checks[1:]])
        self.assertEqual(["0.3mm", MISSING], [r[checks[0].index("用户实测值（未核）")] for r in checks[1:]])
        self.assertEqual([MISSING, "12处"], [r[checks[0].index("用户问题数量")] for r in checks[1:]])
        self.assertTrue(all(r[-1] == "未检 / 待核" for r in checks[1:]))
        reports = next(v for v in tables.values() if "用户报告编号" in v[0])
        self.assertEqual(2, len(reports))
        self.assertEqual("LOT-A", reports[1][reports[0].index("检验批编号")])
        self.assertEqual("GJ-A", reports[1][reports[0].index("用户报告编号")])
        self.assertEqual("本轮未核验", reports[1][-1])
        actions = next(v for k, v in tables.items() if k.startswith("8 不符合"))
        self.assertEqual(["10月20日", MISSING], [r[actions[0].index("计划整改期限")] for r in actions[1:]])
        self.assertEqual(["用户说已完成", MISSING], [r[actions[0].index("用户整改完成记录（未核）")] for r in actions[1:]])
        reviews = next(v for k, v in tables.items() if k.startswith("9 资料闭合"))
        self.assertEqual(MISSING, reviews[1][reviews[0].index("用户实际复查记录（未核）")])
        self.assertEqual("待核", reviews[1][-2])
        self.assertTrue(done["submit_blocked"])

    def test_quality_requires_exact_current_turn_typed_confirmation(self):
        for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
            with self.subTest(fields=fields):
                self.sid = "qualitygate-"+uuid4().hex[:16]
                done, _ = self.post("生成质量检查表，检验批LOT-A。", expert_ids=["quality"], **fields)
                self.assertFalse(done["wrote"])
                self.assertEqual([], done["deliverables"])
                self.assertIn("我明白，将由持证人员签认", done["text"])
        self.no_model.assert_not_called()

    def test_selected_quality_csv_preserves_empty_identity_and_original_bytes(self):
        from packing_assistant.post_drafts.quality import MISSING
        source = "检验批,检验批编号,主控检查项,自检记录,实测,复试报告,报告来源,整改责任,复查状态\n甲批,LOT-A,外观,裂缝,0.3mm,GJ-A,A.pdf,甲班,待复查\n,,,,,GJ-X,,,\n乙批,LOT-B,隐蔽,漏放,,,,乙班,未复查\n"
        data = source.encode("utf-8")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成质量资料.csv", data, "text/csv")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        identifier = uploaded.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        self.assertEqual(data, original.read_bytes())
        md, done = self.markdown("根据所选附件生成质量检查表，缺项保持待核。", "quality", attachments=[identifier], confirm_text="我明白，将由持证人员签认")
        tables = dict(tables_from_md(md))
        checks_title, checks = next((k, v) for k, v in tables.items() if k.startswith("7 通病"))
        self.assertEqual(["LOT-A", MISSING, "LOT-B"], [r[checks[0].index("检验批编号")] for r in checks[1:]])
        self.assertEqual(["0.3mm", MISSING, MISSING], [r[checks[0].index("用户实测值（未核）")] for r in checks[1:]])
        self.assertEqual(checks, self.spreadsheet_rows(done, checks_title))
        self.assertEqual(checks, self.word_table_rows(done, checks[0], title="7 通病防治核对 / 问题记录"))
        reports_title, reports = next((k, v) for k, v in tables.items() if "用户报告编号" in v[0])
        self.assertEqual(["GJ-A", "GJ-X", MISSING], [r[reports[0].index("用户报告编号")] for r in reports[1:]])
        self.assertEqual(["A.pdf", MISSING, MISSING], [r[reports[0].index("报告来源")] for r in reports[1:]])
        self.assertEqual(reports, self.spreadsheet_rows(done, reports_title))
        self.assertEqual(reports, self.word_table_rows(done, reports[0]))
        self.assertEqual(data, original.read_bytes())

    def test_funding_downloads_keep_forecast_actual_balance_and_accounts_separate(self):
        from packing_assistant.post_drafts.funding import MISSING
        text = ("生成资金计划。项目：合成资金A\n计划期：2026-10-06至2026-11-02\n币种：SGD\n"
                "期初余额：300元\n账户：余额户A\n余额来源：bankA.pdf\n"
                "收款：款项A\n金额：1,200.00\n计划日：2026-10-20\n账户：收款户A\n状态：已报未批\n"
                "收款：款项B\n实际收款金额：800.00\n实际到账日：2026-10-22\n账户：收款户B\n来源：回单B.pdf\n"
                "付款：材料款\n金额：200.00\n账户：付款户C\n对象：供方C\n"
                "付款：工资\n金额：100.00\n账户：工资专户\n付款依据：工资表D.xlsx\n"
                "融资方：贷款方E\n授信额度：SGD 8,000\n融资限制：需董事签字\n"
                "审批事项：材料款申请\n审批条件：补合同\n审批状态：待审批")
        md, done = self.markdown(text, "finance-fund")
        tables = dict(tables_from_md(md))
        for title in ("2 期间与余额", "3 编制原则与融资约束", "4 合同价款节点", "5 收入栏", "6 支出栏", "8 风险提示与未分派资料", "9 自检与审批跟踪"):
            rows = tables[title]
            self.assertEqual(rows, self.spreadsheet_rows(done, title))
            self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
        income = tables["5 收入栏"]
        self.assertEqual(["1,200.00", MISSING], [r[income[0].index("预计 / 计划金额原文")] for r in income[1:]])
        self.assertEqual([MISSING, "800.00"], [r[income[0].index("用户实际收付金额（未核）")] for r in income[1:]])
        self.assertEqual(["收款户A", "收款户B"], [r[income[0].index("账户（用户原文）")] for r in income[1:]])
        self.assertTrue(all(r[-1] == "未执行" for r in income[1:]))
        payments = tables["6 支出栏"]
        self.assertEqual(["付款户C", "工资专户"], [r[payments[0].index("账户（用户原文）")] for r in payments[1:]])
        self.assertTrue(all(r[payments[0].index("用户实际收付金额（未核）")] == MISSING for r in payments[1:]))
        self.assertTrue(all(r[-1] == "未执行 /（空）" for r in payments[1:]))
        self.assertIn("无法试算", md)
        self.assertTrue(done["submit_blocked"])

    def test_funding_unknown_balance_and_paid_status_do_not_invent_actual_amount(self):
        from packing_assistant.post_drafts.funding import MISSING
        md, done = self.markdown("生成资金计划。期初余额：300元\n账户：余额户A\n收款：进度款\n金额：200元\n状态：已到账\n实际到账日：2026-10-20\n付款：钢筋款已到账100元", "finance-fund")
        tables = dict(tables_from_md(md))
        balance = tables["2 期间与余额"]
        self.assertEqual(MISSING, balance[1][balance[0].index("用户期初余额（来源原件未核）")])
        income = tables["5 收入栏"]
        self.assertEqual("200元", income[1][income[0].index("预计 / 计划金额原文")])
        self.assertEqual(MISSING, income[1][income[0].index("用户实际收付金额（未核）")])
        self.assertEqual(income, self.spreadsheet_rows(done, "5 收入栏"))
        self.assertEqual(income, self.word_table_rows(done, income[0], title="5 收入栏"))
        payments = tables["6 支出栏"]
        self.assertEqual(MISSING, payments[1][payments[0].index("用户实际收付金额（未核）")])
        orphan = tables["8 风险提示与未分派资料"]
        self.assertEqual("100元", orphan[1][orphan[0].index("用户实际收付金额（未核）")])
        self.assertIn("方向冲突", orphan[1][orphan[0].index("用户收支类别 / 归属待核")])
        self.assertEqual(orphan, self.spreadsheet_rows(done, "8 风险提示与未分派资料"))
        self.assertEqual(orphan, self.word_table_rows(done, orphan[0], title="8 风险提示与未分派资料"))

    def test_selected_funding_csv_preserves_currency_accounts_empty_rows_and_original(self):
        from packing_assistant.post_drafts.funding import MISSING
        source = "收支类别,节点名称,金额,币种,账户,计划日,实际收款金额\n收款,甲款,300,SGD,账户A,2026-10-20,\n付款,乙款,200,CNY,账户B,,\n收款,丙款,,SGD,,,100\n"
        data = source.encode("utf-8")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成资金资料.csv", data, "text/csv")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        identifier = uploaded.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        self.assertEqual(data, original.read_bytes())
        md, done = self.markdown("根据所选附件生成资金计划，币种和缺项保持原样。", "finance-fund", attachments=[identifier])
        tables = dict(tables_from_md(md))
        income = tables["5 收入栏"]
        self.assertEqual(["甲款", "丙款"], [r[0] for r in income[1:]])
        self.assertEqual(["300", MISSING], [r[income[0].index("预计 / 计划金额原文")] for r in income[1:]])
        self.assertEqual([MISSING, "100"], [r[income[0].index("用户实际收付金额（未核）")] for r in income[1:]])
        self.assertEqual(["账户A", MISSING], [r[income[0].index("账户（用户原文）")] for r in income[1:]])
        payments = tables["6 支出栏"]
        self.assertEqual("CNY", payments[1][payments[0].index("币种（用户原文）")])
        self.assertEqual(MISSING, payments[1][payments[0].index("计划日（非收付事实）")])
        for title in ("5 收入栏", "6 支出栏"):
            self.assertEqual(tables[title], self.spreadsheet_rows(done, title))
            self.assertEqual(tables[title], self.word_table_rows(done, tables[title][0], title=title))
        self.assertEqual(data, original.read_bytes())

    def test_equipment_downloads_keep_devices_certificates_people_and_dates_separate(self):
        from packing_assistant.post_drafts.equipment import MISSING
        text = ("生成设备台账。设备名称：塔吊A\n型号：STT293\n机号：D-A\n计划进场：2026-10-20\n"
                "合格证：HG-A\n有效期至：2027-03-01\n司机：张测试\n作业证：OP-A\n维保：每15天一次\n"
                "设备名称：塔吊B\n型号：TC5613\n机号：D-B\n产权：自有\n合格证：未提供")
        md, done = self.markdown(text, "equip", confirm_text="我明白，将由持证人员签认")
        tables = dict(tables_from_md(md))
        for title in ("2 设备清单", "3 进场验收与日期", "4 租赁与台班", "5 维保计划与完成记录", "6 证件与检验台账", "7 故障停用与人员接口"):
            actual = tables[title]
            self.assertEqual(actual, self.spreadsheet_rows(done, title))
            self.assertEqual(actual, self.word_table_rows(done, actual[0], title=title))
        arrivals = tables["3 进场验收与日期"]
        self.assertEqual(["2026-10-20", MISSING], [r[arrivals[0].index("计划进场日（非到场事实）")] for r in arrivals[1:]])
        self.assertTrue(all(r[arrivals[0].index("用户实际进场日（未核）")] == MISSING for r in arrivals[1:]))
        certificates = tables["6 证件与检验台账"]
        self.assertEqual(["D-A", "D-B"], [r[certificates[0].index("设备编号 / 机号")] for r in certificates[1:]])
        self.assertEqual(["HG-A", MISSING], [r[certificates[0].index("用户证件编号")] for r in certificates[1:]])
        self.assertEqual(["2027-03-01", MISSING], [r[certificates[0].index("用户检定期限 / 有效期原文")] for r in certificates[1:]])
        self.assertTrue(all(r[-1] == "本轮未核验" for r in certificates[1:]))
        people_title, people = next((k, v) for k, v in tables.items() if v[0][0] == "设备名称" and "司机 / 操作人员" in v[0])
        self.assertEqual("D-A", people[1][people[0].index("设备编号 / 机号")])
        self.assertEqual("OP-A", people[1][people[0].index("用户作业人员证件编号")])
        self.assertEqual(people, self.spreadsheet_rows(done, people_title))
        self.assertEqual(people, self.word_table_rows(done, people[0]))
        self.assertTrue(done["submit_blocked"])

    def test_equipment_requires_exact_current_turn_typed_confirmation(self):
        for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
            with self.subTest(fields=fields):
                self.sid = "equipgate-"+uuid4().hex[:16]
                done, _ = self.post("生成设备台账，塔吊型号STT293。", expert_ids=["equip"], **fields)
                self.assertFalse(done["wrote"])
                self.assertEqual([], done["deliverables"])
                self.assertIn("我明白，将由持证人员签认", done["text"])
        self.no_model.assert_not_called()

    def test_selected_equipment_csv_preserves_empty_fields_and_original_bytes(self):
        missing = "[A001] / UNSPECIFIED"
        source = "设备名称,型号,机号,合格证,有效期至,证件来源\n塔吊A,STT293,D-A,HG-A,2027-03-01,A.pdf\n塔吊B,TC5613,D-B,,,\n"
        data = source.encode("utf-8")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成设备资料.csv", data, "text/csv")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        identifier = uploaded.json()["files"][0]["id"]
        original, _ = flow.uploads.upload_file(self.sid, identifier)
        self.assertEqual(data, original.read_bytes())
        md, done = self.markdown("根据所选附件生成设备台账，缺项保持待核。", "equip", attachments=[identifier], confirm_text="我明白，将由持证人员签认")
        certificates = dict(tables_from_md(md))["6 证件与检验台账"]
        self.assertEqual(["D-A", "D-B"], [r[certificates[0].index("设备编号 / 机号")] for r in certificates[1:]])
        self.assertEqual(["HG-A", missing], [r[certificates[0].index("用户证件编号")] for r in certificates[1:]])
        self.assertEqual(["2027-03-01", missing], [r[certificates[0].index("用户检定期限 / 有效期原文")] for r in certificates[1:]])
        self.assertEqual(["A.pdf", missing], [r[certificates[0].index("证件来源")] for r in certificates[1:]])
        self.assertEqual(certificates, self.spreadsheet_rows(done, "6 证件与检验台账"))
        self.assertEqual(certificates, self.word_table_rows(done, certificates[0]))
        self.assertEqual(data, original.read_bytes())

    def test_planning_posts_download_scoped_literal_word_and_excel_tables(self):
        from packing_assistant.post_drafts.planning import MISSING
        cases = (
            ("plan-master", "生成总控计划草稿。5#楼基础9月底完，主体封顶甲方卡的是2027-06-30，这栋归三公司干；"
             "6#楼工程量按建施-11算的，工期粗算560天，封顶日子还没给准信，干的队伍也没定。",
             (("3 工作分解", "3 工作分解 / 活动分解"), ("4 一级网络", "4 一级网络与里程碑"))),
            ("plan-lookahead", "生成周计划。活动：3层砌筑\n制约：制约已清\n活动：4层砌筑\n制约：塔吊未到",
             (("3 近细", "3 近细远粗 / 四周工序"), ("4 制约", "4 制约因素 / 前置条件"), ("5 周承诺", "5 周承诺"))),
            ("plan-resource", "生成资源负荷表头。3#楼砌筑下周上瓦工38个，塔吊TC5613吊砖，进场日子定的11月20日；"
             "2#楼二次结构上多少人还没定，只知道要用汽车吊，型号没报，材料是加气块，堆到东侧空地。",
             (("3 劳动力", "3 劳动力负荷表头"), ("4 施工机具", "4 施工机具负荷表头"), ("5 主要材料", "5 主要材料与周转料表头"))),
        )
        for post, text, sections in cases:
            with self.subTest(post=post):
                self.sid = "planning-"+uuid4().hex[:16]
                md, done = self.markdown(text, post)
                tables = dict(tables_from_md(md))
                for prefix, word_title in sections:
                    title = next(t for t in tables if t.startswith(prefix))
                    actual = tables[title]
                    self.assertEqual(actual, self.spreadsheet_rows(done, title))
                    self.assertEqual(actual, self.word_table_rows(done, actual[0], title=word_title))
                if post == "plan-master":
                    actual = next(v for k, v in tables.items() if k.startswith("3 工作分解"))
                    self.assertEqual(["5#楼", "6#楼"], [r[0] for r in actual[1:]])
                    self.assertEqual(["三公司", MISSING], [r[actual[0].index("责任单位")] for r in actual[1:]])
                    self.assertEqual([MISSING, "粗算560天"], [r[actual[0].index("用户工期 / 持续时间口径")] for r in actual[1:]])
                    nodes = tables["4 一级网络与里程碑"]
                    self.assertEqual(["9月底", "2027-06-30", MISSING], [r[nodes[0].index("节点日期原文")] for r in nodes[1:]])
                elif post == "plan-lookahead":
                    promise = tables["5 周承诺"]
                    self.assertEqual(["3层砌筑"], [r[0] for r in promise[1:]])
                    self.assertNotIn("4层砌筑", repr(promise))
                else:
                    crane = tables["4 施工机具负荷表头"]
                    self.assertEqual(["TC5613", MISSING], [r[crane[0].index("规格 / 型号")] for r in crane[1:]])
                    self.assertEqual(["11月20日", MISSING], [r[crane[0].index("计划进场日（非进场事实）")] for r in crane[1:]])
                    self.assertEqual(["3#楼", "2#楼"], [r[crane[0].index("单位工程 / 楼栋")] for r in crane[1:]])
                self.assertTrue(done["submit_blocked"])

    def test_selected_planning_csvs_preserve_row_ownership_and_original_bytes(self):
        missing = "[A001] / UNSPECIFIED"
        cases = (
            ("plan-master", "总控计划草稿", "楼栋,活动,用户工期,责任单位\n3#楼,封顶,18天,甲队\n2#楼,封顶,,\n", "3 工作分解", "3 工作分解 / 活动分解", "用户工期 / 持续时间口径", ["18天", missing]),
            ("plan-lookahead", "周计划", "周次,活动,责任班组,制约\n第1周,作业A,甲队,制约已清\n第2周,作业B,,图纸未发\n", "3 近细", "3 近细远粗 / 四周工序", "责任方 / 班组 / 工长", ["甲队", missing]),
            ("plan-resource", "资源负荷表头", "活动,工种,需用人数,可用量\n3#楼砌筑,木工,20人,8人\n2#楼二次结构,钢筋工,,\n", "3 劳动力", "3 劳动力负荷表头", "实际可用量原文", ["8人", missing]),
        )
        for post, name, source, prefix, word_title, column, expected in cases:
            with self.subTest(post=post):
                self.sid = "planning-file-"+uuid4().hex[:16]
                data = source.encode("utf-8")
                uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                            files={"files": ("合成计划资料.csv", data, "text/csv")})
                self.assertEqual(uploaded.status_code, 200, uploaded.text)
                identifier = uploaded.json()["files"][0]["id"]
                original, _ = flow.uploads.upload_file(self.sid, identifier)
                self.assertEqual(data, original.read_bytes())
                md, done = self.markdown("根据所选附件生成"+name+"，缺失字段保持待核。", post, attachments=[identifier])
                title, actual = next((k, v) for k, v in tables_from_md(md) if k.startswith(prefix))
                self.assertEqual(expected, [r[actual[0].index(column)] for r in actual[1:]])
                self.assertEqual(actual, self.spreadsheet_rows(done, title))
                self.assertEqual(actual, self.word_table_rows(done, actual[0], title=word_title))
                self.assertEqual(data, original.read_bytes())

    def test_procurement_posts_download_scoped_actual_word_and_excel_tables(self):
        missing = "[A001] / UNSPECIFIED"
        cases = (
            ("proc-plan", "生成采购计划。3#楼HRB400E钢筋要87吨，是甲指的宝丰钢铁供货，10月16日前必须进场；"
             "地下室自采的M15砌筑砂浆要121方，需用单编号XQ-2026-042，工程部李强提的。",
             ("甲指", "自采", "4 需用计划来源", "6 提前期倒排")),
            ("proc-compare", "生成比价表。物资：铝板；需要日期：2026-10-23\n供应商：甲商\n"
             "原报价：SGD 38.25/m2\n交期：2026-10-18\n付款：货到付款\n"
             "供应商：乙商\n原报价：USD 29.75/m2\n交期：2026-10-20",
             ("3 询价文件必备栏", "5 比价表")),
            ("proc-vendor", "生成供应商评价表。一个是天成砂石，供机制砂和碎石，执照给了，"
             "业绩说做过京唐港堆场，厂子没去看过；另一个是宏利防水，卷材那家，"
             "上个月张工去他们唐山厂看了产线，业绩和资料都还没交，售后电话倒是留了。",
             ("3 准入", "5 考察", "6 短名单", "7 评价表头", "用户证据位置")),
        )
        for post, text, titles in cases:
            with self.subTest(post=post):
                self.sid = "proc-"+uuid4().hex[:16]
                md, done = self.markdown(text, post)
                tables = dict(tables_from_md(md))
                for title in titles:
                    self.assertEqual(tables[title], self.spreadsheet_rows(done, title))
                    self.assertEqual(tables[title], self.word_table_rows(done, tables[title][0], title=title))
                if post == "proc-plan":
                    self.assertEqual("87吨", tables["甲指"][1][tables["甲指"][0].index("用户数量")])
                    self.assertEqual("121方", tables["自采"][1][tables["自采"][0].index("用户数量")])
                    self.assertTrue(all(row[-2:] == ["UNSPECIFIED", "UNSPECIFIED"] for row in tables["6 提前期倒排"][1:]))
                elif post == "proc-compare":
                    quoted = tables["5 比价表"]
                    self.assertEqual(["2026-10-18", "2026-10-20"], [r[quoted[0].index("到货期")] for r in quoted[1:]])
                    self.assertEqual(["货到付款", missing], [r[quoted[0].index("付款")] for r in quoted[1:]])
                    self.assertEqual(["SGD", "USD"], [r[quoted[0].index("币种")] for r in quoted[1:]])
                else:
                    entry = tables["3 准入"]
                    self.assertEqual(["京唐港堆场", missing], [r[entry[0].index("用户类似业绩")] for r in entry[1:]])
                    visits = tables["5 考察"]
                    self.assertEqual([missing, "张工"], [r[visits[0].index("考察人")] for r in visits[1:]])
                    self.assertTrue(all(r[-2:] == ["待核", "待核"] for r in tables["7 评价表头"][1:]))
                self.assertTrue(done["submit_blocked"])

    def test_quote_revisions_download_as_separate_sources_without_inheriting_payment(self):
        md, done = self.markdown("生成比价表。物资：铝板\n供应商：甲商\n报价来源：BJ-A\n"
                                 "原报价：38元/m2\n付款：货到付款\n供应商：甲商\n"
                                 "报价来源：BJ-B\n原报价：40元/m2", "proc-compare")
        quoted = dict(tables_from_md(md))["5 比价表"]
        self.assertEqual(["BJ-A", "BJ-B"], [r[quoted[0].index("报价来源")] for r in quoted[1:]])
        self.assertEqual(["货到付款", "[A001] / UNSPECIFIED"], [r[quoted[0].index("付款")] for r in quoted[1:]])
        self.assertEqual(quoted, self.spreadsheet_rows(done, "5 比价表"))
        self.assertEqual(quoted, self.word_table_rows(done, quoted[0]))

    def test_lab_posts_download_literal_separate_samples_instruments_and_mix_targets(self):
        missing = "[A001] / UNSPECIFIED"
        cases = (
            ("lab-sample", "生成取样送检清单。第一车HRB500E盘螺Φ8，磅单28.6吨，批号PL0916，"
             "用在1#楼三层楼板；另一车中砂35方，批号没给，见证人监理刘卫东，取样让王师傅去。",
             ("1 封面与取样范围", "2 角色与签认", "3 取样依据与数量待核", "5 送检与报告追踪")),
            ("lab-record", "生成试验台账。样品：砂浆试块\n报告编号：SJ-2026-0921\n"
             "样品：SBS防水卷材\n委托单号：WT-2026-0917\n压力机YL-017上月刚检定完",
             ("2 样品与委托链", "3 试验记录与报告目录", "4 仪器检定   校准", "6 用户证据位置")),
            ("lab-mix", "生成配比报告提纲。地下室外墙C30，抗渗P8，坍落度140±20，水泥用海螺；"
             "砌体M7.5水泥砂浆，稠度没给；砂都是赣江砂场拉的，两边都还没试拌。",
             ("1 用途与性能目标", "4 原材料与复试来源", "5 试配记录与验证", "6 批准资料与施工调整")),
        )
        for post, message, titles in cases:
            with self.subTest(post=post):
                self.sid = "lab-"+uuid4().hex[:16]
                confirmation = {} if post == "lab-record" else {"confirm_text": "我明白，将由持证人员签认"}
                md, done = self.markdown(message, post, **confirmation)
                tables = dict(tables_from_md(md))
                for title in titles:
                    self.assertEqual(tables[title], self.spreadsheet_rows(done, title))
                    word_title = "4 仪器检定 / 校准" if title == "4 仪器检定   校准" else title
                    self.assertEqual(tables[title], self.word_table_rows(done, tables[title][0], title=word_title))
                if post == "lab-sample":
                    actual = tables[titles[0]]
                    self.assertEqual(["28.6吨", "35方"], [r[actual[0].index("用户进场数量")] for r in actual[1:]])
                    self.assertEqual(["PL0916", missing], [r[actual[0].index("材料批号")] for r in actual[1:]])
                    self.assertTrue(all(r[-2:] == ["（空）", "（空）"] for r in tables[titles[1]][1:]))
                    self.assertTrue(all(r[-1] == "UNSPECIFIED" for r in tables[titles[2]][1:]))
                elif post == "lab-record":
                    actual = tables[titles[1]]
                    self.assertEqual(["SJ-2026-0921", missing], [r[actual[0].index("用户报告编号")] for r in actual[1:]])
                    self.assertEqual([missing]*2, [r[actual[0].index("用户仪器编号引用")] for r in actual[1:]])
                    instruments = tables[titles[2]]
                    self.assertEqual("YL-017", instruments[1][0])
                    self.assertEqual(missing, instruments[1][instruments[0].index("用户检定有效期原文")])
                else:
                    actual = tables[titles[0]]
                    self.assertEqual(["C30", "M7.5"], [r[actual[0].index("用户强度等级 / 砂浆等级")] for r in actual[1:]])
                    self.assertEqual(["140±20", missing], [r[actual[0].index("用户坍落度要求")] for r in actual[1:]])
                    self.assertNotIn("140±20mm", md)
                    self.assertTrue(all(r[-2:] == ["UNSPECIFIED", "（空）"] for r in tables[titles[-1]][1:]))
                self.assertTrue(done["submit_blocked"])

    def test_lab_high_risk_posts_require_exact_typed_confirmation(self):
        for post, message in (("lab-sample", "生成取样送检清单。钢筋批号A01。"),
                              ("lab-mix", "生成配比报告提纲。地下室外墙C30。")):
            for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
                with self.subTest(post=post, fields=fields):
                    self.sid = "labgate-"+uuid4().hex[:16]
                    done, _ = self.post(message, expert_ids=[post], **fields)
                    self.assertFalse(done["wrote"], done)
                    self.assertEqual([], done["deliverables"])
                    self.assertIn("我明白，将由持证人员签认", done["text"])
        self.no_model.assert_not_called()

    def test_selected_lab_attachments_have_separate_rows_and_immutable_originals(self):
        cases = (
            ("lab-sample", "取样送检清单", "材料,批号,数量,见证人\n钢筋,A01,62吨,刘测试\n中砂,,35方,\n", "1 封面与取样范围", "材料批号", ["A01", "[A001] / UNSPECIFIED"]),
            ("lab-record", "试验台账", "样品,报告编号,检测机构\n试件A,R-A,机构A\n试件B,,机构B\n", "3 试验记录与报告目录", "用户报告编号", ["R-A", "[A001] / UNSPECIFIED"]),
            ("lab-mix", "配比报告提纲", "部位,强度等级,坍落度,试配记录编号\n池壁A,C35,160±20,PB-A\n砌体B,M7.5,,\n", "1 用途与性能目标", "用户坍落度要求", ["160±20", "[A001] / UNSPECIFIED"]),
        )
        for post, name, source, title, field, expected in cases:
            with self.subTest(post=post):
                self.sid = "labfile-"+uuid4().hex[:16]
                data = source.encode("utf-8")
                uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                            files={"files": ("合成试验资料.csv", data, "text/csv")})
                self.assertEqual(uploaded.status_code, 200, uploaded.text)
                identifier = uploaded.json()["files"][0]["id"]
                original, _ = flow.uploads.upload_file(self.sid, identifier)
                before = original.read_bytes()
                self.assertEqual(data, before)
                confirmation = {} if post == "lab-record" else {"confirm_text": "我明白，将由持证人员签认"}
                md, done = self.markdown("根据所选附件生成"+name+"，缺失字段保持待核。", post,
                                         attachments=[identifier], **confirmation)
                actual = dict(tables_from_md(md))[title]
                self.assertEqual(expected, [r[actual[0].index(field)] for r in actual[1:]])
                self.assertEqual(actual, self.spreadsheet_rows(done, title))
                self.assertEqual(actual, self.word_table_rows(done, actual[0], title=title))
                self.assertEqual(before, original.read_bytes())

    def test_short_compare_commands_download_separate_names_and_only_labelled_prices(self):
        for tail, names, prices in (("甲商 乙商", ["甲商", "乙商"], ["TBD", "TBD"]),
                                   ("甲商 报价3861", ["甲商"], ["3861"]),
                                   ("甲商 3861", ["甲商"], ["TBD"])):
            with self.subTest(tail=tail):
                self.sid = "shortproc-"+uuid4().hex[:16]
                md, done = self.markdown("写一份比价表 钢筋 "+tail, "proc-compare")
                quoted = dict(tables_from_md(md))["5 比价表"]
                self.assertEqual(names, [r[quoted[0].index("供方名称")] for r in quoted[1:]])
                self.assertEqual(prices, [r[quoted[0].index("单价 / 原报价")] for r in quoted[1:]])
                self.assertEqual(quoted, self.spreadsheet_rows(done, "5 比价表"))
                self.assertEqual(quoted, self.word_table_rows(done, quoted[0]))

    def test_selected_procurement_attachment_rows_are_not_merged_or_rewritten(self):
        cases = (
            ("proc-plan", "采购计划", "物资,规格,数量,供应方式\n钢筋,HRB400E,87吨,甲指\n砂浆,M15,121方,自采\n", "4 需用计划来源", "用户数量", ["87吨", "121方"]),
            ("proc-compare", "比价表", "供应商,物资,原报价,币种,到货期\n甲商,铝板,38.25/m2,SGD,2026-10-18\n乙商,铝板,29.75/m2,USD,\n", "5 比价表", "到货期", ["2026-10-18", "[A001] / UNSPECIFIED"]),
            ("proc-vendor", "供应商评价表", "供应商,供货范围,质量资料\n甲商,机制砂,报告A\n乙商,卷材,\n", "7 评价表头", "质量资料", ["报告A", "[A001] / UNSPECIFIED"]),
        )
        for post, name, source, title, field, expected in cases:
            with self.subTest(post=post):
                self.sid = "procfile-"+uuid4().hex[:16]
                data = source.encode("utf-8")
                uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                            files={"files": ("合成采购资料.csv", data, "text/csv")})
                self.assertEqual(uploaded.status_code, 200, uploaded.text)
                original, _ = flow.uploads.upload_file(self.sid, uploaded.json()["files"][0]["id"])
                before = original.read_bytes()
                self.assertEqual(data, before)
                md, done = self.markdown("根据所选附件生成"+name+"，缺失字段保持待核。", post,
                                         attachments=[uploaded.json()["files"][0]["id"]])
                actual = dict(tables_from_md(md))[title]
                self.assertEqual(expected, [r[actual[0].index(field)] for r in actual[1:]])
                self.assertEqual(actual, self.spreadsheet_rows(done, title))
                self.assertEqual(actual, self.word_table_rows(done, actual[0]))
                if post == "proc-compare":
                    self.assertEqual(["合成采购资料.csv"]*2, [r[actual[0].index("输入资料名 / 分组")] for r in actual[1:]])
                elif post == "proc-vendor":
                    proof = dict(tables_from_md(md))["用户证据位置"]
                    self.assertTrue(all("合成采购资料.csv" in r[-1] for r in proof[1:]))
                self.assertEqual(before, original.read_bytes())

    def test_two_selected_quote_files_cannot_share_the_first_supplier_or_material(self):
        documents = (("报价A.txt", "物资：铝板\n供应商：甲商\n原报价：SGD 38.25/m2\n付款：货到付款"),
                     ("报价B.txt", "原报价：40元/m2"))
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                    files=[("files", (name, source.encode("utf-8"), "text/plain")) for name, source in documents])
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        originals = [flow.uploads.upload_file(self.sid, f["id"])[0] for f in uploaded.json()["files"]]
        before = [path.read_bytes() for path in originals]
        md, done = self.markdown("根据两个所选附件生成比价表，归属未说明的字段保持待核。", "proc-compare",
                                 attachments=[f["id"] for f in uploaded.json()["files"]])
        quoted = dict(tables_from_md(md))["5 比价表"]
        missing = "[A001] / UNSPECIFIED"
        for field, expected in (("供方名称", ["甲商", missing]), ("物资 / 标的", ["铝板", missing]),
                                ("付款", ["货到付款", missing]), ("输入资料名 / 分组", ["报价A.txt", "报价B.txt"])):
            self.assertEqual(expected, [r[quoted[0].index(field)] for r in quoted[1:]])
        self.assertEqual(quoted, self.spreadsheet_rows(done, "5 比价表"))
        self.assertEqual(quoted, self.word_table_rows(done, quoted[0]))
        self.assertEqual(before, [path.read_bytes() for path in originals])

    def test_english_tender_uses_shared_http_pipeline_and_extracts_scores(self):
        md, done = self.markdown("Parse the following tender into a requirements table.\n"
                                 "Project: Synthetic Facade A.\nContract duration: 180 calendar days.\n"
                                 "Technical proposal: 35 marks. Programme: 12 marks.\n"
                                 "Leave missing qualifications unspecified. Do not approve submission.", "bid-parse", locale="en")
        self.assertIn("Tender requirements", md)
        self.assertIn("| Contract duration | 180 calendar days |", md)
        self.assertIn("| Technical proposal | 35 marks |", md)
        self.assertIn("| Programme | 12 marks |", md)
        self.assertIn("| Qualifications | Not detected in supplied text |", md)
        self.assertTrue(done["submit_blocked"])

    def test_english_tender_attachment_has_english_export_without_rewriting_source(self):
        text = ("Invitation to Tender\nProject: Synthetic Facade G.\nTender reference: SYN-ATT-010.\n"
                "Contract duration: 140 calendar days.\nTender security: S$35,000.\n")
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                    files=[("files", ("synthetic-tender-en.txt", text.encode(), "text/plain"))])
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        md, _ = self.markdown("Parse the selected tender attachment into a requirements table. "
                              "Leave missing fields unspecified. Do not approve submission.", "bid-parse", locale="en",
                              attachments=[f["id"] for f in uploaded.json()["files"]])
        self.assertIn("Tender requirements", md)
        self.assertIn("| Contract duration | 140 calendar days |", md)
        self.assertIn("| Tender security | S$35,000 |", md)

    def test_instruction_does_not_become_bidder_experience(self):
        md, _ = self.markdown("请解析以下招标正文并生成招标解析表。项目名称：合成幕墙A。"
                              "工期180日历天。未提供的资格、业绩、价格信息保持未检出。", "bid-parse")
        self.assertIn("| 类似业绩 | 未在原文检出 |", md)
        self.assertNotIn("| 类似业绩 | 价格信息保持未检出", md)

    def test_technical_response_prose_is_in_its_business_section_and_downloads(self):
        md, done = self.markdown("请生成技术标目录和响应草稿。\n招标正文：\n"
                                 "项目名称：合成幕墙响应项目；工期180日历天；施工组织设计35分。\n\n"
                                 "投标响应：\n部署工艺：先整理接口图，再核对样板记录。\n"
                                 "质量措施：保存材料检验记录；未知抽检频率保持待填。", "bid-tech")
        body = md.split("## 3 部署工艺", 1)[1].split("## 4 工期资源", 1)[0]
        self.assertIn("先整理接口图", body)
        self.assertIn("来源：", body)
        self.assertTrue(any(f["name"].endswith(".docx") for f in done["deliverables"]))
        self.assertTrue(done["submit_blocked"])

    def test_warehouse_quantity_and_units_reach_actual_cells(self):
        md, _ = self.markdown("请生成材料入库台账。材料名称：铝板；规格：3mm；数量：120张；"
                              "供应商：合成供应商甲；批次：SYN-B001；仓库：一号仓。未知单价保持待填。", "warehouse")
        tables = tables_from_md(md)
        ledger = next(rows for _, rows in tables if "账面结存" in rows[0])
        row = dict(zip(ledger[0], ledger[1]))
        self.assertEqual((row["物资"], row["单位"], row["入库"], row["账面结存"]), ("铝板", "张", "120张", "TBD"))
        self.assertEqual(len(ledger), 2)

    def test_minutes_auto_route_to_document_post_and_keep_decision_separate(self):
        md, done = self.markdown("请生成会议纪要草稿。会议名称：合成幕墙协调会；日期：2026-10-09；"
                                 "地点：东楼301；与会人员：张测试、李测试；"
                                 "决定：李测试在2026-10-12前提供节点图；未提供的信息保持待填。")
        self.assertEqual(done["route"]["expert_ids"], ["admin-doc"])
        self.assertIn("| 时间 | 2026-10-09 |", md)
        self.assertIn("| 李测试在2026-10-12前提供节点图 | 李测试 | 2026-10-12前 |", md)
        self.assertNotIn("节点图；未提供", md)

    def test_reported_minutes_reach_actual_export_and_auto_route(self):
        md, done = self.markdown("请生成会议纪要，10月8日下午在项目部西楼会议室开的材料协调会，"
                                 "项目经理吴海平主持，记录人是资料员林宁，"
                                 "会上定了由材料负责人钱文在10月15日前完成二层材料清点，"
                                 "东门卸货时间没定下来，留到下次例会再议。")
        self.assertEqual(done["route"]["expert_ids"], ["admin-doc"])
        rows = next(rows for title, rows in tables_from_md(md) if title.startswith("已议定事项"))
        self.assertEqual(["二层材料清点", "材料负责人钱文", "10月15日前", ""], rows[1])
        self.assertEqual(rows, self.spreadsheet_rows(done, "已议定事项"))
        self.assertEqual(rows, self.word_table_rows(done, rows[0]))
        info = dict(self.spreadsheet_rows(done, "会议记录要素")[1:])
        self.assertEqual("资料员林宁", info["记录人"])
        self.assertEqual("10月8日下午", info["时间"])

    def test_two_seal_files_keep_per_file_quantities_and_blank_approval_in_download(self):
        md, done = self.markdown("请生成用印审批表，一个是钢筋班组劳动合同，27份，盖合同专用章，"
                                 "审批人写商务经理林安；另一个是施工安全管理协议，11份，"
                                 "盖项目部印章，这个谁批还没定先空着。"
                                 "经办都是我，劳资员周宁，下周二前要用。", "admin-doc")
        rows = next(rows for title, rows in tables_from_md(md) if title.startswith("逐份用印文件核验"))
        self.assertEqual(["钢筋班组劳动合同", "27份", "合同专用章", "商务经理林安", ""], rows[1])
        self.assertEqual(["施工安全管理协议", "11份", "项目部印章"], rows[2][:3])
        self.assertTrue(rows[2][3].startswith("UNSPECIFIED"))
        self.assertEqual("", rows[2][4])
        self.assertEqual(rows, self.spreadsheet_rows(done, "逐份用印文件核验"))
        self.assertEqual(rows, self.word_table_rows(done, rows[0]))

    def test_two_vehicle_requests_keep_second_plate_driver_and_time_unknown_in_download(self):
        md, done = self.markdown("请生成后勤用车清单，周五有两拨用车，第一拨商务车苏E8J613，司机赵文，"
                                 "早上8:10从项目部送甲方徐总他们3个人去苏州高铁站；"
                                 "第二拨下午去诚达钢构厂看构件，安全部王宁带队6个人，"
                                 "来回大概140公里，司机还没定车也没批，几点走等通知。", "admin-office")
        rows = next(rows for title, rows in tables_from_md(md) if title.startswith("逐项用车需求"))
        self.assertEqual(rows, self.spreadsheet_rows(done, "逐项用车需求"))
        self.assertEqual(rows, self.word_table_rows(done, rows[0]))
        first, second = dict(zip(rows[0], rows[1])), dict(zip(rows[0], rows[2]))
        self.assertEqual("苏E8J613", first["拟用车/车牌（待核）"])
        self.assertEqual("赵文", first["拟司机（待核）"])
        self.assertEqual("诚达钢构厂", second["目的地"])
        self.assertEqual("大概140公里", second["用户提供里程（未计算）"])
        for key in ("拟用车/车牌（待核）", "拟司机（待核）", "出发时间"):
            self.assertTrue(second[key].startswith("UNSPECIFIED"))
        self.assertEqual("", second["派车审批"])

    def test_wrong_manual_post_does_not_claim_success_with_a_different_document(self):
        response = self.client.post("/api/chat", json={"message": "请生成会议纪要草稿。会议名称：合成会议",
                                    "session_id": self.sid, "expert_ids": ["admin-office"]})
        self.assertEqual(response.status_code, 200)
        done = next(e["data"] for e in flow.sse_events(response) if e["event"] == "done")
        self.assertFalse(done["ok"])
        self.assertFalse(done["wrote"])
        self.assertEqual(done["deliverables"], [])
        self.assertIn("admin-doc", done["text"])

    def test_unknown_approver_question_is_still_read_only(self):
        response = self.client.post("/api/chat", json={"message": "生成用印审批表，谁批还没定？",
                                    "session_id": self.sid, "expert_ids": ["admin-doc"]})
        self.assertEqual(response.status_code, 200)
        done = next(event["data"] for event in flow.sse_events(response) if event["event"] == "done")
        self.assertTrue(done["ok"])
        self.assertFalse(done["wrote"])
        self.assertEqual([], done["deliverables"])
        self.assertEqual("chat", done["route"]["intent"])
        self.no_model.assert_not_called()

    def test_five_design_posts_download_literal_scoped_word_and_excel_tables(self):
        from packing_assistant.post_drafts.design_services import MISSING
        cases = (
            ("plumbing", "生成给排水设计说明。1#楼给水从学苑路DN150市政管接，水压0.26兆帕；"
             "2#楼从规划三路接，水压还没给。1#楼出户标高-1.75，接市政井W9，井底-2.55；"
             "2#楼出户-1.35，市政井标高没给。", ("水源与市政接驳", "排水与标高衔接"),
             "水源与市政接驳", "用户水压", "0.26兆帕", MISSING),
            ("hvac", "生成暖通设计说明。客房那套是风机盘管带新风，冷负荷1320kW，热负荷还没算出来，"
             "新风量先按17400m3/h；宴会厅是全空气，冷615kW、热408kW，机房在裙房3层K3-08。"
             "两套都出自负荷计算书KT-FH-A08。", ("冷热源方案与负荷资料", "风系统与水系统", "机房、冷却塔与竖向空间"),
             "冷热源方案与负荷资料", "用户热负荷", MISSING, "408kW"),
            ("electrical", "生成电气设计说明。1#联合厂房设备容量大概2250kW，需要系数工艺那边给的0.73，"
             "防雷按二类；5#倒班宿舍才465kW，系数还没人给，防雷类别也没定。"
             "负荷计算书目前只有厂房的，编号JS-DQ-019，宿舍的后补。",
             ("负荷计算书数据登记", "防雷、接地与等电位"), "负荷计算书数据登记", "需要/利用系数", "0.73", MISSING),
            ("fire-protect", "生成消防设计专篇。A座是酒店，高度78.2米，分区面积设计院给的1450平米；"
             "B座是裙房商业，高度22.8米，分区面积还没给。防排烟A座走机械排烟，B座中庭那块是自然排烟窗。"
             "B座消防用水量给排水提的是605立方，A座的还没算出来。",
             ("建筑防火与分区", "防烟排烟", "消防给水与灭火"), "消防给水与灭火", "用户消防水量", MISSING, "605立方"),
            ("steel", "生成钢结构设计说明。办公楼连廊那根钢梁LL-4，跨度17.8米，材料Q355D；"
             "屋顶设备平台的钢柱GZ-7，用的是Q345C，柱脚地脚螺栓4根M32，跨度啥的没有。"
             "甲方说连廊梁刷厚型防火涂料，平台柱的还没定。",
             ("材料与构件规格记录", "构件防火保护", "实体连接界面"), "构件防火保护", "用户防火做法", "厚型防火涂料", MISSING),
        )
        for post, message, titles, target, label, first, second in cases:
            with self.subTest(post=post):
                self.sid = "design-" + post[:8] + "-" + uuid4().hex[:10]
                confirmation = {"confirm_text": "我明白，将由持证人员签认"} if post in {"fire-protect", "steel"} else {}
                md, done = self.markdown(message, post, **confirmation)
                tables = dict(tables_from_md(md))
                for title in titles:
                    rows = next(rows for name, rows in tables.items() if name.startswith(title))
                    self.assertEqual(3, len(rows), (post, title, rows))
                    self.assertEqual(rows, self.spreadsheet_rows(done, title))
                    self.assertEqual(rows, self.word_table_rows(done, rows[0]))
                    for index, header in enumerate(rows[0]):
                        if header.startswith("本稿") or "核验结果" in header:
                            self.assertTrue(all(row[index] == MISSING for row in rows[1:]))
                rows = next(rows for name, rows in tables.items() if name.startswith(target))
                self.assertEqual([first, second], [row[rows[0].index(label)] for row in rows[1:]])
                self.assertTrue(done["submit_blocked"])

    def test_design_narrative_does_not_bypass_high_risk_typed_confirmation(self):
        for post, message in (("fire-protect", "生成消防设计专篇。A座是酒店，高度78.2米。"),
                              ("steel", "生成钢结构设计说明。钢梁LL-4，材料Q355D。")):
            for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
                with self.subTest(post=post, fields=fields):
                    self.sid = "gate-" + uuid4().hex[:16]
                    done, _ = self.post(message, expert_ids=[post], **fields)
                    self.assertFalse(done["wrote"], done)
                    self.assertEqual([], done["deliverables"])
                    self.assertIn("我明白，将由持证人员签认", done["text"])
        self.no_model.assert_not_called()

    def test_four_basic_design_posts_download_literal_hierarchical_word_and_excel_tables(self):
        missing, not_found = "[A001] 待填", "未在原文检出"
        cases = (
            ("architecture", "生成建筑设计说明。1#教学楼，地上5层，总建筑面积5140平米，高23.6米；"
             "2#幼儿园，三层，总建筑面积2380平米，高度还没给。",
             ("1 工程概况与设计范围_2",), "1 工程概况与设计范围_2", "建筑高度", ["23.6米", missing]),
            ("structure", "生成结构设计说明。1#办公楼，辖区：CN，柱子KZ-6用C45；"
             "2#仓库，钢柱GZ-3用Q355D。柱子KZ-6跨度6.7米，基础型式 筏板基础。",
             ("4 荷载与组合", "5 材料", "6 地基基础", "9 构件复核清单"),
             "9 构件复核清单", "用户跨度", ["6.7米", missing]),
            ("geotech", "生成岩土工程勘察简报。孔号 BH13\n辖区：CN\n孔深：33.4m\n水位 1.85m\n"
             "层号：④\nc：11.3kPa\nφ 8.7度\n孔号 BH14\n孔深 34.5m\n水位没给\n"
             "层号：④\nc：12.3kPa\nφ没写",
             ("4 场地与地下水_2", "5 勘探工作量与试验", "6 地层与参数摘录"),
             "6 地层与参数摘录", "φ", ["8.7度", not_found]),
            ("facade", "生成幕墙专业说明。A座有两种幕墙，大堂东立面是全玻幕墙，编号QB-02，"
             "风压2.47；塔楼标准层是单元式的UC-04，风压3.36，水密要1100Pa。"
             "预埋单元式用槽式埋件，全玻的还不知道。",
             ("1 体系选型_2", "3 预埋件与主体结构接口", "4 抗风、气密、水密与层间变位", "8 验收资料"),
             "3 预埋件与主体结构接口", "预埋件", [missing, "槽式埋件"]),
        )
        for post, message, titles, target, label, expected in cases:
            with self.subTest(post=post):
                self.sid = "basic-" + uuid4().hex[:16]
                confirmation = {"confirm_text": "我明白，将由持证人员签认"} if post != "architecture" else {}
                md, done = self.markdown(message, post, **confirmation)
                tables = dict(tables_from_md(md))
                for title in titles:
                    rows = tables[title]
                    self.assertEqual(3, len(rows), (post, title, rows))
                    self.assertEqual(rows, self.spreadsheet_rows(done, title))
                    self.assertEqual(rows, self.word_table_rows(done, rows[0]))
                rows = tables[target]
                self.assertEqual(expected, [row[rows[0].index(label)] for row in rows[1:]])
                if post in {"structure", "geotech"}:
                    self.assertEqual(["CN", "UNSPECIFIED"], [row[-1] for row in rows[1:]])
                if post == "structure":
                    foundations = tables["6 地基基础"]
                    self.assertEqual(["筏板基础", missing], [row[foundations[0].index("用户基础型式")] for row in foundations[1:]])
                    self.assertTrue(all(row[-1] == "UNSPECIFIED" for row in tables["7 计算步骤"][1:]))
                if post == "facade":
                    self.assertTrue(all(row[-2] == "本轮未执行验收" for row in tables["8 验收资料"][1:]))
                    self.assertNotIn("2.47kPa", md)
                self.assertTrue(done["submit_blocked"])

    def test_duplicate_member_reference_downloads_as_unassigned_not_last_building(self):
        md, done = self.markdown("生成结构设计说明。1#办公楼，柱子KZ-6用C45；"
                                 "2#仓库，柱子KZ-6用C35。柱子KZ-6跨度6.7米，没说是哪栋。",
                                 "structure", confirm_text="我明白，将由持证人员签认")
        tables = dict(tables_from_md(md))
        loads = tables["4 荷载与组合"]
        members = tables["9 构件复核清单"]
        self.assertEqual(["1#办公楼", "2#仓库", "[A001] 待填"], [row[0] for row in loads[1:]])
        self.assertEqual(["[A001] 待填", "[A001] 待填", "6.7米"],
                         [row[members[0].index("用户跨度")] for row in members[1:]])
        for title in ("4 荷载与组合", "9 构件复核清单"):
            self.assertEqual(tables[title], self.spreadsheet_rows(done, title))
            self.assertEqual(tables[title], self.word_table_rows(done, tables[title][0]))

    def test_specialty_downloads_keep_two_objects_literal_sources_and_blank_signoffs(self):
        cases = (
            ("landscape", "生成景观说明。南入口广场种银杏，胸径19，共25株；宅间绿地种桂花，冠幅251，覆土1.3米。苗木表是甲方苗表L-28。",
             "苗木与分区来源核对", "数量", ["25株", "UNSPECIFIED"]),
            ("interior", "生成装修做法表。房间：甲卫；地面：砖A；防水材料：材料A；防水上翻：1.9米；节点图号：N-A；节点来源：节点包A；收口责任：张工\n房间：乙卫；地面：砖B；节点图号：N-B；收口责任：李工",
             "房间做法与收口责任", "上翻", ["1.9米", "UNSPECIFIED"]),
            ("intel-weak", "生成弱电说明。系统：视频监控覆盖A楼；点数：213；品牌：甲品牌\n系统：门禁；覆盖：B楼；点数：87\n监控走的桥架规格：201×101\n门禁那边桥架规格：151×76\n数量都是按用户点表DB-28",
             "系统点表与联调边界", "用户数量", ["213", "87"]),
            ("civil-defense", "生成人防专篇提纲。防护单元：F1；平时功能：停车；人防等级：用户核6级；等级批准资料：用户批准包A；柴油电站：不设固定电站，预留移动电站接口\n防护单元：F2；平时功能：储藏",
             "防护单元专业接口核对", "用户等级要求（未核）", ["用户核6级", "UNSPECIFIED"]),
            ("hydraulic", "生成水利工程提纲。水闸：甲闸\n水文断面：H1；水位：8.6m；流量：127.5m3/s；水文报告：水文甲\n勘探孔：Z1；地层：黏土\n水闸：乙闸\n水文断面：H2；水位：7.9m；水文报告：水文乙\n勘探孔：Z2；地层：淤泥",
             "水工参数来源与待计算项", "用户流量", ["127.5m3/s", "UNSPECIFIED"]),
        )
        for post, message, title, field, expected in cases:
            with self.subTest(post=post):
                self.sid = "specialty-" + uuid4().hex[:16]
                confirmation = {"confirm_text": "我明白，将由持证人员签认"} if post in {"civil-defense", "hydraulic"} else {}
                md, done = self.markdown(message, post, **confirmation)
                rows = dict(tables_from_md(md))[title]
                self.assertEqual(3, len(rows), rows)
                self.assertEqual(expected, [r[rows[0].index(field)] for r in rows[1:]])
                self.assertTrue(all(r[-1] == "" for r in rows[1:]))
                self.assertEqual(rows, self.spreadsheet_rows(done, title))
                self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
                self.assertTrue(done["submit_blocked"])

    def test_selected_specialty_csvs_keep_reordered_fields_blank_cells_and_original_bytes(self):
        cases = (
            ("landscape", "数量,苗木来源,冠幅,苗木\n39,苗表A,137,红叶石楠球\n,,251,桂花\n", "苗木与分区来源核对", "数量", 0),
            ("interior", "收口责任,节点来源,防水上翻,房间\n张工,节点A,1.9米,甲卫\n李工,,,乙卫\n", "房间做法与收口责任", "上翻", 5),
            ("intel-weak", "点数,点表来源,品牌,系统,密码\n20,点表A,品牌A,视频监控,synthetic-private-A\n30,,,门禁,synthetic-private-B\n", "系统点表与联调边界", "数量来源", 4),
            ("civil-defense", "数量,型号,防护设备,所属单元\n4,HFM1520,东防护门,F1\n,,西防护门,F2\n", "4 防护设备与孔口登记", "用户数量", 7),
            ("hydraulic", "水位,水文报告,流量,水文断面\n8.63m,水文甲,127.6m3/s,H1\n,水文乙,,H2\n", "3 水文成果与水位流量登记", "用户设计水位", 4),
        )
        for post, source, title, missing_field, _ in cases:
            with self.subTest(post=post):
                self.sid = "specialty-up-" + uuid4().hex[:16]
                raw = source.encode("utf-8")
                uploaded = self.client.post("/api/upload", data={"session_id": self.sid}, files={"files": ("合成专项资料.csv", raw, "text/csv")})
                self.assertEqual(200, uploaded.status_code, uploaded.text)
                identifier = uploaded.json()["files"][0]["id"]
                original, _ = flow.uploads.upload_file(self.sid, identifier)
                self.assertEqual(raw, original.read_bytes())
                confirmation = {"confirm_text": "我明白，将由持证人员签认"} if post in {"civil-defense", "hydraulic"} else {}
                md, done = self.markdown("按所选资料生成本岗内部草稿，缺失字段保持待填。", post, attachments=[identifier], **confirmation)
                rows = dict(tables_from_md(md))[title]
                self.assertEqual(3, len(rows), rows)
                if post == "intel-weak":
                    self.assertNotIn("点表A", rows[2][rows[0].index(missing_field)])
                else:
                    self.assertEqual("UNSPECIFIED", rows[2][rows[0].index(missing_field)])
                self.assertEqual(rows, self.spreadsheet_rows(done, title))
                self.assertEqual(rows, self.word_table_rows(done, rows[0], title=title))
                self.assertNotRegex(md, r"synthetic-private-[AB]")
                self.assertEqual(raw, original.read_bytes())
                self.assertTrue(done["submit_blocked"])

    def test_two_specialty_uploads_do_not_share_node_sources_or_project_basis(self):
        sources = ("项目：甲项目；CN依据：CN-A\n房间：甲卫；防水材料：材料A；节点来源：节点A；收口责任：张工",
                   "项目：乙项目；SG依据：SG-B\n房间：乙卫；地面：砖B")
        raw = [text.encode("utf-8") for text in sources]
        uploaded = self.client.post("/api/upload", data={"session_id": self.sid},
                                    files=[("files", (name, blob, "text/plain")) for name, blob in zip(("甲资料.txt", "乙资料.txt"), raw)])
        self.assertEqual(200, uploaded.status_code, uploaded.text)
        identifiers = [f["id"] for f in uploaded.json()["files"]]
        originals = [flow.uploads.upload_file(self.sid, identifier)[0] for identifier in identifiers]
        md, done = self.markdown("按所选资料生成室内装修内部说明，不要补填。", "interior", attachments=identifiers)
        tables = dict(tables_from_md(md))
        rooms = tables["房间做法与收口责任"]
        self.assertEqual(["甲卫", "乙卫"], [r[0] for r in rooms[1:]])
        self.assertEqual("UNSPECIFIED", rooms[2][rooms[0].index("防水")])
        self.assertEqual("UNSPECIFIED", rooms[2][rooms[0].index("收口责任")])
        self.assertNotIn("节点A", rooms[2][rooms[0].index("节点来源")])
        for title in ("房间做法与收口责任", "各项目与文件依据归属"):
            self.assertEqual(tables[title], self.spreadsheet_rows(done, title))
            self.assertEqual(tables[title], self.word_table_rows(done, tables[title][0], title=title))
        self.assertEqual(raw, [p.read_bytes() for p in originals])

    def test_specialty_weak_credentials_and_private_addresses_do_not_enter_downloads(self):
        md, done = self.markdown('生成弱电说明。系统：门禁\npassword="TOP\n品牌：synthetic-secret-tail\nEND"\n网络安全：隔离10.24.35.46\n点数：20\n品牌：安全品牌', "intel-weak")
        self.assertNotRegex(md, r"synthetic-secret-tail|10\.24\.35\.46")
        rows = dict(tables_from_md(md))["系统点表与联调边界"]
        self.assertEqual("20", rows[1][rows[0].index("用户数量")])
        self.assertEqual("", rows[1][-1])
        self.assertEqual(rows, self.spreadsheet_rows(done, "系统点表与联调边界"))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="系统点表与联调边界"))
        from xml.etree import ElementTree as ET
        for item in done["deliverables"]:
            response = self.client.get("/api/file", params={"path": item["path"]})
            if Path(item["name"]).suffix == ".docx":
                with ZipFile(BytesIO(response.content)) as archive:
                    text = "".join(ET.fromstring(archive.read("word/document.xml")).itertext())
                self.assertNotRegex(text, r"synthetic-secret-tail|10\.24\.35\.46")

    def test_specialty_formula_like_cells_and_encoded_markup_remain_literal(self):
        md, done = self.markdown("生成装修做法表\n|房间|地面|墙面|天花|\n|---|---|---|---|\n|=2+3|x&lt;10 &amp;&amp; A&#124;B|literal &amp;amp;|&lt;b&gt;材料&lt;/b&gt;|", "interior")
        rows = dict(tables_from_md(md))["2 房间功能与饰面做法"]
        self.assertEqual(["=2+3", "x<10 && A|B", "literal &amp;", "<b>材料</b>"], [rows[1][i] for i in (1, 4, 6, 8)])
        self.assertEqual(rows, self.spreadsheet_rows(done, "2 房间功能与饰面做法"))
        self.assertEqual(rows, self.word_table_rows(done, rows[0], title="2 房间功能与饰面做法"))
        import openpyxl
        artifact = next(f for f in done["deliverables"] if f["name"].endswith(".xlsx"))
        book = openpyxl.load_workbook(BytesIO(self.client.get("/api/file", params={"path": artifact["path"]}).content), data_only=False)
        try:
            self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))
        finally:
            book.close()

    def test_specialty_high_risk_posts_require_exact_current_turn_ack(self):
        for post, message in (("civil-defense", "生成人防专篇提纲。防护单元：F1；人防等级：用户核6级"),
                              ("hydraulic", "生成水利提纲。水闸：甲闸；扬程：5.8m")):
            for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
                with self.subTest(post=post, fields=fields):
                    self.sid = "specialty-gate-" + uuid4().hex[:16]
                    done, _ = self.post(message, expert_ids=[post], **fields)
                    self.assertFalse(done["wrote"], done)
                    self.assertEqual([], done["deliverables"])
                    self.assertIn("我明白，将由持证人员签认", done["text"])
        self.no_model.assert_not_called()

    def test_specialty_questions_and_dry_run_do_not_write_or_create_downloads(self):
        for post in ("landscape", "interior", "intel-weak", "civil-defense", "hydraulic"):
            with self.subTest(post=post):
                self.sid = "specialty-chat-" + uuid4().hex[:16]
                done, _ = self.post("这个设计岗位需要哪些资料？先不要生成文件", expert_ids=[post])
                self.assertFalse(done["wrote"], done)
                self.assertEqual([], done["deliverables"])
        self.no_model.assert_not_called()

    def test_basic_design_narrative_requires_exact_high_risk_typed_confirmation(self):
        cases = (("structure", "生成结构设计说明，柱子KZ-6用C45。"),
                 ("geotech", "生成岩土工程勘察简报，ZK4孔深38.7米，第3层是粉砂。"),
                 ("facade", "生成幕墙专业说明，单元式的UC-04，风压3.36。"))
        for post, message in cases:
            for fields in ({}, {"confirm_text": "我明白，将由持证人员签"}, {"confirm_ok": True}):
                with self.subTest(post=post, fields=fields):
                    self.sid = "gate-" + uuid4().hex[:16]
                    done, _ = self.post(message, expert_ids=[post], **fields)
                    self.assertFalse(done["wrote"], done)
                    self.assertEqual([], done["deliverables"])
                    self.assertIn("我明白，将由持证人员签认", done["text"])
        self.no_model.assert_not_called()


if __name__ == "__main__":
    unittest.main()
