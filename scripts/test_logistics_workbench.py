"""Offline logistics acceptance: immutable sources, proposals, handover and solver gates."""
from copy import deepcopy
import hashlib
import io
import json
import os
import struct
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PYTHON_DOTENV_DISABLED"] = "1"
from fastapi import FastAPI
from fastapi.testclient import TestClient
from demo import logistics_api as api
from packing_assistant.logistics import agent, bundle, packing
from packing_assistant.logistics.intake import parse_document
from packing_assistant.logistics.records import LogisticsStore, digest, validate_record, MAX_HISTORY
from packing_assistant.engineering.schedule import ScheduleConflict
from packing_assistant.runtime import cancel

SOURCE = ("package_id,material_id,name,package_count,quantity,units_per_package,unit,length_mm,width_mm,height_mm,dimension_scope,net_kg,gross_kg,weight_scope\n"
          "B01,M01,synthetic box,1,10,10,piece,1000,800,600,package,100,120,package\n").encode()


class Workbench(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = LogisticsStore(self.root)
        for manager in [patch.dict(os.environ, {"CIVIL_SANDBOX_ROOTS": str(self.root), "CIVIL_SANDBOX": "workspace-write"}),
                        patch.object(api, "store", lambda: self.store), patch.object(api, "DOCUMENTS", api.cad.MemoryStore()),
                        patch.object(api, "PROPOSALS", api.cad.MemoryStore())]:
            manager.start()
            self.addCleanup(manager.stop)
        app = FastAPI()
        app.include_router(api.router)
        app.include_router(api.eng.router)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        self.doc = parse_document(SOURCE, "synthetic.csv")

    def create(self):
        return self.store.create("合成箱单", self.doc, SOURCE)

    def post(self, path, body):
        return self.client.post(api.BASE + path, json=body)

    def test_original_and_history_survive_reopen(self):
        original = self.create()
        proposal = agent.propose_command(original["document"], "把 R00001 毛重改为 0.13 t")
        changed = self.store.revise(original["id"], 1, agent.apply_proposal(original["document"], proposal))
        reopened = LogisticsStore(self.root).open(original["id"])
        self.assertEqual(reopened, changed)
        self.assertEqual(reopened["document"]["rows"][0]["gross_kg"], 130)
        self.assertEqual(self.store.open(original["id"], 1)["document"], original["document"])
        self.assertEqual(self.store.source(original["id"])[0], SOURCE)
        self.assertEqual(reopened["document"]["rows"][0]["evidence"]["gross_kg"]["raw"], "120")

    def test_atomic_cancel_keeps_old_file(self):
        project = self.create()
        path = self.store._path(project["id"])
        before = path.read_bytes()
        from demo import projects
        atomic = projects._write_atomic
        def interrupted(path, payload, before_replace=None):
            def cancelled():
                raise cancel.RunCancelled("synthetic cancel immediately before replace")
            return atomic(path, payload, before_replace=cancelled)
        with patch.object(projects, "_write_atomic", interrupted), self.assertRaises(cancel.RunCancelled):
            self.store.confirm(project["id"], 1)
        self.assertEqual(path.read_bytes(), before)

    def test_revision_conflict_and_confirmation_invalidation(self):
        project = self.create()
        confirmed = self.store.confirm(project["id"], 1)
        self.assertTrue(confirmed["confirmed"])
        proposal = agent.propose_command(project["document"], "R00001 毛重改为 130 kg")
        changed = self.store.revise(project["id"], 1, proposal["document"])
        self.assertFalse(changed["confirmed"])
        with self.assertRaises(ScheduleConflict):
            self.store.revise(project["id"], 1, proposal["document"])
        undone = self.store.undo(project["id"], 2)
        self.assertEqual(undone["revision"], 3)
        self.assertEqual(undone["document"], project["document"])
        self.assertFalse(undone["confirmed"])
        self.assertEqual([item["revision"] for item in undone["versions"]], [1, 2])
        self.assertEqual(self.store.open(project["id"], 1)["document"], project["document"])
        self.assertEqual(self.store.open(project["id"], 2)["document"], changed["document"])
        self.assertFalse(undone["can_undo"])
        self.assertIsNone(undone["undo_target_revision"])
        before = self.store._path(project["id"]).read_bytes()
        with self.assertRaisesRegex(ValueError, "没有可撤销"):
            self.store.undo(project["id"], 3)
        self.assertEqual(self.store._path(project["id"]).read_bytes(), before)

    def test_sequential_undo_preserves_all_versions_and_never_toggles(self):
        first = self.create()
        ident = first["id"]
        second_doc = agent.propose_command(first["document"], "R00001 毛重改为 130kg")["document"]
        second = self.store.revise(ident, 1, second_doc)
        third_doc = agent.propose_command(second_doc, "R00001 毛重改为 140kg")["document"]
        self.store.revise(ident, 2, third_doc)
        self.store.confirm(ident, 3)
        fourth = self.store.undo(ident, 3)
        self.assertEqual(fourth["document"], second_doc)
        self.assertEqual(fourth["undo_target_revision"], 1)
        self.assertFalse(fourth["confirmed"])
        # Latest retained history is v3, but the next undo must go to v1.
        response = self.post(f"/projects/{ident}/conversation", {"expected_revision": 4, "message": "撤销上次修改"})
        self.assertEqual(response.status_code, 200, response.text)
        proposal = response.json()
        self.assertEqual(proposal["changes"], [{"field": "revision", "before": 4, "after": 1}])
        response = self.post(f"/proposals/{proposal['proposal_id']}/apply", {"expected_revision": 4})
        self.assertEqual(response.status_code, 200, response.text)
        fifth = response.json()["project"]
        self.assertEqual(fifth["revision"], 5)
        self.assertEqual(fifth["document"], first["document"])
        self.assertFalse(fifth["can_undo"])
        self.assertEqual([item["revision"] for item in fifth["versions"]], [1, 2, 3, 4])
        for revision, document in ((1, first["document"]), (2, second["document"]), (3, third_doc), (4, second_doc)):
            historical = self.store.open(ident, revision)
            self.assertEqual(historical["document"], document)
            self.assertFalse(historical["can_undo"])
            self.assertIsNone(historical["undo_target_revision"])
        reopened = LogisticsStore(self.root).open(ident)
        self.assertEqual(reopened, fifth)

    def test_edit_after_undo_branches_the_undo_stack_without_discarding_history(self):
        first = self.create()
        ident = first["id"]
        second = self.store.revise(ident, 1, agent.propose_command(first["document"], "R00001 毛重改为 130kg")["document"])
        third = self.store.revise(ident, 2, agent.propose_command(second["document"], "R00001 毛重改为 140kg")["document"])
        fourth = self.store.undo(ident, 3)
        fifth = self.store.revise(ident, 4, agent.propose_command(fourth["document"], "R00001 毛重改为 150kg")["document"])
        self.assertEqual(fifth["undo_target_revision"], 4)
        sixth = self.store.undo(ident, 5)
        self.assertEqual(sixth["document"], second["document"])
        self.assertEqual(sixth["undo_target_revision"], 1)
        seventh = self.store.undo(ident, 6)
        self.assertEqual(seventh["document"], first["document"])
        self.assertFalse(seventh["can_undo"])
        self.assertEqual(self.store.open(ident, 3)["document"], third["document"])
        self.assertEqual(self.store.open(ident, 5)["document"], fifth["document"])

    def test_undo_bundle_handoff_retains_audit_history_and_undo_position(self):
        first = self.create()
        ident = first["id"]
        second = self.store.revise(ident, 1, agent.propose_command(first["document"], "R00001 毛重改为 130kg")["document"])
        third = self.store.revise(ident, 2, agent.propose_command(second["document"], "R00001 毛重改为 140kg")["document"])
        fourth = self.store.undo(ident, 3)
        self.store.confirm(ident, 4)
        payload = bundle.export_bundle(self.store._read(ident))
        other_root = self.root / "coworker"
        other_root.mkdir()
        with patch.dict(os.environ, {"CIVIL_SANDBOX_ROOTS": str(other_root)}):
            other = LogisticsStore(other_root)
            imported = other.import_record(bundle.import_bundle(payload))
            self.assertNotEqual(imported["id"], ident)
            self.assertEqual(imported["revision"], 4)
            self.assertFalse(imported["confirmed"])
            self.assertEqual(imported["undo_target_revision"], 1)
            for revision, document in ((1, first["document"]), (2, second["document"]), (3, third["document"])):
                self.assertEqual(other.open(imported["id"], revision)["document"], document)
            reopened = LogisticsStore(other_root).open(imported["id"])
            self.assertEqual(reopened["document"], fourth["document"])
            undone = other.undo(imported["id"], 4)
            self.assertEqual(undone["document"], first["document"])
            self.assertFalse(undone["can_undo"])
            self.assertEqual(other.source(imported["id"])[0], SOURCE)

    def test_legacy_records_migrate_existing_history_without_rewriting_reads(self):
        first = self.create()
        ident = first["id"]
        second = self.store.revise(ident, 1, agent.propose_command(first["document"], "R00001 毛重改为 130kg")["document"])
        record = self.store._read(ident)
        record.pop("undo_stack")
        path = self.store._path(ident)
        path.write_text(json.dumps(dict(record, checksum=digest(record)), ensure_ascii=False), encoding="utf-8")
        raw = path.read_bytes()
        reopened = self.store.open(ident)
        self.assertEqual(reopened["undo_target_revision"], 1)
        self.assertEqual(path.read_bytes(), raw)
        self.assertEqual(self.store.open(ident, 1)["document"], first["document"])
        restored = self.store.undo(ident, 2)
        self.assertEqual(restored["document"], first["document"])
        self.assertEqual(self.store.open(ident, 2)["document"], second["document"])
        self.assertEqual(json.loads(path.read_text(encoding="utf-8"))["undo_stack"], [])
        # An old destructive undo may already have removed every snapshot.
        legacy = self.store._read(ident)
        legacy.pop("undo_stack")
        legacy["history"] = []
        normalized = validate_record(legacy)
        self.assertEqual(normalized["undo_stack"], [])
        imported = self.store.import_record(bundle.import_bundle(bundle.export_bundle(legacy)))
        self.assertFalse(imported["can_undo"])
        self.assertEqual(imported["versions"], [])

    def test_undo_stack_rejects_bad_types_order_duplicates_and_unknown_revisions(self):
        first = self.create()
        second = self.store.revise(first["id"], 1, agent.propose_command(first["document"], "R00001 毛重改为 130kg")["document"])
        self.store.revise(first["id"], 2, agent.propose_command(second["document"], "R00001 毛重改为 140kg")["document"])
        record = self.store._read(first["id"])
        for stack in (None, {}, [True], [1.0], ["1"], [0], [3], [999], [2, 1], [1, 1], [1] * (MAX_HISTORY + 1)):
            bad = deepcopy(record)
            bad["undo_stack"] = stack
            with self.assertRaises(ValueError, msg=str(stack)):
                validate_record(bad)
            with self.assertRaises(ValueError, msg=str(stack)):
                self.store.import_record(bad)
        self.assertEqual(len(self.store.list_projects()), 1)

    def test_bounded_history_prunes_undo_references_and_cancel_does_not_consume_undo(self):
        project = self.create()
        for index in range(MAX_HISTORY + 3):
            doc = agent.propose_command(project["document"], f"R00001 毛重改为 {130 + index}kg")["document"]
            project = self.store.revise(project["id"], project["revision"], doc)
        record = self.store._read(project["id"])
        self.assertEqual(len(record["history"]), MAX_HISTORY)
        self.assertEqual(record["undo_stack"], [item["revision"] for item in record["history"]])
        path = self.store._path(project["id"])
        before = path.read_bytes()
        with patch("demo.projects._write_atomic", side_effect=cancel.RunCancelled("cancel before commit")), self.assertRaises(cancel.RunCancelled):
            self.store.undo(project["id"], project["revision"])
        self.assertEqual(path.read_bytes(), before)
        target = project["undo_target_revision"]
        result = self.store.undo(project["id"], project["revision"])
        self.assertEqual(result["document"], self.store.open(project["id"], target)["document"])
        updated = self.store._read(project["id"])
        self.assertLessEqual(len(updated["history"]), MAX_HISTORY)
        self.assertTrue(set(updated["undo_stack"]) <= {item["revision"] for item in updated["history"]})

    def test_proposal_unknown_forged_and_units(self):
        proposal = agent.propose_command(self.doc, "把 R00001 长度改为 1.2 m；把 R00001 毛重改为 130000 g")
        self.assertEqual(proposal["document"]["rows"][0]["length_mm"], 1200)
        self.assertEqual(proposal["document"]["rows"][0]["gross_kg"], 130)
        self.assertEqual(self.doc["rows"][0]["length_mm"], 1000)
        bad = deepcopy(proposal)
        bad["document"]["rows"][0]["length_mm"] = 9999
        with self.assertRaises(ValueError):
            agent.apply_proposal(self.doc, bad)
        for text in ["R99999 毛重改为 100kg", "猜一下 R00001 长度", "R00001 长度改为 100 kg", "R00001 数量改为 2.5 件"]:
            with self.assertRaises(ValueError, msg=text):
                agent.propose_command(self.doc, text)

    def test_boxes_and_pieces_are_not_interchangeable(self):
        for text in ("R00001 件数改为 2箱", "R00001 箱数改为 20件"):
            with self.assertRaises(ValueError):
                agent.propose_command(self.doc, text)

    def test_model_tools_empty_args_no_write(self):
        project = self.create()
        context = api.project_context(project)
        before = self.store._path(project["id"]).read_bytes()
        self.assertFalse(agent.execute(context, "logistics_propose", {"file_path": "secret"}, "R00001 毛重改为 130kg")["ok"])
        result = agent.execute(context, "logistics_propose", {}, "R00001 毛重改为 130kg")
        self.assertTrue(result["ok"])
        self.assertIn("未修改", result["reply"])
        self.assertEqual(self.store._path(project["id"]).read_bytes(), before)
        self.assertFalse(agent.execute(context, "logistics_undo", {}, "检查")["ok"])

    def test_procurement_matching_never_guesses_names(self):
        other = deepcopy(self.doc)
        other["rows"][0]["quantity"] = 11
        comparison = agent.compare_documents(self.doc, other)
        self.assertEqual(comparison["rows"][0]["differences"]["quantity"], {"current": 10, "comparison": 11})
        other["rows"][0]["material_id"] = "UNSPECIFIED"
        comparison = agent.compare_documents(self.doc, other)
        self.assertEqual(comparison["rows"][0]["status"], "only_current")
        self.assertEqual(comparison["unmatched_comparison"], ["R00001"])

    def test_inspect_catalog_and_human_summary_are_bounded(self):
        project = self.create()
        context = api.project_context(project)
        inspected = agent.execute(context, "logistics_inspect", {}, "检查台账")
        self.assertEqual(inspected["row_catalog"][0]["id"], "R00001")
        self.assertEqual(inspected["row_catalog"][0]["known_values"]["gross_kg"], 120)
        summary = agent.execute(context, "logistics_summarize", {}, "汇总")
        self.assertIn("包装数：1", summary["reply"])
        self.assertIn("毛重：120 kg", summary["reply"])
        self.assertNotIn('"source":', summary["reply"])
        self.assertNotIn("evidence", inspected["row_catalog"][0])
        missing_id = agent.execute(context, "logistics_propose", {}, "把毛重改为120kg")
        self.assertFalse(missing_id["ok"])
        self.assertIn("行号", missing_id["reply"])

    def test_zip_roundtrip_raw_history_and_cleared_confirmation(self):
        project = self.create()
        proposal = agent.propose_command(project["document"], "R00001 毛重改为 130 kg")
        self.store.revise(project["id"], 1, proposal["document"])
        self.store.confirm(project["id"], 2)
        data = bundle.export_bundle(self.store._read(project["id"]))
        imported = self.store.import_record(bundle.import_bundle(data))
        self.assertNotEqual(project["id"], imported["id"])
        self.assertEqual(imported["revision"], 2)
        self.assertFalse(imported["confirmed"])
        self.assertEqual(self.store.source(imported["id"])[0], SOURCE)
        self.assertEqual(self.store.open(imported["id"], 1)["document"], project["document"])

    def test_zip_path_checksum_duplicate_and_broken_rejected(self):
        project = self.create()
        archive = bundle.export_bundle(self.store._read(project["id"]))
        with zipfile.ZipFile(io.BytesIO(archive)) as z:
            files = {i.filename: z.read(i) for i in z.infolist()}
        for kind in ("path", "checksum", "duplicate"):
            out = io.BytesIO()
            with zipfile.ZipFile(out, "w") as z:
                for name, content in files.items():
                    if kind == "path" and name.startswith("sources/"):
                        name = "../escaped.bin"
                    if kind == "checksum" and name.startswith("sources/"):
                        content += b"tampered"
                    z.writestr(name, content)
                if kind == "duplicate":
                    z.writestr("manifest.json", files["manifest.json"])
            with self.assertRaises(ValueError, msg=kind):
                bundle.import_bundle(out.getvalue())
        with self.assertRaises(ValueError):
            bundle.import_bundle(b"bad zip")
        self.assertEqual(len(self.store.list_projects()), 1)

    def test_http_corrupt_deflate_members_rejected_without_changing_projects(self):
        project = self.create()
        original = self.store._path(project["id"]).read_bytes()
        archive = bundle.export_bundle(self.store._read(project["id"]))
        with zipfile.ZipFile(io.BytesIO(archive)) as z:
            offsets = {item.filename: item.header_offset for item in z.infolist()}
        for name, offset in offsets.items():
            with self.subTest(member=name):
                damaged = bytearray(archive)
                name_len, extra_len = struct.unpack_from("<HH", damaged, offset + 26)
                # Invalid DEFLATE block type; central directory remains valid.
                damaged[offset + 30 + name_len + extra_len] = 7
                response = self.client.post(api.BASE + "/import", files={
                    "file": ("broken.zip", bytes(damaged), "application/zip"),
                })
                self.assertEqual(response.status_code, 422, response.text)
                self.assertIn("项目包损坏", response.json()["detail"])
                self.assertEqual(self.store._path(project["id"]).read_bytes(), original)
                self.assertEqual(len(self.store.list_projects()), 1)

    def test_saved_checksum_and_source_hash_rejected(self):
        project = self.create()
        path = self.store._path(project["id"])
        record = json.loads(path.read_text(encoding="utf-8"))
        record["document"]["rows"][0]["quantity"] = 999
        path.write_text(json.dumps(record), encoding="utf-8")
        with self.assertRaises(ValueError):
            self.store.open(project["id"])
        with self.assertRaises(ValueError):
            self.store.create("错误原件", self.doc, b"different")

    def test_xlsx_keeps_literal_formula_text_and_sources(self):
        from openpyxl import load_workbook
        project = self.create()
        project["document"]["rows"][0]["name"] = "=1+1"
        data, mime, filename = bundle.export_ledger(project, "xlsx")
        wb = load_workbook(io.BytesIO(data), data_only=False)
        columns = {cell.value: cell.column for cell in wb["台账"][1]}
        name_cell = wb["台账"].cell(2, columns["name"])
        self.assertEqual(name_cell.value, "=1+1")
        self.assertEqual(name_cell.data_type, "s")
        self.assertIn("来源与修订", wb.sheetnames)

    def test_packaged_uses_loader_only_and_preserves_box_count(self):
        project = self.create()
        with self.assertRaises(ValueError):
            packing.pack(project, "packaged", "40HQ", 2)
        project = self.store.confirm(project["id"], 1)
        plan = {"can_fit": True, "containers_used": 1, "engine": "offline-scripted-3d", "layout": [{"box_id": "R00001-1", "container_no": 1, "position": {"x": 0, "y": 0, "z": 0}, "size": {"dx": 1000, "dy": 800, "dz": 600}}]}
        with patch("packing_assistant.agents.loader.agent_loader", return_value={"container_plan": plan}) as loader, patch("packing_assistant.agents.box_scheme.agent_box_scheme", side_effect=AssertionError("must not repack")):
            result = packing.pack(project, "packaged", "40HQ", 2)
        self.assertTrue(result["ok"])
        self.assertEqual(result["n_boxes"], 1)
        self.assertEqual(loader.call_args.args[0]["boxes"][0]["gross_weight_kg"], 120)
        self.assertFalse(loader.call_args.args[0]["boxes"][0]["allowRotate"])

    def test_missing_scope_quantity_and_repacking_block(self):
        project = self.create()
        project["confirmed"] = True
        for field in ("package_count", "quantity", "gross_kg", "length_mm", "dimension_scope", "weight_scope"):
            bad = deepcopy(project)
            bad["document"]["rows"][0][field] = "UNSPECIFIED"
            with patch("packing_assistant.agents.loader.agent_loader", side_effect=AssertionError("unknown input reached solver")):
                self.assertFalse(packing.pack(bad, "packaged", "40HQ", 2)["ok"], field)
        self.assertFalse(packing.pack(project, "materials", "40HQ", 2)["ok"])

    def test_solver_can_fit_false_and_fallback_not_success(self):
        project = self.create()
        project["confirmed"] = True
        for plan in ({"can_fit": False, "containers_used": 1, "engine": "3d"}, {"can_fit": True, "containers_used": 3, "engine": "3d"}, {"can_fit": True, "containers_used": 1, "engine": "local-1d-fallback"}):
            with patch("packing_assistant.agents.loader.agent_loader", return_value={"container_plan": plan}):
                self.assertFalse(packing.pack(project, "packaged", "40HQ", 2)["ok"])

    def test_real_worker_geometry_no_rotation_or_stack(self):
        project = self.create()
        project["confirmed"] = True
        result = packing.run_pack(project, "packaged", "40HQ", 2)
        self.assertTrue(result["ok"], result)
        self.assertTrue(result["layout_verified"])
        self.assertEqual(result["container_plan"]["layout"][0]["size"], {"dx": 1000, "dy": 800, "dz": 600})
        self.assertEqual(result["container_plan"]["layout"][0]["position"]["z"], 0)
        # This fits a 20GP only if L/W are swapped. Rotation is deliberately unavailable.
        row = project["document"]["rows"][0]
        row.update(length_mm=1500, width_mm=3000, height_mm=500)
        self.assertFalse(packing.run_pack(project, "packaged", "20GP", 1)["ok"])

    def test_layout_verification_rejects_fake_success(self):
        project = self.create()
        project["confirmed"] = True
        boxes = packing.prepare(project, "packaged", "40HQ", 2)["boxes"]
        item = {"box_id": "R00001-1", "container_no": 1, "position": {"x": 0, "y": 0, "z": 0}, "size": {"dx": 1000, "dy": 800, "dz": 600}}
        self.assertTrue(packing.verify_packaged_layout(boxes, {"layout": [item], "containers_used": 1}, "40HQ", 2))
        for variant in (dict(item, size={"dx": 800, "dy": 1000, "dz": 600}), dict(item, position={"x": 0, "y": 0, "z": 600}), dict(item, box_id="unknown")):
            self.assertFalse(packing.verify_packaged_layout(boxes, {"layout": [variant], "containers_used": 1}, "40HQ", 2))

    def test_http_upload_preview_save_conversation_confirm_undo(self):
        response = self.client.post(api.BASE + "/upload", files={"file": ("synthetic.csv", SOURCE, "text/csv")})
        self.assertEqual(response.status_code, 200, response.text)
        upload = response.json()
        self.assertEqual(self.store.list_projects(), [])
        created = self.post("/projects", {"document_id": upload["document_id"], "name": "HTTP合成"})
        self.assertEqual(created.status_code, 200, created.text)
        project = created.json()["project"]
        ident = project["id"]
        suggestion = self.post(f"/projects/{ident}/conversation", {"expected_revision": 1, "message": "R00001 毛重改为 130kg"})
        self.assertEqual(suggestion.status_code, 200, suggestion.text)
        self.assertEqual(self.store.open(ident)["revision"], 1)
        proposal = suggestion.json()
        applied = self.post(f"/proposals/{proposal['proposal_id']}/apply", {"expected_revision": 1})
        self.assertEqual(applied.status_code, 200, applied.text)
        self.assertEqual(self.post(f"/proposals/{proposal['proposal_id']}/apply", {"expected_revision": 1}).status_code, 409)
        self.assertTrue(self.post(f"/projects/{ident}/confirm", {"expected_revision": 2}).json()["project"]["confirmed"])
        self.assertFalse(self.post(f"/projects/{ident}/undo", {"expected_revision": 2}).json()["project"]["confirmed"])
        self.assertEqual(self.client.get(api.BASE + f"/projects/{ident}/source").content, SOURCE)

    def test_http_permission_extra_fields_and_local_only(self):
        project = self.create()
        path = f"/projects/{project['id']}/export"
        self.assertEqual(self.post(path, {"expected_revision": 1, "format": "json", "confirmation": "yes"}).status_code, 403)
        self.assertEqual(self.post(path, {"expected_revision": 1, "format": "json", "confirmation": api.cad.CONFIRMATION, "file_path": "secret"}).status_code, 422)
        response = self.client.get(api.BASE + "/projects", headers={"Origin": "https://evil.invalid"})
        self.assertEqual(response.status_code, 403)
        remote = TestClient(self.client.app, base_url="http://remote.invalid", client=("192.0.2.1", 12345))
        self.addCleanup(remote.close)
        self.assertEqual(remote.get(api.BASE + "/projects").status_code, 403)

    def test_http_precancelled_upload_never_publishes(self):
        ident = "b" * 32
        self.client.post("/api/engineering/operations/" + ident + "/cancel")
        before = self.store.list_projects()
        response = self.client.post(api.BASE + "/upload", headers={"X-CAD-Operation-ID": ident}, files={"file": ("synthetic.csv", SOURCE)})
        self.assertEqual(response.status_code, 499, response.text)
        self.assertEqual(before, self.store.list_projects())

    def test_http_exchange_json_xlsx_and_complete_bundle(self):
        project = self.create()
        ident = project["id"]
        for format in ("json", "xlsx", "zip"):
            response = self.post(f"/projects/{ident}/export", {"expected_revision": 1, "format": format, "confirmation": api.cad.CONFIRMATION})
            self.assertEqual(response.status_code, 200, response.text[:300] if format == "json" else response.status_code)
            endpoint = "/import" if format == "zip" else "/upload"
            imported = self.client.post(api.BASE + endpoint, files={"file": ("export." + format, response.content)})
            self.assertEqual(imported.status_code, 200, imported.text)
            if format == "zip":
                new = imported.json()["project"]
                self.assertNotEqual(new["id"], ident)
                self.assertFalse(new["confirmed"])
                self.assertEqual(self.store.source(new["id"])[0], SOURCE)
            else:
                document = imported.json()["document"]
                self.assertEqual(document["rows"][0]["quantity"], 10)
                self.assertEqual(document["source"]["sha256"], hashlib.sha256(response.content).hexdigest())

    def test_example_label_survives_save(self):
        response = self.client.get(api.BASE + "/example")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["synthetic"])
        project = self.post("/projects", {"document_id": response.json()["document_id"], "name": "合成"}).json()["project"]
        self.assertTrue(self.store.open(project["id"])["document"]["extraction"]["synthetic"])

    def test_cancel_after_atomic_commit_returns_committed_record(self):
        project = self.create()
        from demo import projects
        from threading import Event
        event = Event()
        atomic = projects._write_atomic
        def commit_then_cancel(*args, **kwargs):
            atomic(*args, **kwargs)
            event.set()
        with cancel.scope(event=event), patch.object(projects, "_write_atomic", commit_then_cancel):
            confirmed = self.store.confirm(project["id"], 1)
        self.assertTrue(confirmed["confirmed"])
        self.assertTrue(self.store.open(project["id"])["confirmed"])

    def test_real_worker_timeout_kills_process(self):
        project = self.create()
        project["confirmed"] = True
        with self.assertRaisesRegex(ValueError, "已停止本次进程"):
            packing.run_pack(project, "packaged", "40HQ", 2, timeout=0)


if __name__ == "__main__":
    unittest.main()
