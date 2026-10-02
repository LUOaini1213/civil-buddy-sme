"""Bounded HTTP acceptance for the unified Rust workbench.

Default: test a server configured with model civil-scripted-acceptance.
--live: require the explicitly expected model and provider host in server capabilities.
--token-file: read only a local workbench login token, never a provider key.
--serve-model: run the local deterministic OpenAI-compatible test endpoint.

For a scripted server configure CIVIL_MODEL=civil-scripted-acceptance,
CIVIL_API_BASE=http://127.0.0.1:18889/v1 and a dummy CIVIL_API_KEY in that
server's environment. This script never changes a running server configuration.
"""
from __future__ import annotations

import argparse
from hashlib import sha256
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
import json
import math
from pathlib import Path
import re
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
MODEL = "civil-scripted-acceptance"
TERMINAL = {"completed", "failed", "cancelled", "interrupted"}
# Source-level defaults, not a claim about an arbitrary deployed binary. Prefer
# the actual turn's budget snapshot when one is available.
SERVER_LIMITS = {"total_tokens": 160_000, "task_tokens": 140_000, "max_model_calls": 20,
                 "max_tasks": 5, "max_depth": 1, "timeout_ms": 300_000}
FILES = ["requirements.pdf", "quantities.xlsx", "report.docx"]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
TASK = ("根据 requirements.pdf 中 sample_count 的明确要求，检查 quantities.xlsx 的 Counts!B2 "
        "和 report.docx 中 sample_count 段落，把这两处同时修改为 PDF 要求的数量。"
        "请实际读取三个原件并检索 PDF 要求；每份补丁携带 PDF 的完整原文引用、SHA256和定位，"
        "先预览再保存新副本。保持其他段落、表格、样式、页眉和未修改工作表不变，不覆盖原件。"
        "可委派只读资料子代理核对要求；最后列出两个真实新文件、来源及尚未完成的重算或渲染。")


def package(data):
    with ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def fixture(folder: Path, required=6):
    """Create real synthetic files. The prompt intentionally does not state the required number."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    folder = folder.resolve()
    folder.mkdir(parents=True, exist_ok=True)
    if any((folder / name).exists() for name in FILES):
        raise ValueError("Acceptance fixtures already exist; use a fresh directory")
    if type(required) is not int or not 1 <= required <= 999:
        raise ValueError("Fixture requirement must be a bounded positive integer")
    writer = PdfWriter()
    page = writer.add_blank_page(width=420, height=595)
    font = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 35 530 Td (Synthetic acceptance requirement.) Tj 0 -24 Td (sample_count={required}) Tj 0 -24 Td (Use this value in the specimen register and report.) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    with (folder / FILES[0]).open("wb") as output:
        writer.write(output)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Counts"
    sheet.append(["Metric", "Value", "Comment", "Derived"])
    sheet.append(["sample_count", 4, "Keep this note", "=B2*2"])
    sheet.append(["untouched_count", 99, "Unchanged row"])
    sheet["B2"].font = Font(bold=True, color="003366")
    sheet["B2"].number_format = "0"
    workbook.create_sheet("Unchanged")["A1"] = "Keep this sheet"
    workbook.save(folder / FILES[1])
    workbook.close()
    parts = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/><Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/document.xml": f'<w:document xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body><w:p><w:pPr><w:pStyle w:val="Title"/></w:pPr><w:r><w:t>Specimen report</w:t></w:r></w:p><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>sample_count=4</w:t></w:r></w:p><w:p><w:r><w:t>Unchanged scope paragraph.</w:t></w:r></w:p><w:sectPr><w:headerReference w:type="default" r:id="rIdHeader"/><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>',
        "word/styles.xml": f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="24"/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/></w:style></w:styles>',
        "word/header1.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>Unchanged project header</w:t></w:r></w:p></w:hdr>',
        "word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/><Relationship Id="rIdHeader" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/></Relationships>',
    }
    with ZipFile(folder / FILES[2], "w") as archive:
        for name, value in parts.items():
            archive.writestr(name, '<?xml version="1.0" encoding="UTF-8"?>'+value)
    return {"workspace": str(folder), "files": FILES, "required_sample_count": required,
            "initial_sample_count": 4, "hashes": {name: sha256((folder/name).read_bytes()).hexdigest() for name in FILES}}


def _unwrap(value):
    if isinstance(value, dict) and value.get("ok") is False:
        raise ValueError("Tool failed: "+str(value.get("error")))
    while isinstance(value, dict) and isinstance(value.get("result"), dict):
        value = value["result"]
    return value


def _tool_history(messages):
    calls, records = {}, []
    for message in messages:
        for call in message.get("tool_calls") or []:
            function = call.get("function", {})
            calls[call["id"]] = (function.get("name"), json.loads(function.get("arguments", "{}")))
        if message.get("role") == "tool":
            name, arguments = calls.get(message.get("tool_call_id"), (None, {}))
            try:
                result = json.loads(message.get("content") or "{}")
            except ValueError:
                result = {"ok": False, "error": "Tool returned non-JSON content"}
            records.append({"name": name, "arguments": arguments, "result": result})
    return records


def _calls(items):
    return {"role": "assistant", "content": None, "tool_calls": [
        {"id": "acceptance_"+uuid4().hex[:12], "type": "function", "function": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}}
        for name, args in items]}


def scripted_message(payload):
    """State-free fixture policy: derive edits from real tool outputs, never touch files."""
    messages = payload.get("messages", [])
    records = _tool_history(messages)
    is_child = any(str(m.get("content", "")).startswith("你是只读") and "子代理" in str(m.get("content", "")) for m in messages if m.get("role") == "system")
    if is_child:
        if not records:
            return _calls([("search_sources", {"query": "sample_count"})])
        result = _unwrap(records[-1]["result"])
        return {"role": "assistant", "content": "已读取定位资料；以下是原文检索结果（仍需主代理核对）："+json.dumps(result, ensure_ascii=False)}
    available = {t.get("function", {}).get("name") for t in payload.get("tools", [])}
    if not records:
        first = [("load_skill", {"skill_id": id}) for id in ("doc-pdf", "doc-spreadsheet", "doc-word")]
        first += [("read_file", {"source": FILES[0], "operation": "read", "arguments": {"pages": [1]}}),
                  ("read_file", {"source": FILES[1], "operation": "inspect"}),
                  ("read_file", {"source": FILES[2], "operation": "inspect"}),
                  ("search_sources", {"query": "sample_count"})]
        if "delegate" in available:
            first.append(("delegate", {"tasks": [{"role": "evidence", "goal": "只读检索requirements.pdf中的sample_count要求，给出完整source/hash/locator/quote，不改文件。"}]}))
        return _calls(first)
    for record in records:
        if record["result"].get("ok") is False:
            return {"role": "assistant", "content": "验收工具链有失败，未声称任务完成："+json.dumps(record["result"], ensure_ascii=False)}
    reads = [r for r in records if r["name"] == "read_file"]
    def read(source, operation):
        return next((_unwrap(r["result"]) for r in reversed(reads) if r["arguments"].get("source") == source and r["arguments"].get("operation", "read") == operation), None)
    word_inspect, sheet_inspect = read(FILES[2], "inspect"), read(FILES[1], "inspect")
    pdf_read = read(FILES[0], "read")
    if not all((word_inspect, sheet_inspect, pdf_read)):
        raise ValueError("Scripted provider needs actual inspect/read tool results")
    pdf_text = "\n".join(p["text"] for p in pdf_read["pages"])
    match = re.search(r"sample_count\s*=\s*(\d+)", pdf_text)
    if not match:
        raise ValueError("No source requirement was read from PDF")
    required = int(match.group(1))
    target_node = next(n for n in word_inspect["nodes"] if re.search(r"sample_count\s*=\s*\d+", n["text"]))
    sheet = next(s["name"] for s in sheet_inspect["sheets"] if s["name"] == "Counts")
    if read(FILES[2], "read") is None or read(FILES[1], "read") is None:
        return _calls([("read_file", {"source": FILES[2], "operation": "read", "arguments": {"block_ids": [target_node["id"]]}}),
                       ("read_file", {"source": FILES[1], "operation": "read", "arguments": {"sheet": sheet, "range": "A1:D3"}})])
    word_read, sheet_read = read(FILES[2], "read"), read(FILES[1], "read")
    old_text = next(b["text"] for b in word_read["blocks"] if b["id"] == target_node["id"])
    old_cell = next(c for row in sheet_read["rows"] for c in row if c["cell"] == "B2")
    search = next(_unwrap(r["result"]) for r in reversed(records) if r["name"] == "search_sources")
    evidence = next(h for h in search["hits"] if h.get("source") == FILES[0] and re.search(r"sample_count\s*=\s*"+str(required)+r"\b", h.get("quote", "")))
    word_patch = {"source": FILES[2], "expected_sha256": word_read["source_sha256"], "patches": [
        {"op": "replace_paragraph", "paragraph_id": target_node["id"], "expected_text": old_text,
         "text": re.sub(r"(sample_count\s*=\s*)\d+", lambda m: m.group(1)+str(required), old_text), "evidence": [evidence]}]}
    sheet_patch = {"source": FILES[1], "expected_sha256": sheet_read["source_sha256"], "patches": [
        {"op": "set_cell", "sheet": sheet, "cell": "B2", "expected": {"type": old_cell["type"], "value": old_cell["value"]},
         "value": {"type": "number", "value": required}, "evidence": [evidence]}]}
    plans = [word_patch, sheet_patch]
    previewed = {r["arguments"]["source"] for r in records if r["name"] == "preview_document" and r["result"].get("ok") is True}
    if any(p["source"] not in previewed for p in plans):
        return _calls([("preview_document", p) for p in plans])
    applied = {r["arguments"]["source"] for r in records if r["name"] == "apply_document" and r["result"].get("ok") is True}
    if any(p["source"] not in applied for p in plans):
        return _calls([("apply_document", p) for p in plans if p["source"] not in applied])
    return {"role": "assistant", "content": f"依据PDF原文，已将Word段落与Excel Counts!B2中的sample_count改为{required}，保存两个新副本，原件保留。修改有原件hash和PDF页引用；未进行Excel重算或视觉渲染，不能将结构回读视为工程签认。"}


def model_server(port=18889):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_):
            pass

        def do_GET(self):
            body = json.dumps({"ok": True, "model": MODEL, "kind": "scripted_fixture_provider"}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_POST(self):
            if self.path not in ("/v1/chat/completions", "/chat/completions"):
                self.send_error(404)
                return
            length = int(self.headers.get("Content-Length", "0"))
            if not 0 < length <= 4_000_000:
                self.send_error(413)
                return
            try:
                payload = json.loads(self.rfile.read(length))
                message = scripted_message(payload)
                response = {"id": "scripted-"+uuid4().hex, "object": "chat.completion", "model": MODEL,
                            "choices": [{"index": 0, "message": message, "finish_reason": "tool_calls" if message.get("tool_calls") else "stop"}],
                            "usage": {"prompt_tokens": max(1, length//4), "completion_tokens": max(1, len(json.dumps(message))//4)}}
                body = json.dumps(response, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
            except Exception as exc:
                body = json.dumps({"error": {"type": "scripted_fixture_error", "message": str(exc)}}, ensure_ascii=False).encode()
                self.send_response(422)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
    return ThreadingHTTPServer(("127.0.0.1", port), Handler)


def _product_base(base):
    parsed = urllib.parse.urlsplit(base)
    if (parsed.scheme != "http" or parsed.hostname not in ("127.0.0.1", "localhost", "::1")
            or parsed.username is not None or parsed.password is not None
            or parsed.query or parsed.fragment or parsed.path not in ("", "/")):
        raise ValueError("Acceptance only connects to a loopback HTTP product origin without credentials, path or query")
    return base.rstrip("/")


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        # Even a same-origin redirect can lead to another redirect. Never send
        # the named login credential anywhere except the explicitly chosen URL.
        raise ValueError("HTTP redirect refused by acceptance client")


def http(base, path, body=None, *, binary=False, timeout=30, token=None):
    base = _product_base(base)
    target = urllib.parse.urlsplit(path)
    if (not path.startswith("/") or path.startswith("//") or "\\" in path
            or target.scheme or target.netloc or target.fragment):
        raise ValueError("Acceptance HTTP paths must remain on the selected product origin")
    if token and _redact(path, token) != path:
        raise ValueError("A login credential must not appear in a request URL")
    raw = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = "Bearer "+token
    request = urllib.request.Request(base+path, data=raw, headers=headers)
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}), _NoRedirect())
    try:
        with opener.open(request, timeout=timeout) as response:
            data = response.read(50_000_001)
            if len(data) > 50_000_000:
                raise ValueError("HTTP response exceeds acceptance limit")
            return data if binary else json.loads(data)
    except urllib.error.HTTPError as exc:
        # Do not include an untrusted reflected reason, URL or response body.
        raise ValueError(f"Product HTTP request failed with status {exc.code}") from None
    except urllib.error.URLError:
        raise ConnectionError("Product HTTP request could not be completed") from None


def _login_token(path, workspace, report_dir):
    if path is None:
        return None
    path = Path(path).resolve()
    if any(path.is_relative_to(folder.resolve()) for folder in (workspace, report_dir)):
        raise ValueError("Keep the workbench login token file outside the fixture workspace and reports")
    with path.open("rb") as stream:
        raw = stream.read(4097)
    if len(raw) > 4096:
        raise ValueError("Workbench login token file is too large")
    token = raw.decode("utf-8-sig").strip()
    if not 32 <= len(token) <= 4096 or not all(33 <= ord(c) <= 126 for c in token):
        raise ValueError("Workbench login token must contain at least 32 visible ASCII characters on one line")
    return token


def _redact(text, token):
    if token:
        for value in (token, urllib.parse.quote(token, safe=""), urllib.parse.quote_plus(token),
                      json.dumps(token)[1:-1]):
            text = text.replace(value, "[REDACTED_LOGIN_TOKEN]")
    return text


def _write_json(path, value, token=None):
    text = _redact(json.dumps(value, ensure_ascii=False, indent=2), token)
    path.write_text(text+"\n", encoding="utf-8")
    return json.loads(text)


def original_hashes(workspace, meta):
    checks = {}
    for name, expected in meta["hashes"].items():
        try:
            actual = sha256((workspace/name).read_bytes()).hexdigest()
            checks[name] = {"expected_sha256": expected, "actual_sha256": actual,
                            "status": "unchanged" if actual == expected else "changed"}
        except OSError:
            checks[name] = {"expected_sha256": expected, "actual_sha256": None, "status": "unavailable"}
    return {"checked": True, "unchanged": all(c["status"] == "unchanged" for c in checks.values()), "files": checks}


def _host(value):
    if not isinstance(value, str) or not value or any(c.isspace() for c in value):
        raise ValueError("Expected provider host must be a hostname, without a URL, port or credentials")
    parsed = urllib.parse.urlsplit("https://"+value)
    if parsed.hostname != value.lower() or parsed.port is not None or parsed.username or parsed.path or parsed.query or parsed.fragment:
        raise ValueError("Expected provider host must be a hostname, without a URL, port or credentials")
    return value.lower()


def _scripted(model):
    return isinstance(model, str) and any(word in model.lower() for word in ("scripted", "fixture", "mock", "dummy"))


def model_verification(capabilities, live, expected_model, expected_provider_host):
    models = capabilities.get("models", {})
    actual = models.get("model")
    observed_host = models.get("provider_host")
    if not isinstance(observed_host, str) or not observed_host:
        observed_host = None
    matched = observed_host and (observed_host.lower() == expected_provider_host if live else
                                 observed_host.lower() in ("127.0.0.1", "localhost", "::1"))
    host_status = "unverified" if not observed_host else ("matched" if matched else "mismatch")
    return {"configured": models.get("configured") is True, "requested_mode": "live" if live else "scripted",
            "expected_model": expected_model if live else MODEL, "configured_model": actual,
            "model_status": "matched" if actual == (expected_model if live else MODEL) else "mismatch",
            "scripted_model_detected": _scripted(actual), "expected_provider_host": expected_provider_host,
            "configured_provider_host": observed_host, "provider_host_status": host_status,
            "provider_host_source": "capabilities.models.provider_host" if observed_host else "not_exposed_by_server"}


def usage_report(turn, events):
    result = turn.get("result") if isinstance(turn, dict) else None
    budget = result.get("usage") if isinstance(result, dict) else None
    receipts = [e["data"] for e in events if isinstance(e, dict) and isinstance(e.get("data"), dict)
                and e.get("kind", e.get("event")) == "model"]
    settlements = budget.get("settlements") if isinstance(budget, dict) else None
    records = settlements if isinstance(settlements, list) else [r.get("usage", {}) for r in receipts]
    totals = {kind: {"input_tokens": 0, "output_tokens": 0, "model_calls": 0}
              for kind in ("provider_reported", "estimated", "unknown")}
    for record in records:
        if not isinstance(record, dict):
            continue
        kind = "provider_reported" if record.get("estimated") is False else ("estimated" if record.get("estimated") is True else "unknown")
        for key in totals[kind]:
            value = record.get(key)
            if type(value) is int and value >= 0:
                totals[kind][key] += value
    return {"source": "server_budget_settlements" if isinstance(settlements, list) else "model_events_partial",
            "budget_snapshot": budget, "model_calls": budget.get("model_calls") if isinstance(budget, dict) else None,
            "model_events_observed": len(receipts), "response_models": sorted({r["model"] for r in receipts if isinstance(r.get("model"), str)}),
            "usage_by_measurement": totals, "provider_reported_is_billing_verified": False,
            "currency_cost": None, "currency_budget_enforced_by_script": False}


def validate_outputs(workspace, meta, outputs):
    """Independent reopening: no calls to document worker or its validators."""
    from openpyxl import load_workbook
    required = meta["required_sample_count"]
    checks, details = {}, {}
    checks["originals_unchanged"] = all(sha256((workspace/name).read_bytes()).hexdigest() == sha for name, sha in meta["hashes"].items())
    for kind in ("docx", "xlsx"):
        candidates = [o for o in outputs if Path(o["name"]).suffix.lower() == "."+kind]
        checks[kind+"_artifact_exists"] = bool(candidates)
        valid = []
        for candidate in candidates:
            data = Path(candidate["download_path"]).read_bytes()
            expected_hash = candidate.get("output_sha256")
            hash_ok = bool(expected_hash) and sha256(data).hexdigest() == expected_hash
            before = package((workspace/("report.docx" if kind == "docx" else "quantities.xlsx")).read_bytes())
            after = package(data)
            allowed = {"word/document.xml"} if kind == "docx" else {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}
            preserved = before.keys() == after.keys() and all(before[k] == after[k] for k in before if k not in allowed)
            if kind == "docx":
                root = ET.fromstring(after["word/document.xml"])
                paras = ["".join(n.text or "" for n in p.iter(f"{{{W}}}t")) for p in root.iter(f"{{{W}}}p")]
                exact = "sample_count="+str(required) in paras
                preserved = preserved and "Unchanged scope paragraph." in paras and "Specimen report" in paras
                target = next((p for p in root.iter(f"{{{W}}}p") if "".join(t.text or "" for t in p.iter(f"{{{W}}}t")) == "sample_count="+str(required)), None)
                style = target is not None and target.find(f"{{{W}}}r/{{{W}}}rPr/{{{W}}}b") is not None
            else:
                book = load_workbook(BytesIO(data), data_only=False)
                exact = book["Counts"]["B2"].value == required
                preserved = preserved and book["Counts"]["B3"].value == 99 and book["Counts"]["D2"].value == "=B2*2" and book["Unchanged"]["A1"].value == "Keep this sheet"
                style = book["Counts"]["B2"].font.bold and book["Counts"]["B2"].number_format == "0"
                book.close()
            item = {"artifact": candidate["name"], "required_value": required, "actual_value_matches": exact,
                    "unmodified_parts_preserved": preserved, "target_style_preserved": bool(style), "registered_hash_matches": hash_ok}
            valid.append(all((exact, preserved, style, hash_ok)))
            details.setdefault(kind, []).append(item)
        checks[kind+"_changed_and_preserved"] = any(valid)
    return {"passed": all(checks.values()), "checks": checks, "details": details,
            "rendering": "not_performed", "formula_recalculation": "not_performed", "engineering_approval": "not_evaluated"}


def acceptance(base, workspace, report_dir, *, live=False, timeout=360, token_file=None,
               expected_model=None, expected_provider_host=None, cancel_timeout=15):
    base = _product_base(base)
    workspace, report_dir = Path(workspace).resolve(), Path(report_dir).resolve()
    if not math.isfinite(timeout) or not 0 < timeout <= 3600:
        raise ValueError("Acceptance wait timeout must be positive and at most 3600 seconds")
    if not math.isfinite(cancel_timeout) or not 0 < cancel_timeout <= 120:
        raise ValueError("Cancellation wait must be positive and at most 120 seconds")
    report_dir.mkdir(parents=True, exist_ok=True)
    report = {"mode": "live" if live else "scripted", "passed": False, "live_verified": False,
              "events_file": str(report_dir/"events.jsonl"),
              "authentication": "workbench_login_token_file" if token_file else "none",
              "wait_limits": {"task_wait_seconds": timeout, "cancel_wait_seconds": cancel_timeout,
                              "scope": "client observation and cancellation grace; not a server or currency budget"},
              "server_limits": {"values": SERVER_LIMITS, "source": "repository_source_defaults",
                                "deployed_limits_verified": False, "includes_subagents": True},
              "originals": {"checked": False, "unchanged": None}}
    events, token, meta, turn = [], None, None, None
    after, turn_id, workspace_id, session_id = 0, None, None, None
    start_attempted = False

    def request(path, body=None, **kwargs):
        return http(base, path, body, token=token, **kwargs)

    def poll(deadline):
        nonlocal after, turn
        remaining = deadline-time.monotonic()
        if remaining <= 0:
            raise TimeoutError("Acceptance observation deadline elapsed")
        query = urllib.parse.urlencode({"workspace": workspace_id, "session_id": session_id, "after_seq": after})
        snapshot = request(f"/api/agent/turns/{turn_id}/events?{query}", timeout=min(30, remaining))
        if not isinstance(snapshot, dict) or not isinstance(snapshot.get("turn"), dict) or not isinstance(snapshot.get("events"), list):
            raise ValueError("Invalid product event snapshot")
        for event in snapshot.get("events", []):
            if not isinstance(event, dict) or type(event.get("seq")) is not int or event["seq"] <= after:
                raise ValueError("Product events are not strictly increasing")
            events.append(event)
            after = event["seq"]
        turn = snapshot["turn"]
        report["turn"] = turn
        return turn.get("status") in TERMINAL and after >= turn.get("last_seq", after)

    def cancel_and_drain():
        # A successful cancel POST only acknowledges the request, not shutdown.
        deadline = time.monotonic()+cancel_timeout
        state = {"requested": True, "acknowledged": False, "terminal_observed": False,
                 "last_observed_status": turn.get("status") if turn else "unknown", "still_running": None}
        report["cancellation"] = state
        query = urllib.parse.urlencode({"workspace": workspace_id, "session_id": session_id})
        try:
            response = request(f"/api/agent/turns/{turn_id}/cancel?{query}", {}, timeout=min(30, cancel_timeout))
            state["acknowledged"] = response.get("ok") is True
        except Exception as exc:
            state["request_error"] = {"type": type(exc).__name__, "message": str(exc)}
        while time.monotonic() < deadline:
            try:
                if poll(deadline):
                    state["terminal_observed"] = True
                    break
            except Exception as exc:
                state["observation_error"] = {"type": type(exc).__name__, "message": str(exc)}
                # Continue within the bounded grace if a transient poll fails.
            remaining = deadline-time.monotonic()
            if remaining > 0:
                time.sleep(min(0.25, remaining))
        state["last_observed_status"] = turn.get("status") if turn else "unknown"
        state["still_running"] = (False if state["terminal_observed"] else
                                  True if state["last_observed_status"] in ("running", "cancelling") else None)
        state["shutdown_verified"] = state["terminal_observed"]

    try:
        token = _login_token(token_file, workspace, report_dir)
        meta = fixture(workspace)
        report["fixture"] = meta
        if live:
            if not expected_model or not isinstance(expected_model, str) or _scripted(expected_model):
                raise ValueError("Live acceptance requires an explicit non-scripted --expected-model")
            if not expected_provider_host:
                raise ValueError("Live acceptance requires --expected-provider-host")
            expected_provider_host = _host(expected_provider_host)
        capabilities = request("/api/agent/capabilities")
        # Capability responses are public metadata; never query /api/llm-config
        # (which also returns credential hints) to infer an unavailable host.
        report["capabilities"] = capabilities
        verification = model_verification(capabilities, live, expected_model, expected_provider_host)
        report["model_verification"] = verification
        if not verification["configured"]:
            raise ValueError("The product server has no configured model")
        if live and verification["scripted_model_detected"]:
            raise ValueError("Live acceptance refuses a scripted fixture model")
        if verification["model_status"] != "matched":
            raise ValueError("Configured model does not match --expected-model" if live else
                             "Default scripted acceptance requires model civil-scripted-acceptance; use --live explicitly for a real configured provider")
        if live and verification["provider_host_status"] != "matched":
            raise ValueError("Provider host is unverified: server capabilities do not expose it" if verification["provider_host_status"] == "unverified" else
                             "Configured provider host does not match --expected-provider-host")
        if not live and verification["provider_host_status"] != "matched":
            raise ValueError("Default scripted acceptance requires a verified loopback provider host; external or unverified hosts are refused")
        registered = request("/api/agent/workspaces", {"path": str(workspace)})
        workspace_id = registered["workspace"]["id"]
        session_id = "acceptance_"+uuid4().hex
        # The key permits an operator to find this exact request if the start
        # response is lost. The client deliberately does not retry paid work.
        start_key = "acceptance_"+uuid4().hex
        report.update(workspace_id=workspace_id, session_id=session_id, start_idempotency_key=start_key)
        start_attempted = True
        start = request("/api/agent/turns", {"workspace": workspace_id, "session_id": session_id,
                    "message": TASK, "mode": "model", "sandbox": "workspace-write", "files": FILES,
                    "expert_id": "pm-daily", "idempotency_key": start_key})
        turn_id = start["turn_id"]
        report["turn_id"] = turn_id
        deadline = time.monotonic()+timeout
        while not poll(deadline):
            remaining = deadline-time.monotonic()
            if remaining <= 0:
                raise TimeoutError("Acceptance task exceeded its observation wait; cancellation required")
            time.sleep(min(0.25, remaining))
        outputs = []
        for artifact in (turn.get("result") or {}).get("artifacts", []):
            url = artifact.get("url", "")
            if not url.startswith("/api/agent/artifacts/"):
                raise ValueError("Unexpected artifact URL")
            data = request(url, binary=True)
            if token and token.encode("ascii") in data:
                raise ValueError("Artifact contains a reflected login credential; download was not saved")
            filename = Path(artifact["name"]).name
            if token and _redact(filename, token) != filename:
                raise ValueError("Artifact filename contains a reflected login credential")
            destination = report_dir/"artifacts"
            destination.mkdir(exist_ok=True)
            downloaded = destination/(str(len(outputs)+1)+"-"+filename)
            downloaded.write_bytes(data)
            outputs.append({**artifact, "download_path": str(downloaded)})
        report["artifacts"] = outputs
        report["validation"] = validate_outputs(workspace, meta, outputs)
        names = [e.get("data", {}).get("name") for e in events if e.get("kind", e.get("event")) == "tool_finished"]
        report["tools_observed"] = names
        report["events_monotonic"] = all(events[i]["seq"] < events[i+1]["seq"] for i in range(len(events)-1))
        report["passed"] = turn.get("status") == "completed" and report["validation"]["passed"] and all(n in names for n in ("read_file", "search_sources", "preview_document", "apply_document")) and report["events_monotonic"]
    except (Exception, KeyboardInterrupt) as exc:
        report["error"] = {"type": type(exc).__name__, "message": str(exc) or "Acceptance interrupted"}
    finally:
        if turn_id and (not turn or turn.get("status") not in TERMINAL):
            try:
                cancel_and_drain()
            except (Exception, KeyboardInterrupt) as exc:
                report["cleanup_error"] = {"type": type(exc).__name__, "message": str(exc) or "Cancellation observation interrupted"}
        elif start_attempted and not turn_id:
            report["cancellation"] = {"requested": False, "terminal_observed": False, "still_running": None,
                                       "shutdown_verified": False, "reason": "start_outcome_unknown_no_turn_id; inspect the recorded session and idempotency key"}
        if turn and turn.get("status") in ("failed", "interrupted", "cancelled") and not report.get("error"):
            status = turn["status"]
            result = turn.get("result")
            result = result if isinstance(result, dict) else {}
            reason = next((result[key].strip() for key in ("error", "reason")
                           if isinstance(result.get(key), str) and result[key].strip()), None)
            report["error"] = {"type": "Task"+status.capitalize(),
                               "message": reason or f"The product task ended with status {status}"}
        if meta is not None:
            report["originals"] = original_hashes(workspace, meta)
            report["originals"]["snapshot_is_final"] = bool(turn and turn.get("status") in TERMINAL) or not start_attempted
            # If a turn is still active this only proves the observation time,
            # and the report explicitly refuses to imply future preservation.
            report["passed"] = report["passed"] and report["originals"]["unchanged"]
        report["usage"] = usage_report(turn, events)
        budget = report["usage"]["budget_snapshot"]
        if isinstance(budget, dict) and isinstance(budget.get("limits"), dict):
            report["server_limits"] = {"values": budget["limits"], "source": "turn_result.usage.limits",
                                       "deployed_limits_verified": True, "includes_subagents": True}
        if live:
            usage = report["usage"]
            # These are product-reported receipts, not independent billing or
            # proof that an endpoint operator really ran a particular model.
            report["live_evidence"] = {"scope": "configured host/model and product-reported model events; no independent provider attestation",
                "model_calls_positive": type(usage["model_calls"]) is int and usage["model_calls"] > 0,
                "model_events_present": usage["model_events_observed"] > 0,
                "usage_measurement_recorded": sum(usage["usage_by_measurement"][kind]["model_calls"] for kind in ("provider_reported", "estimated")) > 0,
                "response_models_match": bool(usage["response_models"]) and all(m == expected_model and not _scripted(m) for m in usage["response_models"]),
                "configuration_matched": report.get("model_verification", {}).get("model_status") == "matched" and report.get("model_verification", {}).get("provider_host_status") == "matched"}
            report["live_verified"] = all(report["live_evidence"][key] for key in ("model_calls_positive", "model_events_present", "usage_measurement_recorded", "response_models_match", "configuration_matched"))
            report["passed"] = report["passed"] and report["live_verified"]
        (report_dir/"events.jsonl").write_text(_redact("".join(json.dumps(e, ensure_ascii=False)+"\n" for e in events), token), encoding="utf-8")
        report = _write_json(report_dir/"report.json", report, token)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--serve-model", action="store_true")
    parser.add_argument("--model-port", type=int, default=18889)
    parser.add_argument("--base", default="http://127.0.0.1:8765")
    parser.add_argument("--work-dir", type=Path)
    parser.add_argument("--report-dir", type=Path)
    parser.add_argument("--timeout", type=float, default=360, help="Client task observation wait in seconds; not a server or fee limit")
    parser.add_argument("--cancel-timeout", type=float, default=15, help="Bounded wait after requesting cancellation (maximum 120 seconds)")
    parser.add_argument("--token-file", type=Path, help="Local workbench login token file outside the fixture workspace; never a provider API key")
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--expected-model", help="Required for --live; must exactly match server configuration and model events")
    parser.add_argument("--expected-provider-host", help="Required for --live, e.g. api.deepseek.com; must match capabilities.models.provider_host")
    args = parser.parse_args()
    if args.serve_model:
        server = model_server(args.model_port)
        print(f"Scripted model only: http://127.0.0.1:{server.server_port}/v1; model={MODEL}", flush=True)
        try:
            server.serve_forever()
        except KeyboardInterrupt:
            pass
        finally:
            server.server_close()
        return 0
    directory = (args.work_dir or ROOT/"work"/("unified-acceptance-"+time.strftime("%Y%m%d-%H%M%S")+"-"+uuid4().hex[:6])).resolve()
    reports = (args.report_dir or directory/"acceptance-report").resolve()
    report = acceptance(args.base, directory, reports, live=args.live, timeout=args.timeout,
                        cancel_timeout=args.cancel_timeout, token_file=args.token_file,
                        expected_model=args.expected_model, expected_provider_host=args.expected_provider_host)
    print(json.dumps({"passed": report["passed"], "mode": report["mode"], "report": str(reports/"report.json"), "error": report.get("error")}, ensure_ascii=False), flush=True)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
